"""Tests for the layout and deck rules of scripts/design_lint.py: geometry, pictures, alt text, counts and notes.
"""

import contextlib
import io
import sys
from pathlib import Path
from unittest import mock

from lxml import etree
from pptx.oxml import parse_xml
from pptx.oxml.ns import nsdecls, qn
from pptx.slide import Slide
from pptx.util import Inches

SCRIPTS_DIR = Path(__file__).resolve().parents[1] / "scripts"
for folder in (SCRIPTS_DIR, Path(__file__).resolve().parent):
    if str(folder) not in sys.path:
        sys.path.insert(0, str(folder))

import design_lint as dl  # noqa: E402
from lint_helpers import (  # noqa: E402
    BODY,
    GOOD_TITLE,
    LintCase,
    add_box,
    add_chart,
    add_notes,
    add_picture,
    add_rect,
    add_slide,
    add_source,
    add_table,
    clean_slide,
    damage_blip_without_embed,
    deck_of,
    new_deck,
    set_alt,
)


class InsideSlide(LintCase):
    def probe(self, **place):
        prs = new_deck()
        add_box(add_slide(prs), BODY, name="Probe", **place)
        return self.lint(prs)

    def test_a_shape_beyond_an_edge_is_an_error_and_the_message_names_it(self):
        for place in ({"left": 8.0, "width": 6.0}, {"left": -0.2}, {"top": 6.5, "height": 1.5}, {"top": -0.1}):
            with self.subTest(place=place):
                hit = self.assertFires(self.probe(**place), "inside-slide", "error", "T3")
                self.assertIn("Probe", hit.message)

    def test_a_shape_that_fills_the_slide_exactly_is_inside(self):
        prs = new_deck()
        slide = add_slide(prs)
        panel = add_rect(slide, "Panel", 0, 0, 1, 1)
        panel.left, panel.top, panel.width, panel.height = 0, 0, prs.slide_width, prs.slide_height
        self.assertQuiet(self.lint(prs), "inside-slide")
        panel.width = prs.slide_width + 1
        self.assertFires(self.lint(prs), "inside-slide", "error", "T3")

    def test_a_group_moved_beyond_the_edge_is_an_error(self):
        prs = new_deck()
        slide = add_slide(prs)
        one = add_box(slide, "One", name="One", left=1.0, top=2.0, width=2.0, height=0.5)
        two = add_box(slide, "Two", name="Two", left=1.0, top=3.0, width=2.0, height=0.5)
        group = slide.shapes.add_group_shape([one, two])
        self.assertQuiet(self.lint(prs), "inside-slide")
        group.left = group.left + Inches(12)
        hits = self.of(self.lint(prs), "inside-slide")
        self.assertEqual(sorted(h.severity for h in hits), ["error", "error"])

    def test_a_shape_name_with_a_line_break_gives_a_one_line_message(self):
        prs = new_deck()
        add_box(add_slide(prs), BODY, name="Probe\nsecond line", left=8.0, width=6.0)
        hit = self.only(self.lint(prs), "inside-slide")
        self.assertNotIn("\n", hit.message)
        self.assertIn("Probe second line", hit.message)


class RotatedShapes(LintCase):
    def rotated_deck(self, angle, count=1):
        prs = new_deck()
        slide = add_slide(prs)
        for index in range(count):
            box = add_box(slide, "Axis label", name="Rotated label %d" % index, left=-1.5, top=3.0 + index, width=4.0,
                          height=0.5)
            box.rotation = angle
        return prs

    def test_a_box_rotated_by_270_degrees_has_no_false_inside_error(self):
        findings = self.lint(self.rotated_deck(270))
        self.assertQuiet(findings, "margin")
        hit = self.only(findings, "inside-slide")
        self.assertEqual((hit.slide, hit.severity), (1, "warn"))
        self.assertIn("rotated", hit.message)
        self.assertIn("inside-slide and margin", hit.message)
        self.assertIn("rule skipped", hit.message)
        self.assertIn("Rotated label 0", hit.message)

    def test_two_rotated_boxes_give_one_note_for_the_slide(self):
        findings = self.lint(self.rotated_deck(90, count=2))
        hit = self.only(findings, "inside-slide")
        self.assertIn("Rotated label 0", hit.message)
        self.assertIn("Rotated label 1", hit.message)

    def test_a_box_with_no_rotation_is_still_an_error_beyond_the_edge(self):
        hit = self.assertFires(self.lint(self.rotated_deck(0)), "inside-slide", "error", "T3")
        self.assertNotIn("rotated", hit.message)

    def test_a_shape_in_a_rotated_group_is_skipped(self):
        prs = new_deck()
        slide = add_slide(prs)
        one = add_box(slide, "One", name="One", left=-1.5, top=3.0, width=4.0, height=0.5)
        group = slide.shapes.add_group_shape([one])
        group.rotation = 90
        hit = self.only(self.lint(prs), "inside-slide")
        self.assertEqual(hit.severity, "warn")

    def test_an_unreadable_rotation_does_not_skip_the_rule(self):
        prs = self.rotated_deck(0)
        box = [s for s in list(prs.slides)[0].shapes if s.name == "Rotated label 0"][0]
        box._element.find(qn("p:spPr")).find(qn("a:xfrm")).set("rot", "turn")
        self.assertFires(self.lint(prs), "inside-slide", "error", "T3")


class Margin(LintCase):
    def probe(self, **place):
        prs = new_deck()
        add_box(add_slide(prs), BODY, name="Probe", **place)
        return self.lint(prs)

    def test_a_side_closer_than_half_an_inch_warns(self):
        hit = self.assertFires(self.probe(left=0.3), "margin", "warn", "T3")
        self.assertIn("Probe", hit.message)
        self.assertFires(self.probe(left=5.0, width=8.033), "margin", "warn", "T3")

    def test_exactly_half_an_inch_passes(self):
        self.assertQuiet(self.probe(left=0.5), "margin")
        self.assertQuiet(self.probe(left=4.833, width=8.0), "margin")

    def test_a_body_box_closer_than_half_an_inch_to_the_top_or_the_bottom_warns(self):
        hit = self.assertFires(self.probe(top=0.2, height=0.5), "margin", "warn", "T3")
        self.assertIn("0.20 in from the top edge", hit.message)
        hit = self.assertFires(self.probe(top=6.3, height=0.9), "margin", "warn", "T3")
        self.assertIn("0.30 in from the bottom edge", hit.message)

    def test_exactly_half_an_inch_from_the_top_or_the_bottom_passes(self):
        self.assertQuiet(self.probe(top=0.5, height=0.5), "margin")
        self.assertQuiet(self.probe(top=6.0, height=1.0), "margin")

    def test_the_title_and_the_source_line_keep_their_look_and_need_no_top_or_bottom_margin(self):
        # The look puts the title 0.3 in from the top and the source line 0.25 in from the bottom.
        prs = new_deck()
        slide = add_slide(prs)
        add_source(slide)
        add_box(slide, BODY, name="Footer", top=7.05, height=0.3, left=0.8)
        findings = self.lint(prs)
        self.assertEqual([f.message.split(" is ")[0] for f in self.of(findings, "margin")], ["Footer"])

    def test_a_slide_number_placeholder_needs_no_top_or_bottom_margin(self):
        prs = new_deck()
        slide = add_slide(prs)
        for placeholder in slide.slide_layout.placeholders:
            if "SLIDE_NUMBER" in str(placeholder.placeholder_format.type):
                slide.shapes.clone_placeholder(placeholder)
        shape = [s for s in slide.shapes if s.is_placeholder and "SLIDE_NUMBER" in str(s.placeholder_format.type)][0]
        shape.left, shape.top, shape.width, shape.height = Inches(10), Inches(7.1), Inches(2), Inches(0.3)
        self.assertQuiet(self.lint(prs), "margin")

    def test_a_full_slide_background_picture_is_exempt_and_a_smaller_one_is_not(self):
        def picture(scale):
            prs = new_deck()
            slide = add_slide(prs)
            width, height = int(prs.slide_width * scale), int(prs.slide_height * scale)
            shape = slide.shapes.add_picture(str(self.png("bg.png", 4000, 2250)),
                                             (prs.slide_width - width) // 2, (prs.slide_height - height) // 2,
                                             width, height)
            shape.name = "Background"
            set_alt(shape, "A plain background")
            return self.lint(prs)
        self.assertQuiet(picture(1.0), "margin")
        self.assertQuiet(picture(0.98), "margin")
        self.assertFires(picture(0.97), "margin", "warn", "T3")

    def test_a_full_slide_picture_beyond_the_edge_is_still_an_error(self):
        prs = new_deck()
        slide = add_slide(prs)
        shape = slide.shapes.add_picture(str(self.png("bg.png", 4000, 2250)), -10, 0, prs.slide_width + 20,
                                         prs.slide_height)
        set_alt(shape, "A plain background")
        self.assertFires(self.lint(prs), "inside-slide", "error", "T3")

    def test_the_margin_message_gives_the_limit(self):
        prs = new_deck()
        add_box(add_slide(prs), BODY, name="Probe", left=0.3)
        hit = self.assertFires(self.lint(prs), "margin", "warn", "T3")
        self.assertIn("keep at least 0.5 in", hit.message)
        self.assertIn("0.30 in from the left edge", hit.message)


class AspectRatio(LintCase):
    def sized(self, width, height):
        prs = new_deck(width, height)
        clean_slide(prs)
        return self.lint(prs)

    def test_a_4_to_3_deck_is_an_error_on_the_deck(self):
        hit = self.assertFires(self.sized(10, 7.5), "aspect", "error", "T2", slide=0)
        self.assertIn("16:9", hit.message)

    def test_a_16_to_10_deck_is_an_error(self):
        self.assertFires(self.sized(13.333, 8.333), "aspect", "error", "T2", slide=0)

    def test_16_to_9_decks_pass(self):
        self.assertQuiet(self.sized(13.333, 7.5), "aspect")
        self.assertQuiet(self.sized(12, 6.75), "aspect")

    def test_a_deck_with_no_slide_size_is_an_error(self):
        prs = new_deck()
        clean_slide(prs)
        size = prs._element.find(qn("p:sldSz"))
        size.getparent().remove(size)
        hit = self.assertFires(self.lint(prs), "aspect", "error", "T2", slide=0)
        self.assertIn("no usable slide size", hit.message)


class PictureResolution(LintCase):
    def shown(self, pixels_w, pixels_h, width_in, height_in):
        prs = new_deck()
        slide = add_slide(prs)
        add_picture(slide, self.png("figure.png", pixels_w, pixels_h), left=1.0, top=1.5, width=width_in,
                    height=height_in, alt="A figure of made up data")
        return self.lint(prs)

    def test_144_ppi_passes_and_143_9_ppi_warns(self):
        self.assertQuiet(self.shown(1440, 810, 10, 5.625), "picture-ppi")
        hit = self.assertFires(self.shown(1439, 810, 10, 5.625), "picture-ppi", "warn", "T3")
        self.assertIn("143.9 ppi", hit.message)

    def test_a_small_image_on_a_large_box_warns(self):
        self.assertFires(self.shown(300, 200, 6, 4), "picture-ppi", "warn", "T3")

    def test_the_lower_axis_counts_for_a_stretched_picture(self):
        self.assertFires(self.shown(2000, 400, 10, 5), "picture-ppi", "warn", "T3")

    def test_a_full_slide_picture_gets_the_check(self):
        prs = new_deck()
        slide = add_slide(prs)
        shape = slide.shapes.add_picture(str(self.png("bg.png", 1000, 562)), 0, 0, prs.slide_width,
                                         prs.slide_height)
        set_alt(shape, "A plain background")
        self.assertFires(self.lint(prs), "picture-ppi", "warn", "T3")

    def test_a_picture_with_an_unknown_resolution_gets_a_note(self):
        prs = new_deck()
        damage_blip_without_embed(self, add_slide(prs))
        hits = self.of(self.lint(prs), "picture-ppi")
        self.assertEqual(len(hits), 1)
        self.assertIn("rule skipped", hits[0].message)


class AltText(LintCase):
    def picture(self, alt):
        prs = new_deck()
        slide = add_slide(prs)
        add_picture(slide, self.png("figure.png", 2000, 1200), alt=alt)
        return self.lint(prs)

    def test_the_file_name_that_python_pptx_writes_warns(self):
        hit = self.assertFires(self.picture(None), "alt-text", "warn", "T2")
        self.assertIn("Evidence picture", hit.message)
        self.assertIn("file name", hit.message)

    def test_an_empty_alt_text_warns(self):
        self.assertFires(self.picture(""), "alt-text", "warn", "T2")
        self.assertFires(self.picture("   "), "alt-text", "warn", "T2")

    def test_a_generic_default_warns(self):
        for alt in ("Picture 3", "image.png", "Image", "Content Placeholder 2", "Evidence picture", "SCAN.JPG"):
            with self.subTest(alt=alt):
                self.assertFires(self.picture(alt), "alt-text", "warn", "T2")

    def test_a_real_description_is_quiet(self):
        self.assertQuiet(self.picture("Photo of a made up gel with one band"), "alt-text")

    def test_a_chart_needs_alt_text_too(self):
        def chart(alt):
            prs = new_deck()
            add_chart(add_slide(prs), labels=True, alt=alt)
            return self.lint(prs)
        self.assertFires(chart(""), "alt-text", "warn", "T2")
        self.assertFires(chart("Chart 4"), "alt-text", "warn", "T2")
        self.assertQuiet(chart("Bars of made up hit counts for three targets"), "alt-text")

    DECORATIVE = (
        '<a:extLst %s><a:ext uri="{C183D7F6-B498-43B3-948B-1728B52AA6E4}">'
        '<adec:decorative xmlns:adec="http://schemas.microsoft.com/office/drawing/2017/decorative" val="%s"/>'
        '</a:ext></a:extLst>')

    def mark_decorative(self, shape, value="1"):
        shape._element.xpath("./*[1]/p:cNvPr")[0].append(parse_xml(self.DECORATIVE % (nsdecls("a"), value)))

    def test_a_picture_marked_decorative_needs_no_alt_text(self):
        prs = new_deck()
        picture = add_picture(add_slide(prs), self.png("rule.png", 2000, 1200), alt="")
        self.mark_decorative(picture)
        self.assertQuiet(self.lint(prs), "alt-text")

    def test_a_chart_marked_decorative_needs_no_alt_text(self):
        prs = new_deck()
        frame = add_chart(add_slide(prs), labels=True, alt="")
        self.mark_decorative(frame)
        self.assertQuiet(self.lint(prs), "alt-text")

    def test_a_comment_in_the_picture_properties_does_not_stop_the_check(self):
        for decorative in (True, False):
            with self.subTest(decorative=decorative):
                prs = new_deck()
                picture = add_picture(add_slide(prs), self.png("rule.png", 2000, 1200), alt="")
                node = picture._element.xpath("./*[1]/p:cNvPr")[0]
                node.append(etree.Comment(" written by another tool "))
                if decorative:
                    self.mark_decorative(picture)
                findings = self.lint(prs)
                self.assertEqual(len(self.of(findings, "alt-text")), 0 if decorative else 1)

    def test_a_picture_with_decorative_set_to_false_still_warns(self):
        for value in ("0", "false"):
            with self.subTest(value=value):
                prs = new_deck()
                picture = add_picture(add_slide(prs), self.png("rule.png", 2000, 1200), alt="")
                self.mark_decorative(picture, value)
                self.assertFires(self.lint(prs), "alt-text", "warn", "T2")

    def test_decorative_is_read_when_the_value_is_true(self):
        prs = new_deck()
        picture = add_picture(add_slide(prs), self.png("rule.png", 2000, 1200), alt="")
        self.mark_decorative(picture, "true")
        self.assertQuiet(self.lint(prs), "alt-text")

    def test_a_table_is_not_checked(self):
        prs = new_deck()
        add_table(add_slide(prs))
        self.assertQuiet(self.lint(prs), "alt-text")


class SlideCount(LintCase):
    def test_more_main_slides_than_minutes_warns_on_the_deck(self):
        hit = self.assertFires(self.lint(deck_of([None] * 6), minutes=5), "slide-count", "warn", "T2", slide=0)
        self.assertIn("6 main slides", hit.message)
        self.assertIn("5 minutes", hit.message)

    def test_as_many_main_slides_as_minutes_passes(self):
        self.assertQuiet(self.lint(deck_of([None] * 5), minutes=5), "slide-count")

    def test_title_backup_and_hidden_slides_do_not_count(self):
        prs = deck_of(["title"] + [None] * 4 + ["backup", "backup", None])
        list(prs.slides)[-1]._element.set("show", "0")
        self.assertQuiet(self.lint(prs, minutes=4), "slide-count")

    def test_without_minutes_the_rule_is_skipped_with_a_note_on_the_deck(self):
        hit = self.assertFires(self.lint(deck_of([None] * 6)), "slide-count", "warn", "T2", slide=0)
        self.assertIn("--minutes", hit.message)
        self.assertIn("rule skipped", hit.message)

    def test_the_lab_meeting_profile_allows_one_and_a_half_slides_per_minute(self):
        self.assertQuiet(self.lint(deck_of([None] * 15), minutes=10, profile="lab-meeting"), "slide-count")
        self.assertFires(self.lint(deck_of([None] * 16), minutes=10, profile="lab-meeting"), "slide-count",
                         "warn", "T2", slide=0)

    def test_minutes_can_be_a_fraction(self):
        self.assertQuiet(self.lint(deck_of([None] * 2), minutes=2.5), "slide-count")
        self.assertFires(self.lint(deck_of([None] * 3), minutes=2.5), "slide-count", "warn", "T2", slide=0)

    def test_a_deck_with_only_hidden_slides_is_an_error(self):
        prs = deck_of([None, None])
        for slide in prs.slides:
            slide._element.set("show", "0")
        findings = self.lint(prs, minutes=5)
        self.assertEqual(findings, [dl.Finding(0, "slide-count", "error", "T2", "the deck has no visible slides")])

    def test_a_deck_with_one_visible_slide_is_not(self):
        prs = deck_of([None, None])
        list(prs.slides)[0]._element.set("show", "0")
        self.assertEqual(self.of(self.lint(prs, minutes=5), "slide-count"), [])

    def test_a_centered_title_placeholder_on_a_slide_with_a_kind_does_not_make_it_a_title_slide(self):
        prs = new_deck()
        slide = prs.slides.add_slide(prs.slide_layouts.get_by_name("Title Slide"))
        slide.shapes.title.text_frame.text = "Binder study"
        slide._element.cSld.set("name", "assertion_evidence")
        self.assertFires(self.lint(prs, minutes=0.5), "slide-count", "warn", "T2", slide=0)

    def test_a_slide_with_a_centered_title_placeholder_is_a_title_slide_and_does_not_count(self):
        prs = new_deck()
        slide = prs.slides.add_slide(prs.slide_layouts.get_by_name("Title Slide"))
        slide.shapes.title.text_frame.text = "Binder study"
        for index in range(2):
            clean_slide(prs, title="Binder design succeeds for %d of 15 targets" % index)
        self.assertEqual(self.of(self.lint(prs, minutes=2), "slide-count"), [])
        self.assertFires(self.lint(prs, minutes=1), "slide-count", "warn", "T2", slide=0)

    def test_minutes_must_be_positive_and_finite(self):
        for minutes in (0, -3, float("inf"), float("nan")):
            with self.subTest(minutes=minutes):
                with self.assertRaises(ValueError):
                    self.lint(deck_of([None]), minutes=minutes)

    def test_a_boolean_is_not_a_number_of_minutes(self):
        prs = new_deck()
        clean_slide(prs)
        with self.assertRaises(ValueError):
            self.lint(prs, minutes=True)

    def test_fractional_minutes_reach_the_message(self):
        out = io.StringIO()
        with contextlib.redirect_stdout(out):
            code = dl.main([str(self.save(deck_of([None] * 3))), "--no-notes", "--minutes", "2.5"])
        self.assertEqual(code, 0)
        self.assertIn("3 main slides for 2.5 minutes", out.getvalue())


class TakeHomes(LintCase):
    def with_shapes(self, kind, name, count):
        prs = new_deck()
        slide = add_slide(prs, "Three things to remember from this talk", kind=kind)
        for index in range(count):
            add_box(slide, "Item", name=name, left=0.8 + index * 1.9, width=1.8, height=0.5)
        return self.lint(prs)

    def test_four_claims_warn_and_three_pass(self):
        hit = self.assertFires(self.with_shapes("take_home", "Claim", 4), "take-homes", "warn", "T2")
        self.assertIn("4 claims", hit.message)
        self.assertQuiet(self.with_shapes("take_home", "Claim", 3), "take-homes")

    def test_six_roadmap_sections_warn_and_five_pass(self):
        hit = self.assertFires(self.with_shapes("roadmap", "Section name", 6), "take-homes", "warn", "T2")
        self.assertIn("6 sections", hit.message)
        self.assertQuiet(self.with_shapes("roadmap", "Section name", 5), "take-homes")

    def test_a_slide_with_no_kind_gets_the_generic_rules_only(self):
        self.assertQuiet(self.with_shapes(None, "Claim", 5), "take-homes")
        self.assertQuiet(self.with_shapes(None, "Section name", 6), "take-homes")

    def test_the_limits_come_from_the_style(self):
        prs = new_deck()
        slide = add_slide(prs, kind="take_home")
        for index in range(4):
            add_box(slide, "Item", name="Claim", left=0.8 + index * 1.9, width=1.8, height=0.5)
        project = self.tmp / "project"
        project.mkdir()
        (project / "slide-style-minimal.yaml").write_text("lint:\n  take_homes: {max: 4, tier: T1}\n")
        self.assertQuiet(self.lint(prs, project_dir=str(project)), "take-homes")


class Notes(LintCase):
    def with_notes(self, text, kind=None):
        prs = new_deck()
        slide = clean_slide(prs, kind=kind)
        if text is not None:
            add_notes(slide, text)
        return self.lint(prs, notes=True)

    def test_a_main_slide_with_no_notes_warns(self):
        hit = self.assertFires(self.with_notes(None), "notes", "warn", "T3")
        self.assertIn("no speaker notes", hit.message)

    def test_a_slide_with_a_centered_title_placeholder_needs_no_notes(self):
        prs = new_deck()
        slide = prs.slides.add_slide(prs.slide_layouts.get_by_name("Title Slide"))
        slide.shapes.title.text_frame.text = "Binder study"
        self.assertQuiet(self.lint(prs, notes=True), "notes")

    def test_blank_notes_warn(self):
        self.assertFires(self.with_notes("   "), "notes", "warn", "T3")

    def test_notes_with_text_are_quiet(self):
        self.assertQuiet(self.with_notes("Speak about the assay and the controls."), "notes")

    def test_the_title_slide_and_a_backup_slide_need_no_notes(self):
        self.assertQuiet(self.with_notes(None, kind="title"), "notes")
        self.assertQuiet(self.with_notes(None, kind="backup"), "notes")

    def test_notes_off_switches_both_rules_off(self):
        prs = new_deck()
        slide = clean_slide(prs)
        self.assertFires(self.lint(prs, notes=True), "notes", "warn", "T3")
        self.assertQuiet(self.lint(prs, notes=False), "notes")
        add_notes(slide, GOOD_TITLE + " " + BODY)
        self.assertQuiet(self.lint(prs, notes=False), "notes")

    def redundant(self, shared):
        prs = new_deck()
        slide = add_slide(prs, "Alpha beta gamma delta epsilon zeta eta theta")
        add_box(slide, "iota kappa")
        slide_words = "alpha beta gamma delta epsilon zeta eta theta iota kappa".split()
        add_notes(slide, " ".join(slide_words[:shared]) + " plus some other speech that adds new content")
        return self.lint(prs, notes=True)

    def test_more_than_80_percent_of_the_words_in_the_notes_warns(self):
        self.assertQuiet(self.redundant(8), "notes")
        hit = self.assertFires(self.redundant(9), "notes", "warn", "T3")
        self.assertIn("appear in the notes", hit.message)

    def test_the_comparison_ignores_case_and_punctuation(self):
        prs = new_deck()
        slide = add_slide(prs, "Alpha, beta; gamma. Delta!")
        add_notes(slide, "ALPHA beta GAMMA delta")
        self.assertFires(self.lint(prs, notes=True), "notes", "warn", "T3")

    def test_reading_never_creates_a_notes_slide(self):
        prs = new_deck()
        clean_slide(prs)
        path = self.save(prs)
        with mock.patch.object(Slide, "notes_slide", new_callable=mock.PropertyMock,
                               side_effect=AssertionError("created a notes slide")):
            findings = dl.lint(path, notes=True)
        self.assertEqual(len(self.of(findings, "notes")), 1)

    def test_a_notes_slide_with_no_body_placeholder_counts_as_no_notes(self):
        prs = new_deck()
        slide = clean_slide(prs)
        add_notes(slide, "Some notes")
        for placeholder in list(slide.notes_slide.placeholders):
            if "BODY" in str(placeholder.placeholder_format.type):
                placeholder._element.getparent().remove(placeholder._element)
        self.assertIsNone(slide.notes_slide.notes_text_frame)
        self.assertFires(self.lint(prs, notes=True), "notes", "warn", "T3")
