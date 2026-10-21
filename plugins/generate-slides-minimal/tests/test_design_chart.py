"""Tests for scripts/design_chart.py: the fill reader and the chart reader.
"""

import sys
import unittest
from pathlib import Path

from pptx.chart.data import CategoryChartData, XyChartData
from pptx.dml.color import RGBColor
from pptx.enum.chart import XL_CHART_TYPE
from pptx.oxml import parse_xml
from pptx.oxml.ns import nsdecls
from pptx.util import Inches

SCRIPTS_DIR = Path(__file__).resolve().parents[1] / "scripts"
for folder in (SCRIPTS_DIR, Path(__file__).resolve().parent):
    if str(folder) not in sys.path:
        sys.path.insert(0, str(folder))

import design_chart  # noqa: E402
from lint_helpers import add_chart, add_slide, new_deck  # noqa: E402


class SrgbColor(unittest.TestCase):
    def test_a_color_needs_exactly_one_plain_child(self):
        solid = parse_xml('<a:solidFill %s><a:srgbClr val="112233"/><a:srgbClr val="445566"/></a:solidFill>'
                          % nsdecls("a"))
        self.assertIsNone(design_chart.srgb(solid))
        self.assertIsNone(design_chart.srgb(parse_xml('<a:solidFill %s/>' % nsdecls("a"))))
        self.assertEqual(design_chart.srgb(parse_xml('<a:solidFill %s><a:srgbClr val="1a2b3c"/></a:solidFill>'
                                            % nsdecls("a"))), "#1A2B3C")
        self.assertIsNone(design_chart.srgb(parse_xml('<a:solidFill %s><a:srgbClr val="1a2b"/></a:solidFill>'
                                             % nsdecls("a"))))


class SeriesColors(unittest.TestCase):
    def chart(self, **options):
        prs = new_deck()
        slide = add_slide(prs)
        frame = add_chart(slide, **options)
        return design_chart.Chart(frame, "Chart", "#FFFFFF"), frame

    def test_a_bar_series_reads_its_fill_and_not_its_outline(self):
        chart, _ = self.chart(colors=("1F6FEB",), outline="FFFFFF")
        self.assertEqual([s.colors for s in chart.series], [["#1F6FEB"]])

    def test_a_bar_series_with_two_fills_and_an_outline_lists_both_fills(self):
        chart, frame = self.chart(colors=("1F6FEB",), outline="FFFFFF")
        point = frame.chart.plots[0].series[0].points[1]
        point.format.fill.solid()
        point.format.fill.fore_color.rgb = RGBColor.from_string("D55E00")
        point.format.line.color.rgb = RGBColor.from_string("FFFFFF")
        chart = design_chart.Chart(frame, "Chart", "#FFFFFF")
        self.assertEqual([s.colors for s in chart.series], [["#1F6FEB", "#D55E00"]])

    def test_a_line_series_reads_the_color_of_its_line(self):
        chart, _ = self.chart(colors=("0072B2", "D55E00"), line=True)
        self.assertEqual([s.colors for s in chart.series], [["#0072B2"], ["#D55E00"]])

    def test_a_scatter_series_reads_the_marker_fill_and_the_line(self):
        prs = new_deck()
        slide = add_slide(prs)
        data = XyChartData()
        series = data.add_series("s0")
        for x, y in ((1, 2), (2, 3), (3, 5)):
            series.add_data_point(x, y)
        frame = slide.shapes.add_chart(XL_CHART_TYPE.XY_SCATTER_LINES, Inches(1), Inches(2), Inches(6), Inches(4), data)
        plot_series = frame.chart.plots[0].series[0]
        plot_series.format.line.color.rgb = RGBColor.from_string("0072B2")
        plot_series.marker.format.fill.solid()
        plot_series.marker.format.fill.fore_color.rgb = RGBColor.from_string("D55E00")
        plot_series.marker.format.line.color.rgb = RGBColor.from_string("FFFFFF")
        chart = design_chart.Chart(frame, "Chart", "#FFFFFF")
        self.assertEqual(sorted(chart.series[0].colors), ["#0072B2", "#D55E00"])

    def test_a_bar_with_no_fill_element_and_an_outline_has_no_color_because_the_theme_fills_it(self):
        chart, _ = self.chart(colors=(None,), outline="FFFFFF")
        self.assertEqual(chart.series[0].colors, [])

    def test_the_outline_is_read_only_for_a_line_a_scatter_and_a_radar_series(self):
        for kind in (XL_CHART_TYPE.LINE, XL_CHART_TYPE.RADAR):
            with self.subTest(kind=str(kind)):
                prs = new_deck()
                slide = add_slide(prs)
                data = CategoryChartData()
                data.categories = ["a", "b", "c"]
                data.add_series("s", (1, 2, 3))
                frame = slide.shapes.add_chart(kind, Inches(1), Inches(2), Inches(6), Inches(4), data)
                frame.chart.plots[0].series[0].format.line.color.rgb = RGBColor.from_string("0072B2")
                chart = design_chart.Chart(frame, "Chart", "#FFFFFF")
                self.assertEqual(chart.series[0].colors, ["#0072B2"])

    def test_an_area_series_with_an_outline_and_no_fill_has_no_color(self):
        prs = new_deck()
        slide = add_slide(prs)
        data = CategoryChartData()
        data.categories = ["a", "b", "c"]
        data.add_series("s", (1, 2, 3))
        frame = slide.shapes.add_chart(XL_CHART_TYPE.AREA, Inches(1), Inches(2), Inches(6), Inches(4), data)
        frame.chart.plots[0].series[0].format.line.color.rgb = RGBColor.from_string("FFFFFF")
        chart = design_chart.Chart(frame, "Chart", "#FFFFFF")
        self.assertEqual(chart.series[0].colors, [])

    def test_a_gradient_fill_gives_no_color_even_with_an_outline(self):
        chart, frame = self.chart(colors=("1F6FEB",), outline="FFFFFF")
        series = frame.chart.plots[0].series[0]
        series.format.fill.gradient()
        chart = design_chart.Chart(frame, "Chart", "#FFFFFF")
        self.assertEqual(chart.series[0].colors, [])

    def test_a_scatter_marker_with_a_gradient_gives_no_color_even_with_an_outline(self):
        prs = new_deck()
        slide = add_slide(prs)
        data = XyChartData()
        series = data.add_series("s0")
        for x, y in ((1, 2), (2, 3), (3, 5)):
            series.add_data_point(x, y)
        frame = slide.shapes.add_chart(XL_CHART_TYPE.XY_SCATTER_LINES, Inches(1), Inches(2), Inches(6), Inches(4), data)
        plot_series = frame.chart.plots[0].series[0]
        plot_series.format.line.color.rgb = RGBColor.from_string("0072B2")
        plot_series.marker.format.fill.gradient()
        plot_series.marker.format.line.color.rgb = RGBColor.from_string("FFFFFF")
        chart = design_chart.Chart(frame, "Chart", "#FFFFFF")
        self.assertEqual(chart.series[0].colors, ["#0072B2"])

    def line_chart(self, kind):
        prs = new_deck()
        slide = add_slide(prs)
        data = CategoryChartData()
        data.categories = ["a", "b", "c"]
        data.add_series("s", (1, 2, 3))
        frame = slide.shapes.add_chart(kind, Inches(1), Inches(2), Inches(6), Inches(4), data)
        series = frame.chart.plots[0].series[0]
        series.format.line.color.rgb = RGBColor.from_string("0072B2")
        return frame, series

    def test_a_white_marker_edge_with_no_marker_fill_is_not_a_series_color(self):
        frame, series = self.line_chart(XL_CHART_TYPE.LINE_MARKERS)
        series.marker.format.line.color.rgb = RGBColor.from_string("FFFFFF")
        chart = design_chart.Chart(frame, "Chart", "#FFFFFF")
        self.assertEqual(chart.series[0].colors, ["#0072B2"])

    def test_a_white_edge_on_one_marker_is_not_a_series_color_either(self):
        frame, series = self.line_chart(XL_CHART_TYPE.LINE_MARKERS)
        series.points[1].marker.format.line.color.rgb = RGBColor.from_string("FFFFFF")
        chart = design_chart.Chart(frame, "Chart", "#FFFFFF")
        self.assertEqual(chart.series[0].colors, ["#0072B2"])

    def test_a_line_series_with_a_direct_gradient_or_no_fill_has_no_color_even_with_a_line(self):
        for fill in ("gradient", "background"):
            with self.subTest(fill=fill):
                frame, series = self.line_chart(XL_CHART_TYPE.LINE)
                getattr(series.format.fill, fill)()
                chart = design_chart.Chart(frame, "Chart", "#FFFFFF")
                self.assertEqual(chart.series[0].colors, [])

    def test_a_marker_fill_is_still_a_series_color(self):
        frame, series = self.line_chart(XL_CHART_TYPE.LINE_MARKERS)
        series.marker.format.fill.solid()
        series.marker.format.fill.fore_color.rgb = RGBColor.from_string("D55E00")
        series.marker.format.line.color.rgb = RGBColor.from_string("FFFFFF")
        chart = design_chart.Chart(frame, "Chart", "#FFFFFF")
        self.assertEqual(sorted(chart.series[0].colors), ["#0072B2", "#D55E00"])

    def test_a_no_fill_bar_gives_no_color(self):
        chart, frame = self.chart(colors=("1F6FEB",), outline="FFFFFF")
        frame.chart.plots[0].series[0].format.fill.background()
        chart = design_chart.Chart(frame, "Chart", "#FFFFFF")
        self.assertEqual(chart.series[0].colors, [])
