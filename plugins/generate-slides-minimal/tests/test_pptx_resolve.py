"""Tests for scripts/pptx_resolve.py: the inherited fonts, colors and boxes of a deck.

Most tests build a small deck with python-pptx, so the values come from inheritance.
One class builds a deck with the slide helpers. One class renders a deck with LibreOffice
and compares the resolver sizes with the sizes in the PDF. That class skips when soffice
or pdftotext is missing, and when soffice is installed but cannot render.
"""
import os
import re
import shutil
import signal
import subprocess
import sys
import tempfile
import unittest
from pathlib import Path
from unittest import mock

from lxml import etree
from PIL import Image
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.dml import MSO_THEME_COLOR
from pptx.enum.shapes import MSO_SHAPE, PP_PLACEHOLDER
from pptx.enum.text import MSO_AUTO_SIZE
from pptx.opc.constants import RELATIONSHIP_TARGET_MODE as RTM
from pptx.opc.constants import RELATIONSHIP_TYPE as RT
from pptx.oxml import parse_xml
from pptx.oxml.ns import nsdecls, qn
from pptx.shapes.group import GroupShape
from pptx.util import Inches, Pt

SCRIPTS_DIR = Path(__file__).resolve().parents[1] / "scripts"
if str(SCRIPTS_DIR) not in sys.path:
    sys.path.insert(0, str(SCRIPTS_DIR))

import pptx_resolve as rs  # noqa: E402

EMU_PER_INCH = 914400


# ---------------------------------------------------------------------------
# Helpers
# ---------------------------------------------------------------------------

class DeckCase(unittest.TestCase):
    """A test case with a temporary folder for decks and pictures."""

    def setUp(self):
        self._tmp = tempfile.TemporaryDirectory()
        self.addCleanup(self._tmp.cleanup)
        self.tmp = Path(self._tmp.name)

    def reopen(self, prs):
        """Save the deck and load it again, so the resolver reads a parsed file."""
        path = self.tmp / "deck.pptx"
        prs.save(str(path))
        return Presentation(str(path))

    def png(self, name, width, height):
        path = self.tmp / name
        Image.new("RGB", (width, height), (30, 60, 200)).save(str(path))
        return path


def new_slide(layout_name="Blank", prs=None):
    prs = Presentation() if prs is None else prs
    return prs, prs.slides.add_slide(prs.slide_layouts.get_by_name(layout_name))


def add_box(slide, text, size=None, top=1.0, name=None):
    """Add a text box with one run. Return the shape."""
    box = slide.shapes.add_textbox(Inches(1), Inches(top), Inches(4), Inches(0.6))
    run = box.text_frame.paragraphs[0].add_run()
    run.text = text
    if size is not None:
        run.font.size = Pt(size)
    if name is not None:
        box.name = name
    return box


def shape_named(slide, name):
    for shape in slide.shapes:
        if shape.name == name:
            return shape
    raise AssertionError("no shape named {!r}".format(name))


def resolve(prs, slide, shape, index=0):
    """Resolve the run number `index` of the shape."""
    paragraph, run = list(rs.iter_text_runs(shape))[index]
    return rs.resolve_run(prs, slide, shape, paragraph, run)


def text_body(shape):
    return shape._element.find(qn("p:txBody"))


def list_style_of(shape):
    return text_body(shape).find(qn("a:lstStyle"))


def level_style(level, props_xml):
    """Return an a:lvl{n}pPr element that holds the given a:defRPr XML."""
    tag = "lvl{}pPr".format(level + 1)
    return parse_xml("<a:{0} {1}>{2}</a:{0}>".format(tag, nsdecls("a"), props_xml))


def set_list_style_size(shape, level, hundredths):
    """Set sz on a:lvl{n}pPr/a:defRPr in the list style of a shape. Add the elements when missing."""
    list_style = list_style_of(shape)
    level_element = list_style.find(qn("a:lvl{}pPr".format(level + 1)))
    if level_element is None:
        list_style.append(level_style(level, '<a:defRPr sz="%d"/>' % hundredths))
        return
    props = level_element.find(qn("a:defRPr"))
    if props is None:
        props = etree.SubElement(level_element, qn("a:defRPr"))
    props.set("sz", str(hundredths))


def placeholder_of_type(container, ph_type):
    for placeholder in container.placeholders:
        if placeholder.placeholder_format.type == ph_type:
            return placeholder
    raise AssertionError("no placeholder of type {}".format(ph_type))


def run_props(run, xml):
    """Replace the a:rPr of a run with an element parsed from `xml`."""
    r = run._r
    old = r.find(qn("a:rPr"))
    if old is not None:
        r.remove(old)
    r.insert(0, parse_xml(xml))


def strip_default_text_style(prs, tag="a:solidFill"):
    """Remove every `tag` element from the presentation default text style."""
    for element in prs._element.xpath("p:defaultTextStyle//" + tag):
        element.getparent().remove(element)


def set_theme_fonts(master, major, minor):
    theme = master.part.part_related_by(RT.THEME)
    root = etree.fromstring(theme.blob)
    ns = {"a": "http://schemas.openxmlformats.org/drawingml/2006/main"}
    root.xpath("//a:majorFont/a:latin", namespaces=ns)[0].set("typeface", major)
    root.xpath("//a:minorFont/a:latin", namespaces=ns)[0].set("typeface", minor)
    theme._blob = etree.tostring(root, xml_declaration=True, encoding="UTF-8", standalone=True)


def drop_layout_relationship(slide):
    for rid, rel in list(slide.part.rels.items()):
        if rel.reltype == RT.SLIDE_LAYOUT:
            slide.part.drop_rel(rid)


def group_transform(group):
    """Return the (off, ext, chOff, chExt) elements of a group shape."""
    xfrm = group._element.grpSpPr.find(qn("a:xfrm"))
    return tuple(xfrm.find(qn("a:" + name)) for name in ("off", "ext", "chOff", "chExt"))


# ---------------------------------------------------------------------------
# Resolved
# ---------------------------------------------------------------------------

class ResolvedObject(unittest.TestCase):
    def test_starts_empty(self):
        value = rs.Resolved()
        self.assertEqual(
            (value.size_pt, value.family, value.color, value.bold, value.unresolved, value.sources),
            (None, None, None, None, [], {}),
        )

    def test_keeps_the_values_it_gets(self):
        value = rs.Resolved(size_pt=12.0, family="Arial", color="#112233", bold=True,
                            unresolved=["color"], sources={"size_pt": "run"})
        self.assertEqual((value.size_pt, value.family, value.color, value.bold), (12.0, "Arial", "#112233", True))
        self.assertEqual(value.unresolved, ["color"])
        self.assertEqual(value.sources, {"size_pt": "run"})

    def test_two_objects_do_not_share_their_lists(self):
        first, second = rs.Resolved(), rs.Resolved()
        first.unresolved.append("size_pt")
        first.sources["size_pt"] = "run"
        self.assertEqual((second.unresolved, second.sources), ([], {}))


# ---------------------------------------------------------------------------
# iter_text_runs
# ---------------------------------------------------------------------------

class IterTextRuns(DeckCase):
    def test_yields_each_run_with_its_own_paragraph(self):
        _, slide = new_slide()
        box = slide.shapes.add_textbox(Inches(1), Inches(1), Inches(4), Inches(1))
        first = box.text_frame.paragraphs[0]
        first.add_run().text = "one"
        first.add_run().text = "two"
        second = box.text_frame.add_paragraph()
        second.add_run().text = "three"
        pairs = list(rs.iter_text_runs(box))
        self.assertEqual([run.text for _, run in pairs], ["one", "two", "three"])
        self.assertEqual([paragraph._p for paragraph, _ in pairs], [first._p, first._p, second._p])

    def test_enters_a_group_and_a_group_in_a_group(self):
        _, slide = new_slide()
        outer = slide.shapes.add_group_shape()
        outer.shapes.add_textbox(Inches(1), Inches(1), Inches(2), Inches(1)).text_frame.text = "outer child"
        inner = outer.shapes.add_group_shape()
        inner.shapes.add_textbox(Inches(1), Inches(3), Inches(2), Inches(1)).text_frame.text = "inner child"
        self.assertEqual([run.text for _, run in rs.iter_text_runs(outer)], ["outer child", "inner child"])

    def test_enters_every_cell_of_a_table(self):
        _, slide = new_slide()
        table = slide.shapes.add_table(2, 2, Inches(1), Inches(1), Inches(4), Inches(2))
        table.table.cell(0, 0).text = "a"
        table.table.cell(0, 1).text = "b"
        table.table.cell(1, 1).text = "d"
        self.assertEqual([run.text for _, run in rs.iter_text_runs(table)], ["a", "b", "d"])

    def test_a_picture_has_no_text(self):
        _, slide = new_slide()
        picture = slide.shapes.add_picture(str(self.png("p.png", 20, 10)), Inches(1), Inches(1), width=Inches(2))
        self.assertEqual(list(rs.iter_text_runs(picture)), [])

    def test_a_shape_with_no_text_body_yields_nothing_and_gains_none(self):
        _, slide = new_slide()
        shape = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(1), Inches(1), Inches(2), Inches(1))
        shape._element.remove(text_body(shape))
        self.assertEqual(list(rs.iter_text_runs(shape)), [])
        self.assertEqual(rs.word_count(shape), 0)
        self.assertIsNone(text_body(shape))

    def test_a_run_of_a_shape_in_a_group_resolves(self):
        prs, slide = new_slide()
        group = slide.shapes.add_group_shape()
        child = group.shapes.add_textbox(Inches(1), Inches(1), Inches(2), Inches(1))
        child.text_frame.text = "grouped"
        self.assertEqual(resolve(prs, slide, child).size_pt, 18.0)


# ---------------------------------------------------------------------------
# Size
# ---------------------------------------------------------------------------

class ResolveSize(DeckCase):
    def test_an_explicit_run_size_wins(self):
        prs, slide = new_slide()
        box = add_box(slide, "text", size=20)
        result = resolve(prs, slide, box)
        self.assertEqual(result.size_pt, 20.0)
        self.assertEqual(result.sources["size_pt"], "run")
        self.assertNotIn("size_pt", result.unresolved)

    def test_a_fractional_size_is_kept(self):
        prs, slide = new_slide()
        box = add_box(slide, "text")
        run_props(box.text_frame.paragraphs[0].runs[0], '<a:rPr %s sz="1050"/>' % nsdecls("a"))
        self.assertEqual(resolve(prs, slide, box).size_pt, 10.5)

    def test_the_default_template_gives_the_measured_values(self):
        # The research measured these with python-pptx 0.6.23 and a LibreOffice render.
        prs, slide = new_slide("Title and Content")
        slide.shapes.title.text_frame.text = "headline"
        body = slide.placeholders[1]
        body.text_frame.text = "level zero"
        second = body.text_frame.add_paragraph()
        second.text = "level one"
        second.level = 1
        prs = self.reopen(prs)
        slide = prs.slides[0]
        title, body = slide.shapes.title, slide.placeholders[1]
        self.assertEqual(resolve(prs, slide, title).size_pt, 44.0)
        self.assertEqual(resolve(prs, slide, body, 0).size_pt, 32.0)
        self.assertEqual(resolve(prs, slide, body, 1).size_pt, 28.0)
        self.assertEqual(resolve(prs, slide, title).sources["size_pt"], "master_style")

    def test_a_plain_text_box_gets_18_from_the_default_text_style(self):
        prs, slide = new_slide()
        box = add_box(slide, "plain")
        result = resolve(prs, slide, box)
        self.assertEqual(result.size_pt, 18.0)
        self.assertEqual(result.sources["size_pt"], "default_style")

    def test_the_default_text_style_of_the_deck_is_read(self):
        prs, slide = new_slide()
        box = add_box(slide, "plain")
        prs._element.xpath("p:defaultTextStyle/a:lvl1pPr/a:defRPr")[0].set("sz", "2000")
        self.assertEqual(resolve(prs, slide, box).size_pt, 20.0)

    def test_the_paragraph_default_applies_when_the_run_has_no_size(self):
        # python-pptx writes paragraph.font to a:pPr/a:defRPr. The research found that
        # LibreOffice honors it and renders 30 pt. The resolver follows the brief and the
        # render. The source "paragraph" tells the linter that the value is renderer dependent.
        prs, slide = new_slide()
        box = add_box(slide, "text")
        box.text_frame.paragraphs[0].font.size = Pt(30)
        result = resolve(prs, slide, box)
        self.assertEqual(result.size_pt, 30.0)
        self.assertEqual(result.sources["size_pt"], "paragraph")

    def test_the_paragraph_default_can_be_left_out(self):
        # PowerPoint is said to ignore the paragraph default. The keyword gives that reading.
        prs, slide = new_slide()
        box = add_box(slide, "text")
        font = box.text_frame.paragraphs[0].font
        font.size = Pt(30)
        font.name = "Arial"
        font.color.rgb = RGBColor(1, 2, 3)
        font.bold = True
        paragraph, run = list(rs.iter_text_runs(box))[0]
        with_defaults = rs.resolve_run(prs, slide, box, paragraph, run)
        self.assertEqual((with_defaults.size_pt, with_defaults.family, with_defaults.color, with_defaults.bold),
                         (30.0, "Arial", "#010203", True))
        without = rs.resolve_run(prs, slide, box, paragraph, run, paragraph_defaults=False)
        self.assertEqual((without.size_pt, without.family, without.color, without.bold),
                         (18.0, "Calibri", "#000000", False))
        self.assertEqual(without.sources["size_pt"], "default_style")

    def test_leaving_out_the_paragraph_default_keeps_the_run_values(self):
        prs, slide = new_slide()
        box = add_box(slide, "text", size=12)
        box.text_frame.paragraphs[0].font.size = Pt(30)
        paragraph, run = list(rs.iter_text_runs(box))[0]
        self.assertEqual(rs.resolve_run(prs, slide, box, paragraph, run, paragraph_defaults=False).size_pt, 12.0)

    def test_every_value_of_the_paragraph_default_names_the_paragraph_as_its_source(self):
        prs, slide = new_slide()
        box = add_box(slide, "text")
        font = box.text_frame.paragraphs[0].font
        font.size = Pt(30)
        font.name = "Arial"
        font.color.rgb = RGBColor(1, 2, 3)
        font.bold = True
        result = resolve(prs, slide, box)
        self.assertEqual(result.sources, {"size_pt": "paragraph", "family": "paragraph",
                                          "color": "paragraph", "bold": "paragraph"})

    def test_a_run_size_beats_the_paragraph_default(self):
        prs, slide = new_slide()
        box = add_box(slide, "text", size=12)
        box.text_frame.paragraphs[0].font.size = Pt(30)
        result = resolve(prs, slide, box)
        self.assertEqual((result.size_pt, result.sources["size_pt"]), (12.0, "run"))

    def test_the_list_style_of_the_shape_follows_the_paragraph_level(self):
        prs, slide = new_slide()
        box = add_box(slide, "level zero")
        second = box.text_frame.add_paragraph()
        second.add_run().text = "level one"
        second.level = 1
        lst = list_style_of(box)
        lst.append(level_style(0, '<a:defRPr sz="2200"/>'))
        lst.append(level_style(1, '<a:defRPr sz="1400"/>'))
        first_result, second_result = resolve(prs, slide, box, 0), resolve(prs, slide, box, 1)
        self.assertEqual((first_result.size_pt, second_result.size_pt), (22.0, 14.0))
        self.assertEqual(first_result.sources["size_pt"], "shape")

    def test_a_level_without_its_own_size_falls_through_to_the_next_source(self):
        prs, slide = new_slide()
        box = add_box(slide, "text")
        box.text_frame.paragraphs[0].level = 2
        list_style_of(box).append(level_style(0, '<a:defRPr sz="2200"/>'))
        result = resolve(prs, slide, box)
        self.assertEqual((result.size_pt, result.sources["size_pt"]), (18.0, "default_style"))

    def test_each_source_beats_the_sources_below_it(self):
        prs, slide = new_slide("Title Only")
        title = slide.shapes.title
        title.text_frame.text = "headline"
        layout = slide.slide_layout
        master = layout.slide_master
        steps = [
            (44.0, "master_style"),
            (40.0, "master"),
            (36.0, "layout"),
            (32.0, "shape"),
            (30.0, "paragraph"),
            (28.0, "run"),
        ]

        def check(step):
            result = resolve(prs, slide, title)
            self.assertEqual((result.size_pt, result.sources["size_pt"]), steps[step])

        check(0)
        list_style_of(placeholder_of_type(master, PP_PLACEHOLDER.TITLE)).append(level_style(0, '<a:defRPr sz="4000"/>'))
        check(1)
        list_style_of(layout.placeholders[0]).append(level_style(0, '<a:defRPr sz="3600"/>'))
        check(2)
        list_style_of(title).append(level_style(0, '<a:defRPr sz="3200"/>'))
        check(3)
        title.text_frame.paragraphs[0].font.size = Pt(30)
        check(4)
        title.text_frame.paragraphs[0].runs[0].font.size = Pt(28)
        check(5)

    def test_date_footer_and_slide_number_each_read_their_own_master_placeholder(self):
        prs, slide = new_slide("Title Only")
        master = slide.slide_layout.slide_master
        sizes = {PP_PLACEHOLDER.DATE: 1000, PP_PLACEHOLDER.FOOTER: 1100, PP_PLACEHOLDER.SLIDE_NUMBER: 1400}
        for placeholder in master.placeholders:
            if placeholder.placeholder_format.type in sizes:
                set_list_style_size(placeholder, 0, sizes[placeholder.placeholder_format.type])
        for placeholder in slide.slide_layout.placeholders:
            if placeholder.placeholder_format.type in sizes:
                slide.shapes.clone_placeholder(placeholder)
        found = {}
        for shape in slide.placeholders:
            kind = shape.placeholder_format.type
            if kind in sizes:
                shape.text_frame.text = "text"
                result = resolve(prs, slide, shape)
                found[kind] = (result.size_pt, result.sources["size_pt"])
        self.assertEqual(found, {PP_PLACEHOLDER.DATE: (10.0, "master"), PP_PLACEHOLDER.FOOTER: (11.0, "master"),
                                 PP_PLACEHOLDER.SLIDE_NUMBER: (14.0, "master")})

    def test_a_header_placeholder_reads_the_master_placeholder_of_its_type(self):
        prs, slide = new_slide()
        master = prs.slide_masters[0]._element
        master.find(qn("p:cSld")).find(qn("p:spTree")).append(parse_xml(
            '<p:sp %s><p:nvSpPr><p:cNvPr id="90" name="Header Placeholder 9"/><p:cNvSpPr/>'
            '<p:nvPr><p:ph type="hdr" sz="quarter" idx="5"/></p:nvPr></p:nvSpPr><p:spPr/>'
            '<p:txBody><a:bodyPr/><a:lstStyle><a:lvl1pPr><a:defRPr sz="1700"/></a:lvl1pPr></a:lstStyle>'
            '<a:p/></p:txBody></p:sp>' % nsdecls("a", "p")))
        box = add_box(slide, "header text")
        box._element.find(qn("p:nvSpPr")).find(qn("p:nvPr")).append(
            parse_xml('<p:ph %s type="hdr" idx="5"/>' % nsdecls("p")))
        result = resolve(prs, slide, box)
        self.assertEqual((result.size_pt, result.sources["size_pt"]), (17.0, "master"))

    def test_object_and_subtitle_placeholders_read_the_body_placeholder_of_the_master(self):
        for layout_name in ("Title and Content", "Title Slide"):
            with self.subTest(layout=layout_name):
                prs, slide = new_slide(layout_name)
                set_list_style_size(placeholder_of_type(prs.slide_masters[0], PP_PLACEHOLDER.BODY), 0, 3000)
                body = slide.placeholders[1]
                body.text_frame.text = "text"
                result = resolve(prs, slide, body)
                self.assertEqual((result.size_pt, result.sources["size_pt"]), (30.0, "master"))

    def test_a_subtitle_placeholder_reads_the_body_style_of_the_master(self):
        prs, slide = new_slide("Title Slide")
        subtitle = slide.placeholders[1]
        subtitle.text_frame.text = "a subtitle"
        result = resolve(prs, slide, subtitle)
        self.assertEqual((result.size_pt, result.sources["size_pt"]), (32.0, "master_style"))
        self.assertEqual(result.family, "Calibri")
        # The layout of the default template sets the subtitle color with a tint, which is a transform.
        self.assertIn("color", result.unresolved)

    def test_a_placeholder_that_the_layout_lacks_skips_the_layout_and_reaches_the_master(self):
        prs, slide = new_slide("Title and Content")
        body = slide.placeholders[1]
        body.text_frame.text = "text"
        layout_body = [p for p in slide.slide_layout.placeholders if p.placeholder_format.idx == 1][0]
        set_list_style_size(layout_body, 0, 2600)
        self.assertEqual(resolve(prs, slide, body).size_pt, 26.0)
        body._element.xpath("p:nvSpPr/p:nvPr/p:ph")[0].set("idx", "99")
        result = resolve(prs, slide, body)
        self.assertEqual((result.size_pt, result.sources["size_pt"]), (32.0, "master_style"))

    def test_a_body_placeholder_never_reads_a_title_placeholder_of_the_layout(self):
        prs, slide = new_slide("Title Only")
        set_list_style_size(slide.slide_layout.placeholders[0], 0, 3600)
        box = add_box(slide, "body text")
        box._element.find(qn("p:nvSpPr")).find(qn("p:nvPr")).append(parse_xml('<p:ph %s type="body"/>' % nsdecls("p")))
        result = resolve(prs, slide, box)
        self.assertEqual((result.size_pt, result.sources["size_pt"]), (32.0, "master_style"))

    def test_a_placeholder_reads_the_layout_placeholder_with_its_own_index(self):
        prs, slide = new_slide("Two Content")
        left, right = slide.placeholders[1], slide.placeholders[2]
        left.text_frame.text = "left"
        right.text_frame.text = "right"
        for placeholder in slide.slide_layout.placeholders:
            size = {1: 2400, 2: 2000}.get(placeholder.placeholder_format.idx)
            if size is not None:
                set_list_style_size(placeholder, 0, size)
        self.assertEqual(resolve(prs, slide, left).size_pt, 24.0)
        self.assertEqual(resolve(prs, slide, right).size_pt, 20.0)

    def test_font_scale_multiplies_the_size_of_the_run(self):
        prs, slide = new_slide()
        box = add_box(slide, "text", size=24)
        box.text_frame.auto_size = MSO_AUTO_SIZE.TEXT_TO_FIT_SHAPE
        text_body(box).find(qn("a:bodyPr")).find(qn("a:normAutofit")).set("fontScale", "62500")
        self.assertEqual(resolve(prs, slide, box).size_pt, 15.0)

    def test_font_scale_multiplies_an_inherited_size(self):
        prs, slide = new_slide("Title Only")
        title = slide.shapes.title
        title.text_frame.text = "headline"
        norm = etree.SubElement(text_body(title).find(qn("a:bodyPr")), qn("a:normAutofit"))
        norm.set("fontScale", "50000")
        self.assertEqual(resolve(prs, slide, title).size_pt, 22.0)

    def test_autofit_without_a_font_scale_keeps_the_size(self):
        prs, slide = new_slide()
        box = add_box(slide, "text", size=24)
        box.text_frame.auto_size = MSO_AUTO_SIZE.TEXT_TO_FIT_SHAPE
        self.assertEqual(resolve(prs, slide, box).size_pt, 24.0)

    def test_a_malformed_font_scale_leaves_the_size_unresolved(self):
        prs, slide = new_slide()
        box = add_box(slide, "text", size=24)
        box.text_frame.auto_size = MSO_AUTO_SIZE.TEXT_TO_FIT_SHAPE
        text_body(box).find(qn("a:bodyPr")).find(qn("a:normAutofit")).set("fontScale", "lots")
        result = resolve(prs, slide, box)
        self.assertIsNone(result.size_pt)
        self.assertIn("size_pt", result.unresolved)

    def test_a_run_properties_element_with_no_size_is_not_an_error(self):
        prs, slide = new_slide()
        box = add_box(slide, "text")
        run_props(box.text_frame.paragraphs[0].runs[0], '<a:rPr %s lang="en-US"/>' % nsdecls("a"))
        self.assertEqual(resolve(prs, slide, box).size_pt, 18.0)

    def test_no_size_anywhere_is_unresolved(self):
        prs, slide = new_slide()
        box = add_box(slide, "text")
        strip_default_text_style(prs, "a:defRPr")
        result = resolve(prs, slide, box)
        self.assertIsNone(result.size_pt)
        self.assertIn("size_pt", result.unresolved)
        self.assertNotIn("size_pt", result.sources)

    def test_a_malformed_size_is_unresolved_and_does_not_fall_through(self):
        prs, slide = new_slide()
        for text in ("big", "0", "-300", ""):
            with self.subTest(sz=text):
                box = add_box(slide, "text")
                run_props(box.text_frame.paragraphs[0].runs[0], '<a:rPr %s sz="%s"/>' % (nsdecls("a"), text))
                result = resolve(prs, slide, box)
                self.assertIsNone(result.size_pt)
                self.assertIn("size_pt", result.unresolved)

    def test_a_malformed_paragraph_level_leaves_inherited_values_unresolved(self):
        prs, slide = new_slide()
        box = add_box(slide, "text")
        box.text_frame.paragraphs[0]._p.get_or_add_pPr().set("lvl", "12")
        result = resolve(prs, slide, box)
        self.assertIsNone(result.size_pt)
        self.assertIn("size_pt", result.unresolved)
        box.text_frame.paragraphs[0].runs[0].font.size = Pt(10)
        self.assertEqual(resolve(prs, slide, box).size_pt, 10.0)

    def test_table_cells_use_their_run_and_the_default_text_style(self):
        prs, slide = new_slide()
        table = slide.shapes.add_table(1, 2, Inches(1), Inches(1), Inches(4), Inches(1))
        explicit = table.table.cell(0, 0).text_frame.paragraphs[0].add_run()
        explicit.text = "sized"
        explicit.font.size = Pt(10)
        table.table.cell(0, 1).text = "plain"
        self.assertEqual(resolve(prs, slide, table, 0).size_pt, 10.0)
        self.assertEqual(resolve(prs, slide, table, 1).size_pt, 18.0)

    def test_resolving_a_loaded_deck_gives_the_same_values(self):
        prs, slide = new_slide("Title Only")
        slide.shapes.title.text_frame.text = "headline"
        add_box(slide, "plain")
        prs = self.reopen(prs)
        slide = prs.slides[0]
        self.assertEqual(resolve(prs, slide, slide.shapes.title).size_pt, 44.0)
        self.assertEqual(resolve(prs, slide, slide.shapes[1]).size_pt, 18.0)


# ---------------------------------------------------------------------------
# Color
# ---------------------------------------------------------------------------

class ResolveColor(DeckCase):
    def test_an_explicit_rgb_color_resolves_to_an_upper_case_hex_with_a_hash(self):
        prs, slide = new_slide()
        box = add_box(slide, "text")
        box.text_frame.paragraphs[0].runs[0].font.color.rgb = RGBColor(0x59, 0x59, 0x59)
        result = resolve(prs, slide, box)
        self.assertEqual((result.color, result.sources["color"]), ("#595959", "run"))
        run_props(box.text_frame.paragraphs[0].runs[0],
                  '<a:rPr %s><a:solidFill><a:srgbClr val="ff8800"/></a:solidFill></a:rPr>' % nsdecls("a"))
        self.assertEqual(resolve(prs, slide, box).color, "#FF8800")

    def test_a_theme_color_resolves_through_the_theme_xml(self):
        prs, slide = new_slide()
        box = add_box(slide, "text")
        box.text_frame.paragraphs[0].runs[0].font.color.theme_color = MSO_THEME_COLOR.ACCENT_1
        prs = self.reopen(prs)
        slide = prs.slides[0]
        self.assertEqual(resolve(prs, slide, slide.shapes[0]).color, "#4F81BD")

    def test_a_system_color_resolves_through_its_last_color(self):
        prs, slide = new_slide()
        box = add_box(slide, "text")
        box.text_frame.paragraphs[0].runs[0].font.color.theme_color = MSO_THEME_COLOR.DARK_1
        self.assertEqual(resolve(prs, slide, box).color, "#000000")

    def test_text_and_background_theme_colors_follow_the_color_map(self):
        prs, slide = new_slide()
        box = add_box(slide, "text")
        font = box.text_frame.paragraphs[0].runs[0].font
        font.color.theme_color = MSO_THEME_COLOR.TEXT_1
        self.assertEqual(resolve(prs, slide, box).color, "#000000")
        font.color.theme_color = MSO_THEME_COLOR.BACKGROUND_1
        self.assertEqual(resolve(prs, slide, box).color, "#FFFFFF")
        prs.slide_masters[0]._element.find(qn("p:clrMap")).set("tx1", "lt1")
        font.color.theme_color = MSO_THEME_COLOR.TEXT_1
        self.assertEqual(resolve(prs, slide, box).color, "#FFFFFF")

    def test_a_slide_that_overrides_the_color_map_uses_its_own_map(self):
        prs, slide = new_slide()
        box = add_box(slide, "text")
        override = parse_xml(
            '<p:clrMapOvr %s><a:overrideClrMapping bg1="dk1" tx1="lt1" bg2="dk2" tx2="lt2" accent1="accent1" '
            'accent2="accent2" accent3="accent3" accent4="accent4" accent5="accent5" accent6="accent6" '
            'hlink="hlink" folHlink="folHlink"/></p:clrMapOvr>' % nsdecls("a", "p"))
        old = slide._element.find(qn("p:clrMapOvr"))
        slide._element.replace(old, override)
        self.assertEqual(resolve(prs, slide, box).color, "#FFFFFF")

    def test_a_layout_that_overrides_the_color_map_applies_to_its_slides(self):
        prs, slide = new_slide()
        box = add_box(slide, "text")
        override = parse_xml(
            '<p:clrMapOvr %s><a:overrideClrMapping bg1="dk1" tx1="lt1" bg2="dk2" tx2="lt2" accent1="accent1" '
            'accent2="accent2" accent3="accent3" accent4="accent4" accent5="accent5" accent6="accent6" '
            'hlink="hlink" folHlink="folHlink"/></p:clrMapOvr>' % nsdecls("a", "p"))
        layout = slide.slide_layout._element
        layout.replace(layout.find(qn("p:clrMapOvr")), override)
        self.assertEqual(resolve(prs, slide, box).color, "#FFFFFF")

    def test_an_unset_color_comes_from_the_default_text_style(self):
        prs, slide = new_slide()
        box = add_box(slide, "text")
        result = resolve(prs, slide, box)
        self.assertEqual((result.color, result.sources["color"]), ("#000000", "default_style"))

    def test_text_in_an_autoshape_takes_the_font_reference_of_its_style(self):
        prs, slide = new_slide()
        shape = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(1), Inches(1), Inches(2), Inches(1))
        shape.text_frame.text = "on a shape"
        result = resolve(prs, slide, shape)
        self.assertEqual((result.color, result.sources["color"]), ("#FFFFFF", "shape_style"))

    def test_a_color_set_nowhere_is_unresolved(self):
        prs, slide = new_slide()
        box = add_box(slide, "text")
        strip_default_text_style(prs)
        result = resolve(prs, slide, box)
        self.assertIsNone(result.color)
        self.assertIn("color", result.unresolved)
        self.assertNotIn("color", result.sources)
        self.assertEqual(result.size_pt, 18.0)

    def test_a_color_with_a_transform_is_unresolved(self):
        prs, slide = new_slide()
        for inner in ('<a:schemeClr val="accent1"><a:lumMod val="75000"/></a:schemeClr>',
                      '<a:srgbClr val="336699"><a:alpha val="50000"/></a:srgbClr>'):
            with self.subTest(color=inner):
                box = add_box(slide, "text")
                run_props(box.text_frame.paragraphs[0].runs[0],
                          '<a:rPr %s><a:solidFill>%s</a:solidFill></a:rPr>' % (nsdecls("a"), inner))
                result = resolve(prs, slide, box)
                self.assertIsNone(result.color)
                self.assertIn("color", result.unresolved)

    def test_a_color_kind_that_is_not_scheme_or_rgb_is_unresolved(self):
        prs, slide = new_slide()
        for inner in ('<a:prstClr val="black"/>', '<a:scrgbClr r="0" g="0" b="0"/>',
                      '<a:hslClr hue="0" sat="0" lum="0"/>', '<a:srgbClr val="12"/>',
                      '<a:schemeClr val="phClr"/>', '<a:schemeClr val="nonsense"/>'):
            with self.subTest(color=inner):
                box = add_box(slide, "text")
                run_props(box.text_frame.paragraphs[0].runs[0],
                          '<a:rPr %s><a:solidFill>%s</a:solidFill></a:rPr>' % (nsdecls("a"), inner))
                result = resolve(prs, slide, box)
                self.assertIsNone(result.color)
                self.assertIn("color", result.unresolved)

    def test_a_gradient_text_fill_is_unresolved_and_hides_the_inherited_color(self):
        prs, slide = new_slide()
        box = add_box(slide, "text")
        run_props(box.text_frame.paragraphs[0].runs[0],
                  '<a:rPr %s><a:gradFill><a:gsLst/></a:gradFill></a:rPr>' % nsdecls("a"))
        result = resolve(prs, slide, box)
        self.assertIsNone(result.color)
        self.assertIn("color", result.unresolved)

    def test_a_hyperlink_run_is_unresolved(self):
        prs, slide = new_slide()
        box = add_box(slide, "link")
        box.text_frame.paragraphs[0].runs[0].hyperlink.address = "https://example.org"
        result = resolve(prs, slide, box)
        self.assertIsNone(result.color)
        self.assertIn("color", result.unresolved)

    def test_a_table_cell_without_an_explicit_color_is_unresolved(self):
        prs, slide = new_slide()
        table = slide.shapes.add_table(1, 2, Inches(1), Inches(1), Inches(4), Inches(1))
        table.table.cell(0, 0).text = "styled by the table style"
        colored = table.table.cell(0, 1).text_frame.paragraphs[0].add_run()
        colored.text = "colored"
        colored.font.color.rgb = RGBColor(1, 2, 3)
        plain = resolve(prs, slide, table, 0)
        self.assertIsNone(plain.color)
        self.assertIn("color", plain.unresolved)
        self.assertEqual(resolve(prs, slide, table, 1).color, "#010203")

    def test_a_scheme_color_is_unresolved_when_the_layout_part_is_missing(self):
        prs, slide = new_slide()
        box = add_box(slide, "text")
        box.text_frame.paragraphs[0].runs[0].font.color.theme_color = MSO_THEME_COLOR.ACCENT_1
        drop_layout_relationship(slide)
        result = resolve(prs, slide, box)
        self.assertIsNone(result.color)
        self.assertIn("color", result.unresolved)


# ---------------------------------------------------------------------------
# Family and bold
# ---------------------------------------------------------------------------

class ResolveFamily(DeckCase):
    def test_an_explicit_family_wins(self):
        prs, slide = new_slide()
        box = add_box(slide, "text")
        box.text_frame.paragraphs[0].runs[0].font.name = "Arial"
        result = resolve(prs, slide, box)
        self.assertEqual((result.family, result.sources["family"]), ("Arial", "run"))

    def test_theme_font_references_resolve_through_the_font_scheme(self):
        prs, slide = new_slide("Title Only")
        set_theme_fonts(prs.slide_masters[0], major="Georgia", minor="Verdana")
        slide.shapes.title.text_frame.text = "headline"
        box = add_box(slide, "plain")
        explicit = add_box(slide, "explicit", top=2.0)
        explicit.text_frame.paragraphs[0].runs[0].font.name = "+mj-lt"
        prs = self.reopen(prs)
        slide = prs.slides[0]
        self.assertEqual(resolve(prs, slide, slide.shapes.title).family, "Georgia")
        self.assertEqual(resolve(prs, slide, slide.shapes[1]).family, "Verdana")
        self.assertEqual(resolve(prs, slide, slide.shapes[2]).family, "Georgia")

    def test_text_in_an_autoshape_takes_the_typeface_of_its_font_reference(self):
        prs, slide = new_slide()
        set_theme_fonts(prs.slide_masters[0], major="Georgia", minor="Verdana")
        shape = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(1), Inches(1), Inches(2), Inches(1))
        shape.text_frame.text = "on a shape"
        font_ref = shape._element.find(qn("p:style")).find(qn("a:fontRef"))
        self.assertEqual(resolve(prs, slide, shape).family, "Verdana")
        font_ref.set("idx", "major")
        result = resolve(prs, slide, shape)
        self.assertEqual((result.family, result.sources["family"]), ("Georgia", "shape_style"))

    def test_a_latin_element_with_no_typeface_is_unresolved(self):
        prs, slide = new_slide()
        box = add_box(slide, "text")
        run_props(box.text_frame.paragraphs[0].runs[0], '<a:rPr %s><a:latin/></a:rPr>' % nsdecls("a"))
        self.assertIn("family", resolve(prs, slide, box).unresolved)

    def test_the_default_template_font_is_calibri(self):
        prs, slide = new_slide()
        self.assertEqual(resolve(prs, slide, add_box(slide, "plain")).family, "Calibri")

    def test_a_family_set_nowhere_is_unresolved(self):
        prs, slide = new_slide()
        box = add_box(slide, "plain")
        strip_default_text_style(prs, "a:latin")
        result = resolve(prs, slide, box)
        self.assertIsNone(result.family)
        self.assertIn("family", result.unresolved)

    def test_an_east_asian_theme_reference_in_the_latin_slot_is_unresolved(self):
        prs, slide = new_slide()
        box = add_box(slide, "text")
        box.text_frame.paragraphs[0].runs[0].font.name = "+mn-ea"
        self.assertIn("family", resolve(prs, slide, box).unresolved)


class ResolveBold(DeckCase):
    def test_an_explicit_value_wins_over_an_inherited_one(self):
        prs, slide = new_slide()
        box = add_box(slide, "text")
        list_style_of(box).append(level_style(0, '<a:defRPr b="1"/>'))
        font = box.text_frame.paragraphs[0].runs[0].font
        inherited = resolve(prs, slide, box)
        self.assertEqual((inherited.bold, inherited.sources["bold"]), (True, "shape"))
        font.bold = False
        self.assertEqual(resolve(prs, slide, box).bold, False)
        font.bold = True
        self.assertEqual((resolve(prs, slide, box).bold, resolve(prs, slide, box).sources["bold"]), (True, "run"))

    def test_text_that_nothing_makes_bold_is_not_bold(self):
        prs, slide = new_slide()
        result = resolve(prs, slide, add_box(slide, "text"))
        self.assertEqual((result.bold, result.sources["bold"]), (False, "unset"))
        self.assertNotIn("bold", result.unresolved)

    def test_a_malformed_value_is_unresolved(self):
        prs, slide = new_slide()
        box = add_box(slide, "text")
        run_props(box.text_frame.paragraphs[0].runs[0], '<a:rPr %s b="maybe"/>' % nsdecls("a"))
        result = resolve(prs, slide, box)
        self.assertIsNone(result.bold)
        self.assertIn("bold", result.unresolved)

    def test_the_words_true_and_false_are_valid(self):
        prs, slide = new_slide()
        box = add_box(slide, "text")
        run_props(box.text_frame.paragraphs[0].runs[0], '<a:rPr %s b="true"/>' % nsdecls("a"))
        self.assertIs(resolve(prs, slide, box).bold, True)
        run_props(box.text_frame.paragraphs[0].runs[0], '<a:rPr %s b="false"/>' % nsdecls("a"))
        self.assertIs(resolve(prs, slide, box).bold, False)

    def test_a_table_cell_without_an_explicit_value_is_unresolved(self):
        prs, slide = new_slide()
        table = slide.shapes.add_table(1, 2, Inches(1), Inches(1), Inches(4), Inches(1))
        table.table.cell(0, 0).text = "header text may be bold through the table style"
        bold = table.table.cell(0, 1).text_frame.paragraphs[0].add_run()
        bold.text = "bold"
        bold.font.bold = True
        plain = resolve(prs, slide, table, 0)
        self.assertIsNone(plain.bold)
        self.assertIn("bold", plain.unresolved)
        self.assertIs(resolve(prs, slide, table, 1).bold, True)


# ---------------------------------------------------------------------------
# Broken and foreign decks
# ---------------------------------------------------------------------------

class ForeignDecks(DeckCase):
    def test_a_slide_with_no_shapes_raises_nothing(self):
        prs, slide = new_slide()
        self.assertEqual(list(slide.shapes), [])
        self.assertIsNone(rs.find_title(slide))
        self.assertEqual(rs.background_color(prs, slide), "#FFFFFF")

    def test_a_text_box_on_a_slide_with_no_layout_part_keeps_its_size(self):
        prs, slide = new_slide()
        box = add_box(slide, "text")
        drop_layout_relationship(slide)
        result = resolve(prs, slide, box)
        self.assertEqual(result.size_pt, 18.0)
        self.assertEqual(result.unresolved, ["family", "color"])
        self.assertIsNone(rs.background_color(prs, slide))

    def test_a_master_with_no_theme_part_leaves_the_theme_values_unresolved(self):
        prs, slide = new_slide()
        box = add_box(slide, "text")
        master = slide.slide_layout.slide_master
        for rid, rel in list(master.part.rels.items()):
            if rel.reltype == RT.THEME:
                master.part.drop_rel(rid)
        result = resolve(prs, slide, box)
        self.assertEqual(result.size_pt, 18.0)
        self.assertEqual(result.unresolved, ["family", "color"])
        self.assertIsNone(rs.background_color(prs, slide))

    def test_a_theme_that_is_not_valid_xml_leaves_the_theme_values_unresolved(self):
        for blob in (b"<a:theme><unclosed", b""):
            with self.subTest(blob=blob):
                prs, slide = new_slide()
                box = add_box(slide, "text")
                slide.slide_layout.slide_master.part.part_related_by(RT.THEME)._blob = blob
                result = resolve(prs, slide, box)
                self.assertEqual(result.size_pt, 18.0)
                self.assertEqual(result.unresolved, ["family", "color"])
                self.assertIsNone(rs.background_color(prs, slide))

    def test_a_placeholder_on_a_slide_with_no_layout_part_is_not_guessed(self):
        prs, slide = new_slide("Title Only")
        slide.shapes.title.text_frame.text = "headline"
        drop_layout_relationship(slide)
        result = resolve(prs, slide, slide.shapes.title)
        self.assertIsNone(result.size_pt)
        self.assertIn("size_pt", result.unresolved)
        slide.shapes.title.text_frame.paragraphs[0].runs[0].font.size = Pt(30)
        self.assertEqual(resolve(prs, slide, slide.shapes.title).size_pt, 30.0)

    def test_a_layout_with_no_master_part_is_not_guessed_for_a_placeholder(self):
        prs, slide = new_slide("Title Only")
        slide.shapes.title.text_frame.text = "headline"
        layout = slide.slide_layout
        for rid, rel in list(layout.part.rels.items()):
            if rel.reltype == RT.SLIDE_MASTER:
                layout.part.drop_rel(rid)
        result = resolve(prs, slide, slide.shapes.title)
        self.assertIsNone(result.size_pt)
        self.assertIn("size_pt", result.unresolved)

    def test_a_master_with_no_text_styles_falls_back_to_the_default_text_style(self):
        prs, slide = new_slide("Title Only")
        slide.shapes.title.text_frame.text = "headline"
        master = prs.slide_masters[0]._element
        master.remove(master.find(qn("p:txStyles")))
        self.assertEqual(resolve(prs, slide, slide.shapes.title).size_pt, 18.0)

    def test_resolving_does_not_change_the_deck(self):
        prs, slide = new_slide("Title and Content")
        slide.shapes.title.text_frame.text = "headline"
        slide.placeholders[1].text_frame.text = "body"
        add_box(slide, "plain")
        group = slide.shapes.add_group_shape()
        group.shapes.add_textbox(Inches(1), Inches(1), Inches(1), Inches(1)).text_frame.text = "in group"
        table = slide.shapes.add_table(1, 1, Inches(1), Inches(4), Inches(2), Inches(1))
        table.table.cell(0, 0).text = "cell"
        slide.shapes.add_picture(str(self.png("p.png", 20, 10)), Inches(1), Inches(1), width=Inches(2))
        parts = [slide._element, slide.slide_layout._element, slide.slide_layout.slide_master._element,
                 prs._element]
        before = [etree.tostring(part) for part in parts]
        for shape in slide.shapes:
            for paragraph, run in rs.iter_text_runs(shape):
                rs.resolve_run(prs, slide, shape, paragraph, run)
            rs.shape_box(shape)
            rs.fill_color(shape)
            rs.word_count(shape)
        rs.find_title(slide)
        rs.background_color(prs, slide)
        rs.effective_ppi(slide.shapes[-1])
        self.assertEqual([etree.tostring(part) for part in parts], before)
        self.assertFalse(slide.has_notes_slide)


# ---------------------------------------------------------------------------
# Background and fill
# ---------------------------------------------------------------------------

class BackgroundColor(DeckCase):
    def test_an_explicit_solid_fill_of_the_slide_wins(self):
        prs, slide = new_slide()
        slide.background.fill.solid()
        slide.background.fill.fore_color.rgb = RGBColor(0x14, 0x11, 0x1F)
        self.assertEqual(rs.background_color(prs, self.reopen(prs).slides[0]), "#14111F")

    def test_the_master_reference_to_the_background_style_resolves(self):
        prs, slide = new_slide()
        self.assertEqual(rs.background_color(prs, slide), "#FFFFFF")

    def test_a_master_with_an_explicit_fill_gives_that_fill(self):
        prs, slide = new_slide()
        master = prs.slide_masters[0]
        master.background.fill.solid()
        master.background.fill.fore_color.rgb = RGBColor(0x10, 0x20, 0x30)
        self.assertEqual(rs.background_color(prs, slide), "#102030")

    def test_the_layout_beats_the_master(self):
        prs, slide = new_slide()
        slide.slide_layout.background.fill.solid()
        slide.slide_layout.background.fill.fore_color.rgb = RGBColor(0xAA, 0xBB, 0xCC)
        self.assertEqual(rs.background_color(prs, slide), "#AABBCC")

    def test_the_slide_beats_the_layout(self):
        prs, slide = new_slide()
        slide.slide_layout.background.fill.solid()
        slide.slide_layout.background.fill.fore_color.rgb = RGBColor(0xAA, 0xBB, 0xCC)
        slide.background.fill.solid()
        slide.background.fill.fore_color.theme_color = MSO_THEME_COLOR.ACCENT_2
        self.assertEqual(rs.background_color(prs, slide), "#C0504D")

    def test_a_gradient_is_none_and_does_not_fall_through_to_the_master(self):
        prs, slide = new_slide()
        slide.background.fill.gradient()
        self.assertIsNone(rs.background_color(prs, slide))

    def test_a_pattern_is_none(self):
        prs, slide = new_slide()
        slide.background.fill.patterned()
        self.assertIsNone(rs.background_color(prs, slide))

    def test_a_picture_is_none(self):
        prs, slide = new_slide()
        c_sld = slide._element.find(qn("p:cSld"))
        c_sld.insert(0, parse_xml(
            '<p:bg %s><p:bgPr><a:blipFill><a:blip r:embed="rId99"/></a:blipFill><a:effectLst/></p:bgPr></p:bg>'
            % nsdecls("a", "p", "r")))
        self.assertIsNone(rs.background_color(prs, slide))

    def test_a_background_style_that_is_not_solid_is_none(self):
        prs, slide = new_slide()
        master = prs.slide_masters[0]._element
        master.find(qn("p:cSld")).find(qn("p:bg")).find(qn("p:bgRef")).set("idx", "1003")
        self.assertIsNone(rs.background_color(prs, slide))

    def test_a_slide_whose_layout_has_no_part_gives_none(self):
        prs, slide = new_slide()
        drop_layout_relationship(slide)
        self.assertIsNone(rs.background_color(prs, slide))


class FillColor(DeckCase):
    def test_a_solid_rgb_fill_resolves(self):
        _, slide = new_slide()
        shape = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(1), Inches(1), Inches(1), Inches(1))
        shape.fill.solid()
        shape.fill.fore_color.rgb = RGBColor(0x12, 0x34, 0x56)
        self.assertEqual(rs.fill_color(shape), "#123456")

    def test_a_solid_theme_fill_resolves_through_the_theme(self):
        prs, slide = new_slide()
        shape = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(1), Inches(1), Inches(1), Inches(1))
        shape.fill.solid()
        shape.fill.fore_color.theme_color = MSO_THEME_COLOR.ACCENT_2
        self.assertEqual(rs.fill_color(self.reopen(prs).slides[0].shapes[0]), "#C0504D")

    def test_a_shape_with_no_solid_fill_gives_none(self):
        _, slide = new_slide()
        shape = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(1), Inches(1), Inches(1), Inches(1))
        self.assertIsNone(rs.fill_color(shape))
        shape.fill.background()
        self.assertIsNone(rs.fill_color(shape))
        shape.fill.gradient()
        self.assertIsNone(rs.fill_color(shape))
        shape.fill.patterned()
        self.assertIsNone(rs.fill_color(shape))

    def test_a_text_box_and_a_picture_give_none(self):
        _, slide = new_slide()
        self.assertIsNone(rs.fill_color(add_box(slide, "text")))
        picture = slide.shapes.add_picture(str(self.png("p.png", 20, 10)), Inches(1), Inches(1), width=Inches(2))
        self.assertIsNone(rs.fill_color(picture))

    def test_a_fill_with_a_color_transform_gives_none(self):
        _, slide = new_slide()
        shape = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(1), Inches(1), Inches(1), Inches(1))
        shape.fill.solid()
        shape.fill.fore_color.theme_color = MSO_THEME_COLOR.ACCENT_1
        shape.fill.fore_color.brightness = 0.4
        self.assertIsNone(rs.fill_color(shape))


class FillState(DeckCase):
    def rectangle(self, slide):
        return slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(1), Inches(1), Inches(1), Inches(1))

    def test_an_explicit_solid_fill_is_solid(self):
        _, slide = new_slide()
        shape = self.rectangle(slide)
        shape.fill.solid()
        shape.fill.fore_color.rgb = RGBColor(0x12, 0x34, 0x56)
        self.assertEqual(rs.fill_state(shape), "solid")

    def test_an_explicit_no_fill_is_none(self):
        _, slide = new_slide()
        shape = self.rectangle(slide)
        shape.fill.background()
        self.assertEqual(rs.fill_state(shape), "none")

    def test_a_gradient_a_pattern_and_a_picture_fill_are_unknown(self):
        _, slide = new_slide()
        shape = self.rectangle(slide)
        shape.fill.gradient()
        self.assertEqual(rs.fill_state(shape), "unknown")
        shape.fill.patterned()
        self.assertEqual(rs.fill_state(shape), "unknown")
        sp_pr = shape._element.find(qn("p:spPr"))
        sp_pr.remove(sp_pr.find(qn("a:pattFill")))
        sp_pr.append(parse_xml('<a:blipFill %s><a:blip r:embed="rId9"/></a:blipFill>' % nsdecls("a", "r")))
        self.assertEqual(rs.fill_state(shape), "unknown")

    def test_a_group_fill_is_unknown(self):
        _, slide = new_slide()
        shape = self.rectangle(slide)
        sp_pr = shape._element.find(qn("p:spPr"))
        sp_pr.append(parse_xml("<a:grpFill %s/>" % nsdecls("a")))
        self.assertEqual(rs.fill_state(shape), "unknown")

    def test_an_autoshape_with_only_a_style_fill_is_unknown(self):
        _, slide = new_slide()
        shape = self.rectangle(slide)
        self.assertIsNone(shape._element.find(qn("p:spPr")).find(qn("a:solidFill")))
        self.assertEqual(rs.fill_state(shape), "unknown")

    def test_a_style_fill_reference_with_index_zero_is_none(self):
        _, slide = new_slide()
        shape = self.rectangle(slide)
        shape._element.find(qn("p:style")).find(qn("a:fillRef")).set("idx", "0")
        self.assertEqual(rs.fill_state(shape), "none")

    def test_a_style_fill_reference_with_a_bad_index_is_unknown(self):
        _, slide = new_slide()
        shape = self.rectangle(slide)
        shape._element.find(qn("p:style")).find(qn("a:fillRef")).set("idx", "many")
        self.assertEqual(rs.fill_state(shape), "unknown")

    def test_an_explicit_fill_beats_the_style_fill_reference(self):
        _, slide = new_slide()
        shape = self.rectangle(slide)
        shape.fill.background()
        self.assertEqual(rs.fill_state(shape), "none")

    def test_a_placeholder_with_no_explicit_fill_is_unknown(self):
        _, slide = new_slide("Title Only")
        title = slide.shapes.title
        self.assertEqual(rs.fill_state(title), "unknown")
        layout_title = slide.slide_layout.placeholders[0]
        layout_title._element.find(qn("p:spPr")).append(
            parse_xml('<a:solidFill %s><a:srgbClr val="FF0000"/></a:solidFill>' % nsdecls("a")))
        self.assertEqual(rs.fill_state(title), "unknown")
        self.assertIsNone(rs.fill_color(title))

    def test_a_placeholder_with_an_explicit_fill_keeps_that_state(self):
        _, slide = new_slide("Title Only")
        title = slide.shapes.title
        title.fill.background()
        self.assertEqual(rs.fill_state(title), "none")
        title.fill.solid()
        title.fill.fore_color.rgb = RGBColor(1, 2, 3)
        self.assertEqual(rs.fill_state(title), "solid")

    def test_a_text_box_and_a_picture_are_none(self):
        _, slide = new_slide()
        self.assertEqual(rs.fill_state(add_box(slide, "text")), "none")
        picture = slide.shapes.add_picture(str(self.png("p.png", 20, 10)), Inches(1), Inches(1), width=Inches(2))
        self.assertEqual(rs.fill_state(picture), "none")

    def test_a_solid_fill_with_a_color_transform_is_unknown(self):
        _, slide = new_slide()
        shape = self.rectangle(slide)
        shape.fill.solid()
        shape.fill.fore_color.theme_color = MSO_THEME_COLOR.ACCENT_1
        shape.fill.fore_color.brightness = 0.4
        self.assertEqual(rs.fill_state(shape), "unknown")

    def test_a_solid_theme_fill_without_a_theme_is_unknown(self):
        _, slide = new_slide()
        shape = self.rectangle(slide)
        shape.fill.solid()
        shape.fill.fore_color.theme_color = MSO_THEME_COLOR.ACCENT_1
        self.assertEqual(rs.fill_state(shape), "solid")
        drop_layout_relationship(slide)
        self.assertEqual(rs.fill_state(shape), "unknown")

    def test_the_state_is_solid_exactly_when_fill_color_has_a_value(self):
        _, slide = new_slide()
        cases = []
        shape = self.rectangle(slide)
        cases.append(shape)
        shape = self.rectangle(slide)
        shape.fill.solid()
        shape.fill.fore_color.rgb = RGBColor(1, 2, 3)
        cases.append(shape)
        shape = self.rectangle(slide)
        shape.fill.solid()
        shape.fill.fore_color.theme_color = MSO_THEME_COLOR.ACCENT_1
        shape.fill.fore_color.brightness = -0.25
        cases.append(shape)
        shape = self.rectangle(slide)
        shape.fill.background()
        cases.append(shape)
        shape = self.rectangle(slide)
        shape.fill.gradient()
        cases.append(shape)
        cases.append(add_box(slide, "text"))
        for shape in cases:
            with self.subTest(shape=shape.name):
                self.assertEqual(rs.fill_state(shape) == "solid", rs.fill_color(shape) is not None)

    def test_table_cells_use_the_same_three_words(self):
        prs, slide = new_slide()
        table = slide.shapes.add_table(1, 5, Inches(1), Inches(1), Inches(5), Inches(1))
        cells = [table.table.cell(0, column) for column in range(5)]
        cells[0].fill.solid()
        cells[0].fill.fore_color.rgb = RGBColor(0x11, 0x22, 0x33)
        cells[1].fill.background()
        cells[3].fill.gradient()
        cells[4].fill.solid()
        cells[4].fill.fore_color.theme_color = MSO_THEME_COLOR.ACCENT_1
        self.assertEqual([rs.fill_state(cell) for cell in cells], ["solid", "none", "unknown", "unknown", "solid"])
        self.assertEqual(rs.fill_color(cells[0]), "#112233")
        self.assertEqual(rs.fill_color(cells[4]), "#4F81BD")
        self.assertIsNone(rs.fill_color(cells[2]))

    def test_an_object_that_is_not_a_shape_is_unknown(self):
        self.assertEqual(rs.fill_state(object()), "unknown")
        self.assertIsNone(rs.fill_color(object()))


# ---------------------------------------------------------------------------
# Boxes
# ---------------------------------------------------------------------------

class ShapeBox(DeckCase):
    def test_a_top_level_shape_gives_its_own_box(self):
        _, slide = new_slide()
        box = slide.shapes.add_textbox(Inches(1), Inches(2), Inches(3), Inches(4))
        self.assertEqual(rs.shape_box(box), (Inches(1), Inches(2), Inches(3), Inches(4)))

    def test_a_placeholder_gives_the_box_it_inherits(self):
        _, slide = new_slide("Title Only")
        self.assertEqual(rs.shape_box(slide.shapes.title), (457200, 274638, 8229600, 1143000))

    def test_rotation_is_ignored(self):
        _, slide = new_slide()
        box = slide.shapes.add_textbox(Inches(1), Inches(2), Inches(3), Inches(4))
        box.rotation = 90
        self.assertEqual(rs.shape_box(box), (Inches(1), Inches(2), Inches(3), Inches(4)))

    def test_the_child_of_a_doubled_group_has_double_the_width_and_offset(self):
        _, slide = new_slide()
        group = slide.shapes.add_group_shape()
        first = group.shapes.add_textbox(Inches(1), Inches(1), Inches(2), Inches(1))
        second = group.shapes.add_textbox(Inches(4), Inches(2), Inches(2), Inches(1))
        off, ext, ch_off, ch_ext = group_transform(group)
        self.assertEqual(ext.get("cx"), str(Inches(5)))
        self.assertEqual(ch_ext.get("cx"), str(Inches(5)))
        ext.set("cx", str(int(ext.get("cx")) * 2))
        off.set("x", str(Inches(3)))
        self.assertEqual(rs.shape_box(first), (Inches(3), Inches(1), Inches(4), Inches(1)))
        self.assertEqual(rs.shape_box(second), (Inches(9), Inches(2), Inches(4), Inches(1)))

    def test_a_group_that_scales_height_changes_top_and_height(self):
        _, slide = new_slide()
        group = slide.shapes.add_group_shape()
        child = group.shapes.add_textbox(Inches(1), Inches(1), Inches(2), Inches(1))
        group.shapes.add_textbox(Inches(1), Inches(3), Inches(2), Inches(1))
        _, ext, _, _ = group_transform(group)
        ext.set("cy", str(int(ext.get("cy")) // 2))
        self.assertEqual(rs.shape_box(child), (Inches(1), Inches(1), Inches(2), Inches(0.5)))

    def test_nested_groups_multiply_their_scales(self):
        _, slide = new_slide()
        outer = slide.shapes.add_group_shape()
        inner = outer.shapes.add_group_shape()
        child = inner.shapes.add_textbox(Inches(1), Inches(1), Inches(1), Inches(1))
        _, inner_ext, _, _ = group_transform(inner)
        _, outer_ext, _, _ = group_transform(outer)
        inner_ext.set("cx", str(Inches(2)))
        outer_ext.set("cx", str(Inches(3)))
        left, top, width, height = rs.shape_box(child)
        self.assertEqual((left, width, top, height), (Inches(1), Inches(6), Inches(1), Inches(1)))

    def test_a_group_with_no_child_extent_gives_none(self):
        _, slide = new_slide()
        group = slide.shapes.add_group_shape()
        child = group.shapes.add_textbox(Inches(1), Inches(1), Inches(1), Inches(1))
        _, _, _, ch_ext = group_transform(group)
        ch_ext.set("cx", "0")
        self.assertIsNone(rs.shape_box(child))

    def test_a_shape_with_no_transform_gives_none(self):
        _, slide = new_slide()
        box = slide.shapes.add_textbox(Inches(1), Inches(1), Inches(1), Inches(1))
        sp_pr = box._element.find(qn("p:spPr"))
        sp_pr.remove(sp_pr.find(qn("a:xfrm")))
        self.assertIsNone(rs.shape_box(box))

    def test_a_group_with_no_transform_gives_none_for_its_child(self):
        _, slide = new_slide()
        group = slide.shapes.add_group_shape()
        child = group.shapes.add_textbox(Inches(1), Inches(1), Inches(1), Inches(1))
        grp_pr = group._element.grpSpPr
        grp_pr.remove(grp_pr.find(qn("a:xfrm")))
        self.assertIsNone(rs.shape_box(child))


class EffectivePpi(DeckCase):
    def test_a_cropped_large_picture(self):
        _, slide = new_slide()
        picture = slide.shapes.add_picture(str(self.png("large.png", 3840, 2160)), Inches(1), Inches(1),
                                           width=Inches(4))
        picture.crop_left = 0.25
        picture.crop_right = 0.25
        self.assertAlmostEqual(rs.effective_ppi(picture), 480.0, places=3)

    def test_a_small_picture(self):
        _, slide = new_slide()
        picture = slide.shapes.add_picture(str(self.png("small.png", 200, 100)), Inches(1), Inches(1),
                                           width=Inches(4))
        self.assertAlmostEqual(rs.effective_ppi(picture), 50.0, places=3)

    def test_the_lower_value_of_the_two_axes_counts(self):
        _, slide = new_slide()
        picture = slide.shapes.add_picture(str(self.png("small.png", 200, 100)), Inches(1), Inches(1),
                                           width=Inches(4), height=Inches(4))
        self.assertAlmostEqual(rs.effective_ppi(picture), 25.0, places=3)

    def test_a_group_that_scales_the_picture_changes_the_value(self):
        _, slide = new_slide()
        group = slide.shapes.add_group_shape()
        picture = group.shapes.add_picture(str(self.png("small.png", 200, 100)), Inches(1), Inches(1),
                                           width=Inches(4))
        group.shapes.add_textbox(Inches(1), Inches(4), Inches(4), Inches(1))
        _, ext, _, _ = group_transform(group)
        ext.set("cx", str(int(ext.get("cx")) * 2))
        self.assertAlmostEqual(rs.effective_ppi(picture), 25.0, places=3)

    def test_a_picture_with_no_area_gives_none(self):
        _, slide = new_slide()
        picture = slide.shapes.add_picture(str(self.png("small.png", 200, 100)), Inches(1), Inches(1),
                                           width=Inches(4))
        picture.width = 0
        self.assertIsNone(rs.effective_ppi(picture))

    def test_a_picture_cropped_to_nothing_gives_none(self):
        _, slide = new_slide()
        picture = slide.shapes.add_picture(str(self.png("small.png", 200, 100)), Inches(1), Inches(1),
                                           width=Inches(4))
        picture.crop_left = 0.5
        picture.crop_right = 0.5
        self.assertIsNone(rs.effective_ppi(picture))

    def test_a_picture_whose_image_part_is_missing_gives_none(self):
        _, slide = new_slide()
        picture = slide.shapes.add_picture(str(self.png("small.png", 200, 100)), Inches(1), Inches(1),
                                           width=Inches(4))
        blip = picture._element.xpath(".//a:blip")[0]
        blip.set(qn("r:embed"), "rId99")
        self.assertIsNone(rs.effective_ppi(picture))


# ---------------------------------------------------------------------------
# Title and words
# ---------------------------------------------------------------------------

class FindTitle(DeckCase):
    def test_finds_the_title_placeholder(self):
        _, slide = new_slide("Title Only")
        add_box(slide, "other")
        self.assertIs(rs.find_title(slide)._element, slide.shapes.title._element)

    def test_finds_the_title_slide_placeholder(self):
        _, slide = new_slide("Title Slide")
        self.assertEqual(rs.find_title(slide).name, slide.shapes.title.name)

    def test_finds_a_shape_named_title(self):
        _, slide = new_slide()
        add_box(slide, "other", top=3.0)
        named = add_box(slide, "the title", name="Title")
        self.assertIs(rs.find_title(slide)._element, named._element)

    def test_the_placeholder_wins_over_a_shape_named_title(self):
        _, slide = new_slide("Title Only")
        add_box(slide, "decoy", name="Title")
        self.assertIs(rs.find_title(slide)._element, slide.shapes.title._element)

    def test_a_slide_with_no_title_gives_none(self):
        _, slide = new_slide()
        add_box(slide, "Title text in a box that has another name")
        self.assertIsNone(rs.find_title(slide))

    def test_a_shape_with_no_name_does_not_stop_the_search(self):
        _, slide = new_slide()
        no_name = add_box(slide, "no name", top=0.2)
        no_name._element.find(qn("p:nvSpPr")).find(qn("p:cNvPr")).attrib.pop("name")
        no_properties = add_box(slide, "no properties", top=0.4)
        no_properties._element.remove(no_properties._element.find(qn("p:nvSpPr")))
        self.assertIsNone(rs.find_title(slide))
        named = add_box(slide, "the title", name="Title", top=3.0)
        self.assertIs(rs.find_title(slide)._element, named._element)

    def test_a_body_placeholder_is_not_a_title(self):
        _, slide = new_slide("Title and Content")
        slide.shapes.title._element.getparent().remove(slide.shapes.title._element)
        self.assertIsNone(rs.find_title(slide))


def damage_no_nv_sp_pr(case, prs, slide):
    box = add_box(slide, "no nvSpPr words", top=2.0)
    box._element.remove(box._element.find(qn("p:nvSpPr")))
    return lambda: box.name


def damage_no_name(case, prs, slide):
    box = add_box(slide, "no name words", top=2.0)
    box._element.find(qn("p:nvSpPr")).find(qn("p:cNvPr")).attrib.pop("name")
    return lambda: box.name


def damage_no_sp_pr(case, prs, slide):
    box = add_box(slide, "no spPr words", top=2.0)
    box._element.remove(box._element.find(qn("p:spPr")))
    return lambda: box.left


def damage_transform_without_x(case, prs, slide):
    box = add_box(slide, "no x words", top=2.0)
    box._element.find(qn("p:spPr")).find(qn("a:xfrm")).find(qn("a:off")).attrib.pop("x")
    return lambda: box.left


def damage_no_blip_fill(case, prs, slide):
    picture = slide.shapes.add_picture(str(case.png("p.png", 20, 10)), Inches(1), Inches(3), width=Inches(2))
    picture._element.remove(picture._element.find(qn("p:blipFill")))
    return lambda: picture.image


def damage_bad_crop(case, prs, slide):
    picture = slide.shapes.add_picture(str(case.png("p.png", 20, 10)), Inches(1), Inches(3), width=Inches(2))
    picture._element.find(qn("p:blipFill")).append(parse_xml('<a:srcRect %s l="wide"/>' % nsdecls("a")))
    return lambda: picture.crop_left


def damage_non_numeric_size(case, prs, slide):
    box = add_box(slide, "bad size words", top=2.0)
    run_props(box.text_frame.paragraphs[0].runs[0], '<a:rPr %s sz="big" b="maybe"/>' % nsdecls("a"))
    return lambda: box.text_frame.paragraphs[0].runs[0].font.size


def damage_run_without_text(case, prs, slide):
    box = add_box(slide, "no text element words", top=2.0)
    run = box.text_frame.paragraphs[0].runs[0]
    run._r.remove(run._r.find(qn("a:t")))
    return lambda: run.text


def damage_layout_is_not_a_part(case, prs, slide):
    for rel in slide.part.rels.values():
        if rel.reltype == RT.SLIDE_LAYOUT:
            rel._target_mode = RTM.EXTERNAL
    return lambda: slide.slide_layout


def damage_ragged_table(case, prs, slide):
    table = slide.shapes.add_table(2, 3, Inches(1), Inches(3), Inches(4), Inches(1))
    table.table.cell(0, 0).text = "header words"
    table.table.cell(1, 0).text = "body words"
    row = table._element.findall(".//" + qn("a:tr"))[1]
    row.remove(row.findall(qn("a:tc"))[-1])

    def bite():
        if len(row.findall(qn("a:tc"))) != 3:
            raise ValueError("the row has fewer cells than the table has columns")
    return bite


def damage_no_graphic(case, prs, slide):
    table = slide.shapes.add_table(1, 1, Inches(1), Inches(3), Inches(2), Inches(1))
    table._element.remove(table._element.find(qn("a:graphic")))
    return lambda: table.has_table


def damage_no_table_element(case, prs, slide):
    table = slide.shapes.add_table(1, 1, Inches(1), Inches(3), Inches(2), Inches(1))
    graphic_data = table._element.find(qn("a:graphic")).find(qn("a:graphicData"))
    graphic_data.remove(graphic_data.find(qn("a:tbl")))
    return lambda: list(table.table.rows)


def damage_blip_without_embed(case, prs, slide):
    picture = slide.shapes.add_picture(str(case.png("p.png", 20, 10)), Inches(1), Inches(3), width=Inches(2))
    blip = picture._element.find(qn("p:blipFill")).find(qn("a:blip"))
    blip.attrib.pop(qn("r:embed"))
    return lambda: picture.image


DAMAGES = [
    ("a shape with no nvSpPr", damage_no_nv_sp_pr),
    ("a cNvPr with no name", damage_no_name),
    ("a shape with no spPr", damage_no_sp_pr),
    ("a transform with no x", damage_transform_without_x),
    ("a picture with no blipFill", damage_no_blip_fill),
    ("a picture with a bad crop", damage_bad_crop),
    ("a run with a non-numeric sz", damage_non_numeric_size),
    ("a run with no text element", damage_run_without_text),
    ("a layout relationship that is not a part", damage_layout_is_not_a_part),
    ("a table with a ragged row", damage_ragged_table),
    ("a graphic frame with no a:graphic", damage_no_graphic),
    ("a table frame with no a:tbl", damage_no_table_element),
    ("a picture whose blip has no embed", damage_blip_without_embed),
]


class ForeignDeckAudit(DeckCase):
    """Call every public function of the module on a deck with a defect. No call may raise.

    A package with a layout part that is missing from the file cannot be opened at all.
    So the layout defect is a relationship that is not a part (an external target).
    """

    def deck(self, damage):
        prs, slide = new_slide()
        add_box(slide, "a readable box with words", size=14)
        bite = damage(self, prs, slide)
        return prs, slide, bite

    def failures(self, prs, slide):
        """Return the list of calls that raised, and the number of shapes that the walk visited."""
        found = []
        visited = [0]

        def attempt(label, function, *args, **kwargs):
            try:
                return function(*args, **kwargs)
            except Exception as error:  # the audit reports every kind of error
                found.append("{}: {}: {}".format(label, type(error).__name__, error))
                return None

        def quiet(function):
            try:
                return function()
            except Exception:  # python-pptx may raise on a defect. That is not a call of this module.
                return None

        attempt("background_color", rs.background_color, prs, slide)
        attempt("find_title", rs.find_title, slide)
        attempt("safe_shapes", lambda: list(rs.safe_shapes(slide)))
        report = attempt("safe_shapes_report", rs.safe_shapes_report, slide)
        shapes = [] if report is None else list(report[0])
        index = 0
        while index < len(shapes):
            shape = shapes[index]
            label = "shape {}".format(index)
            index += 1
            visited[0] += 1
            if isinstance(shape, GroupShape):
                attempt(label + " safe_shapes_report", rs.safe_shapes_report, shape)
                shapes.extend(quiet(lambda: list(rs.safe_shapes(shape))) or [])
            for paragraph, run in attempt(label + " iter_text_runs", lambda: list(rs.iter_text_runs(shape))) or []:
                for flag in (True, False):
                    attempt(label + " resolve_run", rs.resolve_run, prs, slide, shape, paragraph, run,
                            paragraph_defaults=flag)
            attempt(label + " shape_box", rs.shape_box, shape)
            attempt(label + " fill_color", rs.fill_color, shape)
            attempt(label + " fill_state", rs.fill_state, shape)
            attempt(label + " word_count", rs.word_count, shape)
            attempt(label + " effective_ppi", rs.effective_ppi, shape)
            for cell in quiet(lambda: [c for row in shape.table.rows for c in row.cells]) or []:
                attempt(label + " cell fill_color", rs.fill_color, cell)
                attempt(label + " cell fill_state", rs.fill_state, cell)
        return found, visited[0]

    def test_each_defect_bites_python_pptx(self):
        # Without this check, a defect that python-pptx tolerates would make the audit pass for nothing.
        for name, damage in DAMAGES:
            with self.subTest(defect=name):
                _, _, bite = self.deck(damage)
                with self.assertRaises(Exception):
                    bite()

    def test_no_public_function_raises_on_a_deck_with_a_defect(self):
        for name, damage in DAMAGES:
            with self.subTest(defect=name):
                prs, slide, _ = self.deck(damage)
                found, visited = self.failures(prs, slide)
                self.assertEqual(found, [])
                self.assertGreaterEqual(visited, 1)

    def test_the_audit_reports_a_call_that_raises(self):
        prs, slide, _ = self.deck(damage_no_name)
        with mock.patch.object(rs, "shape_box", side_effect=ValueError("forced")):
            found, _ = self.failures(prs, slide)
        self.assertTrue(any("shape_box: ValueError: forced" in line for line in found), found)

    def test_a_deck_with_every_defect_at_once_raises_nothing(self):
        prs, slide = new_slide()
        add_box(slide, "a readable box with words", size=14)
        for name, damage in DAMAGES:
            if "layout" not in name:
                damage(self, prs, slide)
        found, visited = self.failures(prs, slide)
        self.assertEqual(found, [])
        self.assertGreaterEqual(visited, 8)


class SafeShapes(DeckCase):
    def slide_with_an_invalid_placeholder_type(self):
        prs, slide = new_slide("Title Only")
        add_box(slide, "text", name="Readable")
        slide.shapes.title._element.find(qn("p:nvSpPr")).find(qn("p:nvPr")).find(qn("p:ph")).set("type", "bogus")
        return prs, slide

    def test_python_pptx_reads_one_shape_at_a_time_on_that_slide(self):
        # python-pptx raises KeyError for the invalid placeholder only. len() and the other index work.
        _, slide = self.slide_with_an_invalid_placeholder_type()
        self.assertEqual(len(slide.shapes), 2)
        with self.assertRaises(KeyError):
            slide.shapes[0]
        self.assertEqual(slide.shapes[1].name, "Readable")
        with self.assertRaises(KeyError):
            list(slide.shapes)

    def test_yields_the_readable_shapes_of_that_slide(self):
        _, slide = self.slide_with_an_invalid_placeholder_type()
        self.assertEqual([shape.name for shape in rs.safe_shapes(slide)], ["Readable"])

    def test_the_report_counts_the_shapes_that_it_could_not_read(self):
        _, slide = self.slide_with_an_invalid_placeholder_type()
        shapes, unreadable = rs.safe_shapes_report(slide)
        self.assertEqual(([shape.name for shape in shapes], unreadable), (["Readable"], 1))
        add_box(slide, "second", top=2.0, name="Second")
        bad = add_box(slide, "another bad one")
        bad._element.find(qn("p:nvSpPr")).find(qn("p:nvPr")).append(parse_xml('<p:ph %s type="bogus2"/>' % nsdecls("p")))
        shapes, unreadable = rs.safe_shapes_report(slide)
        self.assertEqual(([shape.name for shape in shapes], unreadable), (["Readable", "Second"], 2))

    def test_a_slide_that_is_clean_has_a_count_of_zero(self):
        _, slide = new_slide("Title Only")
        add_box(slide, "text")
        shapes, unreadable = rs.safe_shapes_report(slide)
        self.assertEqual((len(shapes), unreadable), (2, 0))
        empty_prs, empty_slide = new_slide()
        self.assertEqual(rs.safe_shapes_report(empty_slide), ([], 0))

    def test_a_container_whose_shapes_cannot_be_counted_gives_no_shapes_and_a_count_of_one(self):
        class CannotCount(object):
            def __len__(self):
                raise KeyError("bogus")

        class Unreadable(object):
            @property
            def shapes(self):
                raise ValueError("bogus")

        class Container(object):
            shapes = CannotCount()

        self.assertEqual(rs.safe_shapes_report(Container()), ([], 1))
        self.assertEqual(rs.safe_shapes_report(Unreadable()), ([], 1))
        self.assertEqual(list(rs.safe_shapes(Container())), [])

    def test_find_title_works_on_the_readable_shapes_of_that_slide(self):
        _, slide = self.slide_with_an_invalid_placeholder_type()
        self.assertIsNone(rs.find_title(slide))
        add_box(slide, "the title", name="Title")
        self.assertEqual(rs.find_title(slide).name, "Title")

    def test_yields_the_shapes_of_a_slide_a_layout_a_master_and_a_group(self):
        _, slide = new_slide("Title and Content")
        group = slide.shapes.add_group_shape()
        group.shapes.add_textbox(Inches(1), Inches(1), Inches(1), Inches(1))
        group.shapes.add_textbox(Inches(2), Inches(2), Inches(1), Inches(1))
        for container in (slide, slide.slide_layout, slide.slide_layout.slide_master, group):
            with self.subTest(container=type(container).__name__):
                expected = [shape.name for shape in container.shapes]
                self.assertTrue(expected)
                self.assertEqual([shape.name for shape in rs.safe_shapes(container)], expected)
                self.assertEqual(rs.safe_shapes_report(container)[1], 0)

    def test_a_wrong_kind_of_object_raises_what_python_pptx_raises(self):
        prs, _ = new_slide()
        for wrong in (object(), None, 42, prs):
            with self.subTest(wrong=type(wrong).__name__):
                with self.assertRaises((AttributeError, TypeError)):
                    list(rs.safe_shapes(wrong))
                with self.assertRaises((AttributeError, TypeError)):
                    rs.safe_shapes_report(wrong)

    def test_a_group_whose_shapes_cannot_be_read_has_no_runs_and_no_words(self):
        _, slide = new_slide()
        group = slide.shapes.add_group_shape()
        group.shapes.add_textbox(Inches(1), Inches(1), Inches(1), Inches(1)).text_frame.text = "hidden words"
        with mock.patch.object(GroupShape, "shapes", new_callable=mock.PropertyMock, side_effect=KeyError("bogus")):
            self.assertEqual(list(rs.iter_text_runs(group)), [])
            self.assertEqual(rs.word_count(group), 0)


class WordCount(DeckCase):
    def test_counts_the_words_of_a_text_box(self):
        _, slide = new_slide()
        box = add_box(slide, "one two  three")
        box.text_frame.add_paragraph().add_run().text = "four"
        self.assertEqual(rs.word_count(box), 4)

    def test_a_word_split_across_runs_counts_once(self):
        _, slide = new_slide()
        box = add_box(slide, "Hel")
        box.text_frame.paragraphs[0].add_run().text = "lo there"
        self.assertEqual(rs.word_count(box), 2)

    def test_a_line_break_separates_words(self):
        _, slide = new_slide()
        box = add_box(slide, "left")
        box.text_frame.paragraphs[0].add_line_break()
        box.text_frame.paragraphs[0].add_run().text = "right"
        self.assertEqual(rs.word_count(box), 2)

    def test_empty_and_blank_runs_count_for_nothing(self):
        _, slide = new_slide()
        box = add_box(slide, "")
        box.text_frame.paragraphs[0].add_run().text = "   "
        box.text_frame.add_paragraph()
        self.assertEqual(rs.word_count(box), 0)

    def test_counts_the_cells_of_a_table(self):
        _, slide = new_slide()
        table = slide.shapes.add_table(2, 2, Inches(1), Inches(1), Inches(4), Inches(2))
        table.table.cell(0, 0).text = "cell one"
        table.table.cell(1, 1).text = "three words here"
        self.assertEqual(rs.word_count(table), 5)

    def test_counts_the_children_of_a_group(self):
        _, slide = new_slide()
        group = slide.shapes.add_group_shape()
        group.shapes.add_textbox(Inches(1), Inches(1), Inches(1), Inches(1)).text_frame.text = "one two"
        group.shapes.add_textbox(Inches(1), Inches(2), Inches(1), Inches(1)).text_frame.text = "three"
        self.assertEqual(rs.word_count(group), 3)

    def test_field_text_counts_as_words_but_a_field_is_not_a_run(self):
        _, slide = new_slide()
        box = add_box(slide, "page ")
        field = parse_xml('<a:fld %s id="{5BCAD085-E8A6-8845-BD4E-CB4CCA059FC4}" type="slidenum">'
                          '<a:rPr lang="en-US"/><a:t>12 of 40</a:t></a:fld>' % nsdecls("a"))
        box.text_frame.paragraphs[0]._p.append(field)
        self.assertEqual(rs.word_count(box), 4)
        self.assertEqual([run.text for _, run in rs.iter_text_runs(box)], ["page "])

    def test_a_picture_has_no_words(self):
        _, slide = new_slide()
        picture = slide.shapes.add_picture(str(self.png("p.png", 20, 10)), Inches(1), Inches(1), width=Inches(2))
        self.assertEqual(rs.word_count(picture), 0)

    def test_a_run_with_no_text_element_is_not_an_error(self):
        _, slide = new_slide()
        box = add_box(slide, "one")
        run = box.text_frame.paragraphs[0].runs[0]
        run._r.remove(run._r.find(qn("a:t")))
        self.assertEqual(rs.word_count(box), 0)


# ---------------------------------------------------------------------------
# Public helpers, fills with inheritance and slides that python-pptx cannot read
# ---------------------------------------------------------------------------

class PublicHelpers(DeckCase):
    def test_find_element_walks_a_path_of_prefixed_names(self):
        _, slide = new_slide("Title Only")
        sp_tree = slide._element.cSld.spTree
        self.assertIs(rs.find_element(slide._element, "p:cSld", "p:spTree"), sp_tree)
        self.assertIsNone(rs.find_element(slide._element, "p:cSld", "p:nothing"))
        self.assertIsNone(rs.find_element(None, "p:cSld"))

    def test_paragraph_text_joins_runs_and_fields_and_marks_a_line_break(self):
        _, slide = new_slide()
        box = add_box(slide, "one")
        paragraph = box.text_frame.paragraphs[0]
        paragraph.add_line_break()
        paragraph.add_run().text = "two"
        paragraph._p.append(parse_xml('<a:fld %s id="{1}" type="slidenum"><a:t>7</a:t></a:fld>' % nsdecls("a")))
        self.assertEqual(rs.paragraph_text(paragraph._p), "one\ntwo7")

    def test_table_cells_lists_the_cells_in_row_order(self):
        _, slide = new_slide()
        table = slide.shapes.add_table(2, 2, Inches(1), Inches(1), Inches(4), Inches(1))
        for index, cell in enumerate(rs.table_cells(table)):
            cell.text = str(index)
        self.assertEqual([cell.text for cell in rs.table_cells(table)], ["0", "1", "2", "3"])
        self.assertEqual(rs.table_cells(add_box(slide, "not a table")), [])

    def test_the_fill_tags_name_the_six_fill_elements(self):
        self.assertEqual(len(rs.FILL_TAGS), 6)
        self.assertIn(qn("a:solidFill"), rs.FILL_TAGS)
        self.assertNotIn(qn("a:ln"), rs.FILL_TAGS)

    def test_placeholder_element_finds_the_ph_of_a_text_shape_a_picture_and_a_table(self):
        prs, slide = new_slide("Title Only")
        self.assertEqual(rs.placeholder_element(slide.shapes.title).get("type"), "title")
        self.assertIsNone(rs.placeholder_element(add_box(slide, "plain")))
        picture_slide = prs.slides.add_slide(prs.slide_layouts.get_by_name("Picture with Caption"))
        placeholder = [p for p in picture_slide.placeholders if p.placeholder_format.type == PP_PLACEHOLDER.PICTURE][0]
        picture = placeholder.insert_picture(str(self.png("p.png", 20, 10)))
        self.assertEqual(picture._element.tag, qn("p:pic"))
        self.assertEqual(rs.placeholder_element(picture).get("type"), "pic")

    def test_placeholder_class_groups_the_types_that_share_a_style(self):
        self.assertEqual(rs.placeholder_class("title"), "title")
        self.assertEqual(rs.placeholder_class("ctrTitle"), "title")
        for ph_type in ("dt", "ftr", "sldNum", "hdr"):
            self.assertEqual(rs.placeholder_class(ph_type), "other")
        for ph_type in ("body", "obj", "subTitle", "pic", "tbl", "chart"):
            self.assertEqual(rs.placeholder_class(ph_type), "body")


class SafeShapesInvalidXml(DeckCase):
    def reloaded(self, damage):
        """Return a slide of a saved and reloaded deck after `damage(slide)`. python-pptx caches the shapes of a slide that it built."""
        prs, _ = new_slide()
        slide = self.reopen(prs).slides[0]
        damage(slide)
        return slide

    def test_a_slide_with_no_shape_tree_gives_no_shapes_and_a_count_of_one(self):
        def damage(slide):
            sp_tree = slide._element.cSld.spTree
            sp_tree.getparent().remove(sp_tree)
        slide = self.reloaded(damage)
        self.assertEqual(rs.safe_shapes_report(slide), ([], 1))
        self.assertEqual(list(rs.safe_shapes(slide)), [])

    def test_a_slide_with_no_common_slide_data_gives_no_shapes_and_a_count_of_one(self):
        slide = self.reloaded(lambda slide: slide._element.remove(slide._element.cSld))
        self.assertEqual(rs.safe_shapes_report(slide), ([], 1))
        self.assertIsNone(rs.find_title(slide))


    def test_a_shape_that_cannot_be_read_by_index_with_invalid_xml_is_counted(self):
        from pptx.exc import InvalidXmlError

        class Shapes(object):
            def __len__(self):
                return 3

            def __getitem__(self, index):
                if index == 1:
                    raise InvalidXmlError("bad shape")
                return "shape %d" % index

        class Container(object):
            shapes = Shapes()

        self.assertEqual(rs.safe_shapes_report(Container()), (["shape 0", "shape 2"], 1))


class SafeShapesCost(unittest.TestCase):
    """python-pptx builds a list of all member elements on every shapes[index], so a read by index costs n squared."""

    class Shapes(object):
        def __init__(self, items, fail_at=None):
            self.items, self.fail_at, self.index_reads = items, fail_at, 0

        def __len__(self):
            return len(self.items)

        def __getitem__(self, index):
            self.index_reads += 1
            if index == self.fail_at:
                raise KeyError("bad shape")
            return self.items[index]

        def __iter__(self):
            for index, item in enumerate(self.items):
                if index == self.fail_at:
                    raise KeyError("bad shape")
                yield item

    class IndexedShapes(object):
        """No __iter__: the sequence protocol reads index 0, 1, 2 and so on, and this fake never ends."""

        def __init__(self, length, fail_at=None):
            self.length, self.fail_at = length, fail_at

        def __len__(self):
            return self.length

        def __getitem__(self, index):
            if index == self.fail_at:
                raise KeyError("bad shape")
            return "shape %d" % index

    def container(self, shapes):
        holder = type("Container", (object,), {})()
        holder.shapes = shapes
        return holder

    def test_a_readable_container_is_read_in_one_pass_with_no_read_by_index(self):
        shapes = self.Shapes(["a", "b", "c"])
        self.assertEqual(rs.safe_shapes_report(self.container(shapes)), (["a", "b", "c"], 0))
        self.assertEqual(shapes.index_reads, 0)

    def test_a_container_with_an_unreadable_shape_falls_back_to_one_read_per_shape(self):
        shapes = self.Shapes(["a", "b", "c"], fail_at=1)
        self.assertEqual(rs.safe_shapes_report(self.container(shapes)), (["a", "c"], 1))
        self.assertEqual(shapes.index_reads, 3)

    def test_a_container_that_cannot_be_iterated_is_read_by_index_up_to_its_length(self):
        shapes = self.IndexedShapes(4)
        self.assertEqual(rs.safe_shapes_report(self.container(shapes)), (["shape 0", "shape 1", "shape 2", "shape 3"], 0))

    def test_a_container_with_no_iteration_and_a_bad_shape_counts_it(self):
        shapes = self.IndexedShapes(3, fail_at=1)
        self.assertEqual(rs.safe_shapes_report(self.container(shapes)), (["shape 0", "shape 2"], 1))

    def test_a_real_slide_is_read_in_one_pass(self):
        _, slide = new_slide()
        for _ in range(40):
            slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(1), Inches(1), Inches(1), Inches(1))
        with mock.patch.object(type(slide.shapes), "__getitem__", side_effect=AssertionError("read by index")):
            shapes, unreadable = rs.safe_shapes_report(slide)
        self.assertEqual((len(shapes), unreadable), (40, 0))


class PicturePlaceholderFill(DeckCase):
    def test_a_picture_placeholder_with_no_fill_element_is_unknown(self):
        prs, _ = new_slide()
        slide = prs.slides.add_slide(prs.slide_layouts.get_by_name("Picture with Caption"))
        placeholder = [p for p in slide.placeholders if p.placeholder_format.type == PP_PLACEHOLDER.PICTURE][0]
        picture = placeholder.insert_picture(str(self.png("p.png", 20, 10)))
        self.assertEqual(rs.fill_state(picture), "unknown")


class InheritedFill(DeckCase):
    def title_slide(self):
        prs, slide = new_slide("Title Only")
        slide.shapes.title.text_frame.text = "A title"
        return prs, slide

    def layout_title(self, slide):
        return placeholder_of_type(slide.slide_layout, PP_PLACEHOLDER.TITLE)

    def master_title(self, slide):
        return placeholder_of_type(slide.slide_layout.slide_master, PP_PLACEHOLDER.TITLE)

    def test_a_placeholder_that_no_level_fills_has_no_fill(self):
        _, slide = self.title_slide()
        title = slide.shapes.title
        self.assertEqual(rs.fill_state(title), "unknown")
        self.assertEqual(rs.inherited_fill_state(slide, title), "none")
        self.assertIsNone(rs.inherited_fill_color(slide, title))

    def test_a_solid_fill_in_the_layout_placeholder_counts(self):
        _, slide = self.title_slide()
        layout = self.layout_title(slide)
        layout.fill.solid()
        layout.fill.fore_color.rgb = RGBColor(0x0B, 0x1F, 0x3A)
        self.assertEqual(rs.inherited_fill_state(slide, slide.shapes.title), "solid")
        self.assertEqual(rs.inherited_fill_color(slide, slide.shapes.title), "#0B1F3A")

    def test_a_solid_fill_in_the_master_placeholder_counts_when_the_layout_sets_none(self):
        _, slide = self.title_slide()
        master = self.master_title(slide)
        master.fill.solid()
        master.fill.fore_color.rgb = RGBColor(0x22, 0x33, 0x44)
        self.assertEqual(rs.inherited_fill_color(slide, slide.shapes.title), "#223344")

    def test_the_layout_wins_over_the_master(self):
        _, slide = self.title_slide()
        for owner, color in ((self.layout_title(slide), RGBColor(1, 2, 3)), (self.master_title(slide), RGBColor(9, 9, 9))):
            owner.fill.solid()
            owner.fill.fore_color.rgb = color
        self.assertEqual(rs.inherited_fill_color(slide, slide.shapes.title), "#010203")

    def test_an_explicit_no_fill_in_the_layout_stops_the_search(self):
        _, slide = self.title_slide()
        self.layout_title(slide).fill.background()
        master = self.master_title(slide)
        master.fill.solid()
        master.fill.fore_color.rgb = RGBColor(9, 9, 9)
        self.assertEqual(rs.inherited_fill_state(slide, slide.shapes.title), "none")

    def test_the_own_fill_of_the_shape_wins(self):
        _, slide = self.title_slide()
        title = slide.shapes.title
        title.fill.solid()
        title.fill.fore_color.rgb = RGBColor(0xAA, 0xBB, 0xCC)
        layout = self.layout_title(slide)
        layout.fill.solid()
        layout.fill.fore_color.rgb = RGBColor(1, 2, 3)
        self.assertEqual(rs.inherited_fill_color(slide, title), "#AABBCC")

    def test_a_gradient_in_the_layout_is_unknown(self):
        _, slide = self.title_slide()
        self.layout_title(slide).fill.gradient()
        self.assertEqual(rs.inherited_fill_state(slide, slide.shapes.title), "unknown")
        self.assertIsNone(rs.inherited_fill_color(slide, slide.shapes.title))

    def test_a_theme_color_in_the_layout_resolves_through_the_theme(self):
        _, slide = self.title_slide()
        layout = self.layout_title(slide)
        layout.fill.solid()
        layout.fill.fore_color.theme_color = MSO_THEME_COLOR.ACCENT_2
        self.assertEqual(rs.inherited_fill_color(slide, slide.shapes.title), "#C0504D")

    def test_a_shape_that_is_not_a_placeholder_gives_its_own_state(self):
        _, slide = new_slide()
        box = add_box(slide, "plain")
        self.assertEqual(rs.inherited_fill_state(slide, box), "none")
        rectangle = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(1), Inches(1), Inches(1), Inches(1))
        rectangle.fill.solid()
        rectangle.fill.fore_color.rgb = RGBColor(5, 6, 7)
        self.assertEqual(rs.inherited_fill_state(slide, rectangle), "solid")
        self.assertEqual(rs.inherited_fill_color(slide, rectangle), "#050607")

    def test_a_body_placeholder_matches_the_layout_placeholder_of_its_class(self):
        prs, _ = new_slide()
        slide = prs.slides.add_slide(prs.slide_layouts.get_by_name("Title and Content"))
        slide.shapes.title.text_frame.text = "t"
        body = slide.placeholders[1]
        layout_body = [p for p in slide.slide_layout.placeholders if p.placeholder_format.idx == 1][0]
        layout_body.fill.solid()
        layout_body.fill.fore_color.rgb = RGBColor(0x10, 0x20, 0x30)
        self.assertEqual(rs.inherited_fill_color(slide, body), "#102030")
        self.assertEqual(rs.inherited_fill_state(slide, slide.shapes.title), "none")

    def test_a_layout_placeholder_of_another_class_with_the_same_index_does_not_count(self):
        prs, _ = new_slide()
        slide = prs.slides.add_slide(prs.slide_layouts.get_by_name("Title and Content"))
        slide.shapes.title.text_frame.text = "t"
        layout_body = [p for p in slide.slide_layout.placeholders if p.placeholder_format.idx == 1][0]
        layout_body.fill.solid()
        layout_body.fill.fore_color.rgb = RGBColor(0x10, 0x20, 0x30)
        layout_body._element.find(qn("p:nvSpPr")).find(qn("p:nvPr")).find(qn("p:ph")).set("idx", "0")
        self.assertEqual(rs.inherited_fill_state(slide, slide.shapes.title), "none")

    def test_a_table_cell_with_no_fill_stays_unknown(self):
        _, slide = new_slide()
        table = slide.shapes.add_table(1, 1, Inches(1), Inches(1), Inches(3), Inches(1))
        cell = table.table.cell(0, 0)
        self.assertEqual(rs.inherited_fill_state(slide, cell), "unknown")
        self.assertIsNone(rs.inherited_fill_color(slide, cell))

    def test_a_slide_with_no_layout_gives_unknown(self):
        _, slide = self.title_slide()
        drop_layout_relationship(slide)
        self.assertEqual(rs.inherited_fill_state(slide, slide.shapes.title), "unknown")


class FillColorOfLayoutShapes(DeckCase):
    def band_in_layout(self, slide, scheme=None, rgb=None):
        band = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, Inches(4), Inches(1))
        band.fill.solid()
        if scheme is not None:
            band.fill.fore_color.theme_color = scheme
        else:
            band.fill.fore_color.rgb = rgb
        element = band._element
        element.getparent().remove(element)
        slide.slide_layout.shapes._spTree.append(element)
        return [s for s in slide.slide_layout.shapes if s._element is element][0]

    def override(self, owner, mapping):
        """Set a color map override on a slide or a layout element."""
        element = owner._element
        old = element.find(qn("p:clrMapOvr"))
        if old is not None:
            element.remove(old)
        attributes = " ".join('{}="{}"'.format(key, value) for key, value in mapping.items())
        element.append(parse_xml("<p:clrMapOvr %s><a:overrideClrMapping %s/></p:clrMapOvr>"
                                 % (nsdecls("p", "a"), attributes + ' bg2="lt2" tx2="dk2" accent1="accent1" accent2="accent2" '
                                    'accent3="accent3" accent4="accent4" accent5="accent5" accent6="accent6" '
                                    'hlink="hlink" folHlink="folHlink"')))

    def test_the_color_map_override_of_the_layout_applies_to_its_shapes(self):
        _, slide = new_slide("Title Only")
        shape = self.band_in_layout(slide, scheme=MSO_THEME_COLOR.BACKGROUND_1)
        self.assertEqual(rs.fill_color(shape), "#FFFFFF")
        self.override(slide.slide_layout, {"bg1": "dk1", "tx1": "lt1"})
        self.assertEqual(rs.fill_color(shape), "#000000")

    def test_the_color_map_override_of_the_slide_applies_to_its_shapes(self):
        _, slide = new_slide("Title Only")
        box = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, Inches(1), Inches(1))
        box.fill.solid()
        box.fill.fore_color.theme_color = MSO_THEME_COLOR.BACKGROUND_1
        self.assertEqual(rs.fill_color(box), "#FFFFFF")
        self.override(slide, {"bg1": "dk1", "tx1": "lt1"})
        self.assertEqual(rs.fill_color(box), "#000000")

    def test_a_theme_color_fill_of_a_layout_shape_resolves(self):
        _, slide = new_slide("Title Only")
        shape = self.band_in_layout(slide, scheme=MSO_THEME_COLOR.ACCENT_2)
        self.assertEqual(rs.fill_color(shape), "#C0504D")
        self.assertEqual(rs.fill_state(shape), "solid")

    def test_a_theme_color_fill_of_a_master_shape_resolves(self):
        _, slide = new_slide("Title Only")
        band = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, Inches(4), Inches(1))
        band.fill.solid()
        band.fill.fore_color.theme_color = MSO_THEME_COLOR.ACCENT_3
        element = band._element
        element.getparent().remove(element)
        master = slide.slide_layout.slide_master
        master.shapes._spTree.append(element)
        shape = [s for s in master.shapes if s._element is element][0]
        self.assertEqual(rs.fill_color(shape), "#9BBB59")

    def test_an_rgb_fill_of_a_layout_shape_resolves(self):
        _, slide = new_slide("Title Only")
        shape = self.band_in_layout(slide, rgb=RGBColor(1, 2, 3))
        self.assertEqual(rs.fill_color(shape), "#010203")


# ---------------------------------------------------------------------------
# A deck from the slide helpers
# ---------------------------------------------------------------------------

class SlideHelpersDeck(DeckCase):
    def test_the_resolver_reads_back_the_values_the_helpers_wrote(self):
        import slide_helpers
        import style as style_mod

        style = style_mod.load_style("peers", "dark")
        prs = slide_helpers.new_presentation(style)
        slide_helpers.add_title_slide(prs, style, "Findings in one line", author="Ada Lovelace")
        prs = self.reopen(prs)
        slide = prs.slides[0]
        title = rs.find_title(slide)
        self.assertEqual(title.name, "Title")
        result = resolve(prs, slide, title)
        self.assertEqual(result.size_pt, 44.0)
        self.assertEqual(result.family, style["typography"]["body_font"])
        self.assertEqual(result.color, style["colors"]["text"].upper())
        self.assertIs(result.bold, False)
        self.assertEqual(result.unresolved, [])
        subtitle = resolve(prs, slide, shape_named(slide, "Subtitle"))
        self.assertEqual(subtitle.size_pt, float(style["typography"]["callout_size"]))
        self.assertEqual(subtitle.color, style["colors"]["muted"].upper())
        self.assertEqual(rs.background_color(prs, slide), style["colors"]["background"].upper())
        self.assertEqual(rs.fill_color(shape_named(slide, "Accent mark")), style["colors"]["accent"].upper())
        self.assertEqual(rs.word_count(title), 4)


# ---------------------------------------------------------------------------
# LibreOffice cross-check
# ---------------------------------------------------------------------------

SOFFICE = shutil.which("soffice")
PDFTOTEXT = shutil.which("pdftotext")
RENDER_TIMEOUT_SECONDS = 120
WORD = re.compile(r'<word xMin="([\d.]+)" yMin="([\d.]+)" xMax="([\d.]+)" yMax="([\d.]+)">([^<]*)</word>')


def run_with_timeout(command, env, seconds):
    """Run a command in its own process group. Kill the whole group on a timeout."""
    process = subprocess.Popen(command, stdout=subprocess.PIPE, stderr=subprocess.PIPE, env=env,
                               start_new_session=True)
    try:
        out, err = process.communicate(timeout=seconds)
    except subprocess.TimeoutExpired:
        os.killpg(process.pid, signal.SIGKILL)
        process.communicate()
        raise
    return process.returncode, out, err


ANCHOR = "AnchorTwenty"
EXPECTED_CROSS_CHECK_SIZES = {
    "TitleWord": 44.0, "BodyZero": 32.0, "BodyOne": 28.0, "RunTwelve": 12.0, "ParaThirty": 30.0,
    ANCHOR: 20.0, "PlainBox": 18.0, "ListStyled": 22.0, "GroupBox": 18.0,
}


def build_cross_check_deck():
    """Return a deck with one word per inheritance case. Each word is its own run."""
    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts.get_by_name("Title and Content"))
    slide.shapes.title.text_frame.text = "TitleWord"
    body = slide.placeholders[1].text_frame
    body.text = "BodyZero"
    one = body.add_paragraph()
    one.text = "BodyOne"
    one.level = 1
    twelve = body.add_paragraph()
    twelve.add_run().text = "RunTwelve"
    twelve.runs[0].font.size = Pt(12)
    thirty = body.add_paragraph()
    thirty.text = "ParaThirty"
    thirty.font.size = Pt(30)
    slide = prs.slides.add_slide(prs.slide_layouts.get_by_name("Blank"))
    add_box(slide, ANCHOR, size=20, top=0.5)
    add_box(slide, "PlainBox", top=1.5)
    listed = add_box(slide, "ListStyled", top=2.5)
    list_style_of(listed).append(level_style(0, '<a:defRPr sz="2200"/>'))
    group = slide.shapes.add_group_shape()
    group.shapes.add_textbox(Inches(1), Inches(3.5), Inches(4), Inches(0.6)).text_frame.text = "GroupBox"
    return prs


def resolved_sizes(prs):
    """Return a dict of run text to the size that the resolver gives the run."""
    sizes = {}
    for slide in prs.slides:
        for shape in slide.shapes:
            for paragraph, run in rs.iter_text_runs(shape):
                sizes[run.text] = rs.resolve_run(prs, slide, shape, paragraph, run).size_pt
    return sizes


class CrossCheckDeck(unittest.TestCase):
    def test_the_resolved_sizes_are_the_ones_we_expect(self):
        self.assertEqual(resolved_sizes(build_cross_check_deck()), EXPECTED_CROSS_CHECK_SIZES)


@unittest.skipUnless(
    SOFFICE and PDFTOTEXT,
    "this check needs soffice (LibreOffice) to render the deck and pdftotext to read the PDF, and one is missing",
)
class LibreOfficeCrossCheck(unittest.TestCase):
    """Compare the resolver sizes with the sizes that LibreOffice draws.

    A word in the PDF is as tall as its font size times a factor that depends on the font.
    All words here use one font family. The test takes the factor from a run with an
    explicit size and checks every other word against it.
    A LibreOffice that is installed but cannot render skips the class. It does not make an error.
    """

    @classmethod
    def setUpClass(cls):
        cls._tmp = tempfile.TemporaryDirectory()
        try:
            cls._render(Path(cls._tmp.name))
        except (RuntimeError, OSError, subprocess.TimeoutExpired) as error:
            cls._tmp.cleanup()
            raise unittest.SkipTest("LibreOffice is installed but could not render the deck: {}".format(error))
        except BaseException:
            cls._tmp.cleanup()
            raise

    @classmethod
    def tearDownClass(cls):
        cls._tmp.cleanup()

    @classmethod
    def _render(cls, folder):
        prs = build_cross_check_deck()
        path = folder / "deck.pptx"
        prs.save(str(path))
        cls.expected = resolved_sizes(prs)
        home = folder / "home"
        home.mkdir()
        env = dict(os.environ, HOME=str(home))
        profile = folder / "profile"
        code, out, err = run_with_timeout(
            [SOFFICE, "-env:UserInstallation=file://{}".format(profile), "--headless", "--norestore",
             "--convert-to", "pdf", "--outdir", str(folder), str(path)],
            env, RENDER_TIMEOUT_SECONDS)
        pdf = folder / "deck.pdf"
        if code != 0 or not pdf.exists():
            raise RuntimeError("soffice failed with code {}: {}".format(code, err.decode("utf-8", "replace")))
        boxes = folder / "boxes.html"
        code, out, err = run_with_timeout([PDFTOTEXT, "-bbox", str(pdf), str(boxes)], env, RENDER_TIMEOUT_SECONDS)
        if code != 0:
            raise RuntimeError("pdftotext failed with code {}: {}".format(code, err.decode("utf-8", "replace")))
        cls.heights = {}
        for match in WORD.finditer(boxes.read_text(encoding="utf-8")):
            cls.heights[match.group(5)] = float(match.group(4)) - float(match.group(2))

    def assert_renders_at_the_resolved_size(self, word):
        self.assertIn(word, self.heights, "the PDF has no word {!r}".format(word))
        factor = self.heights[ANCHOR] / self.expected[ANCHOR]
        size = self.expected[word]
        self.assertIsNotNone(size, word)
        self.assertAlmostEqual(self.heights[word], factor * size, delta=0.03 * factor * size + 0.3,
                               msg="{} resolved to {} pt".format(word, size))

    def test_a_title_placeholder_renders_at_the_resolved_size(self):
        self.assert_renders_at_the_resolved_size("TitleWord")

    def test_body_levels_render_at_the_resolved_sizes(self):
        self.assert_renders_at_the_resolved_size("BodyZero")
        self.assert_renders_at_the_resolved_size("BodyOne")

    def test_an_explicit_run_size_renders_at_the_resolved_size(self):
        self.assert_renders_at_the_resolved_size("RunTwelve")

    def test_a_paragraph_default_renders_at_the_resolved_size(self):
        self.assert_renders_at_the_resolved_size("ParaThirty")

    def test_text_boxes_render_at_the_resolved_sizes(self):
        self.assert_renders_at_the_resolved_size("PlainBox")
        self.assert_renders_at_the_resolved_size("ListStyled")
        self.assert_renders_at_the_resolved_size("GroupBox")


class BrokenLibreOffice(DeckCase):
    """A soffice that is installed but fails makes the cross-check class skip. It gives no error."""

    def script(self, name, body):
        path = self.tmp / name
        path.write_text("#!/bin/sh\n" + body + "\n")
        path.chmod(0o755)
        return str(path)

    def assert_the_class_skips(self, soffice, pdftotext=None, timeout=RENDER_TIMEOUT_SECONDS):
        this_module = sys.modules[__name__]
        with mock.patch.object(this_module, "SOFFICE", soffice), \
                mock.patch.object(this_module, "PDFTOTEXT", pdftotext), \
                mock.patch.object(this_module, "RENDER_TIMEOUT_SECONDS", timeout):
            with self.assertRaises(unittest.SkipTest) as caught:
                LibreOfficeCrossCheck.setUpClass()
        self.assertIn("LibreOffice", str(caught.exception))
        self.assertFalse(os.path.exists(LibreOfficeCrossCheck._tmp.name), "the temporary folder must be gone")

    def test_a_nonzero_exit_skips(self):
        self.assert_the_class_skips(self.script("soffice", "exit 3"))

    def test_an_exit_with_no_pdf_skips(self):
        self.assert_the_class_skips(self.script("soffice", "exit 0"))

    def test_a_timeout_skips(self):
        self.assert_the_class_skips(self.script("soffice", "exec sleep 30"), timeout=1)

    def test_a_missing_program_skips(self):
        self.assert_the_class_skips(str(self.tmp / "no-such-soffice"))

    def test_a_pdftotext_that_fails_skips(self):
        self.assert_the_class_skips(self.script("soffice", 'touch "$7/deck.pdf"'),
                                    pdftotext=self.script("pdftotext", "exit 2"))


if __name__ == "__main__":
    unittest.main()
