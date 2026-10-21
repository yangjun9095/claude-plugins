"""Tests for the text rules of scripts/design_lint.py: titles, words, bullets, objects, lines and fonts.
"""

import sys
from pathlib import Path

from pptx.chart.data import CategoryChartData
from pptx.dml.color import RGBColor
from pptx.enum.chart import XL_CHART_TYPE
from pptx.oxml import parse_xml
from pptx.oxml.ns import nsdecls, qn
from pptx.util import Inches, Pt

SCRIPTS_DIR = Path(__file__).resolve().parents[1] / "scripts"
for folder in (SCRIPTS_DIR, Path(__file__).resolve().parent):
    if str(folder) not in sys.path:
        sys.path.insert(0, str(folder))

import pptx_resolve as rs  # noqa: E402
from lint_helpers import (  # noqa: E402
    BODY,
    GOOD_TITLE,
    LIGHT,
    LintCase,
    add_box,
    add_bullets,
    add_chart,
    add_claim,
    add_rect,
    add_slide,
    add_source,
    add_table,
    clean_slide,
    new_deck,
    set_run,
    sh,
    words,
    write,
)


class TitleSentence(LintCase):
    def titled(self, title, kind=None, **options):
        prs = new_deck()
        add_slide(prs, title, kind=kind)
        return self.of(self.lint(prs, **options), "title-sentence")

    def test_a_topic_label_warns_and_says_so(self):
        hits = self.titled("Results")
        self.assertEqual(len(hits), 1)
        self.assertEqual((hits[0].slide, hits[0].severity, hits[0].tier), (1, "warn", "T3"))
        self.assertIn("topic label", hits[0].message)

    def test_a_title_with_five_words_warns_and_six_words_pass(self):
        five = self.titled("Binders fold in nine designs")
        self.assertEqual(len(five), 1)
        self.assertIn("5 words", five[0].message)
        self.assertIn("6 to 20", five[0].message)
        self.assertEqual(self.titled("Binders fold in nine of twelve"), [])

    def test_a_title_with_twenty_words_passes_and_twenty_one_warns(self):
        self.assertEqual(self.titled("Binders fold " + words(18, "again")), [])
        hits = self.titled("Binders fold " + words(19, "again"))
        self.assertEqual(len(hits), 1)
        self.assertIn("21 words", hits[0].message)

    def test_the_word_limits_come_from_the_style(self):
        prs = new_deck()
        add_slide(prs, "Binders fold in nine designs")
        project = self.tmp / "project"
        project.mkdir()
        (project / "slide-style-minimal.yaml").write_text("lint:\n  title_words: {min: 3, max: 8, tier: T2}\n")
        findings = self.lint(prs, project_dir=str(project))
        self.assertQuiet(findings, "title-sentence")

    def test_a_topic_label_shows_up_when_the_word_floor_is_low(self):
        prs = new_deck()
        add_slide(prs, "Results")
        project = self.tmp / "project"
        project.mkdir()
        (project / "slide-style-minimal.yaml").write_text("lint:\n  title_words: {min: 1, max: 20, tier: T3}\n")
        hit = self.only(self.lint(prs, project_dir=str(project)), "title-sentence")
        self.assertIn("topic label", hit.message)

    def test_a_title_with_no_verb_warns(self):
        hits = self.titled("Overview of the methods and the background material")
        self.assertEqual(len(hits), 1)
        self.assertIn("verb", hits[0].message)

    def test_titles_with_a_verb_pass(self):
        for title in (
            "One thing to remember from this talk",
            "Two things to remember from this talk",
            "Three things to remember from this talk",
            "Accuracy drops when the batch size doubles",
            "Binders fold in nine of twelve designs",
            "Most predicted binders are stable in the assay",
            "Sequencing reads cover every exon of the gene",
            "The model was trained on made up data only",
            "Our method outperforms the baseline on every set",
        ):
            with self.subTest(title=title):
                self.assertEqual(self.titled(title), [])

    def test_titles_with_no_verb_warn(self):
        for title in (
            "A map of cell types across the zebrafish embryo",
            "Overview of the methods and the background material",
            "The results of the first round on the cluster",
        ):
            with self.subTest(title=title):
                self.assertEqual(len(self.titled(title)), 1)

    def test_a_slide_with_a_centered_title_placeholder_is_a_title_slide_even_with_no_kind(self):
        prs = new_deck()
        slide = prs.slides.add_slide(prs.slide_layouts.get_by_name("Title Slide"))
        slide.shapes.title.text_frame.text = "Binder study"
        self.assertQuiet(self.lint(prs), "title-sentence")

    def test_a_title_placeholder_that_is_not_centered_is_checked(self):
        prs = new_deck()
        slide = prs.slides.add_slide(prs.slide_layouts.get_by_name("Title Only"))
        slide.shapes.title.text_frame.text = "Binder study"
        self.assertEqual(len(self.of(self.lint(prs), "title-sentence")), 1)

    def test_the_title_slide_is_exempt(self):
        self.assertEqual(self.titled("Results", kind="title"), [])

    def test_a_big_number_slide_is_not_exempt(self):
        self.assertEqual(len(self.titled("Results", kind="big_number")), 1)

    def test_the_default_take_home_headline_passes(self):
        for count in ("One", "Two", "Three"):
            with self.subTest(count=count):
                self.assertEqual(self.titled(count + " things to remember from this talk", kind="take_home"), [])

    def test_the_title_with_non_ascii_characters_is_read(self):
        self.assertEqual(self.titled("Binders fold at 5 µM in 3 × more designs than α controls"), [])

    def test_sentence_titles_from_a_pilot_pass_the_verb_check(self):
        for title in (
            "Only 11 of 52 test items in our review depend on a study or a standard",
            "Three rules carry the argument of this deck",
            "Peak calls generalize poorly across replicates while counts transfer reliably",
            "The new assay normalizes counts before the first merge",
            "Jobs made a large number concrete by restating it per second",
        ):
            with self.subTest(title=title):
                prs = new_deck()
                add_slide(prs, title)
                self.assertQuiet(self.lint(prs), "title-sentence")

    def test_samples_of_a_title_with_only_an_ed_or_ing_verb_pass(self):
        for title in ("Cells stained with the new dye glowed brighter under blue light",
                      "Washing buffers leaching metals into the eluate",
                      "The new assay normalizes counts across lanes"):
            with self.subTest(title=title):
                prs = new_deck()
                add_slide(prs, title)
                self.assertQuiet(self.lint(prs), "title-sentence")


TITLE_EMPTY = "no title placeholder or shape named Title holds text; add a headline that states the claim"


class TitleEmpty(LintCase):
    def test_a_slide_with_no_title_shape_is_an_error(self):
        prs = new_deck()
        slide = prs.slides.add_slide(prs.slide_layouts.get_by_name("Blank"))
        add_box(slide, BODY)
        hit = self.assertFires(self.lint(prs), "title-empty", "error", "T2")
        self.assertEqual(hit.message, TITLE_EMPTY)

    def test_a_title_in_a_plain_text_box_is_not_read(self):
        # v1 decks and many hand-made decks put the title in a text box with another name.
        prs = new_deck()
        slide = prs.slides.add_slide(prs.slide_layouts.get_by_name("Blank"))
        add_box(slide, GOOD_TITLE, name="TextBox 1", top=0.15, height=0.9, size=28)
        add_box(slide, BODY, top=1.6)
        findings = self.lint(prs)
        self.assertEqual(self.assertFires(findings, "title-empty", "error", "T2").message, TITLE_EMPTY)
        self.assertIn("TextBox 1 is 0.15 in from the top edge", self.assertFires(findings, "margin", "warn", "T3").message)

    def test_an_empty_title_placeholder_is_an_error(self):
        prs = new_deck()
        add_slide(prs, None)
        self.assertFires(self.lint(prs), "title-empty", "error", "T2")

    def test_a_blank_title_is_an_error(self):
        prs = new_deck()
        add_slide(prs, "   ")
        self.assertFires(self.lint(prs), "title-empty", "error", "T2")

    def test_the_title_slide_needs_a_title_too(self):
        prs = new_deck()
        add_slide(prs, None, kind="title")
        self.assertFires(self.lint(prs), "title-empty", "error", "T2")

    def test_a_slide_with_a_title_is_quiet(self):
        prs = new_deck()
        clean_slide(prs)
        self.assertQuiet(self.lint(prs), "title-empty")

    def test_a_shape_named_title_counts_as_the_title(self):
        prs = new_deck()
        slide = prs.slides.add_slide(prs.slide_layouts.get_by_name("Blank"))
        add_box(slide, GOOD_TITLE, name="Title", size=28)
        self.assertQuiet(self.lint(prs), "title-empty")


class TitleUnique(LintCase):
    def deck(self, first, second):
        prs = new_deck()
        clean_slide(prs, title=first)
        clean_slide(prs, title=second)
        return prs

    def test_a_repeated_title_warns_on_the_second_slide(self):
        hit = self.assertFires(self.lint(self.deck(GOOD_TITLE, GOOD_TITLE)), "title-unique", "warn", "T2", slide=2)
        self.assertIn("slide 1", hit.message)

    def test_the_comparison_ignores_case_and_spacing(self):
        findings = self.lint(self.deck(GOOD_TITLE, "  " + GOOD_TITLE.upper().replace(" ", "  ")))
        self.assertFires(findings, "title-unique", "warn", "T2", slide=2)

    def test_different_titles_are_quiet(self):
        self.assertQuiet(self.lint(self.deck(GOOD_TITLE, "Binder design fails for the last 3 targets")),
                         "title-unique")

    def test_a_hidden_slide_does_not_count(self):
        prs = self.deck(GOOD_TITLE, GOOD_TITLE)
        list(prs.slides)[0]._element.set("show", "0")
        self.assertQuiet(self.lint(prs), "title-unique")

    def test_empty_titles_are_not_repeats(self):
        prs = new_deck()
        add_slide(prs, None)
        add_slide(prs, None)
        self.assertQuiet(self.lint(prs), "title-unique")


class WordsPerSlide(LintCase):
    def with_words(self, count, kind=None, **options):
        prs = new_deck()
        slide = add_slide(prs, kind=kind)
        add_box(slide, words(count), height=5.0)
        return self.lint(prs, **options)

    def test_25_words_pass_and_26_warn(self):
        self.assertQuiet(self.with_words(25), "words-per-slide")
        hit = self.assertFires(self.with_words(26), "words-per-slide", "warn", "T3")
        self.assertIn("26 words", hit.message)

    def test_40_words_do_not_error_and_41_do(self):
        self.assertFires(self.with_words(40), "words-per-slide", "warn", "T3")
        self.assertFires(self.with_words(41), "words-per-slide", "error", "T3")

    def test_the_title_the_source_and_chart_labels_are_not_counted(self):
        prs = new_deck()
        slide = add_slide(prs, kind="assertion_evidence")
        add_box(slide, words(25))
        add_source(slide, words(30))
        frame = add_chart(slide, labels=True)
        # The chart part holds text of its own: the category names and the series name. The words rule skips it.
        names = [node.text for node in frame.chart._chartSpace.xpath(".//c:pt/c:v")]
        self.assertTrue({"a", "b", "c", "s0"} <= set(names), names)
        self.assertEqual(rs.word_count(frame), 0)
        self.assertQuiet(self.lint(prs), "words-per-slide")

    def test_an_unnamed_text_box_at_the_bottom_counts_as_the_source_line(self):
        prs = new_deck()
        slide = add_slide(prs)
        add_box(slide, words(25))
        add_source(slide, words(30), name="Footer note")
        self.assertQuiet(self.lint(prs), "words-per-slide")

    def test_an_unnamed_text_box_at_the_top_counts(self):
        prs = new_deck()
        slide = add_slide(prs)
        add_box(slide, words(25))
        add_box(slide, words(3), name="Extra", top=2.9)
        self.assertFires(self.lint(prs), "words-per-slide", "warn", "T3")

    def test_table_words_count(self):
        prs = new_deck()
        slide = add_slide(prs)
        add_table(slide, rows=3, cols=3, text=words(3))
        self.assertFires(self.lint(prs), "words-per-slide", "warn", "T3")

    def test_a_group_counts_once(self):
        prs = new_deck()
        slide = add_slide(prs)
        one = add_box(slide, words(13), name="One")
        two = add_box(slide, words(13), name="Two", top=3.0)
        slide.shapes.add_group_shape([one, two])
        hit = self.assertFires(self.lint(prs), "words-per-slide", "warn", "T3")
        self.assertIn("26 words", hit.message)

    def test_a_take_home_slide_is_skipped(self):
        prs = new_deck()
        slide = add_slide(prs, "Three things to remember from this talk", kind="take_home")
        for index in range(3):
            add_box(slide, words(12), name="Claim", top=1.6 + index * 1.5, height=0.6)
            add_box(slide, words(7), name="Baseline", top=2.2 + index * 1.5, height=0.5, size=20)
        self.assertEqual(sum(len(w.split()) for w in [words(12), words(7)]) * 3, 57)
        self.assertQuiet(self.lint(prs), "words-per-slide")

    def test_the_profiles_change_the_numbers(self):
        self.assertFires(self.with_words(41, profile="lab-meeting"), "words-per-slide", "warn", "T3")
        self.assertQuiet(self.with_words(40, profile="lab-meeting"), "words-per-slide")
        self.assertFires(self.with_words(61, profile="lab-meeting"), "words-per-slide", "error", "T3")
        self.assertFires(self.with_words(16, profile="general"), "words-per-slide", "warn", "T3")
        self.assertFires(self.with_words(26, profile="general"), "words-per-slide", "error", "T3")

    def test_a_shape_named_source_is_the_source_line_wherever_it_sits(self):
        prs = new_deck()
        slide = add_slide(prs)
        add_box(slide, words(25))
        add_box(slide, words(30), name="Source", top=3.0)
        self.assertQuiet(self.lint(prs), "words-per-slide")


class TakeHomeWords(LintCase):
    def take_home(self, claim_words, baseline_words):
        prs = new_deck()
        slide = add_slide(prs, "One thing to remember from this talk", kind="take_home")
        add_box(slide, words(claim_words), name="Claim", top=1.6, height=0.6)
        add_box(slide, words(baseline_words), name="Baseline", top=2.4, height=0.6, size=20)
        return self.lint(prs)

    def test_a_claim_of_16_words_warns_and_15_pass(self):
        hit = self.assertFires(self.take_home(16, 5), "take-home-words", "warn", "T3")
        self.assertIn("claim", hit.message)
        self.assertIn("16 words", hit.message)
        self.assertQuiet(self.take_home(15, 5), "take-home-words")

    def test_a_baseline_of_16_words_warns_and_15_pass(self):
        hit = self.assertFires(self.take_home(5, 16), "take-home-words", "warn", "T3")
        self.assertIn("baseline", hit.message)
        self.assertQuiet(self.take_home(5, 15), "take-home-words")

    def test_other_slides_are_not_checked(self):
        prs = new_deck()
        slide = add_slide(prs, kind="assertion_evidence")
        add_box(slide, words(16), name="Claim")
        self.assertQuiet(self.lint(prs), "take-home-words")


TAKE_HOME_TITLE = "Three things to remember from this talk"
SECONDARY_NOUN = "secondary text such as call-outs, tables and chart text"


class TakeHomeClaimBox(LintCase):
    """The helper puts each claim and its baseline in one Claim shape: the claim is paragraph 0, the baseline paragraph 1."""

    def deck(self, rows, kind="take_home"):
        """Return a deck with one slide. Each row is (claim, baseline, options of add_claim)."""
        prs = new_deck()
        slide = add_slide(prs, TAKE_HOME_TITLE, kind=kind)
        for index, (claim, baseline, options) in enumerate(rows):
            add_claim(slide, claim, baseline, top=1.8 + index * 1.5, **options)
        return prs

    def test_a_claim_paragraph_of_16_words_warns_and_15_pass(self):
        hit = self.assertFires(self.lint(self.deck([(words(16), words(5), {})])), "take-home-words", "warn", "T3")
        self.assertEqual(hit.message, "a claim has 16 words; keep at most 15")
        self.assertQuiet(self.lint(self.deck([(words(15), words(15), {})])), "take-home-words")

    def test_a_baseline_paragraph_of_16_words_warns_and_15_pass(self):
        hit = self.assertFires(self.lint(self.deck([(words(5), words(16), {})])), "take-home-words", "warn", "T3")
        self.assertEqual(hit.message, "a baseline has 16 words; keep at most 15")
        self.assertQuiet(self.lint(self.deck([(words(5), words(15), {})])), "take-home-words")

    def test_the_word_limits_come_from_the_style(self):
        project = self.tmp / "project"
        project.mkdir()
        (project / "slide-style-minimal.yaml").write_text("lint:\n  take_home_words: {max_claim: 4, max_baseline: 6}\n")
        findings = self.lint(self.deck([(words(5), words(7), {}), (words(4), words(6), {})]), project_dir=str(project))
        self.assertEqual([h.message for h in self.of(findings, "take-home-words")],
                         ["a claim has 5 words; keep at most 4", "a baseline has 7 words; keep at most 6"])

    def test_the_baseline_paragraph_is_secondary_text_and_the_claim_paragraph_primary(self):
        def body_font(claim_pt, baseline_pt, kind="take_home"):
            prs = self.deck([("A claim", "A baseline", {"size": claim_pt, "baseline_font": {"size": baseline_pt}})],
                            kind=kind)
            return [(h.severity, h.message) for h in self.of(self.lint(prs), "body-font")]
        self.assertEqual(body_font(24, 20), [])
        self.assertEqual(body_font(24, 19.9), [
            ("warn", "Claim (baseline) is 19.9 pt; use at least 20 pt for " + SECONDARY_NOUN)])
        self.assertEqual(body_font(24, 17.9), [("error", "Claim (baseline) is 17.9 pt; use at least 18 pt")])
        self.assertEqual(body_font(20, 20), [("warn", "Claim is 20 pt; use at least 24 pt for body text")])
        # A slide with no kind has no take-home parts, so paragraph 1 is primary text.
        self.assertEqual(body_font(24, 20, kind=None), [("warn", "Claim is 20 pt; use at least 24 pt for body text")])

    def test_a_baseline_paragraph_with_low_contrast_is_named_in_the_finding(self):
        hit = self.only(self.lint(self.deck([("A claim", "A baseline", {"baseline_font": {"color": "BBBBBB"}})])),
                        "text-contrast")
        self.assertEqual((hit.severity, hit.message), (
            "error", "Claim (baseline): text color #BBBBBB on #FFFFFF has a contrast of 1.92:1; use at least 3:1"))

    def test_each_paragraph_of_a_claim_shape_counts_its_own_lines(self):
        # A line of 11.28 in holds 81.2 characters at 20 pt and 67.7 at 24 pt, at 0.5 em per character.
        quiet = self.deck([("x" * 135, "x" * 162, {})])
        self.assertQuiet(self.lint(quiet), "lines-per-paragraph")
        hit = self.assertFires(self.lint(self.deck([("x" * 135, "x" * 163, {})])), "lines-per-paragraph", "warn", "T2")
        self.assertTrue(hit.message.startswith("Claim (baseline) has a paragraph of an estimated 3 lines"), hit.message)
        hit = self.assertFires(self.lint(self.deck([("x" * 136, "x" * 10, {})])), "lines-per-paragraph", "warn", "T2")
        self.assertTrue(hit.message.startswith("Claim has a paragraph of an estimated 3 lines"), hit.message)

    def test_the_fonts_of_both_paragraphs_count(self):
        rows = [("A claim", "A baseline", {}), ("A claim", "A baseline", {"baseline_font": {"size": 22}})]
        hit = self.assertFires(self.lint(self.deck(rows)), "fonts-per-slide", "warn", "T2")
        self.assertEqual(hit.message, "the slide uses 4 sizes; keep at most 3")
        rows = [("A claim", "A baseline", {"baseline_font": {"family": "Courier New"}}),
                ("A claim", "A baseline", {"family": "Georgia"})]
        hit = self.assertFires(self.lint(self.deck(rows)), "fonts-per-slide", "warn", "T2")
        self.assertEqual(hit.message, "the slide uses 3 font families; keep at most 2")

    def test_three_claim_shapes_with_their_marks_are_three_objects(self):
        prs = self.deck([("A claim", "A baseline", {})] * 3)
        slide = prs.slides[0]
        for index in range(3):
            add_rect(slide, "Claim mark", 0.8, 1.92 + index * 1.5, 0.18, 0.18)
        findings = self.lint(prs, minutes=12)
        self.assertQuiet(findings, "objects")
        self.assertEqual(findings, [])

    def test_a_take_home_slide_of_the_helper_reads_as_claims_and_baselines(self):
        prs = sh.new_presentation(LIGHT)
        sh.add_take_home_slide(prs, LIGHT, [(words(16), words(16)), ("A claim of six words here", "A baseline")])
        findings = self.lint(prs, minutes=12)
        self.assertEqual([(f.rule, f.message) for f in findings], [
            ("take-home-words", "a claim has 16 words; keep at most 15"),
            ("take-home-words", "a baseline has 16 words; keep at most 15")])


def add_old_take_home(prs):
    """Add a take-home slide in the structure of earlier decks: a Claim shape and a Baseline shape per claim.

    Seeded defects: a claim of 16 words, a baseline of 16 words, a Claim shape with a second paragraph
    at 20 pt, and a baseline at 19 pt in a pale grey.
    """
    slide = add_slide(prs, TAKE_HOME_TITLE, kind="take_home")
    rows = [(words(16), words(5), {}), (words(6), words(16), {}), (words(6), words(5), {"size": 19, "color": "BBBBBB"})]
    for index, (claim, baseline, font) in enumerate(rows):
        top = 1.8 + index * 1.5
        add_rect(slide, "Claim mark", 0.8, top + 0.12, 0.18, 0.18)
        box = add_box(slide, claim, name="Claim", left=1.25, top=top, width=11.28, height=0.6)
        if index == 2:
            run = box.text_frame.add_paragraph().add_run()
            run.text = "A second paragraph"
            set_run(run, size=20)
        base = {"size": 20, "color": "6B7280"}
        base.update(font)
        add_box(slide, baseline, name="Baseline", left=1.25, top=top + 0.62, width=11.28, height=0.5, **base)
    return slide


class TakeHomeOldStructure(LintCase):
    """A slide with a Baseline shape keeps the rules of earlier decks: a Claim shape is all claim."""

    # The findings of the linter for add_old_take_home before the Claim shape held its baseline.
    EXPECTED = [
        (1, "take-home-words", "warn", "T3", "a claim has 16 words; keep at most 15"),
        (1, "take-home-words", "warn", "T3", "a baseline has 16 words; keep at most 15"),
        (1, "body-font", "warn", "T2", "Claim is 20 pt; use at least 24 pt for body text"),
        (1, "body-font", "warn", "T2", "Baseline is 19 pt; use at least 20 pt for " + SECONDARY_NOUN),
        (1, "fonts-per-slide", "warn", "T2", "the slide uses 4 sizes; keep at most 3"),
        (1, "text-contrast", "error", "T1",
         "Baseline: text color #BBBBBB on #FFFFFF has a contrast of 1.92:1; use at least 3:1"),
    ]

    def test_an_old_structure_deck_lints_as_before(self):
        prs = new_deck()
        add_old_take_home(prs)
        self.assertEqual([tuple(f) for f in self.lint(prs, minutes=12)], self.EXPECTED)

    def test_a_second_paragraph_of_a_claim_stays_claim_text_when_a_baseline_shape_exists(self):
        prs = new_deck()
        slide = add_slide(prs, TAKE_HOME_TITLE, kind="take_home")
        add_claim(slide, words(5), words(16), baseline_font={"size": 20})
        add_box(slide, "A baseline", name="Baseline", left=1.25, top=3.2, width=11.28, height=0.5, size=20)
        findings = self.lint(prs)
        self.assertEqual([h.message for h in self.of(findings, "take-home-words")],
                         ["a claim has 21 words; keep at most 15"])
        self.assertEqual([h.message for h in self.of(findings, "body-font")],
                         ["Claim is 20 pt; use at least 24 pt for body text"])


class Bullets(LintCase):
    def test_five_bullets_warn_and_four_pass(self):
        prs = new_deck()
        add_bullets(add_slide(prs), 5)
        hit = self.assertFires(self.lint(prs), "bullets", "warn", "T2")
        self.assertIn("5 bullets", hit.message)
        prs = new_deck()
        add_bullets(add_slide(prs), 4)
        self.assertQuiet(self.lint(prs), "bullets")

    def test_two_lists_warn_and_one_passes(self):
        prs = new_deck()
        slide = add_slide(prs)
        add_bullets(slide, 3, name="First")
        add_bullets(slide, 3, name="Second", top=4.0)
        hit = self.assertFires(self.lint(prs), "bullets", "warn", "T2")
        self.assertIn("2 bulleted lists", hit.message)
        prs = new_deck()
        add_bullets(add_slide(prs), 3)
        self.assertQuiet(self.lint(prs), "bullets")

    def test_paragraphs_with_no_bullet_mark_are_not_a_list(self):
        prs = new_deck()
        add_bullets(add_slide(prs), 6, bullets=False)
        self.assertQuiet(self.lint(prs), "bullets")

    def test_a_single_bullet_is_not_a_list(self):
        prs = new_deck()
        slide = add_slide(prs)
        add_bullets(slide, 1, name="First")
        add_bullets(slide, 1, name="Second", top=4.0)
        self.assertQuiet(self.lint(prs), "bullets")

    def test_a_body_placeholder_inherits_its_bullets(self):
        prs = new_deck()
        slide = prs.slides.add_slide(prs.slide_layouts.get_by_name("Title and Content"))
        slide.shapes.title.text_frame.text = GOOD_TITLE
        frame = slide.placeholders[1].text_frame
        frame.text = "First point"
        for index in range(5):
            frame.add_paragraph().text = "Point %d" % index
        self.assertFires(self.lint(prs), "bullets", "warn", "T2")

    def test_a_body_placeholder_with_no_bullet_set_by_the_paragraphs_is_quiet(self):
        prs = new_deck()
        slide = prs.slides.add_slide(prs.slide_layouts.get_by_name("Title and Content"))
        slide.shapes.title.text_frame.text = GOOD_TITLE
        frame = slide.placeholders[1].text_frame
        frame.text = "First point"
        for index in range(5):
            frame.add_paragraph().text = "Point %d" % index
        for paragraph in frame.paragraphs:
            paragraph._p.get_or_add_pPr().append(parse_xml("<a:buNone %s/>" % nsdecls("a")))
        self.assertQuiet(self.lint(prs), "bullets")

    def test_the_limit_comes_from_the_style(self):
        prs = new_deck()
        add_bullets(add_slide(prs), 5)
        project = self.tmp / "project"
        project.mkdir()
        (project / "slide-style-minimal.yaml").write_text("lint:\n  bullets_per_list: {max: 6, tier: T2}\n")
        self.assertQuiet(self.lint(prs, project_dir=str(project)), "bullets")


class Objects(LintCase):
    def with_boxes(self, count, **options):
        prs = new_deck()
        slide = add_slide(prs)
        for index in range(count):
            add_box(slide, "Box %d" % index, name="Box %d" % index, top=1.4 + index * 0.7, height=0.5)
        return self.lint(prs, **options)

    def test_seven_objects_warn_and_six_pass(self):
        hit = self.assertFires(self.with_boxes(7), "objects", "warn", "T2")
        self.assertIn("7 objects", hit.message)
        self.assertQuiet(self.with_boxes(6), "objects")

    def test_the_title_and_the_source_line_are_not_objects(self):
        prs = new_deck()
        slide = add_slide(prs)
        for index in range(6):
            add_box(slide, "Box %d" % index, name="Box %d" % index, top=1.4 + index * 0.7, height=0.5)
        add_source(slide)
        self.assertQuiet(self.lint(prs), "objects")

    def test_the_general_profile_allows_four(self):
        self.assertFires(self.with_boxes(5, profile="general"), "objects", "warn", "T2")
        self.assertQuiet(self.with_boxes(4, profile="general"), "objects")

    def test_a_group_counts_as_one_object(self):
        prs = new_deck()
        slide = add_slide(prs)
        boxes = [add_box(slide, "Box %d" % i, name="Box %d" % i, top=1.4 + i * 0.5, height=0.4) for i in range(8)]
        slide.shapes.add_group_shape(boxes)
        self.assertQuiet(self.lint(prs), "objects")

    def test_a_group_of_groups_counts_as_one_object(self):
        prs = new_deck()
        slide = add_slide(prs)
        inner = []
        for index in range(5):
            box = add_box(slide, "Box %d" % index, name="Box %d" % index, top=1.4 + index * 0.6, height=0.5)
            inner.append(slide.shapes.add_group_shape([box]))
        slide.shapes.add_group_shape(inner)
        self.assertQuiet(self.lint(prs, profile="general"), "objects")

    def test_the_parts_of_a_roadmap_count_as_one_object_per_section(self):
        prs = new_deck()
        slide = add_slide(prs, "Four steps take us from design to assay", kind="roadmap")
        for index in range(5):
            left = 0.8 + index * 2.4
            add_rect(slide, "Section mark", left, 2.9, 2.2, 0.03)
            add_box(slide, "Step", name="Section name", left=left, top=3.1, width=2.2, height=0.6, size=32)
            add_box(slide, "Run it", name="Section line", left=left, top=4.0, width=2.2, height=1.0)
        self.assertQuiet(self.lint(prs), "objects")

    def test_the_parts_of_a_take_home_count_as_one_object_per_claim(self):
        prs = new_deck()
        slide = add_slide(prs, "Three things to remember from this talk", kind="take_home")
        for index in range(3):
            add_rect(slide, "Claim mark", 0.8, 1.8 + index * 1.5, 0.18, 0.18)
            add_box(slide, "A claim", name="Claim", left=1.25, top=1.7 + index * 1.5, height=0.5)
            add_box(slide, "A baseline", name="Baseline", left=1.25, top=2.3 + index * 1.5, height=0.5, size=20)
        self.assertQuiet(self.lint(prs), "objects")

    def test_a_connector_is_not_an_object(self):
        prs = new_deck()
        slide = add_slide(prs)
        for index in range(6):
            add_box(slide, "Box %d" % index, name="Box %d" % index, top=1.4 + index * 0.7, height=0.5)
        slide.shapes.add_connector(1, Inches(1), Inches(6.5), Inches(5), Inches(6.5))
        self.assertQuiet(self.lint(prs), "objects")

    def test_the_baselines_of_four_claims_are_not_counted(self):
        prs = new_deck()
        slide = add_slide(prs, "Four things to remember from this talk", kind="take_home")
        for index in range(4):
            add_box(slide, "A claim", name="Claim", left=0.8 + index * 2.9, width=2.8, height=0.5)
            add_box(slide, "A baseline", name="Baseline", left=0.8 + index * 2.9, top=2.3, width=2.8, height=0.5,
                    size=20)
        self.assertQuiet(self.lint(prs), "objects")


class LinesPerParagraph(LintCase):
    def with_text(self, chars, size=24, width=4.0, wrap=True):
        prs = new_deck()
        slide = add_slide(prs)
        box = add_box(slide, "x" * chars, width=width, height=3.0, size=size)
        box.text_frame.word_wrap = wrap
        return self.lint(prs)

    def test_three_estimated_lines_warn_and_two_pass(self):
        # 24 pt over 4 in: one line holds 24 characters at 0.5 em.
        self.assertQuiet(self.with_text(48), "lines-per-paragraph")
        hit = self.assertFires(self.with_text(49), "lines-per-paragraph", "warn", "T2")
        self.assertIn("estimated", hit.message)
        self.assertIn("3 lines", hit.message)

    def test_the_estimate_grows_with_the_size(self):
        self.assertFires(self.with_text(40, size=36), "lines-per-paragraph", "warn", "T2")

    def test_a_box_that_does_not_wrap_is_one_line(self):
        self.assertQuiet(self.with_text(200, wrap=False), "lines-per-paragraph")

    def test_a_title_of_three_estimated_lines_warns(self):
        prs = new_deck()
        add_slide(prs, "Binders fold " + words(28, "again"))
        self.assertFires(self.lint(prs), "lines-per-paragraph", "warn", "T2")

    def test_the_insets_of_the_box_count(self):
        prs = new_deck()
        slide = add_slide(prs)
        box = add_box(slide, "x" * 48, width=4.0, height=3.0)
        box.text_frame.margin_left = box.text_frame.margin_right = Inches(0.5)
        self.assertFires(self.lint(prs), "lines-per-paragraph", "warn", "T2")

    def test_a_superscript_digit_in_an_inset_does_not_crash(self):
        prs = new_deck()
        slide = add_slide(prs)
        box = add_box(slide, "x" * 400, width=4.0, height=3.0)
        box.text_frame._txBody.find(qn("a:bodyPr")).set("lIns", "\u00b2")
        self.assertIsInstance(self.lint(prs), list)

    def test_a_paragraph_with_an_unresolved_size_is_skipped(self):
        prs = new_deck()
        slide = add_slide(prs)
        box = add_box(slide, "x" * 400, width=4.0, height=3.0)
        box.text_frame.paragraphs[0].runs[0]._r.get_or_add_rPr().set("sz", "big")
        self.assertQuiet(self.lint(prs), "lines-per-paragraph")

    def test_a_line_break_starts_a_new_estimated_line(self):
        prs = new_deck()
        slide = add_slide(prs)
        box = add_box(slide, "x" * 13, width=4.0, height=3.0)
        paragraph = box.text_frame.paragraphs[0]
        for _ in range(2):
            paragraph.add_line_break()
            run = paragraph.add_run()
            run.text = "x" * 13
            set_run(run)
        self.assertFires(self.lint(prs), "lines-per-paragraph", "warn", "T2")


class BodyFont(LintCase):
    def sized(self, size, name="Body", kind=None, **options):
        prs = new_deck()
        slide = add_slide(prs, kind=kind)
        add_box(slide, "Twelve binders fold well", name=name, size=size)
        return self.of(self.lint(prs, **options), "body-font")

    def test_primary_text_at_24_pt_passes_and_at_23_9_pt_warns(self):
        self.assertEqual(self.sized(24), [])
        hits = self.sized(23.9)
        self.assertEqual([(h.slide, h.severity, h.tier) for h in hits], [(1, "warn", "T2")])
        self.assertIn("23.9 pt", hits[0].message)

    def test_a_run_at_18_pt_does_not_error_and_one_at_17_9_pt_does(self):
        self.assertEqual([h.severity for h in self.sized(18)], ["warn"])
        hits = self.sized(17.9)
        self.assertEqual([(h.severity, h.tier) for h in hits], [("error", "T2")])

    def test_a_callout_at_20_pt_passes_and_at_19_9_pt_warns(self):
        self.assertEqual(self.sized(20, name="Callout", kind="assertion_evidence"), [])
        hits = self.sized(19.9, name="Callout", kind="assertion_evidence")
        self.assertEqual([(h.severity, h.tier) for h in hits], [("warn", "T2")])
        self.assertEqual([h.severity for h in self.sized(17.9, name="Callout", kind="assertion_evidence")], ["error"])

    def test_a_primary_run_at_20_pt_warns(self):
        self.assertEqual([h.severity for h in self.sized(20)], ["warn"])

    def test_a_callout_on_a_slide_with_no_kind_is_primary(self):
        self.assertEqual([h.severity for h in self.sized(20, name="Callout")], ["warn"])

    def test_a_subtitle_is_secondary(self):
        self.assertEqual(self.sized(20, name="Subtitle", kind="title"), [])
        self.assertEqual([h.severity for h in self.sized(19.9, name="Subtitle", kind="title")], ["warn"])

    def test_a_take_home_baseline_is_secondary_and_another_baseline_is_primary(self):
        self.assertEqual(self.sized(20, name="Baseline", kind="take_home"), [])
        self.assertEqual([h.severity for h in self.sized(19.9, name="Baseline", kind="take_home")], ["warn"])
        self.assertEqual([h.severity for h in self.sized(20, name="Baseline", kind="big_number")], ["warn"])

    def test_table_cells_are_secondary(self):
        def table(size):
            prs = new_deck()
            add_table(add_slide(prs), size=size)
            return self.of(self.lint(prs), "body-font")
        self.assertEqual(table(20), [])
        self.assertEqual([h.severity for h in table(19.9)], ["warn"])
        self.assertEqual([h.severity for h in table(17.9)], ["error"])

    def test_chart_text_is_secondary(self):
        def chart(size):
            prs = new_deck()
            add_chart(add_slide(prs), size=size)
            return self.of(self.lint(prs), "body-font")
        self.assertEqual(chart(20), [])
        hits = chart(19.9)
        self.assertEqual([h.severity for h in hits], ["warn"])
        self.assertIn("chart", hits[0].message)
        self.assertEqual([h.severity for h in chart(17.9)], ["error"])

    def test_the_big_number_is_exempt_on_a_big_number_slide(self):
        self.assertEqual(self.sized(10, name="Number", kind="big_number"), [])
        self.assertEqual([h.severity for h in self.sized(10, name="Number")], ["error"])

    def test_the_source_line_needs_12_pt_and_never_the_body_floor(self):
        self.assertEqual(self.sized(12, name="Source"), [])
        hits = self.sized(11.9, name="Source")
        self.assertEqual([(h.severity, h.tier) for h in hits], [("warn", "T3")])
        self.assertIn("source line", hits[0].message)
        self.assertEqual([h.severity for h in self.sized(8, name="Source")], ["warn"])

    def test_an_unnamed_text_box_at_the_bottom_is_checked_as_the_source_line(self):
        prs = new_deck()
        slide = add_slide(prs)
        add_source(slide, size=11, name="Footer note")
        hits = self.of(self.lint(prs), "body-font")
        self.assertEqual([h.tier for h in hits], ["T3"])

    def test_the_title_is_not_a_body_run(self):
        prs = new_deck()
        add_slide(prs, size=20)
        self.assertQuiet(self.lint(prs), "body-font")

    def test_the_profiles_change_the_floors(self):
        self.assertEqual(self.sized(18, profile="lab-meeting"), [])
        self.assertEqual([h.severity for h in self.sized(17.9, profile="lab-meeting")], ["error"])
        self.assertEqual([h.severity for h in self.sized(27, profile="general")], ["warn"])
        self.assertEqual(self.sized(28, profile="general"), [])
        self.assertEqual([h.severity for h in self.sized(23, name="Callout", kind="assertion_evidence",
                                                         profile="general")], ["warn"])

    def test_a_size_that_cannot_be_resolved_is_skipped_with_one_note_for_the_slide(self):
        prs = new_deck()
        slide = add_slide(prs)
        for index in range(2):
            box = add_box(slide, "Twelve binders fold well", name="Box %d" % index, top=1.6 + index * 1.5)
            box.text_frame.paragraphs[0].runs[0]._r.get_or_add_rPr().set("sz", "big")
        hits = self.of(self.lint(prs), "body-font")
        self.assertEqual(len(hits), 1)
        self.assertEqual((hits[0].slide, hits[0].severity, hits[0].tier), (1, "warn", "T2"))
        self.assertIn("size unresolved", hits[0].message)
        self.assertIn("rule skipped", hits[0].message)

    def test_a_size_from_paragraph_defaults_is_skipped_and_the_note_says_why(self):
        prs = new_deck()
        slide = add_slide(prs)
        box = slide.shapes.add_textbox(Inches(0.8), Inches(1.6), Inches(8), Inches(1))
        box.name = "Body"
        paragraph = box.text_frame.paragraphs[0]
        paragraph.font.size = Pt(10)
        run = paragraph.add_run()
        run.text = "Twelve binders fold well"
        run.font.color.rgb = RGBColor.from_string("111111")
        hits = self.of(self.lint(prs), "body-font")
        self.assertEqual(len(hits), 1)
        self.assertEqual((hits[0].severity, hits[0].tier), ("warn", "T2"))
        self.assertIn("paragraph defaults", hits[0].message)
        self.assertIn("PowerPoint may ignore", hits[0].message)
        self.assertIn("rule skipped", hits[0].message)

    def test_a_run_with_no_text_is_not_checked(self):
        prs = new_deck()
        slide = add_slide(prs)
        box = add_box(slide, " ", size=10)
        self.assertEqual(box.text_frame.paragraphs[0].runs[0].text, " ")
        self.assertQuiet(self.lint(prs), "body-font")

    def test_a_chart_legend_counts_as_chart_text(self):
        prs = new_deck()
        frame = add_chart(add_slide(prs), colors=("6D4BD6", "0072B2"), labels=True)
        frame.chart.has_legend = True
        frame.chart.legend.font.size = Pt(10)
        hits = self.of(self.lint(prs), "body-font")
        self.assertEqual([h.severity for h in hits], ["error"])

    def test_a_hidden_axis_does_not_count_as_chart_text(self):
        def axis_text(visible):
            prs = new_deck()
            frame = add_chart(add_slide(prs), labels=True)
            frame.chart.value_axis.visible = visible
            frame.chart.value_axis.tick_labels.font.size = Pt(8)
            return self.of(self.lint(prs), "body-font")
        self.assertEqual(axis_text(False), [])
        self.assertEqual([h.severity for h in axis_text(True)], ["error"])

    def test_a_chart_from_python_pptx_with_no_font_set_runs(self):
        prs = new_deck()
        slide = add_slide(prs)
        data = CategoryChartData()
        data.categories = ["a", "b"]
        data.add_series("s", (1, 2))
        slide.shapes.add_chart(XL_CHART_TYPE.COLUMN_CLUSTERED, Inches(1), Inches(2), Inches(6), Inches(4), data)
        findings = self.lint(prs)
        # python-pptx writes a default text size of 18 pt and no colors: the size is read, the colors are skipped.
        self.assertEqual([f.severity for f in self.of(findings, "body-font")], ["warn"])
        self.assertIn("18 pt", self.only(findings, "body-font").message)
        self.assertIn("chart text color unresolved", self.only(findings, "text-contrast").message)
        self.assertIn("series color comes from the theme", self.only(findings, "mark-contrast").message)

    def test_the_unresolved_size_of_a_big_number_gets_no_note(self):
        prs = new_deck()
        slide = add_slide(prs, kind="big_number")
        box = add_box(slide, "33%", name="Number", size=96)
        box.text_frame.paragraphs[0].runs[0]._r.get_or_add_rPr().set("sz", "big")
        self.assertQuiet(self.lint(prs), "body-font")

    def test_a_slide_number_placeholder_is_not_body_text(self):
        prs = new_deck()
        slide = add_slide(prs)
        add_box(slide, BODY)
        for placeholder in slide.slide_layout.placeholders:
            if "SLIDE_NUMBER" in str(placeholder.placeholder_format.type):
                slide.shapes.clone_placeholder(placeholder)
        numbers = [shape for shape in slide.shapes if shape.is_placeholder and shape.placeholder_format.idx != 0
                   and "SLIDE_NUMBER" in str(shape.placeholder_format.type)]
        self.assertEqual(len(numbers), 1)
        write(numbers[0], "7", size=10, color="111111")
        findings = self.lint(prs)
        self.assertQuiet(findings, "body-font")
        self.assertQuiet(findings, "objects")
        self.assertQuiet(findings, "fonts-per-slide")

    def test_a_superscript_digit_in_a_chart_font_size_does_not_crash(self):
        prs = new_deck()
        frame = add_chart(add_slide(prs), labels=True)
        for props in frame.chart._chartSpace.iter(qn("a:defRPr")):
            props.set("sz", "\u00b2")
        hits = self.of(self.lint(prs), "body-font")
        self.assertEqual(len(hits), 1)
        self.assertIn("size unresolved", hits[0].message)

    def test_a_chart_title_counts_as_chart_text(self):
        prs = new_deck()
        frame = add_chart(add_slide(prs), labels=True)
        frame.chart.has_title = True
        frame.chart.chart_title.text_frame.text = "Hits"
        frame.chart.chart_title.text_frame.paragraphs[0].runs[0].font.size = Pt(10)
        self.assertEqual([h.severity for h in self.of(self.lint(prs), "body-font")], ["error"])

    def test_data_labels_that_show_nothing_do_not_count_as_chart_text(self):
        def labels(show):
            prs = new_deck()
            frame = add_chart(add_slide(prs))
            plot = frame.chart.plots[0]
            plot.has_data_labels = True
            plot.data_labels.font.size = Pt(8)
            plot.data_labels.show_value = show
            return self.of(self.lint(prs), "body-font")
        self.assertEqual(labels(False), [])
        self.assertEqual([h.severity for h in labels(True)], ["error"])

    def test_an_unknown_kind_counts_as_no_kind(self):
        prs = new_deck()
        slide = add_slide(prs, kind="bogus")
        add_box(slide, "Twelve binders fold", name="Callout", left=9.2, width=3.3, size=20)
        self.assertEqual([h.severity for h in self.of(self.lint(prs), "body-font")], ["warn"])


class TitleFont(LintCase):
    def titled(self, size, kind=None):
        prs = new_deck()
        add_slide(prs, kind=kind, size=size)
        return self.of(self.lint(prs), "title-font")

    def test_28_pt_passes_and_27_9_pt_warns(self):
        self.assertEqual(self.titled(28), [])
        hits = self.titled(27.9)
        self.assertEqual([(h.slide, h.severity, h.tier) for h in hits], [(1, "warn", "T2")])
        self.assertIn("27.9 pt", hits[0].message)

    def test_the_title_slide_is_exempt(self):
        self.assertEqual(self.titled(20, kind="title"), [])

    def test_a_title_size_that_cannot_be_resolved_is_skipped_with_a_note(self):
        prs = new_deck()
        slide = add_slide(prs)
        slide.shapes.title.text_frame.paragraphs[0].runs[0]._r.get_or_add_rPr().set("sz", "big")
        hits = self.of(self.lint(prs), "title-font")
        self.assertEqual(len(hits), 1)
        self.assertIn("rule skipped", hits[0].message)


class FontsPerSlide(LintCase):
    def with_sizes(self, sizes, families=("Avenir Light",)):
        prs = new_deck()
        slide = add_slide(prs)
        for index, size in enumerate(sizes):
            add_box(slide, "Twelve binders fold well", name="Box %d" % index, top=1.5 + index * 0.9, height=0.8,
                    size=size, family=families[index % len(families)])
        return self.lint(prs)

    def test_four_sizes_warn_and_three_pass(self):
        self.assertQuiet(self.with_sizes([24, 20]), "fonts-per-slide")
        hit = self.assertFires(self.with_sizes([24, 20, 16]), "fonts-per-slide", "warn", "T2")
        self.assertIn("4 sizes", hit.message)

    def test_three_families_warn_and_two_pass(self):
        self.assertQuiet(self.with_sizes([24, 24], ("Avenir Light", "PT Mono")), "fonts-per-slide")
        hit = self.assertFires(self.with_sizes([24, 24, 24], ("Avenir Light", "PT Mono", "Arial")),
                               "fonts-per-slide", "warn", "T2")
        self.assertIn("3 font families", hit.message)

    def test_the_source_line_and_chart_text_do_not_count(self):
        prs = new_deck()
        slide = add_slide(prs, kind="assertion_evidence")
        add_box(slide, BODY, size=24)
        add_box(slide, "Twelve binders fold", name="Callout", left=9.2, width=3.3, size=20)
        add_source(slide)
        add_chart(slide, size=18)
        self.assertQuiet(self.lint(prs), "fonts-per-slide")
