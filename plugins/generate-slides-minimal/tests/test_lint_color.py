"""Tests for the color rules of scripts/design_lint.py: text contrast, mark contrast and red and green.
"""

import sys
from pathlib import Path
from unittest import mock

from pptx.dml.color import RGBColor
from pptx.enum.dml import MSO_THEME_COLOR
from pptx.oxml import parse_xml
from pptx.oxml.ns import nsdecls, qn
from pptx.util import Inches

SCRIPTS_DIR = Path(__file__).resolve().parents[1] / "scripts"
for folder in (SCRIPTS_DIR, Path(__file__).resolve().parent):
    if str(folder) not in sys.path:
        sys.path.insert(0, str(folder))

import design_color  # noqa: E402
from lint_helpers import (  # noqa: E402
    BODY,
    LintCase,
    add_band,
    add_box,
    add_chart,
    add_picture,
    add_rect,
    add_slide,
    add_table,
    clean_slide,
    make_png,
    new_deck,
    write,
)


class TextContrast(LintCase):
    def colored(self, color, size=24, background="FFFFFF", bold=False, **options):
        prs = new_deck()
        slide = add_slide(prs, background=background, color="111111" if background == "FFFFFF" else "F3F0FA")
        add_box(slide, "Twelve binders fold well", color=color, size=size, bold=bold)
        return self.of(self.lint(prs, **options), "text-contrast")

    def test_dark_text_on_white_is_quiet(self):
        self.assertEqual(self.colored("111111"), [])

    def test_a_ratio_below_3_is_an_error(self):
        hits = self.colored("CCCCCC")
        self.assertEqual([(h.slide, h.severity, h.tier) for h in hits], [(1, "error", "T1")])
        self.assertIn("Body", hits[0].message)
        self.assertIn("1.6", hits[0].message)

    def test_small_text_below_4_5_warns(self):
        hits = self.colored("888888", size=12)
        self.assertEqual([(h.severity, h.tier) for h in hits], [("warn", "T1")])

    def test_large_text_passes_at_3(self):
        self.assertEqual(self.colored("888888", size=24), [])
        self.assertEqual(self.colored("888888", size=18), [])
        self.assertEqual(len(self.colored("888888", size=17.9)), 1)

    def test_bold_text_of_14_pt_counts_as_large(self):
        self.assertEqual(self.colored("888888", size=14, bold=True), [])
        self.assertEqual(len(self.colored("888888", size=14)), 1)
        self.assertEqual(len(self.colored("888888", size=13.9, bold=True)), 1)

    def test_the_real_pair_around_4_5(self):
        self.assertEqual(self.colored("767676", size=12), [])  # 4.54
        self.assertEqual(len(self.colored("777777", size=12)), 1)  # 4.48

    def contrast_at(self, ratio):
        with mock.patch.object(design_color, "contrast_ratio", return_value=ratio):
            return [h for h in self.colored("111111", size=12) if "Body" in h.message]

    def test_exactly_4_5_passes_and_4_499_warns(self):
        self.assertEqual(self.contrast_at(4.5), [])
        self.assertEqual([h.severity for h in self.contrast_at(4.499)], ["warn"])

    def test_exactly_3_warns_and_2_999_is_an_error(self):
        self.assertEqual([h.severity for h in self.contrast_at(3.0)], ["warn"])
        self.assertEqual([h.severity for h in self.contrast_at(2.999)], ["error"])

    def test_the_slide_background_counts(self):
        self.assertEqual([h.severity for h in self.colored("111111", background="14111F")], ["error"])
        self.assertEqual(self.colored("F3F0FA", background="14111F"), [])

    def test_a_solid_fill_of_the_shape_counts_instead(self):
        def on_black(color):
            prs = new_deck()
            slide = add_slide(prs)
            box = add_box(slide, "Twelve binders fold well", color=color)
            box.fill.solid()
            box.fill.fore_color.rgb = RGBColor.from_string("000000")
            return self.of(self.lint(prs), "text-contrast")
        self.assertEqual(on_black("FFFFFF"), [])
        self.assertEqual([h.severity for h in on_black("111111")], ["error"])

    def test_a_gradient_fill_is_skipped_with_one_note_for_the_slide(self):
        prs = new_deck()
        slide = add_slide(prs)
        for index in range(2):
            box = add_box(slide, "Twelve binders fold well", name="Box %d" % index, top=1.6 + index * 1.5)
            box.fill.gradient()
        hits = self.of(self.lint(prs), "text-contrast")
        self.assertEqual(len(hits), 1)
        self.assertEqual((hits[0].slide, hits[0].severity, hits[0].tier), (1, "warn", "T1"))
        self.assertIn("fill unknown", hits[0].message)
        self.assertIn("rule skipped", hits[0].message)

    def test_text_over_a_picture_is_skipped_with_a_note(self):
        prs = new_deck()
        slide = add_slide(prs)
        add_picture(slide, self.png("photo.png", 2000, 1200), left=1.0, top=1.6, width=5.0, height=3.0, alt="A photo")
        add_box(slide, "Twelve binders fold well", color="FFFFFF", left=1.2, top=1.8, width=4.0, height=1.0)
        hits = self.of(self.lint(prs), "text-contrast")
        self.assertEqual(len(hits), 1)
        self.assertEqual((hits[0].severity, hits[0].tier), ("warn", "T1"))
        self.assertIn("picture", hits[0].message)
        self.assertIn("rule skipped", hits[0].message)

    def test_text_beside_a_picture_is_checked(self):
        prs = new_deck()
        slide = add_slide(prs)
        add_picture(slide, self.png("photo.png", 2000, 1200), left=7.0, top=1.6, width=5.0, height=3.0, alt="A photo")
        add_box(slide, "Twelve binders fold well", color="CCCCCC", left=0.8, top=1.8, width=4.0, height=1.0)
        self.assertEqual([h.severity for h in self.of(self.lint(prs), "text-contrast")], ["error"])

    def test_a_title_placeholder_with_no_fill_looks_through_to_the_background(self):
        prs = new_deck()
        clean_slide(prs)
        self.assertEqual(self.of(self.lint(prs), "text-contrast"), [])

    def test_a_fill_in_the_layout_placeholder_counts(self):
        prs = new_deck()
        clean_slide(prs)
        layout = list(prs.slides)[0].slide_layout
        placeholder = [p for p in layout.placeholders if p.placeholder_format.type is not None
                       and "TITLE" in str(p.placeholder_format.type)][0]
        placeholder.fill.solid()
        placeholder.fill.fore_color.rgb = RGBColor.from_string("000000")
        hits = self.of(self.lint(prs), "text-contrast")
        self.assertEqual([h.severity for h in hits], ["error"])
        self.assertIn("Title", hits[0].message)

    def test_table_cells_count_against_their_own_fill(self):
        def cell_fill(color, fill):
            prs = new_deck()
            slide = add_slide(prs)
            frame = add_table(slide, rows=1, cols=1, color=color)
            cell = frame.table.cell(0, 0)
            if fill is None:
                cell._tc.tcPr.remove(cell._tc.tcPr[0])
            else:
                cell.fill.solid()
                cell.fill.fore_color.rgb = RGBColor.from_string(fill)
            return self.of(self.lint(prs), "text-contrast")
        self.assertEqual(cell_fill("FFFFFF", "000000"), [])
        self.assertEqual([h.severity for h in cell_fill("111111", "000000")], ["error"])
        hits = cell_fill("111111", None)
        self.assertEqual(len(hits), 1)
        self.assertIn("fill unknown", hits[0].message)

    def test_a_color_that_cannot_be_resolved_is_skipped_with_a_note(self):
        prs = new_deck()
        slide = add_slide(prs)
        box = add_box(slide, "Twelve binders fold well")
        fill = box.text_frame.paragraphs[0].runs[0]._r.rPr.find(qn("a:solidFill"))
        for child in list(fill):
            fill.remove(child)
        fill.append(parse_xml('<a:schemeClr %s val="accent1"><a:lumMod val="50000"/></a:schemeClr>' % nsdecls("a")))
        hits = self.of(self.lint(prs), "text-contrast")
        self.assertEqual(len(hits), 1)
        self.assertIn("color unresolved", hits[0].message)
        self.assertIn("rule skipped", hits[0].message)

    def test_a_color_from_paragraph_defaults_is_skipped_and_the_note_says_why(self):
        prs = new_deck()
        slide = add_slide(prs)
        box = slide.shapes.add_textbox(Inches(0.8), Inches(1.6), Inches(8), Inches(1))
        box.name = "Body"
        paragraph = box.text_frame.paragraphs[0]
        paragraph.font.color.rgb = RGBColor.from_string("CCCCCC")
        run = paragraph.add_run()
        run.text = "Twelve binders fold well"
        run._r.get_or_add_rPr().set("sz", "2400")
        hits = self.of(self.lint(prs), "text-contrast")
        self.assertEqual(len(hits), 1)
        self.assertEqual(hits[0].severity, "warn")
        self.assertIn("paragraph defaults", hits[0].message)

    def test_a_size_that_cannot_be_resolved_still_gets_the_error_check_only(self):
        prs = new_deck()
        slide = add_slide(prs)
        box = add_box(slide, "Twelve binders fold well", color="888888")
        box.text_frame.paragraphs[0].runs[0]._r.get_or_add_rPr().set("sz", "big")
        self.assertQuiet(self.lint(prs), "text-contrast")
        box.text_frame.paragraphs[0].runs[0].font.color.rgb = RGBColor.from_string("CCCCCC")
        self.assertEqual([h.severity for h in self.of(self.lint(prs), "text-contrast")], ["error"])

    def test_a_table_cell_with_no_color_of_its_own_is_skipped_with_a_note(self):
        prs = new_deck()
        frame = add_table(add_slide(prs), rows=1, cols=1)
        run = frame.table.cell(0, 0).text_frame.paragraphs[0].runs[0]
        run._r.rPr.remove(run._r.rPr.find(qn("a:solidFill")))
        hits = self.of(self.lint(prs), "text-contrast")
        self.assertEqual(len(hits), 1)
        self.assertIn("color unresolved", hits[0].message)

    def test_a_gradient_slide_background_is_skipped_with_one_note(self):
        prs = new_deck()
        slide = add_slide(prs, background=None)
        slide.background.fill.gradient()
        add_box(slide, BODY)
        hits = self.of(self.lint(prs), "text-contrast")
        self.assertEqual(len(hits), 1)
        self.assertIn("slide background unresolved", hits[0].message)
        self.assertIn("rule skipped", hits[0].message)

    def title_placeholder_in_master(self, master_fill):
        prs = new_deck()
        clean_slide(prs)
        master = list(prs.slides)[0].slide_layout.slide_master
        title = [p for p in master.placeholders if "TITLE" in str(p.placeholder_format.type)][0]
        title.fill.solid()
        title.fill.fore_color.rgb = RGBColor.from_string(master_fill)
        return self.of(self.lint(prs), "text-contrast")

    def test_a_fill_in_the_master_placeholder_counts(self):
        self.assertEqual([h.severity for h in self.title_placeholder_in_master("000000")], ["error"])
        self.assertEqual(self.title_placeholder_in_master("FFFFFF"), [])

    def test_a_placeholder_with_a_style_fill_is_unknown(self):
        prs = new_deck()
        slide = clean_slide(prs)
        title = slide.shapes.title._element
        title.find(qn("p:spPr")).addnext(parse_xml(
            '<p:style %s><a:lnRef idx="0"><a:scrgbClr r="0" g="0" b="0"/></a:lnRef>'
            '<a:fillRef idx="1"><a:scrgbClr r="0" g="0" b="0"/></a:fillRef>'
            '<a:effectRef idx="0"><a:scrgbClr r="0" g="0" b="0"/></a:effectRef>'
            '<a:fontRef idx="minor"/></p:style>' % nsdecls("p", "a")))
        hits = self.of(self.lint(prs), "text-contrast")
        self.assertEqual(len(hits), 1)
        self.assertIn("fill unknown", hits[0].message)


class TextBackground(LintCase):
    """The color behind text comes from the shapes drawn below it, nearest first, then from the slide."""

    def slide_with_title(self, title_color="FFFFFF"):
        prs = new_deck()
        slide = add_slide(prs, color=title_color)
        add_box(slide, BODY, top=2.0)
        return prs, slide

    def contrast(self, prs):
        return self.of(self.lint(prs), "text-contrast")

    def test_a_white_title_over_a_dark_band_of_the_layout_is_quiet(self):
        prs, slide = self.slide_with_title()
        add_band(slide, "layout")
        self.assertEqual(self.contrast(prs), [])

    def test_a_white_title_over_a_dark_band_of_the_master_is_quiet(self):
        prs, slide = self.slide_with_title()
        add_band(slide, "master")
        self.assertEqual(self.contrast(prs), [])

    def test_a_band_with_a_theme_color_resolves_through_the_theme(self):
        prs, slide = self.slide_with_title()
        add_band(slide, "layout", theme=MSO_THEME_COLOR.TEXT_2)
        self.assertEqual(self.contrast(prs), [])

    def test_a_white_title_over_a_band_in_a_layout_group_is_quiet(self):
        prs, slide = self.slide_with_title()
        add_band(slide, "layout", groups=1)
        self.assertEqual(self.contrast(prs), [])

    def test_a_white_title_over_a_band_in_a_master_group_is_quiet(self):
        prs, slide = self.slide_with_title()
        add_band(slide, "master", groups=1)
        self.assertEqual(self.contrast(prs), [])

    def test_a_white_title_over_a_band_in_nested_groups_is_quiet(self):
        prs, slide = self.slide_with_title()
        add_band(slide, "layout", groups=3)
        self.assertEqual(self.contrast(prs), [])

    def test_the_transform_of_a_layout_group_moves_the_band(self):
        prs, slide = self.slide_with_title()
        group = add_band(slide, "layout", groups=1)
        group.top = group.top + Inches(3)  # the band no longer lies behind the title
        hits = self.contrast(prs)
        self.assertEqual([h.severity for h in hits], ["error"])
        self.assertIn("#FFFFFF on #FFFFFF", hits[0].message)

    def test_a_dark_title_over_a_band_in_a_layout_group_is_an_error(self):
        prs, slide = self.slide_with_title(title_color="111111")
        add_band(slide, "layout", groups=1)
        self.assertEqual([h.severity for h in self.contrast(prs)], ["error"])

    def test_a_band_with_a_gradient_in_a_layout_group_is_skipped_with_a_note(self):
        prs, slide = self.slide_with_title()
        add_band(slide, "layout", groups=2, gradient=True)
        hits = self.contrast(prs)
        self.assertEqual(len(hits), 1)
        self.assertEqual((hits[0].severity, hits[0].tier), ("warn", "T1"))
        self.assertIn("text over a filled shape", hits[0].message)
        self.assertIn("Title", hits[0].message)

    def test_a_group_in_a_layout_that_the_slide_hides_does_not_count(self):
        prs, slide = self.slide_with_title()
        add_band(slide, "layout", groups=1)
        slide._element.set("showMasterSp", "0")
        self.assertEqual([h.severity for h in self.contrast(prs)], ["error"])

    def test_the_nearest_of_two_layout_bands_decides(self):
        for first, second, color in (("0B1F3A", "F2F2F2", "111111"), ("F2F2F2", "0B1F3A", "FFFFFF")):
            with self.subTest(first=first, second=second):
                prs, slide = self.slide_with_title(title_color=color)
                add_band(slide, "layout", rgb=first, groups=1)
                add_band(slide, "layout", rgb=second)
                self.assertEqual(self.contrast(prs), [])

    def test_a_layout_band_is_nearer_than_a_master_band(self):
        prs, slide = self.slide_with_title(title_color="111111")
        add_band(slide, "master", rgb="0B1F3A")
        add_band(slide, "layout", rgb="F2F2F2")
        self.assertEqual(self.contrast(prs), [])

    def break_group(self, group):
        """Give a group a zero child extent, so that python-pptx cannot map its children to the slide."""
        group._element.find(qn("p:grpSpPr")).find(qn("a:xfrm")).find(qn("a:chExt")).set("cx", "0")

    def test_a_layout_group_with_an_unusable_box_and_a_dark_band_gives_a_note_and_no_false_error(self):
        prs, slide = self.slide_with_title()
        group = add_band(slide, "layout", groups=1)
        self.break_group(group)
        hits = self.contrast(prs)
        self.assertEqual(len(hits), 1)
        self.assertEqual((hits[0].severity, hits[0].tier), ("warn", "T1"))
        self.assertIn("text over a filled shape", hits[0].message)

    def test_a_layout_group_with_an_unusable_box_and_no_fill_covers_nothing(self):
        prs, slide = self.slide_with_title(title_color="111111")
        group = add_band(slide, "layout", groups=1)
        group.shapes[0].fill.background()
        self.break_group(group)
        self.assertEqual(self.contrast(prs), [])

    def test_a_slide_group_with_an_unusable_box_below_the_text_gives_a_note(self):
        prs, slide = self.slide_with_title()
        rect = add_rect(slide, "Dark", 0.0, 0.0, 13.333, 1.4, "0B1F3A")
        group = slide.shapes.add_group_shape([rect])
        self.break_group(group)
        slide.shapes.title._element.addprevious(group._element)
        hits = self.contrast(prs)
        self.assertEqual(len(hits), 1)
        self.assertIn("text over a filled shape", hits[0].message)

    def test_a_picture_with_an_unusable_box_below_the_text_gives_a_note(self):
        prs, slide = self.slide_with_title()
        picture = add_picture(slide, make_png(self.tmp, "p.png", 2000, 1200), left=0.0, top=0.0, width=13.0, height=1.4,
                              alt="A photo")
        picture._element.find(qn("p:spPr")).find(qn("a:xfrm")).find(qn("a:ext")).set("cx", "notanumber")
        slide.shapes.title._element.addprevious(picture._element)
        hits = self.contrast(prs)
        self.assertEqual(len(hits), 1)
        self.assertIn("text over a filled shape", hits[0].message)
        self.assertIn("rule skipped", hits[0].message)

    def test_a_dark_title_over_a_dark_band_is_an_error(self):
        prs, slide = self.slide_with_title(title_color="111111")
        add_band(slide, "layout")
        hits = self.contrast(prs)
        self.assertEqual([(h.severity, h.tier) for h in hits], [("error", "T1")])
        self.assertIn("Title", hits[0].message)

    def test_a_band_that_the_slide_hides_does_not_count(self):
        prs, slide = self.slide_with_title()
        add_band(slide, "layout")
        slide._element.set("showMasterSp", "0")
        hits = self.contrast(prs)
        self.assertEqual([h.severity for h in hits], ["error"])
        self.assertIn("#FFFFFF on #FFFFFF", hits[0].message)

    def test_a_master_band_that_the_layout_hides_does_not_count(self):
        prs, slide = self.slide_with_title()
        add_band(slide, "master")
        slide.slide_layout._element.set("showMasterSp", "0")
        self.assertEqual([h.severity for h in self.contrast(prs)], ["error"])

    def test_a_layout_band_still_counts_when_the_layout_hides_the_master(self):
        prs, slide = self.slide_with_title()
        add_band(slide, "layout")
        slide.slide_layout._element.set("showMasterSp", "0")
        self.assertEqual(self.contrast(prs), [])

    def test_a_band_that_does_not_hold_the_text_is_not_a_background(self):
        prs, slide = self.slide_with_title()
        add_band(slide, "layout", height=0.8)
        self.assertEqual([h.severity for h in self.contrast(prs)], ["error"])

    def card_deck(self, text_color, card="0B1F3A", top=1.8, height=0.8, card_height=1.2):
        prs = new_deck()
        slide = add_slide(prs)
        add_rect(slide, "Card", 0.8, 1.6, 8.0, card_height, card)
        add_box(slide, "Twelve binders fold well", name="On the card", color=text_color, left=1.0, top=top,
                width=7.5, height=height)
        return prs, slide

    def test_white_text_over_a_dark_rectangle_is_quiet(self):
        prs, _ = self.card_deck("FFFFFF")
        self.assertEqual(self.contrast(prs), [])

    def test_dark_text_over_a_dark_rectangle_is_an_error(self):
        prs, _ = self.card_deck("111111")
        hits = self.contrast(prs)
        self.assertEqual([h.severity for h in hits], ["error"])
        self.assertIn("On the card", hits[0].message)
        self.assertIn("#0B1F3A", hits[0].message)

    def test_dark_text_over_a_light_card_is_quiet(self):
        prs, _ = self.card_deck("111111", card="F2F2F2")
        self.assertEqual(self.contrast(prs), [])

    def test_text_that_sticks_out_of_the_card_is_checked_against_the_slide(self):
        prs, _ = self.card_deck("FFFFFF", top=2.0, height=1.5)
        self.assertEqual([h.severity for h in self.contrast(prs)], ["error"])

    def test_text_over_a_gradient_card_is_skipped_with_one_note(self):
        prs, slide = self.card_deck("FFFFFF")
        card = [s for s in slide.shapes if s.name == "Card"][0]
        card.fill.gradient()
        add_box(slide, "Another line on the card", name="Second", color="FFFFFF", left=1.0, top=2.3, width=7.0,
                height=0.2)
        hits = self.contrast(prs)
        self.assertEqual(len(hits), 1)
        self.assertEqual((hits[0].severity, hits[0].tier), ("warn", "T1"))
        self.assertIn("text over a filled shape", hits[0].message)
        self.assertIn("rule skipped", hits[0].message)

    def test_the_nearest_shape_below_decides(self):
        prs, slide = self.slide_with_title(title_color="111111")
        add_band(slide, "layout")
        add_rect(slide, "Light card", 0.8, 0.3, 11.7, 0.95, "FFFFFF")
        title = slide.shapes.title._element
        slide.shapes._spTree.remove(title)
        slide.shapes._spTree.append(title)
        self.assertEqual(self.contrast(prs), [])

    def test_a_rectangle_with_no_fill_does_not_hide_the_band(self):
        prs, slide = self.slide_with_title()
        add_band(slide, "layout")
        frame = add_rect(slide, "Outline", 0.8, 0.3, 11.7, 0.95, "FFFFFF")
        frame.fill.background()
        title = slide.shapes.title._element
        slide.shapes._spTree.remove(title)
        slide.shapes._spTree.append(title)
        self.assertEqual(self.contrast(prs), [])

    def test_a_layout_placeholder_that_the_slide_does_not_use_is_not_drawn(self):
        prs = new_deck()
        slide = prs.slides.add_slide(prs.slide_layouts.get_by_name("Title and Content"))
        slide.background.fill.solid()
        slide.background.fill.fore_color.rgb = RGBColor.from_string("FFFFFF")
        write(slide.shapes.title, "Binder design succeeds for 12 of 15 targets", size=28, color="111111")
        body = slide.placeholders[1]
        body._element.getparent().remove(body._element)
        layout_body = [p for p in slide.slide_layout.placeholders if p.placeholder_format.idx == 1][0]
        layout_body.fill.solid()
        layout_body.fill.fore_color.rgb = RGBColor.from_string("0B1F3A")
        add_box(slide, "Twelve binders fold well", name="In the empty area", color="111111", left=0.7, top=1.9,
                width=6.0, height=1.0)
        self.assertEqual(self.contrast(prs), [])

    def test_the_nearest_of_two_cards_decides(self):
        for first, second, color in (("0B1F3A", "F2F2F2", "111111"), ("F2F2F2", "0B1F3A", "FFFFFF")):
            with self.subTest(first=first, second=second):
                prs = new_deck()
                slide = add_slide(prs)
                add_rect(slide, "Far card", 0.6, 1.5, 9.0, 1.6, first)
                add_rect(slide, "Near card", 0.8, 1.6, 8.0, 1.2, second)
                add_box(slide, "Twelve binders fold well", name="On the cards", color=color, left=1.0, top=1.8,
                        width=7.5, height=0.8)
                self.assertEqual(self.contrast(prs), [])

    def test_a_slide_that_hides_background_graphics_by_the_word_false_does_too(self):
        prs, slide = self.slide_with_title()
        add_band(slide, "layout")
        slide._element.set("showMasterSp", "false")
        self.assertEqual([h.severity for h in self.contrast(prs)], ["error"])

    def test_a_text_box_with_its_own_solid_fill_ignores_the_slide(self):
        prs = new_deck()
        slide = add_slide(prs)
        box = add_box(slide, "Twelve binders fold well", color="FFFFFF")
        box.fill.solid()
        box.fill.fore_color.rgb = RGBColor.from_string("000000")
        self.assertEqual(self.contrast(prs), [])

    def test_a_rectangle_drawn_above_the_text_is_not_a_background(self):
        prs = new_deck()
        slide = add_slide(prs)
        add_box(slide, "Twelve binders fold well", name="Under", color="FFFFFF", left=1.0, top=1.8, width=7.5,
                height=0.8)
        add_rect(slide, "Cover", 0.8, 1.6, 8.0, 1.2, "0B1F3A")
        self.assertEqual([h.severity for h in self.contrast(prs)], ["error"])

    def test_table_text_over_a_dark_card_is_checked_against_the_card(self):
        prs = new_deck()
        slide = add_slide(prs)
        add_rect(slide, "Card", 0.8, 1.6, 7.7, 1.6, "0B1F3A")
        add_table(slide, rows=1, cols=1, color="FFFFFF")
        self.assertEqual(self.contrast(prs), [])

    def test_a_card_with_a_picture_below_the_text_is_a_picture_note(self):
        prs = new_deck()
        slide = add_slide(prs)
        add_picture(slide, make_png(self.tmp, "photo.png", 2000, 1200), left=0.8, top=1.6, width=8.0, height=3.0,
                    alt="A photo")
        add_box(slide, "Twelve binders fold well", color="FFFFFF", left=1.0, top=1.8, width=4.0, height=1.0)
        hits = self.contrast(prs)
        self.assertEqual(len(hits), 1)
        self.assertIn("text over a picture", hits[0].message)


class ChartTextContrast(LintCase):
    def dark_chart(self, **chart_options):
        prs = new_deck()
        slide = add_slide(prs, background="14111F", color="F3F0FA")
        add_chart(slide, colors=("9D7BFF",), labels=True, text_color="F3F0FA", **chart_options)
        return self.lint(prs)

    def test_light_labels_on_a_white_plot_area_fire(self):
        hits = self.of(self.dark_chart(plot_fill="FFFFFF"), "text-contrast")
        self.assertEqual([(h.severity, h.tier) for h in hits], [("error", "T1")])
        self.assertIn("chart", hits[0].message)

    def test_the_same_chart_with_a_transparent_plot_area_is_quiet(self):
        self.assertQuiet(self.dark_chart(plot_fill="none", chart_fill="none"), "text-contrast")

    def test_a_chart_with_no_fill_elements_counts_against_the_slide(self):
        self.assertQuiet(self.dark_chart(), "text-contrast")

    def test_a_white_chart_area_counts_too(self):
        hits = self.of(self.dark_chart(plot_fill="none", chart_fill="FFFFFF"), "text-contrast")
        self.assertEqual([h.severity for h in hits], ["error"])

    def test_dark_labels_on_a_dark_slide_fire(self):
        prs = new_deck()
        slide = add_slide(prs, background="14111F", color="F3F0FA")
        add_chart(slide, colors=("9D7BFF",), labels=True, text_color="111111")
        self.assertEqual([h.severity for h in self.of(self.lint(prs), "text-contrast")], ["error"])

    def test_the_plot_area_wins_over_the_chart_area(self):
        prs = new_deck()
        slide = add_slide(prs, background="14111F", color="F3F0FA")
        add_chart(slide, colors=("9D7BFF",), labels=True, text_color="F3F0FA", plot_fill="14111F",
                  chart_fill="FFFFFF")
        self.assertQuiet(self.lint(prs), "text-contrast")

    def test_a_chart_text_color_with_a_transform_is_unresolved(self):
        prs = new_deck()
        frame = add_chart(add_slide(prs), labels=True)
        for color in frame.chart._chartSpace.xpath("c:txPr//a:srgbClr"):
            color.append(parse_xml('<a:alpha %s val="50000"/>' % nsdecls("a")))
        hits = self.of(self.lint(prs), "text-contrast")
        self.assertEqual(len(hits), 1)
        self.assertIn("chart text color unresolved", hits[0].message)


class MarkContrast(LintCase):
    def bars(self, color="9CA3AF", background="FFFFFF", prepare=None, **chart_options):
        prs = new_deck()
        slide = add_slide(prs, background=background)
        frame = add_chart(slide, colors=(color,), **chart_options)
        if prepare is not None:
            prepare(frame.chart)
        return self.lint(prs)

    def test_grey_bars_with_no_labels_fire(self):
        hit = self.assertFires(self.bars(), "mark-contrast", "warn", "T1")
        self.assertIn("9CA3AF", hit.message)

    def test_grey_bars_with_data_labels_are_exempt(self):
        self.assertQuiet(self.bars(labels=True), "mark-contrast")

    def test_accent_bars_pass(self):
        self.assertQuiet(self.bars("6D4BD6"), "mark-contrast")

    def test_labels_of_the_series_exempt_it(self):
        def values(chart):
            chart.plots[0].series[0].data_labels.show_value = True

        def names(chart):
            labels = chart.plots[0].series[0].data_labels
            labels.show_value = False
            labels.show_series_name = True

        def one_point(chart):
            chart.plots[0].series[0].points[2].data_label.text_frame.text = "last"

        for prepare in (values, names, one_point):
            with self.subTest(prepare=prepare.__name__):
                self.assertQuiet(self.bars(prepare=prepare), "mark-contrast")

    def test_labels_that_show_nothing_do_not_exempt_the_series(self):
        def silent(chart):
            chart.plots[0].has_data_labels = True
            labels = chart.plots[0].series[0].data_labels
            labels.show_value = False
        self.assertFires(self.bars(prepare=silent), "mark-contrast", "warn", "T1")

    def test_a_series_with_labels_beside_a_series_with_none(self):
        prs = new_deck()
        slide = add_slide(prs)
        frame = add_chart(slide, colors=("9CA3AF", "CCCCCC"), line=True)
        frame.chart.plots[0].series[0].points[2].data_label.text_frame.text = "first"
        hit = self.assertFires(self.lint(prs), "mark-contrast", "warn", "T1")
        self.assertIn("CCCCCC", hit.message)

    def test_the_ratio_boundary(self):
        with mock.patch.object(design_color, "contrast_ratio", return_value=3.0):
            self.assertQuiet(self.bars(), "mark-contrast")
        with mock.patch.object(design_color, "contrast_ratio", return_value=2.999):
            self.assertFires(self.bars(), "mark-contrast", "warn", "T1")

    def test_the_context_color_of_the_dark_theme_passes(self):
        self.assertQuiet(self.bars("6E6788", background="14111F"), "mark-contrast")

    def test_marks_count_against_the_plot_area_fill(self):
        prs = new_deck()
        slide = add_slide(prs, background="14111F", color="F3F0FA")
        add_chart(slide, colors=("F3F0FA",), text_color="111111", plot_fill="FFFFFF")
        self.assertFires(self.lint(prs), "mark-contrast", "warn", "T1")

    def test_series_colors_from_the_theme_are_skipped_with_a_note(self):
        hit = self.assertFires(self.bars(None), "mark-contrast", "warn", "T1")
        self.assertIn("rule skipped", hit.message)

    def test_a_filled_shape_in_the_evidence_area_is_a_mark(self):
        def shape(color, top=2.0):
            prs = new_deck()
            add_rect(add_slide(prs), "Mark", 1.0, top, 2.0, 0.5, color)
            return self.lint(prs)
        self.assertFires(shape("DDDDDD"), "mark-contrast", "warn", "T1")
        self.assertQuiet(shape("6D4BD6"), "mark-contrast")
        self.assertQuiet(shape("DDDDDD", top=0.2), "mark-contrast")

    def test_a_light_card_behind_a_text_box_is_not_a_data_mark(self):
        prs = new_deck()
        slide = add_slide(prs)
        add_rect(slide, "Card", 0.8, 1.6, 8.0, 1.2, "EEEEEE")
        add_box(slide, "Twelve binders fold well", name="On the card", left=1.0, top=1.8, width=7.5, height=0.8)
        self.assertQuiet(self.lint(prs), "mark-contrast")

    def test_a_light_rectangle_beside_a_text_box_is_still_a_mark(self):
        prs = new_deck()
        slide = add_slide(prs)
        add_rect(slide, "Bar", 0.8, 1.6, 4.0, 1.2, "EEEEEE")
        add_box(slide, "Twelve binders fold well", name="Beside", left=6.0, top=1.8, width=4.0, height=0.8)
        self.assertFires(self.lint(prs), "mark-contrast", "warn", "T1")

    def test_a_filled_shape_that_holds_text_is_a_label_and_not_a_mark(self):
        prs = new_deck()
        slide = add_slide(prs)
        shape = add_rect(slide, "Label", 1.0, 2.0, 3.0, 0.8, "DDDDDD")
        write(shape, "Twelve", color="111111")
        self.assertQuiet(self.lint(prs), "mark-contrast")

    def test_the_ratio_boundary_of_a_filled_shape(self):
        def shape():
            prs = new_deck()
            add_rect(add_slide(prs), "Mark", 1.0, 2.0, 2.0, 0.5, "DDDDDD")
            return self.lint(prs)
        with mock.patch.object(design_color, "contrast_ratio", return_value=3.0):
            self.assertQuiet(shape(), "mark-contrast")
        with mock.patch.object(design_color, "contrast_ratio", return_value=2.999):
            self.assertFires(shape(), "mark-contrast", "warn", "T1")


class SeriesOutline(LintCase):
    def test_a_white_outline_around_blue_bars_is_not_a_white_series(self):
        prs = new_deck()
        add_chart(add_slide(prs), colors=("1F6FEB",), outline="FFFFFF")
        self.assertEqual(self.of(self.lint(prs), "mark-contrast"), [])

    def test_a_white_outline_on_bars_that_the_theme_fills_is_a_skip_and_not_a_white_series(self):
        prs = new_deck()
        add_chart(add_slide(prs), colors=(None,), outline="FFFFFF")
        hit = self.only(self.lint(prs), "mark-contrast")
        self.assertIn("series color comes from the theme", hit.message)
        self.assertNotIn("1.00", hit.message)

    def test_a_line_series_with_a_white_marker_edge_gives_no_white_series(self):
        prs = new_deck()
        frame = add_chart(add_slide(prs), colors=("0072B2",), line=True)
        frame.chart.plots[0].series[0].marker.format.line.color.rgb = RGBColor.from_string("FFFFFF")
        self.assertEqual(self.of(self.lint(prs), "mark-contrast"), [])

    def test_a_pale_fill_with_a_dark_outline_still_fires(self):
        prs = new_deck()
        add_chart(add_slide(prs), colors=("DDDDDD",), outline="111111")
        hit = self.only(self.lint(prs), "mark-contrast")
        self.assertIn("#DDDDDD", hit.message)

    def test_a_green_outline_around_a_red_and_a_blue_series_is_not_a_red_green_pair(self):
        prs = new_deck()
        add_chart(add_slide(prs), colors=("FF0000", "0072B2"), outline="00FF00", labels=True)
        self.assertEqual(self.of(self.lint(prs), "red-green"), [])


class RedGreen(LintCase):
    def test_the_premise_of_the_pairs(self):
        self.assertTrue(design_color.red_green_conflict("#FF0000", "#00FF00"))
        self.assertFalse(design_color.red_green_conflict("#D55E00", "#009E73"))
        self.assertFalse(design_color.red_green_conflict("#7A0000", "#80FF80"))

    def chart(self, colors):
        prs = new_deck()
        add_chart(add_slide(prs), colors=colors, labels=True)
        return self.lint(prs)

    def test_red_and_green_series_fire_and_the_message_names_the_heuristic(self):
        hit = self.assertFires(self.chart(("FF0000", "00FF00")), "red-green", "warn", "T2")
        self.assertIn("hue heuristic", hit.message)
        self.assertIn("FF0000", hit.message)

    def test_the_okabe_ito_vermilion_and_bluish_green_pair_is_not_flagged(self):
        self.assertQuiet(self.chart(("D55E00", "009E73")), "red-green")

    def test_a_pair_with_a_luminance_ratio_of_3_passes(self):
        self.assertQuiet(self.chart(("7A0000", "80FF80")), "red-green")

    def test_one_series_is_quiet(self):
        self.assertQuiet(self.chart(("FF0000",)), "red-green")

    def test_red_and_green_shapes_on_one_slide_fire(self):
        prs = new_deck()
        slide = add_slide(prs)
        add_rect(slide, "Bad", 1.0, 2.0, 2.0, 0.5, "FF0000")
        add_rect(slide, "Good", 4.0, 2.0, 2.0, 0.5, "00FF00")
        self.assertFires(self.lint(prs), "red-green", "warn", "T2")

    def test_red_and_green_shapes_with_enough_contrast_are_quiet(self):
        prs = new_deck()
        slide = add_slide(prs)
        add_rect(slide, "Bad", 1.0, 2.0, 2.0, 0.5, "7A0000")
        add_rect(slide, "Good", 4.0, 2.0, 2.0, 0.5, "80FF80")
        self.assertQuiet(self.lint(prs), "red-green")
