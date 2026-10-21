"""Tests for scripts/design_lint.py as a whole: the result type, unusual decks, profiles,
read errors, a deck with all six layouts, defective decks and the command line.
"""

import contextlib
import gc
import io
import os
import subprocess
import sys
import time
import weakref
from pathlib import Path
from unittest import mock

from pptx import Presentation
from pptx.oxml.ns import qn
from pptx.util import Inches, Pt

SCRIPTS_DIR = Path(__file__).resolve().parents[1] / "scripts"
for folder in (SCRIPTS_DIR, Path(__file__).resolve().parent):
    if str(folder) not in sys.path:
        sys.path.insert(0, str(folder))

import design_lint as dl  # noqa: E402
import slide_helpers as sh  # noqa: E402
import style as style_mod  # noqa: E402
from lint_helpers import (  # noqa: E402
    BODY,
    DAMAGES,
    DARK,
    GOOD_TITLE,
    LIGHT,
    LintCase,
    ROW,
    SCRIPT,
    add_box,
    add_group_of_boxes,
    add_chart,
    add_notes,
    add_slide,
    add_source,
    build_all_layouts,
    clean_slide,
    new_deck,
    words,
)


class ResultType(LintCase):
    def test_finding_has_five_named_fields(self):
        finding = dl.Finding(3, "bullets", "warn", "T2", "text")
        self.assertEqual(dl.Finding._fields, ("slide", "rule", "severity", "tier", "message"))
        self.assertEqual(tuple(finding), (3, "bullets", "warn", "T2", "text"))

    def test_lint_returns_a_list_of_findings(self):
        prs = new_deck()
        add_slide(prs, "Results")
        findings = self.lint(prs)
        self.assertIsInstance(findings, list)
        self.assertEqual([(f.slide, f.rule, f.severity, f.tier) for f in findings],
                         [(0, "slide-count", "warn", "T2"), (1, "title-sentence", "warn", "T3")])
        self.assertTrue(all(isinstance(f, dl.Finding) for f in findings))

    def test_a_clean_deck_gives_no_finding_at_all(self):
        prs = new_deck()
        slide = add_slide(prs, kind="assertion_evidence")
        add_box(slide, BODY)
        add_box(slide, "Twelve binders fold", name="Callout", left=9.2, width=3.3, size=20)
        add_source(slide)
        add_notes(slide, "Speak about the assay and the controls in your own words.")
        self.assertEqual(self.lint(prs, notes=True, minutes=5), [])

    def test_the_slide_kinds_match_the_helpers(self):
        self.assertEqual(tuple(dl.KINDS), tuple(sh.KINDS))

    def test_reading_does_not_change_the_file(self):
        prs = new_deck()
        clean_slide(prs)
        path = self.save(prs)
        before = path.read_bytes()
        first = dl.lint(path, notes=True)
        self.assertEqual(path.read_bytes(), before)
        self.assertEqual(dl.lint(path, notes=True), first)

    def test_findings_are_ordered_by_slide(self):
        prs = new_deck()
        clean_slide(prs, title="Results")
        clean_slide(prs, title="Methods")
        findings = self.lint(prs, minutes=1)
        self.assertEqual([f.slide for f in findings], sorted(f.slide for f in findings))
        self.assertEqual(findings[0].slide, 0)

    def test_lint_takes_a_string_or_a_path(self):
        prs = new_deck()
        clean_slide(prs)
        path = self.save(prs)
        self.assertEqual(dl.lint(str(path), notes=False), dl.lint(path, notes=False))


class UnusualDecks(LintCase):
    def test_a_deck_with_no_slides_is_one_deck_level_error(self):
        findings = self.lint(new_deck(), minutes=5)
        self.assertEqual(findings, [dl.Finding(0, "slide-count", "error", "T2", "the deck has no slides")])

    def test_a_deck_with_only_a_title_slide_runs(self):
        prs = new_deck()
        slide = add_slide(prs, "A made up talk about binder design", kind="title", size=44)
        add_box(slide, "Ann Lee", name="Subtitle", top=4.4, size=20, color="6B7280")
        findings = self.lint(prs, notes=True, minutes=5)
        self.assertEqual([f.severity for f in findings if f.severity == "error"], [])
        self.assertQuiet(findings, "slide-count")

    def test_a_hidden_slide_is_skipped_and_numbering_keeps_its_place(self):
        prs = new_deck()
        add_slide(prs, None)
        clean_slide(prs, title="Results")
        list(prs.slides)[0]._element.set("show", "0")
        findings = self.lint(prs)
        self.assertQuiet(findings, "title-empty")
        self.assertFires(findings, "title-sentence", "warn", "T3", slide=2)

    def test_a_slide_with_an_unreadable_shape_gets_a_note_and_no_false_error(self):
        prs = new_deck()
        slide = add_slide(prs)
        add_box(slide, words(41), height=5.0)
        ph = slide.shapes.title._element.find(qn("p:nvSpPr")).find(qn("p:nvPr")).find(qn("p:ph"))
        ph.set("type", "bogus")
        findings = self.lint(prs)
        hit = self.assertFires(findings, "unreadable-shape", "warn", "T3")
        self.assertIn("1 shape(s) could not be read", hit.message)
        self.assertQuiet(findings, "title-empty")
        self.assertFires(findings, "words-per-slide", "error", "T3")

    def reloaded_findings(self, damage, **options):
        prs = new_deck()
        slide = add_slide(prs)
        add_box(slide, BODY)
        damage(slide)
        return self.lint(prs, **options)

    def test_a_slide_with_no_shape_tree_gets_the_note_and_no_error(self):
        def damage(slide):
            tree = slide._element.cSld.spTree
            tree.getparent().remove(tree)
        findings = self.reloaded_findings(damage)
        hit = self.assertFires(findings, "unreadable-shape", "warn", "T3")
        self.assertIn("1 shape(s) could not be read", hit.message)
        self.assertQuiet(findings, "title-empty")

    def test_a_slide_with_no_common_slide_data_gets_the_note_and_no_error(self):
        findings = self.reloaded_findings(lambda slide: slide._element.remove(slide._element.cSld))
        self.assertFires(findings, "unreadable-shape", "warn", "T3")
        self.assertQuiet(findings, "title-empty")

    def test_hidden_slides_and_an_unreadable_slide_do_not_make_the_deck_empty(self):
        prs = new_deck()
        clean_slide(prs)
        slide = list(prs.slides)[0]
        slide._element.set("show", "0")
        path = self.save(prs)
        with mock.patch.object(dl, "_slides", return_value=([(1, slide)], [2])):
            findings = dl.lint(path, notes=False, minutes=5)
        self.assertEqual(self.of(findings, "slide-count"), [])
        self.assertFires(findings, "unreadable-shape", "warn", "T3", slide=2)

    def test_a_deck_from_another_tool_runs_and_reports_the_skipped_rules(self):
        prs = Presentation()
        slide = prs.slides.add_slide(prs.slide_layouts.get_by_name("Title and Content"))
        slide.shapes.title.text_frame.text = GOOD_TITLE
        slide.placeholders[1].text_frame.text = "Twelve of fifteen binders fold."
        second = prs.slides.add_slide(prs.slide_layouts.get_by_name("Title Only"))
        second.shapes.title.text_frame.text = "Twelve of fifteen binders fold well"
        box = second.shapes.add_textbox(Inches(1), Inches(2), Inches(6), Inches(1))
        box.text_frame.paragraphs[0].font.size = Pt(10)
        box.text_frame.paragraphs[0].add_run().text = "Small text from paragraph defaults"
        findings = self.lint(prs, notes=True, minutes=3)
        self.assertFires(findings, "aspect", "error", "T2", slide=0)
        skipped = [f for f in findings if "rule skipped" in f.message]
        self.assertTrue(any(f.slide == 2 and f.rule == "body-font" and "paragraph defaults" in f.message
                            for f in skipped), findings)
        self.assertEqual([f for f in findings if f.slide == 1 and f.rule in ("body-font", "text-contrast")], [])

    def test_a_deck_from_another_tool_with_inherited_sizes_resolves_them(self):
        prs = Presentation()
        slide = prs.slides.add_slide(prs.slide_layouts.get_by_name("Title and Content"))
        slide.shapes.title.text_frame.text = GOOD_TITLE
        slide.placeholders[1].text_frame.text = "Twelve of fifteen binders fold."
        findings = self.lint(prs)
        self.assertEqual(self.of(findings, "body-font"), [])
        self.assertEqual(self.of(findings, "title-font"), [])

    def test_a_slide_that_python_pptx_cannot_read_is_counted(self):
        class Collection(object):
            def __len__(self):
                return 3

            def __getitem__(self, index):
                if index == 1:
                    raise KeyError("rId9")
                return "slide %d" % index

        class Deck(object):
            slides = Collection()

        self.assertEqual(dl._slides(Deck()), ([(1, "slide 0"), (3, "slide 2")], [2]))


class Profiles(LintCase):
    def deck(self):
        prs = new_deck()
        slide = add_slide(prs)
        add_box(slide, words(30), size=22, height=5.0)
        return prs

    def test_each_profile_loads_and_changes_the_numbers(self):
        seen = {}
        for profile in style_mod.PROFILES:
            findings = self.lint(self.deck(), profile=profile)
            seen[profile] = (
                [f.severity for f in self.of(findings, "words-per-slide")],
                [f.severity for f in self.of(findings, "body-font")],
            )
        self.assertEqual(seen, {
            "peers": (["warn"], ["warn"]),
            "lab-meeting": ([], []),
            "general": (["error"], ["warn"]),
        })

    def test_the_dark_theme_gives_the_same_findings_as_the_light_theme(self):
        prs = new_deck()
        slide = add_slide(prs, "Results")
        add_box(slide, words(30), size=22, height=5.0)
        light = self.lint(prs, theme="light", minutes=1)
        self.assertTrue(light)
        self.assertEqual(self.lint(prs, theme="dark", minutes=1), light)

    def test_an_unknown_profile_or_theme_raises(self):
        prs = new_deck()
        clean_slide(prs)
        path = self.save(prs)
        with self.assertRaises(ValueError):
            dl.lint(path, profile="nobody")
        with self.assertRaises(ValueError):
            dl.lint(path, theme="blue")

    def test_the_project_file_changes_the_numbers_and_the_tier(self):
        project = self.tmp / "project"
        project.mkdir()
        (project / "slide-style-minimal.yaml").write_text(
            "lint:\n  words_per_slide: {warn: 50, error: 80, tier: T1}\n  text_contrast: {tier: T2}\n")
        prs = new_deck()
        slide = add_slide(prs)
        add_box(slide, words(60), height=5.0, color="CCCCCC")
        findings = self.lint(prs, project_dir=str(project))
        self.assertFires(findings, "words-per-slide", "warn", "T1")
        self.assertFires(findings, "text-contrast", "error", "T2")

    def test_a_missing_project_directory_raises(self):
        prs = new_deck()
        clean_slide(prs)
        with self.assertRaises(ValueError):
            self.lint(prs, project_dir=str(self.tmp / "nowhere"))


class StyleValues(LintCase):
    """A project style file can hold a value that the rules cannot use. lint() names the key and raises ValueError."""

    def project(self, text):
        folder = self.tmp / "project"
        folder.mkdir(exist_ok=True)
        (folder / "slide-style-minimal.yaml").write_text(text)
        return str(folder)

    def deck_path(self):
        prs = new_deck()
        clean_slide(prs)
        return self.save(prs)

    def test_a_threshold_that_is_not_a_number_raises_and_names_the_key(self):
        for text, key in (
            ("lint:\n  words_per_slide: {warn: many}\n", "lint.words_per_slide.warn"),
            ("lint:\n  body_font: {error_below: '18'}\n", "lint.body_font.error_below"),
            ("lint:\n  text_contrast: {large_text_pt: [18]}\n", "lint.text_contrast.large_text_pt"),
            ("lint:\n  notes: {redundancy_above: .nan}\n", "lint.notes.redundancy_above"),
            ("lint:\n  picture_ppi: {warn_below: .inf}\n", "lint.picture_ppi.warn_below"),
            ("layout:\n  content_top: tall\n", "layout.content_top"),
        ):
            with self.subTest(key=key):
                with self.assertRaises(ValueError) as caught:
                    dl.lint(self.deck_path(), project_dir=self.project(text))
                self.assertIn(key, str(caught.exception))

    def test_a_boolean_threshold_raises(self):
        with self.assertRaises(ValueError) as caught:
            dl.lint(self.deck_path(), project_dir=self.project("lint:\n  words_per_slide: {warn: true}\n"))
        self.assertIn("lint.words_per_slide.warn", str(caught.exception))

    def test_a_tier_that_is_not_t1_t2_or_t3_raises(self):
        for value in ("T4", "1", "true"):
            with self.subTest(value=value):
                with self.assertRaises(ValueError) as caught:
                    dl.lint(self.deck_path(), project_dir=self.project("lint:\n  bullets_per_list: {tier: %s}\n" % value))
                self.assertIn("lint.bullets_per_list.tier", str(caught.exception))

    def test_a_lint_key_that_the_linter_does_not_read_is_ignored_with_a_note(self):
        for text, key in (("lint:\n  my_own_rule: {warn: 3}\n", "my_own_rule"), ("lint:\n  my_own_rule: 5\n", "my_own_rule"),
                          ("lint:\n  my_own_rule: {tier: T9}\n", "my_own_rule"),
                          ("lint:\n  words_per_slides: {warn: 10, error: 20}\n", "words_per_slides")):
            with self.subTest(key=key, text=text):
                findings = dl.lint(self.deck_path(), project_dir=self.project(text), notes=False, minutes=5)
                self.assertEqual(findings, [dl.Finding(0, "style-key", "warn", "T3",
                                                       "ignored unknown lint key '{}'".format(key))])

    def test_each_unknown_key_gets_its_own_note_in_order(self):
        text = "lint:\n  zeta_rule: {warn: 1}\n  alpha_rule: {warn: 2}\n"
        findings = dl.lint(self.deck_path(), project_dir=self.project(text), notes=False, minutes=5)
        self.assertEqual([f.message for f in findings],
                         ["ignored unknown lint key 'alpha_rule'", "ignored unknown lint key 'zeta_rule'"])

    def test_a_typo_does_not_apply_its_limit_and_the_right_key_does(self):
        prs = new_deck()
        add_box(add_slide(prs), words(30), height=5.0)
        path = self.save(prs)
        typo = dl.lint(path, project_dir=self.project("lint:\n  words_per_slides: {warn: 50, error: 80}\n"), notes=False,
                       minutes=5)
        self.assertEqual(len(self.of(typo, "words-per-slide")), 1)
        self.assertEqual(len(self.of(typo, "style-key")), 1)
        right = dl.lint(path, project_dir=self.project("lint:\n  words_per_slide: {warn: 50, error: 80}\n"), notes=False,
                        minutes=5)
        self.assertEqual(self.of(right, "words-per-slide"), [])
        self.assertEqual(self.of(right, "style-key"), [])

    def test_the_known_lint_keys_give_no_note(self):
        for profile in style_mod.PROFILES:
            with self.subTest(profile=profile):
                self.assertEqual(self.of(dl.lint(self.deck_path(), profile=profile, minutes=5), "style-key"), [])

    def test_the_note_shows_on_the_command_line_as_a_deck_row(self):
        out, err = io.StringIO(), io.StringIO()
        with contextlib.redirect_stdout(out), contextlib.redirect_stderr(err):
            code = dl.main([str(self.deck_path()), "--no-notes", "--minutes", "5", "--project-dir",
                            self.project("lint:\n  words_per_slides: {warn: 10}\n")])
        self.assertEqual(code, 0)
        self.assertIn("deck: style-key: warn: T3: ignored unknown lint key 'words_per_slides'", out.getvalue())

    def test_lint_keys_of_mixed_types_get_a_note_each_and_the_deck_is_linted(self):
        # YAML reads the keys 1, true and ~ as a number, a boolean and None. The notes sort the keys by their text.
        # A Python dict holds 1 and True as one key, so each file has only one of them.
        for text, keys in (("lint:\n  1: {warn: 3}\n  foo: {warn: 2}\n  ~: 5\n", ("1", "None", "foo")),
                           ("lint:\n  foo: {warn: 2}\n  true: 4\n", ("True", "foo"))):
            with self.subTest(keys=keys):
                out, err = io.StringIO(), io.StringIO()
                with contextlib.redirect_stdout(out), contextlib.redirect_stderr(err):
                    code = dl.main([str(self.deck_path()), "--no-notes", "--minutes", "5", "--project-dir",
                                    self.project(text)])
                self.assertEqual((code, err.getvalue()), (0, ""))
                rows = [row for row in out.getvalue().splitlines() if "style-key" in row]
                self.assertEqual(rows, ["deck: style-key: warn: T3: ignored unknown lint key '{}'".format(key)
                                        for key in keys])
                self.assertIn("errors: 0", out.getvalue())

    def test_every_lint_key_that_the_linter_reads_needs_a_tier(self):
        for key in ("source_font", "lists_per_slide", "roadmap_sections", "title_unique", "red_green", "alt_text",
                    "title_words", "margin", "notes"):
            with self.subTest(key=key):
                with self.assertRaises(ValueError) as caught:
                    dl.lint(self.deck_path(), project_dir=self.project("lint:\n  %s: {tier: T9}\n" % key))
                self.assertIn("lint.%s.tier" % key, str(caught.exception))

    def test_a_tier_of_a_key_that_the_linter_reads_is_still_checked(self):
        with self.assertRaises(ValueError) as caught:
            dl.lint(self.deck_path(), project_dir=self.project("lint:\n  inside_slide: {tier: T9}\n"))
        self.assertIn("lint.inside_slide.tier", str(caught.exception))

    def test_a_number_that_the_rules_can_use_passes(self):
        text = "lint:\n  words_per_slide: {warn: 30.5, error: 50}\n  slide_count: {per_minute: 2}\n"
        self.assertIsInstance(dl.lint(self.deck_path(), project_dir=self.project(text)), list)

    def test_the_command_line_gives_exit_2_and_one_line_for_a_bad_style_value(self):
        out, err = io.StringIO(), io.StringIO()
        with contextlib.redirect_stdout(out), contextlib.redirect_stderr(err):
            code = dl.main([str(self.deck_path()), "--project-dir",
                            self.project("lint:\n  words_per_slide: {warn: many}\n")])
        self.assertEqual((code, out.getvalue()), (2, ""))
        self.assertEqual(len(err.getvalue().splitlines()), 1)
        self.assertIn("lint.words_per_slide.warn", err.getvalue())

    def test_the_peers_profile_and_the_others_pass_the_check(self):
        for profile in style_mod.PROFILES:
            for theme in style_mod.THEMES:
                with self.subTest(profile=profile, theme=theme):
                    self.assertIsInstance(dl.lint(self.deck_path(), profile=profile, theme=theme), list)


class ReadErrors(LintCase):
    def test_a_missing_file_raises_value_error(self):
        with self.assertRaises(ValueError) as caught:
            dl.lint(self.tmp / "missing.pptx")
        self.assertIn("missing.pptx", str(caught.exception))

    def test_a_file_that_is_not_a_deck_raises_value_error(self):
        path = self.tmp / "notes.pptx"
        path.write_text("this is not a zip file")
        with self.assertRaises(ValueError):
            dl.lint(path)

    def test_a_folder_raises_value_error(self):
        with self.assertRaises(ValueError):
            dl.lint(self.tmp)


class AllSixLayouts(LintCase):
    def test_the_deck_uses_every_layout(self):
        prs = build_all_layouts(LIGHT, self.tmp)
        self.assertEqual({sh.slide_kind(slide) for slide in prs.slides}, set(sh.KINDS))

    def test_the_light_deck_has_no_finding_under_peers(self):
        findings = dl.lint(self.save(build_all_layouts(LIGHT, self.tmp)), profile="peers", minutes=10)
        self.assertEqual(findings, [])

    def test_the_dark_deck_has_no_finding_under_peers(self):
        findings = dl.lint(self.save(build_all_layouts(DARK, self.tmp)), profile="peers", minutes=10,
                           theme="dark")
        self.assertEqual(findings, [])

    def test_the_deck_has_no_finding_under_the_other_profiles(self):
        path = self.save(build_all_layouts(LIGHT, self.tmp))
        for profile in ("lab-meeting", "general"):
            with self.subTest(profile=profile):
                findings = dl.lint(path, profile=profile, minutes=10)
                self.assertEqual([f for f in findings if f.severity == "error"], [])


class MasterGroupCost(LintCase):
    """A master or a layout with a large group is read once per call of lint()."""

    def deck(self, slides, shapes, where="master"):
        prs = new_deck()
        first = None
        for index in range(slides):
            slide = clean_slide(prs, title="Binder design succeeds for %d of 15 targets" % index)
            first = first or slide
        add_group_of_boxes(first, shapes, where)
        return prs

    def count_reads(self, prs, runs=1):
        path = self.save(prs)
        calls = []
        real = dl._leaves

        def spy(container):
            calls.append(type(container).__name__)
            return real(container)

        with mock.patch.object(dl, "_leaves", side_effect=spy):
            for _ in range(runs):
                dl.lint(path, minutes=100, notes=False)
        return calls

    def test_the_master_and_the_layout_are_read_once_for_eight_slides(self):
        calls = self.count_reads(self.deck(8, 20))
        self.assertEqual(calls.count("SlideMaster"), 1)
        self.assertEqual(calls.count("SlideLayout"), 1)
        self.assertEqual(calls.count("Slide"), 8)

    def test_a_layout_group_is_read_once_too(self):
        calls = self.count_reads(self.deck(6, 20, where="layout"))
        self.assertEqual((calls.count("SlideMaster"), calls.count("SlideLayout")), (1, 1))

    def test_each_call_of_lint_reads_them_again_because_the_cache_belongs_to_the_call(self):
        calls = self.count_reads(self.deck(3, 5), runs=2)
        self.assertEqual((calls.count("SlideMaster"), calls.count("SlideLayout")), (2, 2))
        # A cache that outlives the call keeps the parts, and the parts keep the deck.
        # So gc must free the deck after the call.
        path, opened, real = self.save(self.deck(3, 5)), [], dl._open

        def spy(source):
            prs = real(source)
            opened.append(weakref.ref(prs))
            return prs

        with mock.patch.object(dl, "_open", new=spy):
            dl.lint(path, minutes=100, notes=False)
        gc.collect()
        self.assertEqual(len(opened), 1)
        self.assertIsNone(opened[0](), "the deck of the call is still alive after lint() returned")

    def test_two_layouts_are_read_once_each(self):
        prs = new_deck()
        for name in ("Title Only", "Blank", "Title Only", "Blank"):
            slide = prs.slides.add_slide(prs.slide_layouts.get_by_name(name))
            slide.shapes.add_textbox(Inches(1), Inches(2), Inches(3), Inches(1)).text_frame.text = "words here"
        calls = self.count_reads(prs)
        self.assertEqual((calls.count("SlideMaster"), calls.count("SlideLayout")), (1, 2))

    def test_a_large_master_group_lints_in_a_few_seconds(self):
        prs = self.deck(25, 1000)
        path = self.save(prs)
        started = time.time()
        findings = dl.lint(path, minutes=100, notes=False)
        # On this machine the code with the cache took 1.35 s, and 7.5 s with the cache off. The old code took 12.8 s.
        # The counter test above guards the cache. This bound catches a gross slowdown and leaves room for a loaded machine.
        self.assertLess(time.time() - started, 10.0)
        self.assertEqual(self.of(findings, "text-contrast"), [])


class GeneralProfile(LintCase):
    """A deck that the helpers build from the general style, with content inside the general budgets."""

    def findings(self):
        general = style_mod.load_style("general", "light")
        path = self.save(build_all_layouts(general, self.tmp))
        return dl.lint(path, profile="general", minutes=10)

    def test_the_helpers_hold_the_content_in_the_larger_font(self):
        general = style_mod.load_style("general", "light")
        prs = build_all_layouts(general, self.tmp)
        self.assertEqual({sh.slide_kind(slide) for slide in prs.slides}, set(sh.KINDS))

    def test_the_deck_meets_the_general_budgets_on_all_six_layouts(self):
        self.assertEqual(self.findings(), [])

    def test_the_same_deck_under_peers_has_no_finding(self):
        general = style_mod.load_style("general", "light")
        path = self.save(build_all_layouts(general, self.tmp))
        self.assertEqual(dl.lint(path, profile="peers", minutes=10), [])


class CommandLine(LintCase):
    def run_cli(self, *args, **env):
        environment = dict(os.environ)
        environment.update(env)
        environment.pop("PYTHONIOENCODING", None)
        return subprocess.run([sys.executable, "-B", str(SCRIPT)] + [str(a) for a in args],
                              stdout=subprocess.PIPE, stderr=subprocess.PIPE, universal_newlines=True,
                              env=environment)

    def main(self, *args):
        out, err = io.StringIO(), io.StringIO()
        with contextlib.redirect_stdout(out), contextlib.redirect_stderr(err):
            code = dl.main([str(a) for a in args])
        return code, out.getvalue(), err.getvalue()

    def clean_path(self):
        prs = new_deck()
        slide = clean_slide(prs)
        add_notes(slide, "Speak about the assay and the controls.")
        return self.save(prs)

    def error_path(self):
        prs = new_deck()
        slide = add_slide(prs, None)
        add_box(slide, BODY)
        return self.save(prs)

    def test_a_clean_deck_exits_0_and_prints_the_summary(self):
        done = self.run_cli(self.clean_path(), "--minutes", "5")
        self.assertEqual((done.returncode, done.stderr), (0, ""))
        self.assertEqual(done.stdout.splitlines(), ["errors: 0", "warnings: 0"])

    def test_a_deck_with_an_error_exits_1_and_prints_rows_in_the_format(self):
        done = self.run_cli(self.error_path(), "--no-notes", "--minutes", "5")
        self.assertEqual(done.returncode, 1)
        lines = done.stdout.splitlines()
        self.assertEqual(lines[-2:], ["errors: 1", "warnings: 0"])
        rows = lines[:-2]
        self.assertEqual(len(rows), 1)
        self.assertRegex(rows[0], ROW)
        self.assertTrue(rows[0].startswith("slide 1: title-empty: error: T2: "), rows[0])

    def test_a_deck_level_row_starts_with_deck(self):
        prs = new_deck(10, 7.5)
        slide = add_slide(prs)
        slide.shapes.title.width = Inches(8.4)
        add_box(slide, BODY)
        code, out, err = self.main(self.save(prs), "--no-notes", "--minutes", "5")
        self.assertEqual((code, err), (1, ""))
        self.assertEqual(out.splitlines()[0].split(": ")[:4], ["deck", "aspect", "error", "T2"])
        self.assertRegex(out.splitlines()[0], ROW)

    def test_a_slide_row_starts_with_the_slide_number(self):
        prs = new_deck()
        clean_slide(prs)
        add_slide(prs, "Results")
        code, out, err = self.main(self.save(prs), "--no-notes", "--minutes", "5")
        self.assertEqual((code, err), (0, ""))
        self.assertIn("slide 2: title-sentence: warn: T3: ", out)

    def test_format_finding_gives_the_slide_row_and_the_deck_row(self):
        self.assertEqual(dl.format_finding(dl.Finding(3, "margin", "warn", "T3", "a  message\non two lines")),
                         "slide 3: margin: warn: T3: a message on two lines")
        self.assertEqual(dl.format_finding(dl.Finding(0, "aspect", "error", "T2", "the slide ratio is 1.333")),
                         "deck: aspect: error: T2: the slide ratio is 1.333")
        message = dl.format_finding(dl.Finding(1, "notes", "warn", "T3", "w" * 500)).split(": ", 4)[4]
        self.assertEqual((len(message), message[-3:]), (400, "..."))

    def test_main_prints_each_finding_with_format_finding(self):
        path = self.error_path()
        code, out, _ = self.main(path, "--minutes", "5")
        rows = out.splitlines()[:-2]
        self.assertTrue(rows)
        self.assertEqual(rows, [dl.format_finding(f) for f in dl.lint(path, minutes=5)])
        with mock.patch.object(dl, "format_finding", side_effect=lambda finding: "ROW " + finding.rule):
            _, out, _ = self.main(path, "--minutes", "5")
        self.assertEqual(out.splitlines()[:-2], ["ROW " + row.split(": ")[1] for row in rows])

    def test_warnings_alone_exit_0(self):
        code, out, err = self.main(self.clean_path())
        self.assertEqual(code, 0)
        self.assertIn("deck: slide-count: warn: T2: ", out)
        self.assertIn("warnings: 1", out)

    def test_a_missing_file_exits_2_with_one_line_on_stderr(self):
        done = self.run_cli(self.tmp / "missing.pptx")
        self.assertEqual((done.returncode, done.stdout), (2, ""))
        self.assertEqual(len(done.stderr.splitlines()), 1)
        self.assertIn("missing.pptx", done.stderr)

    def test_a_file_that_is_not_a_deck_exits_2(self):
        path = self.tmp / "text.pptx"
        path.write_text("not a deck")
        code, out, err = self.main(path)
        self.assertEqual((code, out), (2, ""))
        self.assertEqual(len(err.splitlines()), 1)

    def test_usage_errors_exit_2_with_one_line_on_stderr(self):
        path = self.clean_path()
        for args in ((), (path, "--profile", "nobody"), (path, "--theme", "blue"), (path, "--minutes", "abc"),
                     (path, "--minutes", "0"), (path, "--minutes", "-1"), (path, "--minutes", "inf"),
                     (path, "--minutes", "1e400"), (path, "--minutes", "nan"), (path, "--bogus"),
                     (path, "--project-dir", self.tmp / "nowhere")):
            with self.subTest(args=args):
                code, out, err = self.main(*args)
                self.assertEqual((code, out), (2, ""))
                self.assertEqual(len(err.splitlines()), 1, err)

    def test_an_unexpected_error_is_reported_with_its_type_and_exits_2(self):
        with mock.patch.object(dl, "lint", side_effect=RuntimeError("boom")):
            code, out, err = self.main(self.clean_path())
        self.assertEqual((code, out), (2, ""))
        self.assertEqual(err.splitlines(), ["design_lint.py: error: RuntimeError: boom"])

    def test_help_exits_0(self):
        code, out, err = self.main("--help")
        self.assertEqual((code, err), (0, ""))
        self.assertIn("--minutes", out)

    def test_the_help_and_the_module_say_that_the_theme_does_not_change_the_findings(self):
        code, out, err = self.main("--help")
        self.assertIn("no effect on the findings", " ".join(out.split()))
        self.assertIn("The theme has no effect on the findings", " ".join(dl.__doc__.split()))

    def test_the_theme_does_not_change_the_findings(self):
        path = self.clean_path()
        self.assertEqual(self.main(path, "--minutes", "5"), self.main(path, "--minutes", "5", "--theme", "dark"))

    def test_main_reads_sys_argv_by_default(self):
        out = io.StringIO()
        with mock.patch.object(sys, "argv", ["design_lint.py", str(self.clean_path()), "--minutes", "5"]):
            with contextlib.redirect_stdout(out):
                self.assertEqual(dl.main(), 0)
        self.assertIn("errors: 0", out.getvalue())

    def test_no_notes_switches_the_notes_rules_off(self):
        prs = new_deck()
        clean_slide(prs)
        path = self.save(prs)
        self.assertIn("notes: warn", self.main(path, "--minutes", "5")[1])
        self.assertNotIn("notes: warn", self.main(path, "--minutes", "5", "--no-notes")[1])

    def test_the_profile_and_the_theme_options_apply(self):
        prs = new_deck()
        slide = add_slide(prs)
        add_box(slide, words(30), height=5.0)
        path = self.save(prs)
        self.assertIn("words-per-slide: warn", self.main(path, "--no-notes")[1])
        self.assertNotIn("words-per-slide", self.main(path, "--no-notes", "--profile", "lab-meeting")[1])
        self.assertEqual(self.main(path, "--no-notes", "--theme", "dark")[0], 0)

    def test_the_project_directory_option_applies(self):
        project = self.tmp / "project"
        project.mkdir()
        (project / "slide-style-minimal.yaml").write_text("lint:\n  words_per_slide: {warn: 50, error: 80}\n")
        prs = new_deck()
        add_box(add_slide(prs), words(30), height=5.0)
        path = self.save(prs)
        self.assertNotIn("words-per-slide", self.main(path, "--no-notes", "--project-dir", project)[1])

    def test_no_render_message_without_the_option(self):
        self.assertNotIn("render", self.main(self.clean_path(), "--minutes", "5")[1])

    def test_non_ascii_text_does_not_crash_under_the_c_locale(self):
        prs = new_deck()
        slide = add_slide(prs, "Binders fold at 5 µM in 3 × more designs than α controls")
        add_box(slide, "µ × α", name="Probe µ × α", left=8.0, width=6.0)
        add_notes(slide, "Notes with µ and × and α in them.")
        # Python 3.7 and later coerce the C locale to UTF-8 and turn on the UTF-8 mode. Both variables switch that off,
        # so that stdout stays ASCII on every version and the encoder is the code under test.
        done = self.run_cli(self.save(prs), "--minutes", "5", LC_ALL="C", LANG="C", PYTHONCOERCECLOCALE="0",
                            PYTHONUTF8="0")
        self.assertIn(done.returncode, (0, 1))
        self.assertNotIn("Traceback", done.stderr)
        self.assertEqual(done.stderr, "")
        self.assertIn("inside-slide", done.stdout)
        self.assertTrue(all(ord(char) < 128 for char in done.stdout), done.stdout)
        self.assertIn("Probe ? ? ?", done.stdout)

    def test_a_lone_surrogate_is_replaced_on_a_strict_ascii_stream(self):
        stream = io.TextIOWrapper(io.BytesIO(), encoding="ascii")
        dl._write(stream, "aµ\ud800b\n")
        self.assertEqual(stream.buffer.getvalue(), b"a??b\n")

    def test_a_stream_with_no_buffer_gets_the_text_as_it_is(self):
        stream = io.StringIO()
        dl._write(stream, "aµb\n")
        self.assertEqual(stream.getvalue(), "aµb\n")

    def test_a_path_with_undecodable_bytes_exits_2_and_never_crashes(self):
        done = self.run_cli(self.tmp / "bad\udcffname.pptx", LC_ALL="C", LANG="C")
        self.assertEqual(done.returncode, 2)
        self.assertNotIn("Traceback", done.stderr)
        self.assertEqual(len(done.stderr.splitlines()), 1)

    def test_a_shape_name_with_a_line_break_stays_on_one_row(self):
        prs = new_deck()
        add_box(add_slide(prs), BODY, name="Probe\nsecond line", left=8.0, width=6.0)
        code, out, err = self.main(self.save(prs), "--no-notes", "--minutes", "5")
        self.assertEqual(code, 1)
        for row in out.splitlines()[:-2]:
            self.assertRegex(row, ROW)

    def test_a_series_name_with_a_line_break_stays_on_one_row(self):
        prs = new_deck()
        slide = add_slide(prs)
        frame = add_chart(slide, colors=("DDDDDD",))
        frame.chart._chartSpace.xpath(".//c:ser/c:tx//c:v")[0].text = "first\nsecond"
        out, err = io.StringIO(), io.StringIO()
        with contextlib.redirect_stdout(out), contextlib.redirect_stderr(err):
            dl.main([str(self.save(prs)), "--no-notes", "--minutes", "5"])
        rows = out.getvalue().splitlines()[:-2]
        self.assertTrue([row for row in rows if "mark-contrast" in row])
        for row in rows:
            self.assertRegex(row, ROW)


class DefectiveDecks(LintCase):
    def deck(self, damage):
        prs = new_deck()
        slide = add_slide(prs, kind="assertion_evidence")
        add_box(slide, "a readable box with words", size=14)
        add_notes(slide, "Notes with words")
        damage(self, slide)
        return prs

    def test_each_defect_gives_a_list_and_the_readable_box_is_still_checked(self):
        for name, damage in DAMAGES:
            with self.subTest(defect=name):
                findings = self.lint(self.deck(damage), notes=True, minutes=5)
                self.assertTrue(all(isinstance(f, dl.Finding) for f in findings))
                self.assertTrue(self.of(findings, "body-font"), findings)

    def test_a_deck_with_every_defect_at_once_gives_a_list(self):
        prs = new_deck()
        slide = add_slide(prs, kind="assertion_evidence")
        add_box(slide, "a readable box with words", size=14)
        for name, damage in DAMAGES:
            damage(self, slide)
        findings = self.lint(prs, notes=True, minutes=5)
        self.assertTrue(self.of(findings, "body-font"))
        self.assertTrue(any("could not be read" in f.message for f in findings))
