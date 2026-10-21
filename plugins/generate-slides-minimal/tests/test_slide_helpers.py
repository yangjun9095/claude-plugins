"""Tests for scripts/slide_helpers.py: six layouts, speaker notes and alt text."""
import copy
import datetime
import sys
import tempfile
import unittest
from pathlib import Path

from PIL import Image
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.dml import MSO_COLOR_TYPE, MSO_FILL
from pptx.enum.shapes import PP_PLACEHOLDER
from pptx.enum.text import MSO_ANCHOR, MSO_AUTO_SIZE
from pptx.oxml.ns import qn
from pptx.util import Inches, Pt

SCRIPTS_DIR = Path(__file__).resolve().parents[1] / "scripts"
if str(SCRIPTS_DIR) not in sys.path:
    sys.path.insert(0, str(SCRIPTS_DIR))

import slide_helpers as sh  # noqa: E402
import style as style_mod  # noqa: E402

EMU_PER_INCH = 914400.0
LIGHT = style_mod.load_style("peers", "light")
DARK = style_mod.load_style("peers", "dark")

TABLE_STYLE_ID = "{2D5ABB26-0587-4C30-8999-92F81FD0307C}"
EMU_TOLERANCE = 2  # Inches() truncates, so two lengths that should touch can differ by one EMU.


def inches(emu):
    return emu / EMU_PER_INCH


def size_pt(style, key):
    """Return the font size of a typography key as a pptx length."""
    return Pt(style["typography"][key])


def lay(style, key):
    """Return a layout value of the style, in inches, as a pptx length."""
    return Inches(style["layout"][key])


def make_png(path, width, height, color=(200, 60, 60)):
    Image.new("RGB", (width, height), color).save(str(path))
    return path


def rgb(hex_text):
    return RGBColor.from_string(hex_text.lstrip("#").upper())


class TmpDirCase(unittest.TestCase):
    def setUp(self):
        self._tmp = tempfile.TemporaryDirectory()
        self.addCleanup(self._tmp.cleanup)
        self.tmp = Path(self._tmp.name)


def blank_slide(style=LIGHT):
    prs = sh.new_presentation(style)
    return prs, prs.slides.add_slide(prs.slide_layouts[6])


class NewPresentation(unittest.TestCase):
    def test_size_comes_from_the_style_and_there_are_no_slides(self):
        prs = sh.new_presentation(LIGHT)
        self.assertEqual(prs.slide_width, Inches(13.333))
        self.assertEqual(prs.slide_height, Inches(7.5))
        self.assertEqual(len(prs.slides), 0)

    def test_a_style_with_another_size_changes_the_slide(self):
        style = copy.deepcopy(LIGHT)
        style["layout"]["slide_width"] = 10.0
        style["layout"]["slide_height"] = 5.625
        prs = sh.new_presentation(style)
        self.assertEqual((prs.slide_width, prs.slide_height), (Inches(10.0), Inches(5.625)))


class NewPresentationMetadata(TmpDirCase):
    """The python-pptx template names an author and holds dates from 2013. The helper must clear them."""

    def test_no_field_holds_the_name_of_the_template_author(self):
        properties = sh.new_presentation(LIGHT).core_properties
        self.assertNotIn("Canny", properties.last_modified_by)
        self.assertEqual(properties.author, "")
        self.assertEqual(properties.last_modified_by, "")

    def test_the_dates_are_the_current_utc_time(self):
        before = datetime.datetime.utcnow() - datetime.timedelta(seconds=1)
        properties = sh.new_presentation(LIGHT).core_properties
        after = datetime.datetime.utcnow() + datetime.timedelta(seconds=1)
        for name in ("created", "modified"):
            with self.subTest(field=name):
                value = getattr(properties, name)
                self.assertLessEqual(abs(value - datetime.datetime.utcnow()), datetime.timedelta(minutes=1))
                self.assertTrue(before <= value <= after, value)

    def test_the_text_fields_are_empty_and_the_revision_is_1(self):
        properties = sh.new_presentation(LIGHT).core_properties
        for name in ("title", "subject", "keywords", "comments", "category"):
            with self.subTest(field=name):
                self.assertEqual(getattr(properties, name), "")
        self.assertEqual(properties.revision, 1)

    def test_the_cleared_fields_survive_a_save_and_a_load(self):
        path = self.tmp / "deck.pptx"
        sh.new_presentation(LIGHT).save(str(path))
        properties = Presentation(str(path)).core_properties
        self.assertEqual((properties.author, properties.last_modified_by, properties.comments), ("", "", ""))
        self.assertLess(abs(properties.created - datetime.datetime.utcnow()), datetime.timedelta(minutes=1))

    def test_a_caller_can_set_a_field_after_the_helper(self):
        prs = sh.new_presentation(LIGHT)
        prs.core_properties.title = "Made-up title"
        prs.core_properties.author = "A. Author"
        self.assertEqual((prs.core_properties.title, prs.core_properties.author), ("Made-up title", "A. Author"))


class SlideKind(TmpDirCase):
    def test_round_trip_for_every_kind(self):
        for kind in ("title", "assertion_evidence", "backup", "big_number", "roadmap", "take_home"):
            _, slide = blank_slide()
            self.assertIsNone(sh.slide_kind(slide))
            sh.set_slide_kind(slide, kind)
            self.assertEqual(sh.slide_kind(slide), kind)

    def test_the_kind_is_the_name_attribute_of_the_slide_data(self):
        _, slide = blank_slide()
        sh.set_slide_kind(slide, "roadmap")
        self.assertEqual(slide._element.cSld.get("name"), "roadmap")

    def test_an_unknown_kind_raises_and_names_the_valid_kinds(self):
        _, slide = blank_slide()
        with self.assertRaises(ValueError) as caught:
            sh.set_slide_kind(slide, "poster")
        self.assertIn("poster", str(caught.exception))
        self.assertIn("take_home", str(caught.exception))
        self.assertIsNone(sh.slide_kind(slide))

    def test_the_kind_survives_save_and_reload(self):
        prs, slide = blank_slide()
        sh.set_slide_kind(slide, "big_number")
        out = self.tmp / "kind.pptx"
        prs.save(str(out))
        self.assertEqual(sh.slide_kind(Presentation(str(out)).slides[0]), "big_number")


class Notes(TmpDirCase):
    def test_notes_text_is_written_when_given(self):
        _, slide = blank_slide()
        sh.set_notes(slide, "Say this aloud.")
        self.assertTrue(slide.has_notes_slide)
        self.assertEqual(slide.notes_slide.notes_text_frame.text, "Say this aloud.")

    def test_none_and_empty_text_never_create_a_notes_slide(self):
        for text in (None, ""):
            _, slide = blank_slide()
            sh.set_notes(slide, text)
            self.assertFalse(slide.has_notes_slide)

    def test_text_that_is_not_a_string_raises(self):
        _, slide = blank_slide()
        with self.assertRaises(ValueError):
            sh.set_notes(slide, 42)
        self.assertFalse(slide.has_notes_slide)

    def test_notes_survive_save_and_reload(self):
        prs, slide = blank_slide()
        sh.set_notes(slide, "Kept.")
        out = self.tmp / "notes.pptx"
        prs.save(str(out))
        self.assertEqual(Presentation(str(out)).slides[0].notes_slide.notes_text_frame.text, "Kept.")


class AltText(unittest.TestCase):
    def descr(self, shape):
        return shape._element.xpath("./*[1]/p:cNvPr")[0].get("descr")

    def test_alt_text_goes_into_the_descr_attribute_for_each_shape_kind(self):
        _, slide = blank_slide()
        box = slide.shapes.add_textbox(Inches(1), Inches(1), Inches(2), Inches(1))
        table = slide.shapes.add_table(2, 2, Inches(1), Inches(3), Inches(3), Inches(1))
        for shape in (box, table):
            sh.set_alt_text(shape, "A short description.")
            self.assertEqual(self.descr(shape), "A short description.")

    def test_alt_text_replaces_an_earlier_value(self):
        _, slide = blank_slide()
        box = slide.shapes.add_textbox(Inches(1), Inches(1), Inches(2), Inches(1))
        sh.set_alt_text(box, "first")
        sh.set_alt_text(box, "second")
        self.assertEqual(self.descr(box), "second")

    def test_alt_text_that_is_not_a_string_raises(self):
        _, slide = blank_slide()
        box = slide.shapes.add_textbox(Inches(1), Inches(1), Inches(2), Inches(1))
        with self.assertRaises(ValueError):
            sh.set_alt_text(box, None)


class FitPicture(TmpDirCase):
    def test_a_wide_image_fills_the_box_width_and_centers_vertically(self):
        _, slide = blank_slide()
        path = make_png(self.tmp / "wide.png", 1600, 900)
        pic, _ = sh.fit_picture(slide, path, 1, 2, 4, 4)
        self.assertAlmostEqual(inches(pic.width), 4.0, places=4)
        self.assertAlmostEqual(inches(pic.height), 2.25, places=3)
        self.assertAlmostEqual(inches(pic.left), 1.0, places=4)
        self.assertAlmostEqual(inches(pic.top), 2.0 + (4 - 2.25) / 2, places=3)

    def test_a_tall_image_fills_the_box_height_and_centers_horizontally(self):
        _, slide = blank_slide()
        path = make_png(self.tmp / "tall.png", 900, 1600)
        pic, _ = sh.fit_picture(slide, path, 1, 2, 4, 4)
        self.assertAlmostEqual(inches(pic.height), 4.0, places=4)
        self.assertAlmostEqual(inches(pic.width), 2.25, places=3)
        self.assertAlmostEqual(inches(pic.left), 1.0 + (4 - 2.25) / 2, places=3)
        self.assertAlmostEqual(inches(pic.top), 2.0, places=4)

    def test_aspect_ratio_stays_within_half_a_percent_and_the_picture_stays_in_the_box(self):
        _, slide = blank_slide()
        box = (Inches(1.3), Inches(0.7), Inches(7.7), Inches(5.2))
        for width, height in ((1600, 900), (900, 1600), (1000, 1000), (3840, 2160), (333, 777), (50, 4000)):
            path = make_png(self.tmp / "p{}x{}.png".format(width, height), width, height)
            pic, _ = sh.fit_picture(slide, path, *box)
            self.assertLess(abs(pic.width / float(pic.height) / (width / float(height)) - 1), 0.005)
            self.assertGreaterEqual(pic.left, box[0])
            self.assertGreaterEqual(pic.top, box[1])
            self.assertLessEqual(pic.left + pic.width, box[0] + box[2])
            self.assertLessEqual(pic.top + pic.height, box[1] + box[3])

    def test_a_3840_pixel_image_in_a_10_inch_box_has_384_ppi(self):
        _, slide = blank_slide()
        path = make_png(self.tmp / "big.png", 3840, 2160)
        _, ppi = sh.fit_picture(slide, path, 1, 1, 10, 7)
        self.assertAlmostEqual(ppi, 384.0, places=3)

    def test_ppi_is_pixels_over_the_inches_of_the_fitted_picture(self):
        _, slide = blank_slide()
        path = make_png(self.tmp / "tall.png", 900, 1600)
        pic, ppi = sh.fit_picture(slide, path, 1, 1, 4, 4)
        crop = 0.0
        self.assertAlmostEqual(ppi, 900 * (1 - crop) / inches(pic.width), places=3)
        self.assertAlmostEqual(ppi, 1600 / 4.0, places=3)

    def test_inches_values_and_python_pptx_lengths_agree(self):
        _, slide = blank_slide()
        path = make_png(self.tmp / "a.png", 1600, 900)
        by_number, ppi_number = sh.fit_picture(slide, path, 1.5, 2, 6, 3)
        by_length, ppi_length = sh.fit_picture(slide, path, Inches(1.5), Inches(2), Inches(6), Inches(3))
        self.assertEqual((by_number.left, by_number.top, by_number.width, by_number.height),
                         (by_length.left, by_length.top, by_length.width, by_length.height))
        self.assertAlmostEqual(ppi_number, ppi_length, places=6)

    def test_the_file_name_never_becomes_the_alt_text(self):
        _, slide = blank_slide()
        path = make_png(self.tmp / "secret-name.png", 400, 300)
        pic, _ = sh.fit_picture(slide, path, 1, 1, 4, 3)
        descr = pic._element.xpath("./*[1]/p:cNvPr")[0].get("descr")
        self.assertNotIn("secret-name", descr or "")

    def test_a_missing_path_raises(self):
        _, slide = blank_slide()
        for value in (None, ""):
            with self.assertRaises(ValueError) as caught:
                sh.fit_picture(slide, value, 1, 1, 4, 3)
            self.assertIn("path", str(caught.exception))

    def test_an_empty_path_says_empty(self):
        _, slide = blank_slide()
        with self.assertRaises(ValueError) as caught:
            sh.fit_picture(slide, "", 1, 1, 4, 3)
        self.assertIn("empty", str(caught.exception))

    def test_a_folder_and_a_missing_file_raise_with_a_specific_message(self):
        _, slide = blank_slide()
        with self.assertRaises(ValueError) as caught:
            sh.fit_picture(slide, self.tmp, 1, 1, 4, 3)
        self.assertIn("folder", str(caught.exception))
        with self.assertRaises(ValueError) as caught:
            sh.fit_picture(slide, self.tmp / "none.png", 1, 1, 4, 3)
        self.assertIn("does not exist", str(caught.exception))

    def test_an_unreadable_image_raises(self):
        _, slide = blank_slide()
        text_file = self.tmp / "fake.png"
        text_file.write_text("this is not an image")
        truncated = self.tmp / "cut.png"
        full = make_png(self.tmp / "full.png", 400, 300)
        truncated.write_bytes(full.read_bytes()[:60])
        for bad in (text_file, truncated):
            with self.assertRaises(ValueError) as caught:
                sh.fit_picture(slide, bad, 1, 1, 4, 3)
            self.assertIn("cannot read the image", str(caught.exception))
        self.assertEqual(len(slide.shapes), 0)

    def test_an_extreme_aspect_ratio_still_gives_a_picture_with_a_positive_size(self):
        _, slide = blank_slide()
        path = make_png(self.tmp / "thin.png", 1, 30000)
        pic, ppi = sh.fit_picture(slide, path, 1, 1, 4, 3)
        self.assertGreater(pic.width, 0)
        self.assertGreater(ppi, 0)

    def test_a_format_that_python_pptx_cannot_embed_raises_before_anything_is_added(self):
        _, slide = blank_slide()
        path = self.tmp / "image.ppm"
        Image.new("RGB", (20, 10)).save(str(path))
        with self.assertRaises(ValueError) as caught:
            sh.fit_picture(slide, path, 1, 1, 4, 3)
        self.assertIn("PPM", str(caught.exception))
        self.assertIn("PNG", str(caught.exception))
        self.assertEqual(len(slide.shapes), 0)

    def test_a_box_with_no_area_raises(self):
        _, slide = blank_slide()
        path = make_png(self.tmp / "a.png", 400, 300)
        with self.assertRaises(ValueError):
            sh.fit_picture(slide, path, 1, 1, 0, 3)


# ---------------------------------------------------------------------------
# The six layouts
# ---------------------------------------------------------------------------

HEADLINE = "The new filter removed most noisy reads from every sample in the pilot"
ROADMAP_SECTIONS = [
    ("Aim", "State the question"),
    ("Method", "Describe the pipeline"),
    ("Result", "Show the main effect"),
    ("Limits", "Name the open points"),
]
CLAIMS = [
    ("The filter keeps signal and drops noise", "Effect size 0.50 against the control"),
    ("Report one number per sample", "Median of three runs"),
    ("Set the budget by audience", "Group A takes 4 items, group B takes 12"),
]


def build_title(prs, style, tmp):
    return sh.add_title_slide(prs, style, "What the evidence says about slides for scientific talks",
                              author="A. Author", affiliation="Biohub", date="October 2026",
                              notes="Open with the claim.")


def build_big_number(prs, style, tmp):
    return sh.add_big_number_slide(prs, style, HEADLINE, "42", "items per second",
                                   "3 million items a day in the invented example",
                                   source="Invented example data",
                                   notes="Divide the daily count by the seconds in a day.")


def build_roadmap(prs, style, tmp):
    return sh.add_roadmap_slide(prs, style, HEADLINE, ROADMAP_SECTIONS, notes="Four parts.")


def build_take_home(prs, style, tmp):
    return sh.add_take_home_slide(prs, style, CLAIMS, notes="Three claims.")


PICTURE_HEADLINE = "Labels next to the curves make the three groups easy to tell apart"
TABLE_HEADLINE = "The groups differ in how many items they allow, so the tool sets a budget"
BAR_HEADLINE = "Most items in the invented review rest on one kind of source"
LINE_HEADLINE = "The high dose raised the response above the other two doses within six days"
BACKUP_HEADLINE = "Each rule of the checker lists the source of its limit"
PICTURE_ALT = "Illustration with made-up data: three curves over time. The high-dose curve is violet."
TABLE_ROWS = [["Group", "Items per slide"], ["Group A", "at most 4"], ["Group B", "about 12 and one figure"],
              ["Group C", "up to 30"]]
TWO_CALLOUTS = ["Highlight one series, grey the rest", "The counts are invented"]
CHART_ALT = "Chart of invented values for three groups."


def evidence_picture(tmp, name="figure.png", alt=PICTURE_ALT):
    return {"kind": "picture", "path": str(make_png(tmp / name, 1600, 900)), "alt": alt}


def evidence_table(highlight=(2, 1), rows=TABLE_ROWS):
    return {"kind": "table", "rows": rows, "highlight": highlight}


def evidence_bar(highlight="Group A", **extra):
    data = {"kind": "chart", "chart": "bar", "categories": ["Group A", "Group B", "Group C"],
            "series": {"Items": [12, 30, 8]}, "highlight": highlight, "alt": CHART_ALT}
    data.update(extra)
    return data


def evidence_line(highlight="High dose", **extra):
    data = {"kind": "chart", "chart": "line", "categories": ["Day 0", "Day 2", "Day 4", "Day 6"],
            "series": {"Control": [0.18, 0.2, 0.22, 0.24], "Low dose": [0.18, 0.3, 0.4, 0.5],
                       "High dose": [0.18, 0.4, 0.62, 0.9]}, "highlight": highlight, "alt": CHART_ALT}
    data.update(extra)
    return data


def evidence_slide(prs, style, evidence, headline=BAR_HEADLINE, **kwargs):
    return sh.add_assertion_evidence_slide(prs, style, headline, evidence, **kwargs)


def build_picture(prs, style, tmp):
    return evidence_slide(prs, style, evidence_picture(tmp), PICTURE_HEADLINE, callouts=TWO_CALLOUTS,
                          source="Illustration with made-up data.", notes="The figure is made up.")


def build_table(prs, style, tmp):
    return evidence_slide(prs, style, evidence_table(), TABLE_HEADLINE, callouts=["Group A allows the fewest items"],
                          source="Invented example data.", notes="Three groups, three limits.")


def build_bar(prs, style, tmp):
    return evidence_slide(prs, style, evidence_bar(vmax=36), BAR_HEADLINE, callouts=TWO_CALLOUTS,
                          source="Invented counts for the test.", notes="Group B is the largest.")


def build_line(prs, style, tmp):
    return evidence_slide(prs, style, evidence_line(), LINE_HEADLINE, callouts=["Three doses, six days"],
                          source="Made-up data for illustration.", notes="Name each line.")


def build_backup(prs, style, tmp):
    return evidence_slide(prs, style, evidence_table(highlight=None), BACKUP_HEADLINE,
                          callouts=["Rule and source in one row"],
                          source="Invented example data.", notes="Backup slide.", backup=True)


# Each entry: (layout name, builder, expected slide kind, expected headline text).
LAYOUTS = [
    ("title", build_title, "title", "What the evidence says about slides for scientific talks"),
    ("big_number", build_big_number, "big_number", HEADLINE),
    ("roadmap", build_roadmap, "roadmap", HEADLINE),
    ("take_home", build_take_home, "take_home", "Three things to remember from this talk"),
    ("picture", build_picture, "assertion_evidence", PICTURE_HEADLINE),
    ("table", build_table, "assertion_evidence", TABLE_HEADLINE),
    ("bar", build_bar, "assertion_evidence", BAR_HEADLINE),
    ("line", build_line, "assertion_evidence", LINE_HEADLINE),
    ("backup", build_backup, "backup", BACKUP_HEADLINE),
]


def build_all(style, tmp):
    prs = sh.new_presentation(style)
    return prs, [(name, builder(prs, style, tmp), kind, headline) for name, builder, kind, headline in LAYOUTS]


def iter_runs(slide):
    """Yield (shape, run) for every text run in the shapes and in the table cells."""
    for shape in slide.shapes:
        frames = []
        if shape.has_text_frame:
            frames.append(shape.text_frame)
        if shape.has_table:
            frames.extend(cell.text_frame for row in shape.table.rows for cell in row.cells)
        for frame in frames:
            for paragraph in frame.paragraphs:
                for run in paragraph.runs:
                    yield shape, run


def shapes_named(slide, name):
    return [shape for shape in slide.shapes if shape.name == name]


class EveryLayout(TmpDirCase):
    def each(self):
        for style_name, style in (("light", LIGHT), ("dark", DARK)):
            prs, built = build_all(style, self.tmp)
            for name, slide, kind, headline in built:
                with self.subTest(style=style_name, layout=name):
                    yield style, prs, name, slide, kind, headline

    def test_every_shape_lies_inside_the_slide_and_keeps_the_margin(self):
        for style, prs, name, slide, kind, headline in self.each():
            margin = Inches(style["layout"]["margin"])
            for shape in slide.shapes:
                self.assertGreaterEqual(shape.left, margin, shape.name)
                self.assertLessEqual(shape.left + shape.width, prs.slide_width - margin, shape.name)
                self.assertGreaterEqual(shape.top, 0, shape.name)
                self.assertLessEqual(shape.top + shape.height, prs.slide_height, shape.name)

    def test_every_run_has_an_explicit_name_size_color_and_bold_false(self):
        for style, prs, name, slide, kind, headline in self.each():
            runs = list(iter_runs(slide))
            self.assertGreater(len(runs), 1)
            for shape, run in runs:
                font = run.font
                self.assertEqual(font.name, style["typography"]["body_font"], shape.name)
                self.assertIsNotNone(font.size, shape.name)
                self.assertIs(font.bold, False, shape.name)
                self.assertEqual(font.color.type, MSO_COLOR_TYPE.RGB, shape.name)

    def test_every_text_frame_wraps_without_margins_or_autofit(self):
        for style, prs, name, slide, kind, headline in self.each():
            for shape in slide.shapes:
                if not shape.has_text_frame or not shape.text_frame.text:
                    continue
                frame = shape.text_frame
                self.assertEqual((frame.margin_left, frame.margin_right, frame.margin_top, frame.margin_bottom),
                                 (0, 0, 0, 0), shape.name)
                self.assertTrue(frame.word_wrap, shape.name)
                self.assertEqual(frame.auto_size, MSO_AUTO_SIZE.NONE, shape.name)

    def test_the_title_shape_is_the_title_placeholder_and_holds_the_headline(self):
        for style, prs, name, slide, kind, headline in self.each():
            titles = shapes_named(slide, "Title")
            self.assertEqual(len(titles), 1)
            self.assertIs(slide.shapes.title._element, titles[0]._element)
            self.assertEqual(titles[0].placeholder_format.type, PP_PLACEHOLDER.TITLE)
            self.assertEqual(titles[0].text_frame.text, headline)

    def test_a_content_title_sits_at_the_top_in_28_pt_with_nothing_under_it(self):
        for style, prs, name, slide, kind, headline in self.each():
            if name == "title":
                continue
            title = slide.shapes.title
            self.assertLessEqual(inches(title.top), 0.35)
            self.assertEqual(title.top, Inches(style["layout"]["title_top"]))
            self.assertEqual(title.height, Inches(style["layout"]["title_height"]))
            self.assertEqual(title.left, Inches(style["layout"]["margin"]))
            self.assertEqual(title.text_frame.vertical_anchor, MSO_ANCHOR.TOP)
            self.assertEqual({run.font.size for run in title.text_frame.paragraphs[0].runs},
                             {size_pt(style, "title_size")})
            between = [s.name for s in slide.shapes if s is not title and s.top >= title.top + title.height
                       and s.top < Inches(style["layout"]["content_top"]) - EMU_TOLERANCE]
            self.assertEqual(between, [])

    def test_the_title_slide_title_uses_44_pt(self):
        for style, prs, name, slide, kind, headline in self.each():
            if name == "title":
                sizes = {run.font.size for run in slide.shapes.title.text_frame.paragraphs[0].runs}
                self.assertEqual(sizes, {Pt(sh.TITLE_SLIDE_PT)})
                self.assertEqual(sh.TITLE_SLIDE_PT, 44)

    def test_the_slide_kind_round_trips_for_each_layout(self):
        for style, prs, name, slide, kind, headline in self.each():
            self.assertEqual(sh.slide_kind(slide), kind)

    def test_notes_are_set_when_given(self):
        for style, prs, name, slide, kind, headline in self.each():
            self.assertTrue(slide.has_notes_slide)
            self.assertTrue(slide.notes_slide.notes_text_frame.text)

    def test_notes_are_absent_when_none(self):
        prs = sh.new_presentation(LIGHT)
        slides = [
            sh.add_title_slide(prs, LIGHT, "A title"),
            sh.add_big_number_slide(prs, LIGHT, HEADLINE, "42", "items per second", "a baseline"),
            sh.add_roadmap_slide(prs, LIGHT, HEADLINE, ROADMAP_SECTIONS),
            sh.add_take_home_slide(prs, LIGHT, CLAIMS),
            evidence_slide(prs, LIGHT, evidence_bar()),
        ]
        for slide in slides:
            self.assertFalse(slide.has_notes_slide)

    def test_the_background_comes_from_the_style(self):
        for style, prs, name, slide, kind, headline in self.each():
            fill = slide.background.fill
            self.assertEqual(fill.type, MSO_FILL.SOLID)
            self.assertEqual(fill.fore_color.rgb, rgb(style["colors"]["background"]))
        self.assertNotEqual(LIGHT["colors"]["background"], DARK["colors"]["background"])

    def test_the_dark_theme_sets_a_dark_background_on_every_slide(self):
        prs, built = build_all(DARK, self.tmp)
        for name, slide, kind, headline in built:
            color = slide.background.fill.fore_color.rgb
            self.assertLess(max(color), 0x40, name)

    def test_no_shape_is_named_accent_line(self):
        for style, prs, name, slide, kind, headline in self.each():
            self.assertNotIn("Accent line", [shape.name for shape in slide.shapes])

    def test_the_text_colors_come_from_the_style(self):
        for style, prs, name, slide, kind, headline in self.each():
            allowed = {rgb(style["colors"][key]) for key in ("text", "muted", "accent_ink")}
            for shape, run in iter_runs(slide):
                self.assertIn(run.font.color.rgb, allowed, shape.name)
            self.assertEqual(slide.shapes.title.text_frame.paragraphs[0].runs[0].font.color.rgb,
                             rgb(style["colors"]["text"]))

    def test_the_deck_keeps_kinds_and_notes_after_save_and_reload(self):
        prs, built = build_all(LIGHT, self.tmp)
        out = self.tmp / "deck.pptx"
        prs.save(str(out))
        reloaded = Presentation(str(out))
        self.assertEqual(len(reloaded.slides), len(built))
        for (name, slide, kind, headline), again in zip(built, reloaded.slides):
            self.assertEqual(sh.slide_kind(again), kind)
            self.assertEqual(again.notes_slide.notes_text_frame.text, slide.notes_slide.notes_text_frame.text)
            self.assertEqual(again.shapes.title.text_frame.text, headline)

    def test_a_blank_headline_raises_and_adds_no_slide(self):
        prs = sh.new_presentation(LIGHT)
        with self.assertRaises(ValueError):
            sh.add_big_number_slide(prs, LIGHT, "  ", "42", "items per second", "a baseline")
        with self.assertRaises(ValueError):
            sh.add_title_slide(prs, LIGHT, "")
        self.assertEqual(len(prs.slides), 0)

    def test_the_free_layouts_stay_in_the_content_area(self):
        for style, prs, name, slide, kind, headline in self.each():
            if name == "title":
                continue
            top = Inches(style["layout"]["content_top"])
            bottom = Inches(style["layout"]["content_bottom"])
            for shape in slide.shapes:
                if shape.name in ("Title", "Source"):
                    continue
                self.assertGreaterEqual(shape.top, top - EMU_TOLERANCE, shape.name)
                self.assertLessEqual(shape.top + shape.height, bottom + EMU_TOLERANCE, shape.name)

    def test_two_styles_that_differ_in_body_size_give_different_sizes(self):
        big = copy.deepcopy(LIGHT)
        big["typography"]["body_size"] = 30
        sizes = []
        for style in (LIGHT, big):
            prs = sh.new_presentation(style)
            slide = sh.add_big_number_slide(prs, style, HEADLINE, "42", "items per second", "a baseline")
            sizes.append(shapes_named(slide, "Baseline")[0].text_frame.paragraphs[0].runs[0].font.size)
        self.assertEqual(sizes, [size_pt(LIGHT, "body_size"), Pt(30)])

    def test_two_styles_that_differ_in_color_and_content_top_give_different_looks(self):
        other = copy.deepcopy(LIGHT)
        other["colors"]["accent_ink"] = "#B00020"
        other["layout"]["content_top"] = 2.0
        numbers = []
        for style in (LIGHT, other):
            prs = sh.new_presentation(style)
            slide = sh.add_big_number_slide(prs, style, HEADLINE, "42", "items per second", "a baseline")
            numbers.append(shapes_named(slide, "Number")[0])
        self.assertEqual(numbers[0].text_frame.paragraphs[0].runs[0].font.color.rgb, rgb("#6D4BD6"))
        self.assertEqual(numbers[1].text_frame.paragraphs[0].runs[0].font.color.rgb, rgb("#B00020"))
        self.assertAlmostEqual(inches(numbers[1].top - numbers[0].top), 0.5, places=3)

    def test_a_style_that_lacks_a_value_raises_and_names_it(self):
        broken = copy.deepcopy(LIGHT)
        del broken["colors"]["accent"]
        prs = sh.new_presentation(broken)
        with self.assertRaises(ValueError) as caught:
            sh.add_roadmap_slide(prs, broken, HEADLINE, ROADMAP_SECTIONS)
        self.assertIn("colors.accent", str(caught.exception))
        self.assertEqual(len(prs.slides), 0)

    def test_a_style_with_a_bad_color_raises_before_any_slide_is_added(self):
        broken = copy.deepcopy(LIGHT)
        broken["colors"]["muted"] = "grey"
        prs = sh.new_presentation(broken)
        with self.assertRaises(ValueError):
            sh.add_take_home_slide(prs, broken, CLAIMS)
        self.assertEqual(len(prs.slides), 0)


class TitleSlide(TmpDirCase):
    def test_the_shapes_and_their_names(self):
        prs = sh.new_presentation(LIGHT)
        slide = build_title(prs, LIGHT, self.tmp)
        self.assertEqual([s.name for s in slide.shapes], ["Title", "Accent mark", "Subtitle"])

    def test_the_accent_mark_is_a_short_thin_rule_in_the_accent_color(self):
        for style in (LIGHT, DARK):
            prs = sh.new_presentation(style)
            slide = build_title(prs, style, self.tmp)
            mark = shapes_named(slide, "Accent mark")[0]
            self.assertAlmostEqual(inches(mark.width), 1.4, places=3)
            self.assertEqual(mark.height, Inches(style["layout"]["rule_width"]))
            self.assertEqual(mark.fill.fore_color.rgb, rgb(style["colors"]["accent"]))
            self.assertEqual(mark.line.fill.type, MSO_FILL.BACKGROUND)
            self.assertEqual(mark.left, Inches(style["layout"]["margin"]))
            self.assertGreaterEqual(mark.top, slide.shapes.title.top + slide.shapes.title.height)

    def test_the_subtitle_holds_one_muted_paragraph_per_given_line(self):
        prs = sh.new_presentation(LIGHT)
        slide = build_title(prs, LIGHT, self.tmp)
        subtitle = shapes_named(slide, "Subtitle")[0]
        self.assertEqual([p.text for p in subtitle.text_frame.paragraphs], ["A. Author", "Biohub", "October 2026"])
        for run in (r for p in subtitle.text_frame.paragraphs for r in p.runs):
            self.assertEqual(run.font.size, size_pt(LIGHT, "callout_size"))
            self.assertEqual(run.font.color.rgb, rgb(LIGHT["colors"]["muted"]))

    def test_empty_lines_are_skipped_and_no_lines_means_no_subtitle(self):
        prs = sh.new_presentation(LIGHT)
        only_date = sh.add_title_slide(prs, LIGHT, "A title", date="2026")
        self.assertEqual([p.text for p in shapes_named(only_date, "Subtitle")[0].text_frame.paragraphs], ["2026"])
        bare = sh.add_title_slide(prs, LIGHT, "A title")
        self.assertEqual(shapes_named(bare, "Subtitle"), [])

    def test_a_subtitle_line_that_is_not_a_string_raises(self):
        prs = sh.new_presentation(LIGHT)
        with self.assertRaises(ValueError):
            sh.add_title_slide(prs, LIGHT, "A title", author=7)
        self.assertEqual(len(prs.slides), 0)


class BigNumberSlide(TmpDirCase):
    def setUp(self):
        super().setUp()
        self.prs = sh.new_presentation(LIGHT)
        self.slide = build_big_number(self.prs, LIGHT, self.tmp)

    def run_of(self, name):
        return shapes_named(self.slide, name)[0].text_frame.paragraphs[0].runs[0]

    def test_the_shape_names(self):
        self.assertEqual([s.name for s in self.slide.shapes],
                         ["Title", "Number", "Number label", "Divider", "Baseline", "Source"])

    def test_sizes_and_colors_follow_the_style(self):
        number = self.run_of("Number")
        self.assertEqual((number.font.size, number.font.color.rgb),
                         (size_pt(LIGHT, "big_number_size"), rgb(LIGHT["colors"]["accent_ink"])))
        label = self.run_of("Number label")
        self.assertEqual((label.font.size, label.font.color.rgb),
                         (size_pt(LIGHT, "label_size"), rgb(LIGHT["colors"]["text"])))
        baseline = self.run_of("Baseline")
        self.assertEqual((baseline.font.size, baseline.font.color.rgb),
                         (size_pt(LIGHT, "body_size"), rgb(LIGHT["colors"]["text"])))
        source = self.run_of("Source")
        self.assertEqual((source.font.size, source.font.color.rgb),
                         (size_pt(LIGHT, "source_size"), rgb(LIGHT["colors"]["muted"])))

    def test_the_number_sits_left_of_the_divider_and_the_baseline_right_of_it(self):
        number, divider, baseline = (shapes_named(self.slide, n)[0] for n in ("Number", "Divider", "Baseline"))
        self.assertLessEqual(number.left + number.width, divider.left)
        self.assertLessEqual(divider.left + divider.width, baseline.left)
        self.assertEqual(divider.width, Inches(LIGHT["layout"]["rule_width"]))
        self.assertEqual(divider.fill.fore_color.rgb, rgb(LIGHT["colors"]["accent"]))
        label = shapes_named(self.slide, "Number label")[0]
        self.assertLessEqual(number.top + number.height, label.top + EMU_TOLERANCE)

    def test_the_source_line_sits_below_the_content_area(self):
        source = shapes_named(self.slide, "Source")[0]
        self.assertGreaterEqual(source.top, Inches(LIGHT["layout"]["content_bottom"]))
        self.assertEqual(source.text_frame.text, "Invented example data")

    def test_no_source_means_no_source_shape(self):
        slide = sh.add_big_number_slide(self.prs, LIGHT, HEADLINE, "42", "items per second", "a baseline")
        self.assertEqual(shapes_named(slide, "Source"), [])

    def test_a_number_too_long_for_its_box_raises_and_adds_no_slide(self):
        count = len(self.prs.slides)
        with self.assertRaises(ValueError) as caught:
            sh.add_big_number_slide(self.prs, LIGHT, HEADLINE, "1,234,567,890", "items", "a baseline")
        self.assertIn("1,234,567,890", str(caught.exception))
        self.assertEqual(len(self.prs.slides), count)

    def test_a_short_number_with_a_unit_fits(self):
        for number in ("58", "$1.2M", "4.5x", "98.6%"):
            sh.add_big_number_slide(self.prs, LIGHT, HEADLINE, number, "label", "a baseline")


class RoadmapSlide(TmpDirCase):
    def test_the_shape_names_and_counts(self):
        prs = sh.new_presentation(LIGHT)
        slide = build_roadmap(prs, LIGHT, self.tmp)
        names = [s.name for s in slide.shapes]
        for name in ("Section mark", "Section name", "Section line"):
            self.assertEqual(names.count(name), 4)
        self.assertEqual(names.count("Title"), 1)

    def test_it_rejects_2_and_6_sections_and_accepts_3_to_5(self):
        prs = sh.new_presentation(LIGHT)
        for count in (0, 1, 2, 6, 7):
            with self.assertRaises(ValueError) as caught:
                sh.add_roadmap_slide(prs, LIGHT, HEADLINE, [("S{}".format(i), "line") for i in range(count)])
            self.assertIn("3 to 5", str(caught.exception))
        self.assertEqual(len(prs.slides), 0)
        for count in (3, 4, 5):
            sh.add_roadmap_slide(prs, LIGHT, HEADLINE, [("S{}".format(i), "line") for i in range(count)])
        self.assertEqual(len(prs.slides), 3)

    def test_a_section_that_is_not_a_name_and_line_pair_raises(self):
        prs = sh.new_presentation(LIGHT)
        for sections in ([("A", "x"), ("B", "y"), "C"], [("A", "x"), ("B", "y"), ("C",)],
                         [("A", "x"), ("B", "y"), ("C", "")]):
            with self.assertRaises(ValueError):
                sh.add_roadmap_slide(prs, LIGHT, HEADLINE, sections)
        self.assertEqual(len(prs.slides), 0)

    def test_the_columns_have_equal_width_do_not_overlap_and_fit_the_margins(self):
        for count in (3, 4, 5):
            prs = sh.new_presentation(LIGHT)
            slide = sh.add_roadmap_slide(prs, LIGHT, HEADLINE, [("S{}".format(i), "line") for i in range(count)])
            for name in ("Section mark", "Section name", "Section line"):
                shapes = shapes_named(slide, name)
                self.assertEqual(len({s.width for s in shapes}), 1)
                for left, right in zip(shapes, shapes[1:]):
                    self.assertLess(left.left + left.width, right.left)
                self.assertGreaterEqual(shapes[0].left, lay(LIGHT, "margin"))
                self.assertLessEqual(shapes[-1].left + shapes[-1].width, prs.slide_width - lay(LIGHT, "margin"))

    def test_the_text_follows_the_style(self):
        prs = sh.new_presentation(DARK)
        slide = build_roadmap(prs, DARK, self.tmp)
        name_run = shapes_named(slide, "Section name")[0].text_frame.paragraphs[0].runs[0]
        line_run = shapes_named(slide, "Section line")[0].text_frame.paragraphs[0].runs[0]
        self.assertEqual(name_run.font.color.rgb, rgb(DARK["colors"]["accent_ink"]))
        self.assertGreater(name_run.font.size, line_run.font.size)
        self.assertEqual(line_run.font.size, size_pt(DARK, "body_size"))
        self.assertEqual(shapes_named(slide, "Section mark")[0].fill.fore_color.rgb, rgb(DARK["colors"]["accent"]))
        self.assertEqual(shapes_named(slide, "Section mark")[0].height, Inches(DARK["layout"]["rule_width"]))


    def names(self, *names):
        return [(name, "a line of text") for name in names]

    def test_a_name_that_needs_more_than_one_line_raises_and_names_the_section_and_the_limit(self):
        prs = sh.new_presentation(LIGHT)
        cases = [  # (names, the culprit, the limit in characters for that column count)
            (("Aim", "Method", "Result", "Limits", "Introduction"), "Introduction", "9 characters"),
            (("Aim", "Background", "Result", "Limits", "Tool"), "Background", "9 characters"),
            (("Aim", "Method", "Result", "Limits", "Conclusion"), "Conclusion", "9 characters"),
            (("Aim", "Results and discussion", "Result", "Limits"), "Results and discussion", "12 characters"),
        ]
        for names, culprit, limit in cases:
            with self.subTest(culprit=culprit):
                with self.assertRaises(ValueError) as caught:
                    sh.add_roadmap_slide(prs, LIGHT, HEADLINE, self.names(*names))
                for text in (culprit, limit, "one line", "Shorten"):
                    self.assertIn(text, str(caught.exception))
        self.assertEqual(len(prs.slides), 0)

    def test_the_longest_name_that_fits_passes_and_one_more_letter_raises(self):
        prs = sh.new_presentation(LIGHT)
        five = ("Aim", "Method", "Result", "Limits")
        sh.add_roadmap_slide(prs, LIGHT, HEADLINE, self.names(*five, "Reasoning"))  # 9 letters
        with self.assertRaises(ValueError):
            sh.add_roadmap_slide(prs, LIGHT, HEADLINE, self.names(*five, "Discussion"))  # 10 letters
        three = ("Aim", "Method")
        sh.add_roadmap_slide(prs, LIGHT, HEADLINE, self.names(*three, "x" * 16))
        with self.assertRaises(ValueError):
            sh.add_roadmap_slide(prs, LIGHT, HEADLINE, self.names(*three, "x" * 17))
        self.assertEqual(len(prs.slides), 2)

    def test_short_names_pass_with_five_sections(self):
        prs = sh.new_presentation(LIGHT)
        slide = sh.add_roadmap_slide(prs, LIGHT, HEADLINE, self.names("Aim", "Method", "Result", "Limits", "Tool"))
        self.assertEqual(len(shapes_named(slide, "Section name")), 5)

    def test_the_limit_follows_the_style(self):
        big = copy.deepcopy(LIGHT)
        big["typography"]["body_size"] = 40
        names = self.names("Aim", "Method", "Result", "Limits", "Results")
        sh.add_roadmap_slide(sh.new_presentation(LIGHT), LIGHT, HEADLINE, names)
        with self.assertRaises(ValueError):
            sh.add_roadmap_slide(sh.new_presentation(big), big, HEADLINE, names)


class TakeHomeSlide(TmpDirCase):
    def test_the_shape_names_and_counts(self):
        prs = sh.new_presentation(LIGHT)
        slide = build_take_home(prs, LIGHT, self.tmp)
        names = [s.name for s in slide.shapes]
        for name in ("Claim", "Claim mark"):
            self.assertEqual(names.count(name), 3)
        self.assertEqual(names.count("Title"), 1)
        self.assertEqual(names.count("Baseline"), 0, "the baseline is a paragraph of its Claim shape")

    def test_each_claim_shape_holds_the_claim_and_then_its_baseline(self):
        prs = sh.new_presentation(LIGHT)
        slide = build_take_home(prs, LIGHT, self.tmp)
        for shape, (claim, baseline) in zip(shapes_named(slide, "Claim"), CLAIMS):
            self.assertEqual([p.text for p in shape.text_frame.paragraphs], [claim, baseline])

    def test_it_rejects_0_and_4_claims_and_accepts_1_to_3(self):
        prs = sh.new_presentation(LIGHT)
        for count in (0, 4, 5):
            with self.assertRaises(ValueError) as caught:
                sh.add_take_home_slide(prs, LIGHT, [("Claim {}".format(i), "baseline") for i in range(count)])
            self.assertIn("1 to 3", str(caught.exception))
        self.assertEqual(len(prs.slides), 0)
        for count in (1, 2, 3):
            sh.add_take_home_slide(prs, LIGHT, [("Claim {}".format(i), "baseline") for i in range(count)])
        self.assertEqual(len(prs.slides), 3)

    def test_a_claim_that_is_not_a_claim_and_baseline_pair_raises(self):
        prs = sh.new_presentation(LIGHT)
        for claims in (["only a string"], [("claim",)], [("claim", "")], [("", "baseline")]):
            with self.assertRaises(ValueError):
                sh.add_take_home_slide(prs, LIGHT, claims)
        self.assertEqual(len(prs.slides), 0)

    def test_the_default_headline_counts_the_claims_and_a_given_headline_replaces_it(self):
        prs = sh.new_presentation(LIGHT)
        titles = []
        for count in (1, 2, 3):
            slide = sh.add_take_home_slide(prs, LIGHT, CLAIMS[:count])
            titles.append(slide.shapes.title.text_frame.text)
        self.assertEqual(titles, ["One thing to remember from this talk", "Two things to remember from this talk",
                                  "Three things to remember from this talk"])
        for title in titles:
            self.assertGreaterEqual(len(title.split()), 6)
        slide = sh.add_take_home_slide(prs, LIGHT, CLAIMS, headline="Three rules carry the argument of this deck")
        self.assertEqual(slide.shapes.title.text_frame.text, "Three rules carry the argument of this deck")
        with self.assertRaises(ValueError):
            sh.add_take_home_slide(prs, LIGHT, CLAIMS, headline=" ")

    def test_claims_and_baselines_use_the_body_and_callout_sizes_and_colors(self):
        prs = sh.new_presentation(LIGHT)
        slide = build_take_home(prs, LIGHT, self.tmp)
        paragraphs = shapes_named(slide, "Claim")[0].text_frame.paragraphs
        claim, baseline = paragraphs[0].runs[0], paragraphs[1].runs[0]
        self.assertEqual((claim.font.size, claim.font.color.rgb),
                         (size_pt(LIGHT, "body_size"), rgb(LIGHT["colors"]["text"])))
        self.assertEqual((baseline.font.size, baseline.font.color.rgb),
                         (size_pt(LIGHT, "callout_size"), rgb(LIGHT["colors"]["muted"])))
        mark = shapes_named(slide, "Claim mark")[0]
        self.assertEqual(mark.fill.fore_color.rgb, rgb(LIGHT["colors"]["accent"]))
        self.assertEqual(mark.width, mark.height)

    def assert_no_overlap(self, slide):
        boxes = sorted(shapes_named(slide, "Claim"), key=lambda s: s.top)
        for upper, lower in zip(boxes, boxes[1:]):
            self.assertLessEqual(upper.top + upper.height, lower.top)

    def test_a_claim_of_71_characters_leaves_no_gap_before_its_baseline(self):
        # The estimate gives this claim two lines at 24 pt, but a renderer can set it on one line.
        # In one top-anchored box with no paragraph space, the baseline then follows the first line.
        claim = "The 16-hour schedule raised the mean height of seedlings by 40 percent."
        self.assertEqual(len(claim), 71)
        text_in = LIGHT["layout"]["slide_width"] - 2 * LIGHT["layout"]["margin"] - sh.CLAIM_INDENT
        body, small = LIGHT["typography"]["body_size"], LIGHT["typography"]["callout_size"]
        self.assertEqual(sh._estimate_lines(claim, text_in, body), 2)
        prs = sh.new_presentation(LIGHT)
        slide = sh.add_take_home_slide(prs, LIGHT, [(claim, "Baseline: the 8-hour schedule"), ("Next claim", "base")])
        self.assertEqual(shapes_named(slide, "Baseline"), [])
        first, second = shapes_named(slide, "Claim")
        frame = first.text_frame
        self.assertEqual([p.text for p in frame.paragraphs], [claim, "Baseline: the 8-hour schedule"])
        self.assertEqual(frame.vertical_anchor, MSO_ANCHOR.TOP)
        self.assertEqual(frame.margin_top, 0)
        for paragraph in frame.paragraphs:
            p_pr = paragraph._p.find(qn("a:pPr"))
            for tag in ("a:spcBef", "a:spcAft", "a:lnSpc"):
                self.assertIsNone(None if p_pr is None else p_pr.find(qn(tag)), tag)
        # The box holds two claim lines and one baseline line, and the next claim starts below it.
        self.assertGreaterEqual(inches(first.height) + 1e-6, 2 * sh._line_height(body) + sh._line_height(small))
        self.assertGreaterEqual(second.top, first.top + first.height)
        mark = shapes_named(slide, "Claim mark")[0]
        self.assertAlmostEqual(mark.top, first.top + Inches(sh.CLAIM_MARK_OFFSET), delta=EMU_TOLERANCE)

    def test_the_claims_keep_their_vertical_rhythm(self):
        # Each claim box starts a fixed gap below the box above it. Claims of one line each step by
        # 1.4533 in, the step of the earlier layout with a Claim shape and a Baseline shape.
        prs = sh.new_presentation(LIGHT)
        slide = build_take_home(prs, LIGHT, self.tmp)
        boxes = shapes_named(slide, "Claim")
        top = LIGHT["layout"]["content_top"] - sh.FREE_FORM_REFERENCE_TOP + sh.CLAIM_TOP
        self.assertEqual(boxes[0].top, Inches(top))
        for upper, lower in zip(boxes, boxes[1:]):
            gap = inches(lower.top - upper.top - upper.height)
            self.assertAlmostEqual(gap, sh.CLAIM_BLOCK_GAP, places=4)
            self.assertAlmostEqual(inches(lower.top - upper.top), 1.4533, places=3)

    def test_one_line_claims_do_not_overlap(self):
        prs = sh.new_presentation(LIGHT)
        slide = build_take_home(prs, LIGHT, self.tmp)
        self.assert_no_overlap(slide)
        marks = shapes_named(slide, "Claim mark")
        self.assertEqual([m.top for m in marks], sorted(m.top for m in marks))

    def test_a_long_claim_gets_a_taller_box_and_pushes_the_next_claim_down(self):
        long_claim = "A claim of fifteen words runs over a single line of the slide and needs a second line here"
        prs = sh.new_presentation(LIGHT)
        short = sh.add_take_home_slide(prs, LIGHT, [("Short claim", "base"), ("Next claim", "base")])
        tall = sh.add_take_home_slide(prs, LIGHT, [(long_claim, "base"), ("Next claim", "base")])
        self.assertGreater(shapes_named(tall, "Claim")[0].height, shapes_named(short, "Claim")[0].height)
        self.assertGreater(shapes_named(tall, "Claim")[1].top, shapes_named(short, "Claim")[1].top)
        self.assert_no_overlap(tall)
        self.assertEqual(shapes_named(tall, "Claim mark")[1].top - shapes_named(tall, "Claim")[1].top,
                         shapes_named(short, "Claim mark")[1].top - shapes_named(short, "Claim")[1].top)

    def test_claims_that_need_more_room_than_the_slide_has_raise_and_add_no_slide(self):
        very_long = " ".join(["word"] * 30)
        prs = sh.new_presentation(LIGHT)
        with self.assertRaises(ValueError) as caught:
            sh.add_take_home_slide(prs, LIGHT, [(very_long, very_long)] * 3)
        self.assertIn("room", str(caught.exception))
        self.assertEqual(len(prs.slides), 0)


def chart_of(slide):
    return shapes_named(slide, "Evidence chart")[0].chart


def count_names(slide):
    names = {}
    for shape in slide.shapes:
        names[shape.name] = names.get(shape.name, 0) + 1
    return names


class AssertionEvidenceSlide(TmpDirCase):
    def setUp(self):
        super().setUp()
        self.prs = sh.new_presentation(LIGHT)

    def fails(self, evidence=None, headline=BAR_HEADLINE, text=None, **kwargs):
        """Assert that the call raises ValueError, adds no slide, and the message holds `text`."""
        count = len(self.prs.slides)
        with self.assertRaises(ValueError) as caught:
            sh.add_assertion_evidence_slide(self.prs, LIGHT, headline, evidence if evidence is not None
                                            else evidence_bar(), **kwargs)
        self.assertEqual(len(self.prs.slides), count)
        if text:
            self.assertIn(text, str(caught.exception))

    def test_the_shape_names_of_a_slide_with_a_picture_and_two_callouts(self):
        slide = build_picture(self.prs, LIGHT, self.tmp)
        self.assertEqual(count_names(slide), {"Title": 1, "Evidence picture": 1, "Callout rule": 2, "Callout": 2,
                                              "Source": 1})

    def test_backup_true_gives_the_backup_kind_and_the_default_gives_assertion_evidence(self):
        self.assertEqual(sh.slide_kind(build_table(self.prs, LIGHT, self.tmp)), "assertion_evidence")
        self.assertEqual(sh.slide_kind(build_backup(self.prs, LIGHT, self.tmp)), "backup")

    def test_the_slide_is_returned_and_added_last(self):
        slide = build_bar(self.prs, LIGHT, self.tmp)
        self.assertIs(self.prs.slides[len(self.prs.slides) - 1]._element, slide._element)

    # ---- evidence box ----------------------------------------------------------------------
    def test_every_kind_of_evidence_starts_at_content_top_and_stays_in_the_evidence_box(self):
        box_left = Inches(LIGHT["layout"]["margin"])
        box_top = Inches(LIGHT["layout"]["content_top"])
        box_w = Inches(LIGHT["layout"]["evidence_width"])
        box_h = Inches(LIGHT["layout"]["evidence_height"])
        for builder, name in ((build_picture, "Evidence picture"), (build_table, "Evidence table"),
                              (build_bar, "Evidence chart"), (build_line, "Evidence chart")):
            slide = builder(sh.new_presentation(LIGHT), LIGHT, self.tmp)
            shape = shapes_named(slide, name)[0]
            self.assertGreaterEqual(shape.left, box_left - EMU_TOLERANCE, name)
            self.assertLessEqual(shape.left + shape.width, box_left + box_w + EMU_TOLERANCE, name)
            self.assertGreaterEqual(shape.top, box_top - EMU_TOLERANCE, name)
            self.assertLessEqual(shape.top + shape.height, box_top + box_h + EMU_TOLERANCE, name)
            if name != "Evidence picture":
                self.assertAlmostEqual(shape.top, box_top, delta=EMU_TOLERANCE)
                self.assertAlmostEqual(shape.width, box_w, delta=EMU_TOLERANCE)

    # ---- picture ---------------------------------------------------------------------------
    def test_a_picture_gets_its_alt_text_and_never_the_file_name(self):
        slide = build_picture(self.prs, LIGHT, self.tmp)
        picture = shapes_named(slide, "Evidence picture")[0]
        self.assertEqual(picture._element.nvPicPr.cNvPr.get("descr"), PICTURE_ALT)
        self.assertNotIn("figure.png", PICTURE_ALT)

    def test_the_descr_of_a_picture_is_always_the_given_alt_text(self):
        evidence = evidence_picture(self.tmp, name="never-use-me.png", alt="A red box on a white page.")
        slide = evidence_slide(self.prs, LIGHT, evidence)
        descr = shapes_named(slide, "Evidence picture")[0]._element.nvPicPr.cNvPr.get("descr")
        self.assertEqual(descr, "A red box on a white page.")

    def test_a_picture_or_chart_without_usable_alt_text_raises_and_adds_no_slide(self):
        for label, build in (("picture", lambda: evidence_picture(self.tmp)), ("chart", evidence_bar)):
            for bad in ("missing", None, "", "   ", 7):
                evidence = build()
                if bad == "missing":
                    del evidence["alt"]
                else:
                    evidence["alt"] = bad
                with self.subTest(kind=label, alt=bad):
                    self.fails(evidence, text="{} alt text".format(label))

    def test_picture_errors_name_the_problem_and_add_no_slide(self):
        folder = self.tmp / "a-folder"
        folder.mkdir()
        fake = self.tmp / "fake.png"
        fake.write_text("not an image")
        other_format = self.tmp / "other.ppm"
        Image.new("RGB", (20, 10)).save(str(other_format))
        cases = [
            ({"kind": "picture", "alt": "x"}, "no path"),
            ({"kind": "picture", "path": None, "alt": "x"}, "no path"),
            ({"kind": "picture", "path": "", "alt": "x"}, "empty"),
            ({"kind": "picture", "path": str(folder), "alt": "x"}, "folder"),
            ({"kind": "picture", "path": str(self.tmp / "gone.png"), "alt": "x"}, "does not exist"),
            ({"kind": "picture", "path": str(fake), "alt": "x"}, "cannot read the image"),
            ({"kind": "picture", "path": str(other_format), "alt": "x"}, "has the format"),
        ]
        for evidence, text in cases:
            with self.subTest(text=text):
                self.fails(evidence, text=text)

    # ---- errors in the evidence and the callouts ---------------------------------------------
    def test_an_unknown_evidence_kind_raises_and_lists_the_valid_kinds(self):
        self.fails({"kind": "video"}, text="video")
        self.fails({"kind": "video"}, text="picture, table, chart")

    def test_evidence_that_is_not_a_dict_with_a_kind_raises(self):
        self.fails("a picture", text="kind")
        self.fails({"path": "x.png"}, text="kind")

    def test_three_callouts_raise(self):
        self.fails(evidence_bar(), callouts=["a", "b", "c"], text="at most 2 callouts")

    def test_a_callout_that_is_not_a_string_raises(self):
        self.fails(evidence_bar(), callouts=["a", 7], text="callout")
        self.fails(evidence_bar(), callouts=[""], text="callout")

    def test_a_blank_headline_or_a_bad_source_raises(self):
        self.fails(evidence_bar(), headline="", text="headline")
        self.fails(evidence_bar(), source=3, text="source")
        self.fails(evidence_bar(), notes=3, text="notes")

    # ---- callouts and source -----------------------------------------------------------------
    def test_a_callout_is_a_plain_text_box_with_an_accent_rule_at_its_left(self):
        for style in (LIGHT, DARK):
            slide = build_picture(sh.new_presentation(style), style, self.tmp)
            rules, texts = shapes_named(slide, "Callout rule"), shapes_named(slide, "Callout")
            self.assertEqual(len(rules), 2)
            for rule, text in zip(rules, texts):
                self.assertEqual(rule.width, Inches(style["layout"]["rule_width"]))
                self.assertEqual(rule.left, Inches(style["layout"]["callout_left"]))
                self.assertEqual(rule.fill.fore_color.rgb, rgb(style["colors"]["accent"]))
                self.assertEqual(rule.line.fill.type, MSO_FILL.BACKGROUND)
                self.assertGreater(text.left, rule.left + rule.width)
                self.assertEqual(text.top, rule.top)
                self.assertEqual(text.fill.type, MSO_FILL.BACKGROUND)
                run = text.text_frame.paragraphs[0].runs[0]
                self.assertEqual(run.font.size, size_pt(style, "callout_size"))
                self.assertEqual(run.font.color.rgb, rgb(style["colors"]["text"]))
            self.assertLess(texts[0].top, texts[1].top)
            self.assertLessEqual(texts[0].top + texts[0].height, texts[1].top)
            self.assertEqual([t.text_frame.text for t in texts], TWO_CALLOUTS)

    def test_no_callouts_means_no_callout_shapes_and_one_string_means_one_callout(self):
        bare = evidence_slide(self.prs, LIGHT, evidence_bar())
        self.assertEqual(shapes_named(bare, "Callout") + shapes_named(bare, "Callout rule"), [])
        one = evidence_slide(self.prs, LIGHT, evidence_bar(), callouts="A single callout")
        self.assertEqual([s.text_frame.text for s in shapes_named(one, "Callout")], ["A single callout"])

    def test_the_source_line_uses_the_source_size_and_the_muted_color(self):
        for style in (LIGHT, DARK):
            slide = build_bar(sh.new_presentation(style), style, self.tmp)
            source = shapes_named(slide, "Source")[0]
            run = source.text_frame.paragraphs[0].runs[0]
            self.assertEqual((run.font.size, run.font.color.rgb),
                             (size_pt(style, "source_size"), rgb(style["colors"]["muted"])))
            self.assertGreaterEqual(source.top, Inches(style["layout"]["content_bottom"]))
        bare = evidence_slide(self.prs, LIGHT, evidence_bar())
        self.assertEqual(shapes_named(bare, "Source"), [])


class TableEvidence(TmpDirCase):
    def build(self, style=LIGHT, **kwargs):
        prs = sh.new_presentation(style)
        slide = evidence_slide(prs, style, evidence_table(**kwargs))
        return slide, shapes_named(slide, "Evidence table")[0]

    def cells(self, shape):
        return [[cell for cell in row.cells] for row in shape.table.rows]

    def test_the_table_style_and_the_header_and_band_flags(self):
        _, shape = self.build()
        tbl_pr = shape.table._tbl.tblPr
        self.assertEqual(tbl_pr.find(qn("a:tableStyleId")).text, TABLE_STYLE_ID)
        self.assertEqual(tbl_pr.get("firstRow"), "0")
        self.assertEqual(tbl_pr.get("bandRow"), "0")

    def test_the_rows_and_columns_match_the_data_and_the_widths_fill_the_evidence_width(self):
        _, shape = self.build()
        self.assertEqual((len(shape.table.rows), len(shape.table.columns)), (4, 2))
        self.assertEqual(sum(column.width for column in shape.table.columns), lay(LIGHT, "evidence_width"))
        self.assertEqual(shape.width, lay(LIGHT, "evidence_width"))
        self.assertEqual([[c.text_frame.text for c in row] for row in self.cells(shape)], TABLE_ROWS)

    def test_the_column_widths_add_up_to_the_table_width_when_the_division_leaves_a_remainder(self):
        # Text lengths 6, 6 and 7 split 7040880 EMU into three parts that sum to 2 EMU less.
        prs = sh.new_presentation(LIGHT)
        slide = evidence_slide(prs, LIGHT, evidence_table(rows=[["Abcdef", "Ghijkl", "Mnopqrs"]], highlight=None))
        shape = shapes_named(slide, "Evidence table")[0]
        self.assertEqual(sum(column.width for column in shape.table.columns), lay(LIGHT, "evidence_width"))

    def test_a_column_with_longer_text_gets_a_wider_column(self):
        _, shape = self.build()
        first, second = shape.table.columns[0].width, shape.table.columns[1].width
        self.assertGreater(second, first)

    def test_every_cell_run_uses_at_least_the_callout_size_and_is_explicit(self):
        for style in (LIGHT, DARK):
            _, shape = self.build(style)
            for row in self.cells(shape):
                for cell in row:
                    for run in cell.text_frame.paragraphs[0].runs:
                        self.assertGreaterEqual(run.font.size, Pt(style["typography"]["callout_size"]))
                        self.assertIs(run.font.bold, False)
                        self.assertEqual(run.font.name, style["typography"]["body_font"])

    def test_the_header_row_has_an_accent_rule_under_it_as_the_first_child_before_the_fill(self):
        for style in (LIGHT, DARK):
            _, shape = self.build(style)
            for cell in self.cells(shape)[0]:
                tc_pr = cell._tc.tcPr
                self.assertEqual(tc_pr[0].tag, qn("a:lnB"))
                self.assertEqual(tc_pr[0].get("w"), str(int(Pt(1.5))))
                self.assertEqual(tc_pr[0].find(qn("a:solidFill")).find(qn("a:srgbClr")).get("val"),
                                 style["colors"]["accent"].lstrip("#").upper())
                self.assertEqual(tc_pr[1].tag, qn("a:noFill"))

    def test_the_body_rows_have_no_border(self):
        _, shape = self.build()
        for row in self.cells(shape)[1:]:
            for cell in row:
                self.assertIsNone(cell._tc.tcPr.find(qn("a:lnB")))

    def test_the_highlighted_row_has_accent_text_and_no_fill(self):
        for style in (LIGHT, DARK):
            _, shape = self.build(style, highlight=(2, 1))
            accent_ink, text = rgb(style["colors"]["accent_ink"]), rgb(style["colors"]["text"])
            rows = self.cells(shape)
            for cell in rows[2]:
                self.assertEqual(cell.text_frame.paragraphs[0].runs[0].font.color.rgb, accent_ink)
            for index in (1, 3):
                for cell in rows[index]:
                    self.assertEqual(cell.text_frame.paragraphs[0].runs[0].font.color.rgb, text)
            for row in rows:
                for cell in row:
                    self.assertEqual(cell.fill.type, MSO_FILL.BACKGROUND)

    def test_the_header_text_is_accent_colored_and_no_highlight_leaves_the_body_plain(self):
        _, shape = self.build(highlight=None)
        rows = self.cells(shape)
        accent_ink, text = rgb(LIGHT["colors"]["accent_ink"]), rgb(LIGHT["colors"]["text"])
        self.assertEqual({c.text_frame.paragraphs[0].runs[0].font.color.rgb for c in rows[0]}, {accent_ink})
        self.assertEqual({c.text_frame.paragraphs[0].runs[0].font.color.rgb for row in rows[1:] for c in row},
                         {text})

    def test_numbers_are_written_as_text(self):
        prs = sh.new_presentation(LIGHT)
        slide = evidence_slide(prs, LIGHT, evidence_table(rows=[["Step", "Share"], ["One", 0.5], ["Two", 3]],
                                                          highlight=None))
        shape = shapes_named(slide, "Evidence table")[0]
        self.assertEqual([c.text_frame.text for c in self.cells(shape)[1]], ["One", "0.5"])
        self.assertEqual([c.text_frame.text for c in self.cells(shape)[2]], ["Two", "3"])

    def test_a_table_needs_no_alt_text_and_gets_none(self):
        prs = sh.new_presentation(LIGHT)
        for alt in (None, "Words that a screen reader would read."):
            evidence = evidence_table()
            if alt:
                evidence["alt"] = alt
            shape = shapes_named(evidence_slide(prs, LIGHT, evidence), "Evidence table")[0]
            self.assertIsNone(shape._element.nvGraphicFramePr.cNvPr.get("descr"))

    def test_a_table_that_does_not_fit_or_is_malformed_raises_and_adds_no_slide(self):
        prs = sh.new_presentation(LIGHT)
        many = [["Row", "Value"]] + [["r{}".format(i), "v"] for i in range(11)]
        bad = [
            (many, (2, 1), "rows"),
            ([], None, "rows"),
            ([["A", "B"], ["x"]], None, "same number"),
            ([["A", "B"], "xy"], None, "row"),
            ([[]], None, "column"),
            (TABLE_ROWS, (9, 0), "highlight"),
            (TABLE_ROWS, (1, 5), "highlight"),
            (TABLE_ROWS, "row 2", "highlight"),
        ]
        for rows, highlight, text in bad:
            with self.subTest(text=text, highlight=highlight):
                with self.assertRaises(ValueError) as caught:
                    evidence_slide(prs, LIGHT, evidence_table(highlight=highlight, rows=rows))
                self.assertIn(text, str(caught.exception))
        self.assertEqual(len(prs.slides), 0)

    def test_a_table_of_nine_rows_fits(self):
        rows = [["Row", "Value"]] + [["r{}".format(i), "v"] for i in range(8)]
        slide, shape = self.build(rows=rows, highlight=None)
        self.assertLessEqual(shape.top + shape.height,
                             lay(LIGHT, "content_top") + lay(LIGHT, "evidence_height") + EMU_TOLERANCE)


class ChartEvidence(TmpDirCase):
    def bar(self, style=LIGHT, **extra):
        prs = sh.new_presentation(style)
        slide = evidence_slide(prs, style, evidence_bar(**extra))
        return slide, chart_of(slide)

    def line(self, style=LIGHT, **extra):
        prs = sh.new_presentation(style)
        slide = evidence_slide(prs, style, evidence_line(**extra))
        return slide, chart_of(slide)

    def point_colors(self, chart):
        series = chart.plots[0].series[0]
        return [series.points[i].format.fill.fore_color.rgb for i in range(len(list(chart.plots[0].categories)))]

    # ---- bar chart with one series -------------------------------------------------------------
    def test_a_bar_chart_is_a_clustered_column_chart_without_legend_title_axis_or_gridlines(self):
        from pptx.enum.chart import XL_CHART_TYPE
        _, chart = self.bar()
        self.assertEqual(chart.chart_type, XL_CHART_TYPE.COLUMN_CLUSTERED)
        self.assertFalse(chart.has_legend)
        self.assertFalse(chart.has_title)
        self.assertFalse(chart.value_axis.visible)
        self.assertFalse(chart.value_axis.has_major_gridlines)
        self.assertFalse(chart.category_axis.has_major_gridlines)
        self.assertTrue(chart.category_axis.visible)

    def test_the_highlighted_bar_is_accent_and_the_rest_are_context(self):
        for style in (LIGHT, DARK):
            accent, context = rgb(style["colors"]["accent"]), rgb(style["colors"]["context"])
            for highlight in ("Group B", 1):
                _, chart = self.bar(style, highlight=highlight)
                self.assertEqual(self.point_colors(chart), [context, accent, context])

    def test_no_highlight_colors_every_bar_with_the_accent(self):
        _, chart = self.bar(highlight=None)
        self.assertEqual(set(self.point_colors(chart)), {rgb(LIGHT["colors"]["accent"])})

    def test_a_bar_chart_shows_data_labels_with_the_caller_format(self):
        _, chart = self.bar()
        plot = chart.plots[0]
        self.assertTrue(plot.has_data_labels)
        self.assertEqual(plot.data_labels.number_format, "General")
        self.assertFalse(plot.data_labels.number_format_is_linked)
        self.assertTrue(plot.data_labels.show_value)
        _, formatted = self.bar(fmt='0"%"')
        self.assertEqual(formatted.plots[0].data_labels.number_format, '0"%"')

    def test_vmax_sets_the_top_of_the_hidden_value_axis_and_bars_start_at_zero(self):
        _, chart = self.bar(vmax=42)
        self.assertEqual(chart.value_axis.maximum_scale, 42)
        self.assertEqual(chart.value_axis.minimum_scale, 0)
        _, auto = self.bar()
        self.assertIsNone(auto.value_axis.maximum_scale)

    def test_bars_with_a_negative_value_leave_the_axis_minimum_to_the_application(self):
        _, chart = self.bar(series={"Change": [-3, 5, 2]}, highlight=None)
        self.assertIsNone(chart.value_axis.minimum_scale)

    def test_chart_text_uses_the_style_font_sizes_and_colors(self):
        for style in (LIGHT, DARK):
            _, chart = self.bar(style)
            text = rgb(style["colors"]["text"])
            self.assertEqual(chart.font.size, size_pt(style, "callout_size"))
            self.assertEqual(chart.category_axis.tick_labels.font.size, size_pt(style, "callout_size"))
            self.assertEqual(chart.plots[0].data_labels.font.size, size_pt(style, "body_size"))
            for font in (chart.font, chart.category_axis.tick_labels.font, chart.plots[0].data_labels.font):
                self.assertEqual(font.name, style["typography"]["body_font"])
                self.assertEqual(font.color.rgb, text)
                self.assertIs(font.bold, False)

    def test_every_run_property_in_the_chart_xml_is_explicit(self):
        for builder in (self.bar, self.line):
            _, chart = builder()
            nodes = chart._chartSpace.xpath(".//a:defRPr | .//a:rPr")
            self.assertGreater(len(nodes), 2)
            for node in nodes:
                self.assertIsNotNone(node.get("sz"))
                self.assertEqual(node.get("b"), "0")
                self.assertIsNotNone(node.find(qn("a:latin")))
                self.assertIsNotNone(node.find(qn("a:solidFill")).find(qn("a:srgbClr")))

    def assert_transparent(self, sp):
        self.assertEqual(sp.tag, qn("c:spPr"))
        self.assertEqual([child.tag for child in sp], [qn("a:noFill"), qn("a:ln")])
        self.assertEqual([child.tag for child in sp[1]], [qn("a:noFill")])

    def test_the_chart_area_and_the_plot_area_are_transparent_in_both_places(self):
        for builder in (self.bar, self.line):
            for style in (LIGHT, DARK):
                _, chart = builder(style)
                space = chart._chartSpace
                children = list(space)
                after_chart = children[[c.tag for c in children].index(qn("c:chart")) + 1]
                self.assert_transparent(after_chart)
                plot_area = space.find(qn("c:chart")).find(qn("c:plotArea"))
                self.assert_transparent(plot_area[len(plot_area) - 1])
                self.assertEqual(len(space.xpath("c:spPr")), 1)
                self.assertEqual(len(plot_area.xpath("c:spPr")), 1)

    def test_the_chart_xml_keeps_the_schema_order(self):
        _, chart = self.line()
        tags = [child.tag for child in chart._chartSpace]
        self.assertLess(tags.index(qn("c:chart")), tags.index(qn("c:spPr")))
        self.assertLess(tags.index(qn("c:spPr")), tags.index(qn("c:txPr")))
        plot_area = chart._chartSpace.find(qn("c:chart")).find(qn("c:plotArea"))
        self.assertEqual(plot_area[0].tag, qn("c:layout"))
        self.assertEqual(len(plot_area.findall(qn("c:layout"))), 1)

    def test_the_chart_gets_the_given_alt_text(self):
        slide, _ = self.bar(alt="Bars: group B holds the most items.")
        frame = shapes_named(slide, "Evidence chart")[0]
        self.assertEqual(frame._element.nvGraphicFramePr.cNvPr.get("descr"), "Bars: group B holds the most items.")

    # ---- line chart with one series --------------------------------------------------------------
    def test_a_one_series_line_chart_marks_the_key_point_in_the_accent_color(self):
        _, chart = self.line(series={"Response": [0.2, 0.4, 0.9, 0.5]}, highlight="Day 4")
        series = chart.plots[0].series[0]
        accent, context = rgb(LIGHT["colors"]["accent"]), rgb(LIGHT["colors"]["context"])
        fills = [series.points[i].marker.format.fill.fore_color.rgb for i in range(4)]
        self.assertEqual(fills, [context, context, accent, context])
        self.assertGreater(series.points[2].marker.size, series.points[0].marker.size)
        self.assertEqual(series.format.line.color.rgb, context)
        self.assertEqual(series.format.line.width, Pt(sh.CONTEXT_LINE_PT))
        self.assertTrue(chart.plots[0].has_data_labels)
        self.assertTrue(chart.plots[0].data_labels.show_value)
        self.assertFalse(chart.has_legend)

    def test_a_one_series_line_chart_without_highlight_is_all_accent(self):
        _, chart = self.line(series={"Response": [0.2, 0.4, 0.9, 0.5]}, highlight=None)
        series = chart.plots[0].series[0]
        self.assertEqual(series.format.line.color.rgb, rgb(LIGHT["colors"]["accent"]))
        self.assertEqual({series.points[i].marker.format.fill.fore_color.rgb for i in range(4)},
                         {rgb(LIGHT["colors"]["accent"])})

    # ---- several series --------------------------------------------------------------------------
    def test_a_line_chart_with_several_series_has_no_legend_and_a_thick_accent_series(self):
        _, chart = self.line()
        self.assertFalse(chart.has_legend)
        by_name = {series.name: series for series in chart.plots[0].series}
        accent, context = rgb(LIGHT["colors"]["accent"]), rgb(LIGHT["colors"]["context"])
        self.assertEqual(by_name["High dose"].format.line.color.rgb, accent)
        for name in ("Control", "Low dose"):
            self.assertEqual(by_name[name].format.line.color.rgb, context)
            self.assertGreater(by_name["High dose"].format.line.width, by_name[name].format.line.width)

    def test_each_series_is_labeled_at_its_last_point_with_its_name_and_nowhere_else(self):
        _, chart = self.line()
        plot = chart.plots[0]
        self.assertTrue(plot.has_data_labels)
        self.assertFalse(plot.data_labels.show_value)
        for series in plot.series:
            labels = series._element.xpath("c:dLbls/c:dLbl")
            self.assertEqual(len(labels), 1)
            self.assertEqual(labels[0].xpath("c:idx")[0].get("val"), "3")
            self.assertEqual(series.points[3].data_label.text_frame.text, series.name)
            self.assertEqual(series._element.xpath("c:dLbls/c:showVal")[0].get("val"), "0")
            run = series.points[3].data_label.text_frame.paragraphs[0].runs[0]
            self.assertEqual(run.font.size, size_pt(LIGHT, "callout_size"))
            self.assertIs(run.font.bold, False)

    def test_the_highlighted_series_label_is_accent_ink_and_the_others_use_the_text_color(self):
        _, chart = self.line(DARK)
        for series in chart.plots[0].series:
            run = series.points[3].data_label.text_frame.paragraphs[0].runs[0]
            expected = DARK["colors"]["accent_ink"] if series.name == "High dose" else DARK["colors"]["text"]
            self.assertEqual(run.font.color.rgb, rgb(expected))

    def test_a_line_chart_with_several_series_leaves_room_at_the_right_for_the_labels(self):
        _, chart = self.line()
        layout = chart._chartSpace.find(qn("c:chart")).find(qn("c:plotArea")).find(qn("c:layout"))
        values = {node.tag.split("}")[1]: node.get("val") for node in layout.find(qn("c:manualLayout"))}
        self.assertEqual(values["layoutTarget"], "inner")
        self.assertLess(float(values["x"]) + float(values["w"]), 0.9)
        self.assertGreater(float(values["w"]), 0.4)
        self.assertLessEqual(float(values["y"]) + float(values["h"]), 0.92)

    def test_a_bar_chart_with_several_series_colors_the_named_series(self):
        prs = sh.new_presentation(LIGHT)
        evidence = evidence_bar(highlight="After", series={"Before": [3, 4, 5], "After": [5, 7, 9]})
        chart = chart_of(evidence_slide(prs, LIGHT, evidence))
        by_name = {series.name: series for series in chart.plots[0].series}
        self.assertEqual(by_name["After"].format.fill.fore_color.rgb, rgb(LIGHT["colors"]["accent"]))
        self.assertEqual(by_name["Before"].format.fill.fore_color.rgb, rgb(LIGHT["colors"]["context"]))
        self.assertFalse(chart.has_legend)
        self.assertEqual(by_name["After"].points[2].data_label.text_frame.text, "After")
        self.assertEqual(by_name["Before"].points[2].data_label.text_frame.text, "Before")

    def test_several_series_without_a_highlight_use_the_data_palette_colors_that_show_on_the_background(self):
        import design_color
        for style in (LIGHT, DARK):
            _, chart = self.line(style, highlight=None)
            colors = [series.format.line.color.rgb for series in chart.plots[0].series]
            self.assertEqual(len(set(colors)), 3)
            palette = {rgb(c) for c in style["colors"]["data_palette"]}
            for color in colors:
                self.assertIn(color, palette)
                ratio = design_color.contrast_ratio("#" + str(color), style["colors"]["background"])
                self.assertGreaterEqual(ratio, 3.0)

    # ---- errors ----------------------------------------------------------------------------------
    def fails(self, text, kind="line", **extra):
        prs = sh.new_presentation(LIGHT)
        evidence = evidence_line(**extra) if kind == "line" else evidence_bar(**extra)
        with self.assertRaises(ValueError) as caught:
            evidence_slide(prs, LIGHT, evidence)
        self.assertIn(text, str(caught.exception))
        self.assertEqual(len(prs.slides), 0)

    def test_a_series_with_the_wrong_length_raises_and_names_the_series(self):
        self.fails("'Low dose'", series={"Control": [1, 2, 3, 4], "Low dose": [1, 2, 3]}, highlight="Control")
        self.fails("4 categories", series={"Control": [1, 2, 3, 4], "Low dose": [1, 2, 3]}, highlight="Control")

    def test_an_unknown_chart_type_raises_and_lists_the_valid_types(self):
        self.fails("pie", chart="pie")
        self.fails("bar, line", chart="pie")

    def test_missing_or_empty_data_raises(self):
        self.fails("categories", categories=[])
        self.fails("series", series={})
        self.fails("series", series=[1, 2, 3, 4])
        self.fails("number", series={"Control": [1, 2, "x", 4]}, highlight=None)
        self.fails("number", series={"Control": [1, 2, True, 4]}, highlight=None)

    def test_a_highlight_that_matches_nothing_raises_and_lists_the_choices(self):
        self.fails("Nothing", highlight="Nothing")
        self.fails("High dose", highlight="Nothing")
        self.fails("series name", highlight=2)
        self.fails("highlight", kind="bar", highlight="Missing")
        self.fails("highlight", kind="bar", highlight=3)
        self.fails("highlight", kind="bar", highlight=-1)

    def test_a_palette_with_no_usable_color_raises_before_any_slide_is_added(self):
        style = copy.deepcopy(LIGHT)
        style["colors"]["data_palette"] = ["#FFFF99", "#FFFFCC"]
        prs = sh.new_presentation(style)
        with self.assertRaises(ValueError) as caught:
            evidence_slide(prs, style, evidence_line(highlight=None))
        self.assertIn("data_palette", str(caught.exception))
        self.assertEqual(len(prs.slides), 0)

    def test_fmt_for_a_chart_with_several_series_raises_and_one_series_accepts_it(self):
        self.fails("fmt", fmt="0.0")
        self.fails("fmt", fmt="General")
        self.fails("fmt", kind="bar", series={"Before": [3, 4, 5], "After": [5, 7, 9]}, highlight="After", fmt="0")
        _, chart = self.bar(fmt="0.0")
        self.assertEqual(chart.plots[0].data_labels.number_format, "0.0")

    def test_a_vmax_below_the_largest_value_raises(self):
        self.fails("vmax", kind="bar", vmax=20)

    def test_a_bar_chart_accepts_numeric_categories_as_text(self):
        prs = sh.new_presentation(LIGHT)
        slide = evidence_slide(prs, LIGHT, evidence_bar(categories=[2019, 2020, 2021], highlight="2020"))
        self.assertEqual(list(chart_of(slide).plots[0].categories), ["2019", "2020", "2021"])


# ---------------------------------------------------------------------------
# Input that a call must reject without leaving a half-built slide
# ---------------------------------------------------------------------------

class FailedCallsLeaveNothing(TmpDirCase):
    def assert_fails(self, call, text, prs=None):
        """Assert that call(prs) raises ValueError that holds `text` and adds no slide."""
        prs = prs or sh.new_presentation(LIGHT)
        before = len(prs.slides)
        with self.assertRaises(ValueError) as caught:
            call(prs)
        self.assertIn(text, str(caught.exception))
        self.assertEqual(len(prs.slides), before)
        return str(caught.exception)


BAD = "x\ud800"  # a lone surrogate: Python holds it, but no file format can save it
WORD = "x" * 40


class LoneSurrogate(FailedCallsLeaveNothing):
    def test_each_text_argument_of_each_layout_raises_and_adds_no_slide(self):
        title, big, road, home, evid = sh.add_title_slide, sh.add_big_number_slide, sh.add_roadmap_slide, \
            sh.add_take_home_slide, sh.add_assertion_evidence_slide
        cases = [
            ("title", lambda p: title(p, LIGHT, BAD)),
            ("author", lambda p: title(p, LIGHT, "T", author=BAD)),
            ("affiliation", lambda p: title(p, LIGHT, "T", affiliation=BAD)),
            ("date", lambda p: title(p, LIGHT, "T", date=BAD)),
            ("notes", lambda p: title(p, LIGHT, "T", notes=BAD)),
            ("headline", lambda p: big(p, LIGHT, BAD, "42", "items", "base")),
            ("number", lambda p: big(p, LIGHT, HEADLINE, BAD, "items", "base")),
            ("label", lambda p: big(p, LIGHT, HEADLINE, "42", BAD, "base")),
            ("baseline", lambda p: big(p, LIGHT, HEADLINE, "42", "items", BAD)),
            ("source", lambda p: big(p, LIGHT, HEADLINE, "42", "items", "base", source=BAD)),
            ("notes", lambda p: big(p, LIGHT, HEADLINE, "42", "items", "base", notes=BAD)),
            ("message", lambda p: road(p, LIGHT, BAD, ROADMAP_SECTIONS)),
            ("section name", lambda p: road(p, LIGHT, HEADLINE, [(BAD, "line")] + ROADMAP_SECTIONS[1:])),
            ("section line", lambda p: road(p, LIGHT, HEADLINE, [("Aim", BAD)] + ROADMAP_SECTIONS[1:])),
            ("notes", lambda p: road(p, LIGHT, HEADLINE, ROADMAP_SECTIONS, notes=BAD)),
            ("claim", lambda p: home(p, LIGHT, [(BAD, "base")])),
            ("baseline", lambda p: home(p, LIGHT, [("claim", BAD)])),
            ("headline", lambda p: home(p, LIGHT, CLAIMS, headline=BAD)),
            ("notes", lambda p: home(p, LIGHT, CLAIMS, notes=BAD)),
            ("headline", lambda p: evid(p, LIGHT, BAD, evidence_bar())),
            ("callout", lambda p: evid(p, LIGHT, BAR_HEADLINE, evidence_bar(), callouts=[BAD])),
            ("source", lambda p: evid(p, LIGHT, BAR_HEADLINE, evidence_bar(), source=BAD)),
            ("notes", lambda p: evid(p, LIGHT, BAR_HEADLINE, evidence_bar(), notes=BAD)),
            ("table cell", lambda p: evid(p, LIGHT, BAR_HEADLINE, evidence_table(rows=[["A", "B"], [BAD, "y"]]))),
            ("category", lambda p: evid(p, LIGHT, BAR_HEADLINE, evidence_bar(categories=["A", BAD, "C"]))),
            ("series name", lambda p: evid(p, LIGHT, BAR_HEADLINE, evidence_bar(series={BAD: [1, 2, 3]}))),
            ("number format", lambda p: evid(p, LIGHT, BAR_HEADLINE, evidence_bar(fmt=BAD))),
            ("chart alt text", lambda p: evid(p, LIGHT, BAR_HEADLINE, evidence_bar(alt=BAD))),
            ("picture alt text", lambda p: evid(p, LIGHT, BAR_HEADLINE, evidence_picture(self.tmp, alt=BAD))),
        ]
        for text, call in cases:
            with self.subTest(item=text):
                message = self.assert_fails(call, text)
                self.assertIn("lone surrogate", message)

    def test_the_reviewer_examples_raise_value_error_and_not_unicode_error(self):
        self.assert_fails(lambda p: sh.add_title_slide(p, LIGHT, "bad \ud800"), "title")
        self.assert_fails(lambda p: sh.add_big_number_slide(p, LIGHT, HEADLINE, "42", "items", "base", notes=BAD),
                          "notes")
        self.assert_fails(lambda p: sh.add_assertion_evidence_slide(p, LIGHT, BAR_HEADLINE, evidence_bar(),
                                                                    callouts=[BAD]), "callout")

    def test_set_notes_and_set_alt_text_reject_a_lone_surrogate_and_change_nothing(self):
        _, slide = blank_slide()
        with self.assertRaises(ValueError):
            sh.set_notes(slide, BAD)
        self.assertFalse(slide.has_notes_slide)
        box = slide.shapes.add_textbox(0, 0, 100, 100)
        with self.assertRaises(ValueError) as caught:
            sh.set_alt_text(box, BAD)
        self.assertIn("alt text", str(caught.exception))
        self.assertIn("lone surrogate", str(caught.exception))
        self.assertIsNone(box._element.xpath("./*[1]/p:cNvPr")[0].get("descr"))

    def test_a_path_is_not_checked_for_surrogates(self):
        # A file name that is not valid UTF-8 reaches Python as a string with a surrogate.
        # The path check must keep working for it.
        _, slide = blank_slide()
        with self.assertRaises(ValueError) as caught:
            sh.fit_picture(slide, str(self.tmp / "missing-\udcff.png"), 1, 1, 4, 3)
        self.assertIn("does not exist", str(caught.exception))


class ControlCharacters(FailedCallsLeaveNothing):
    """XML 1.0 forbids most control characters. lxml raises on them only after the helper has added the slide."""

    CHARACTERS = ("\x00", "\x08", "\x0b", "\x0c", "\x1f", "\ufffe", "\uffff")

    def cases(self, bad):
        evid = sh.add_assertion_evidence_slide
        return [
            ("category", lambda p: evid(p, LIGHT, BAR_HEADLINE, evidence_bar(highlight=None, categories=["A", bad, "C"]))),
            ("series name", lambda p: evid(p, LIGHT, BAR_HEADLINE, evidence_bar(highlight=None, series={bad: [1, 2, 3]}))),
            ("series name", lambda p: evid(p, LIGHT, LINE_HEADLINE, evidence_line(
                highlight=None, series={"A": [1, 2, 3, 4], bad: [2, 3, 4, 5]}))),
            ("chart alt text", lambda p: evid(p, LIGHT, BAR_HEADLINE, evidence_bar(alt=bad))),
            ("picture alt text", lambda p: evid(p, LIGHT, BAR_HEADLINE, evidence_picture(self.tmp, alt=bad))),
            ("table cell", lambda p: evid(p, LIGHT, BAR_HEADLINE, evidence_table(highlight=None, rows=[["A", "B"], [bad, "y"]]))),
        ]

    def test_a_control_character_in_each_path_raises_before_the_slide_and_names_the_character(self):
        for character in self.CHARACTERS:
            for text, call in self.cases("x" + character):
                with self.subTest(character=repr(character), item=text):
                    message = self.assert_fails(call, text)
                    self.assertIn("U+{:04X}".format(ord(character)), message)

    def test_the_deck_saves_after_a_rejected_call(self):
        prs = sh.new_presentation(LIGHT)
        for text, call in self.cases("x\x00"):
            with self.assertRaises(ValueError):
                call(prs)
        evidence_slide(prs, LIGHT, evidence_bar())
        prs.save(str(self.tmp / "after.pptx"))
        self.assertEqual(len(Presentation(str(self.tmp / "after.pptx")).slides), 1)

    def test_tab_newline_and_carriage_return_stay_allowed(self):
        prs = sh.new_presentation(LIGHT)
        evidence_slide(prs, LIGHT, evidence_table(highlight=None, rows=[["A", "B"], ["x\ty", "z\r\nw"]]))
        evidence_slide(prs, LIGHT, evidence_bar(alt="Bars of\tthree groups.\n"))
        prs.save(str(self.tmp / "allowed.pptx"))
        self.assertEqual(len(prs.slides), 2)

    def test_set_alt_text_rejects_a_control_character_and_changes_nothing(self):
        _, slide = blank_slide()
        box = slide.shapes.add_textbox(0, 0, 100, 100)
        with self.assertRaises(ValueError) as caught:
            sh.set_alt_text(box, "x\x0b")
        self.assertIn("U+000B", str(caught.exception))
        self.assertIsNone(box._element.xpath("./*[1]/p:cNvPr")[0].get("descr"))


class TextThatCannotFit(FailedCallsLeaveNothing):
    """Text that needs more lines than its box holds raises ValueError and adds no slide."""

    @staticmethod
    def words(count, length):
        return " ".join(["x" * length] * count)

    def headline_calls(self):
        return {
            "assertion-evidence": lambda text: (lambda p: sh.add_assertion_evidence_slide(
                p, LIGHT, text, evidence_bar())),
            "big number": lambda text: (lambda p: sh.add_big_number_slide(p, LIGHT, text, "42", "items", "base")),
            "roadmap": lambda text: (lambda p: sh.add_roadmap_slide(p, LIGHT, text, ROADMAP_SECTIONS)),
            "take-home": lambda text: (lambda p: sh.add_take_home_slide(p, LIGHT, CLAIMS, headline=text)),
        }

    def test_a_headline_of_three_lines_raises_and_one_of_two_lines_passes(self):
        for name, make in self.headline_calls().items():
            with self.subTest(layout=name):
                make(self.words(2, 40))(sh.new_presentation(LIGHT))  # two lines at the title size
                message = self.assert_fails(make(self.words(3, 40)), "headline")
                self.assertIn("Shorten", message)

    def test_the_title_of_the_title_slide_may_take_two_lines_at_44_pt_and_not_three(self):
        sh.add_title_slide(sh.new_presentation(LIGHT), LIGHT, self.words(2, 30))
        # Three words of 25 letters take two lines at 28 pt and three lines at 44 pt.
        message = self.assert_fails(lambda p: sh.add_title_slide(p, LIGHT, self.words(3, 25)), "title")
        self.assertIn("Shorten", message)

    def test_a_callout_may_take_two_lines_and_not_three(self):
        # Two lines is the maximum of the lines-per-paragraph rule of the linter.
        prs = sh.new_presentation(LIGHT)
        sh.add_assertion_evidence_slide(prs, LIGHT, BAR_HEADLINE, evidence_bar(), callouts=[self.words(2, 20)])
        message = self.assert_fails(lambda p: sh.add_assertion_evidence_slide(
            p, LIGHT, BAR_HEADLINE, evidence_bar(), callouts=["ok", self.words(3, 20)]), "callout", prs=prs)
        self.assertIn("Shorten", message)
        self.assertIn("2", message)

    def test_words_that_fill_a_line_exactly_stay_on_that_line(self):
        # A callout line holds 23 characters, so two words of 11 letters fill it exactly.
        prs = sh.new_presentation(LIGHT)
        slide = sh.add_assertion_evidence_slide(prs, LIGHT, BAR_HEADLINE, evidence_bar(), callouts=[self.words(4, 11)])
        self.assertEqual(len(shapes_named(slide, "Callout")), 1)

    def test_the_limits_follow_the_style(self):
        text = " ".join(["alpha beta gamma delta"] * 5)  # two lines at 28 pt, three lines at 44 pt
        sh.add_big_number_slide(sh.new_presentation(LIGHT), LIGHT, text, "42", "items", "base")
        big = copy.deepcopy(LIGHT)
        big["typography"]["title_size"] = 44
        self.assert_fails(lambda p: sh.add_big_number_slide(p, big, text, "42", "items", "base"), "headline",
                          prs=sh.new_presentation(big))
        wide = copy.deepcopy(LIGHT)
        wide["typography"]["callout_size"] = 30
        callout = self.words(4, 10)  # two lines at 20 pt, four lines at 30 pt
        sh.add_assertion_evidence_slide(sh.new_presentation(LIGHT), LIGHT, BAR_HEADLINE, evidence_bar(),
                                        callouts=[callout])
        self.assert_fails(lambda p: sh.add_assertion_evidence_slide(p, wide, BAR_HEADLINE, evidence_bar(),
                                                                    callouts=[callout]),
                          "callout", prs=sh.new_presentation(wide))

    def test_the_default_take_home_headline_fits_for_every_claim_count(self):
        for count in (1, 2, 3):
            sh.add_take_home_slide(sh.new_presentation(LIGHT), LIGHT, CLAIMS[:count])

    def test_a_long_table_cell_is_not_checked(self):
        rows = [["Group", "Note"], ["Group A", " ".join(["word"] * 60)]]
        slide = evidence_slide(sh.new_presentation(LIGHT), LIGHT, evidence_table(rows=rows, highlight=None))
        self.assertEqual(len(shapes_named(slide, "Evidence table")), 1)


class WhitespaceNotes(TmpDirCase):
    def test_notes_with_only_whitespace_count_as_empty(self):
        for text in ("   ", "\n", " \t\r\n "):
            _, slide = blank_slide()
            sh.set_notes(slide, text)
            self.assertFalse(slide.has_notes_slide, repr(text))
            prs = sh.new_presentation(LIGHT)
            slides = [
                sh.add_title_slide(prs, LIGHT, "A title", notes=text),
                sh.add_big_number_slide(prs, LIGHT, HEADLINE, "42", "items", "base", notes=text),
                sh.add_roadmap_slide(prs, LIGHT, HEADLINE, ROADMAP_SECTIONS, notes=text),
                sh.add_take_home_slide(prs, LIGHT, CLAIMS, notes=text),
                evidence_slide(prs, LIGHT, evidence_bar(), notes=text),
            ]
            for slide in slides:
                self.assertFalse(slide.has_notes_slide, repr(text))

    def test_a_source_with_only_whitespace_adds_no_source_shape(self):
        prs = sh.new_presentation(LIGHT)
        slide = sh.add_big_number_slide(prs, LIGHT, HEADLINE, "42", "items", "base", source="  \n ")
        self.assertEqual(shapes_named(slide, "Source"), [])
        self.assertEqual(shapes_named(evidence_slide(prs, LIGHT, evidence_bar(), source=" "), "Source"), [])

    def test_notes_with_text_keep_their_surrounding_whitespace(self):
        _, slide = blank_slide()
        sh.set_notes(slide, "  Say this.\n")
        self.assertEqual(slide.notes_slide.notes_text_frame.text, "  Say this.\n")


class EdgeCases(TmpDirCase):
    def test_a_table_with_only_a_header_row_builds(self):
        prs = sh.new_presentation(LIGHT)
        slide = evidence_slide(prs, LIGHT, evidence_table(rows=[["Group", "Items"]], highlight=None))
        shape = shapes_named(slide, "Evidence table")[0]
        self.assertEqual(len(shape.table.rows), 1)
        for cell in shape.table.rows[0].cells:
            self.assertEqual(cell._tc.tcPr[0].tag, qn("a:lnB"))

    def test_a_table_with_one_column_builds_and_the_column_fills_the_width(self):
        prs = sh.new_presentation(LIGHT)
        slide = evidence_slide(prs, LIGHT, evidence_table(rows=[["Group"], ["Group A"]], highlight=(1, 0)))
        shape = shapes_named(slide, "Evidence table")[0]
        self.assertEqual(shape.table.columns[0].width, lay(LIGHT, "evidence_width"))

    def test_charts_with_one_category_build(self):
        specs = [
            ("bar", {"Items": [3]}, "Only"),
            ("line", {"Items": [3]}, "Only"),
            ("bar", {"Before": [3], "After": [5]}, "After"),
            ("line", {"Before": [3], "After": [5]}, "After"),
        ]
        for kind, series, highlight in specs:
            with self.subTest(kind=kind, series=len(series)):
                prs = sh.new_presentation(LIGHT)
                evidence = {"kind": "chart", "chart": kind, "categories": ["Only"], "series": series,
                            "highlight": highlight, "alt": CHART_ALT}
                chart = chart_of(evidence_slide(prs, LIGHT, evidence))
                self.assertEqual(list(chart.plots[0].categories), ["Only"])
                if len(series) > 1:
                    for one in chart.plots[0].series:
                        self.assertEqual(one.points[0].data_label.text_frame.text, one.name)

    def test_non_ascii_text_survives_in_the_title_notes_callout_and_table_cell(self):
        text = "µ × α 中"  # micro sign, multiplication sign, alpha, a CJK character
        prs = sh.new_presentation(LIGHT)
        title_slide = sh.add_title_slide(prs, LIGHT, "Title " + text, author="Author " + text,
                                         notes="Notes " + text)
        rows = [["Name " + text, "B"], ["Cell " + text, "x"]]
        content = evidence_slide(prs, LIGHT, evidence_table(rows=rows, highlight=(1, 0)), "Headline " + text,
                                 callouts=["Callout " + text], source="Source " + text, notes="Notes " + text)
        out = self.tmp / "unicode.pptx"
        prs.save(str(out))
        first, second = Presentation(str(out)).slides
        self.assertEqual(first.shapes.title.text_frame.text, "Title " + text)
        self.assertEqual(shapes_named(first, "Subtitle")[0].text_frame.text, "Author " + text)
        self.assertEqual(first.notes_slide.notes_text_frame.text, "Notes " + text)
        self.assertEqual(second.shapes.title.text_frame.text, "Headline " + text)
        self.assertEqual(shapes_named(second, "Callout")[0].text_frame.text, "Callout " + text)
        self.assertEqual(shapes_named(second, "Source")[0].text_frame.text, "Source " + text)
        self.assertEqual(second.notes_slide.notes_text_frame.text, "Notes " + text)
        table = shapes_named(second, "Evidence table")[0].table
        self.assertEqual([c.text_frame.text for row in table.rows for c in row.cells],
                         ["Name " + text, "B", "Cell " + text, "x"])
