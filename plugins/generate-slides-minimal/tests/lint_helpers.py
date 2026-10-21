"""Helpers for the tests of the design linter: decks built with python-pptx, and a test case base class.
"""
import re
import sys
import tempfile
import unittest
from pathlib import Path

from PIL import Image
from pptx import Presentation
from pptx.chart.data import CategoryChartData
from pptx.dml.color import RGBColor
from pptx.enum.chart import XL_CHART_TYPE
from pptx.enum.shapes import MSO_SHAPE
from pptx.oxml import parse_xml
from pptx.oxml.ns import nsdecls, qn
from pptx.util import Inches, Pt

SCRIPTS_DIR = Path(__file__).resolve().parents[1] / "scripts"
if str(SCRIPTS_DIR) not in sys.path:
    sys.path.insert(0, str(SCRIPTS_DIR))

import design_lint as dl  # noqa: E402
import slide_helpers as sh  # noqa: E402
import style as style_mod  # noqa: E402


SCRIPT = SCRIPTS_DIR / "design_lint.py"


DARK = style_mod.load_style("peers", "dark")


LIGHT = style_mod.load_style("peers", "light")


GOOD_TITLE = "Binder design succeeds for 12 of 15 targets"  # 8 words, the verb is "succeeds"


BODY = "Twelve of fifteen designed binders fold and bind the target"  # 10 words


ROW = re.compile(r"^(slide \d+|deck): [a-z-]+: (error|warn): T[123]: \S.*$")


def words(count, word="word"):
    return " ".join([word] * count)


def new_deck(width=13.333, height=7.5):
    prs = Presentation()
    prs.slide_width, prs.slide_height = Inches(width), Inches(height)
    return prs


def set_run(run, size=24, color="111111", family="Avenir Light", bold=False):
    """Give a run explicit font values. The size goes in as hundredths of a point, so 17.9 stays 17.9."""
    run.font.name = family
    run.font.bold = bold
    run.font.color.rgb = RGBColor.from_string(color)
    run._r.get_or_add_rPr().set("sz", str(int(round(size * 100))))


def write(shape, text, **font):
    """Add one run of text to the first paragraph of a shape and return the run."""
    run = shape.text_frame.paragraphs[0].add_run()
    run.text = text
    set_run(run, **font)
    return run


def add_slide(prs, title=GOOD_TITLE, kind=None, background="FFFFFF", **title_font):
    """Add a Title Only slide. A title of None leaves the title placeholder empty."""
    slide = prs.slides.add_slide(prs.slide_layouts.get_by_name("Title Only"))
    if background is not None:
        slide.background.fill.solid()
        slide.background.fill.fore_color.rgb = RGBColor.from_string(background)
    if kind is not None:
        slide._element.cSld.set("name", kind)
    box = slide.shapes.title
    box.name = "Title"
    box.left, box.top, box.width, box.height = Inches(0.8), Inches(0.3), Inches(11.7), Inches(0.95)
    if title is not None:
        font = {"size": 28}
        font.update(title_font)
        write(box, title, **font)
    return slide


def add_box(slide, text, name="Body", left=0.8, top=1.6, width=8.0, height=1.2, **font):
    """Add a text box with no insets and one run of text."""
    box = slide.shapes.add_textbox(Inches(left), Inches(top), Inches(width), Inches(height))
    box.name = name
    frame = box.text_frame
    frame.word_wrap = True
    frame.margin_left = frame.margin_right = frame.margin_top = frame.margin_bottom = 0
    write(box, text, **font)
    return box


def add_claim(slide, claim, baseline, top=1.8, height=1.2, baseline_font=None, **claim_font):
    """Add a Claim shape of the take-home slide: the claim in paragraph 0 and its baseline in paragraph 1.

    The baseline is 20 pt in the muted grey unless `baseline_font` says otherwise.
    """
    box = add_box(slide, claim, name="Claim", left=1.25, top=top, width=11.28, height=height, **claim_font)
    font = {"size": 20, "color": "6B7280"}
    font.update(baseline_font or {})
    run = box.text_frame.add_paragraph().add_run()
    run.text = baseline
    set_run(run, **font)
    return box


def add_source(slide, text="Source: made up data", size=12, color="6B7280", name="Source"):
    return add_box(slide, text, name=name, left=0.8, top=6.95, width=11.7, height=0.3, size=size, color=color)


def add_rect(slide, name, left, top, width, height, color="6D4BD6"):
    shape = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(left), Inches(top), Inches(width), Inches(height))
    shape.name = name
    shape.fill.solid()
    shape.fill.fore_color.rgb = RGBColor.from_string(color)
    shape.line.fill.background()
    return shape


def add_band(slide, where="layout", rgb="0B1F3A", theme=None, left=0.0, top=0.0, width=13.333, height=1.4,
             groups=0, gradient=False):
    """Add a filled rectangle to the layout or the master of a slide. It sits below every shape of the slide.

    `groups` wraps the rectangle in that many nested groups. `gradient` gives it a gradient fill.
    Return the outermost shape.
    """
    band = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(left), Inches(top), Inches(width), Inches(height))
    band.name = "Band"
    band.fill.solid()
    if theme is not None:
        band.fill.fore_color.theme_color = theme
    else:
        band.fill.fore_color.rgb = RGBColor.from_string(rgb)
    if gradient:
        band.fill.gradient()
    band.line.fill.background()
    outer = band
    for _ in range(groups):
        outer = slide.shapes.add_group_shape([outer])
    element = outer._element
    element.getparent().remove(element)
    owner = slide.slide_layout if where == "layout" else slide.slide_layout.slide_master
    owner.shapes._spTree.append(element)
    return outer


def add_group_of_boxes(slide, count, where="master"):
    """Move a group of `count` small rectangles to the master or the layout of the slide. They lie in the top corner."""
    boxes = [slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0.05 + (i % 50) * 0.04), Inches(0.05 + (i // 50) * 0.01),
                                    Inches(0.03), Inches(0.03)) for i in range(count)]
    group = slide.shapes.add_group_shape(boxes)
    element = group._element
    element.getparent().remove(element)
    owner = slide.slide_layout if where == "layout" else slide.slide_layout.slide_master
    owner.shapes._spTree.append(element)
    return group


def add_bullets(slide, count, name="List", top=1.6, bullets=True, size=24):
    box = slide.shapes.add_textbox(Inches(0.8), Inches(top), Inches(8), Inches(2))
    box.name = name
    frame = box.text_frame
    frame.word_wrap = True
    for index in range(count):
        paragraph = frame.paragraphs[0] if index == 0 else frame.add_paragraph()
        if bullets:
            paragraph._p.get_or_add_pPr().append(parse_xml('<a:buChar %s char="*"/>' % nsdecls("a")))
        run = paragraph.add_run()
        run.text = "Point %d of the list" % (index + 1)
        set_run(run, size=size)
    return box


def add_table(slide, size=20, rows=2, cols=2, text="cell text", color="111111", name="Evidence table"):
    frame = slide.shapes.add_table(rows, cols, Inches(0.8), Inches(1.6), Inches(7.7), Inches(0.75 * rows))
    frame.name = name
    for row in range(rows):
        for col in range(cols):
            cell = frame.table.cell(row, col)
            cell.fill.background()
            run = cell.text_frame.paragraphs[0].add_run()
            run.text = text
            set_run(run, size=size, color=color)
    return frame


def add_picture(slide, path, left=1.0, top=1.6, width=5.0, height=3.0, alt=None, name="Evidence picture"):
    shape = slide.shapes.add_picture(str(path), Inches(left), Inches(top), Inches(width), Inches(height))
    shape.name = name
    if alt is not None:
        shape._element.xpath("./*[1]/p:cNvPr")[0].set("descr", alt)
    return shape


def set_alt(shape, text):
    shape._element.xpath("./*[1]/p:cNvPr")[0].set("descr", text)


def make_png(folder, name, width, height):
    path = Path(str(folder)) / name
    Image.new("RGB", (width, height), (30, 60, 200)).save(str(path))
    return path


def add_chart(slide, colors=("6D4BD6",), labels=False, line=False, size=20, text_color="111111",
              plot_fill=None, chart_fill=None, alt="A chart of made up values", name="Evidence chart",
              outline=None):
    """Add a chart with explicit series colors. `plot_fill` and `chart_fill` are hex colors or "none".

    `outline` is the hex color of a 1 pt line around each bar of a bar chart.
    """
    data = CategoryChartData()
    data.categories = ["a", "b", "c"]
    for index in range(len(colors)):
        data.add_series("s%d" % index, (1 + index, 2 + index, 3 + index))
    kind = XL_CHART_TYPE.LINE if line else XL_CHART_TYPE.COLUMN_CLUSTERED
    frame = slide.shapes.add_chart(kind, Inches(0.8), Inches(1.6), Inches(7.7), Inches(4.5), data)
    frame.name = name
    set_alt(frame, alt)
    chart = frame.chart
    chart.has_legend = False
    chart.font.size = Pt(size)
    chart.font.color.rgb = RGBColor.from_string(text_color)
    plot = chart.plots[0]
    for series, color in zip(plot.series, colors):
        if line:
            if color is not None:
                series.format.line.color.rgb = RGBColor.from_string(color)
            continue
        if color is not None:
            series.format.fill.solid()
            series.format.fill.fore_color.rgb = RGBColor.from_string(color)
        if outline is not None:
            series.format.line.color.rgb = RGBColor.from_string(outline)
            series.format.line.width = Pt(1)
    if labels:
        plot.has_data_labels = True
    space = chart._chartSpace
    for owner, fill in ((space.find(qn("c:chart")).find(qn("c:plotArea")), plot_fill), (space, chart_fill)):
        if fill is None:
            continue
        inner = "<a:noFill/>" if fill == "none" else '<a:solidFill><a:srgbClr val="%s"/></a:solidFill>' % fill
        properties = parse_xml("<c:spPr %s>%s</c:spPr>" % (nsdecls("c", "a"), inner))
        if owner is space:
            space.find(qn("c:chart")).addnext(properties)
        else:
            owner.append(properties)
    return frame


def add_notes(slide, text):
    slide.notes_slide.notes_text_frame.text = text


def clean_slide(prs, kind=None, title=GOOD_TITLE):
    """Add a slide that no rule flags: a sentence title and one 24 pt body box."""
    slide = add_slide(prs, title, kind=kind)
    add_box(slide, BODY)
    return slide


class LintCase(unittest.TestCase):
    """A test case with a temporary folder, and helpers that save a deck and lint it."""

    def setUp(self):
        self._tmp = tempfile.TemporaryDirectory()
        self.addCleanup(self._tmp.cleanup)
        self.tmp = Path(self._tmp.name)

    def save(self, prs, name="deck.pptx"):
        path = self.tmp / name
        prs.save(str(path))
        return path

    def lint(self, prs, **options):
        options.setdefault("notes", False)
        return dl.lint(self.save(prs), **options)

    def png(self, name, width, height):
        return make_png(self.tmp, name, width, height)

    def of(self, findings, rule, slide=None):
        return [f for f in findings if f.rule == rule and (slide is None or f.slide == slide)]

    def assertFires(self, findings, rule, severity, tier, slide=1, count=1):
        hits = self.of(findings, rule, slide)
        self.assertEqual(len(hits), count, "expected {} finding(s) of {} on slide {}, got {}".format(
            count, rule, slide, findings))
        for hit in hits:
            self.assertEqual((hit.severity, hit.tier), (severity, tier), hit)
        return hits[0]

    def assertQuiet(self, findings, rule):
        self.assertEqual(self.of(findings, rule), [], "{} should stay quiet".format(rule))

    def only(self, findings, rule):
        """Return the one finding of a rule."""
        hits = self.of(findings, rule)
        self.assertEqual(len(hits), 1, "expected one {} finding, got {}".format(rule, findings))
        return hits[0]


def deck_of(kinds):
    prs = new_deck()
    for index, kind in enumerate(kinds):
        clean_slide(prs, kind=kind, title="Binder design succeeds for %d of 15 targets" % index)
    return prs


def build_all_layouts(style, folder):
    """Build one deck with all six layouts and in-budget content. Return the Presentation."""
    prs = sh.new_presentation(style)
    photo = make_png(folder, "photo.png", 2400, 1350)
    sh.add_title_slide(prs, style, "A made up talk about binder design", "Ann Lee", "Made up lab", "2026-10-20",
                       notes="Open with the question that the audience came with.")
    sh.add_assertion_evidence_slide(
        prs, style, "Binder design succeeds for 12 of 15 targets",
        {"kind": "chart", "chart": "bar", "categories": ["A", "B", "C"], "series": {"Hits": [4, 12, 7]},
         "highlight": "B", "alt": "Bar chart of made up hit counts for three targets"},
        callouts=["Twelve of fifteen targets bind"], source="Source: made up data",
        notes="Walk along the bars from left to right and pause on the second bar.")
    sh.add_assertion_evidence_slide(
        prs, style, "Designed binders outperform the control after round two",
        {"kind": "chart", "chart": "line", "categories": ["R1", "R2", "R3"],
         "series": {"Design": [1, 4, 9], "Control": [1, 1, 2]}, "highlight": "Design",
         "alt": "Lines of made up hit counts for design and control over three rounds"},
        source="Source: made up data", notes="Point at the gap that opens in round two.")
    sh.add_assertion_evidence_slide(
        prs, style, "Hit rate varies across the three targets",
        {"kind": "table", "rows": [["Target", "Hits", "Rate"], ["A", "4", "27%"], ["B", "12", "80%"],
                                    ["C", "7", "47%"]], "highlight": (2, 0)},
        callouts=["Target B leads"], source="Source: made up data", notes="Read the highlighted row aloud.")
    sh.add_assertion_evidence_slide(
        prs, style, "The gel shows one clear band for each design",
        {"kind": "picture", "path": str(photo), "alt": "Photo of a made up gel with one band per lane"},
        source="Source: made up data", notes="Say which lane holds the ladder.")
    sh.add_big_number_slide(
        prs, style, "One in three designs binds the target", "33%", "of designs bind",
        "Controls bind in fewer than one in twenty runs", source="Source: made up data",
        notes="Explain how a hit is counted.")
    sh.add_roadmap_slide(
        prs, style, "Four steps take us from design to assay",
        [("Design", "Pick targets"), ("Build", "Make the genes"), ("Test", "Run the assay"),
         ("Learn", "Read the hits")], notes="Name each step in one short sentence.")
    sh.add_take_home_slide(
        prs, style,
        [("Binder design works for most targets", "Twelve of fifteen targets bind"),
         ("Controls stay silent", "Zero of fifteen controls bind"),
         ("The assay scales to a plate", "Ninety six designs fit one plate")],
        notes="Repeat the three claims and stop.")
    sh.add_assertion_evidence_slide(
        prs, style, "The assay controls bind nothing in fifteen runs",
        {"kind": "table", "rows": [["Run", "Hits"], ["1", "0"], ["2", "0"]]}, source="Source: made up data",
        backup=True)
    return prs


def run_props(run, xml):
    """Replace the a:rPr of a run with an element parsed from `xml`."""
    old = run._r.find(qn("a:rPr"))
    if old is not None:
        run._r.remove(old)
    run._r.insert(0, parse_xml(xml))


def damage_no_nv_sp_pr(case, slide):
    box = add_box(slide, "no nvSpPr words", top=2.0)
    box._element.remove(box._element.find(qn("p:nvSpPr")))


def damage_no_name(case, slide):
    box = add_box(slide, "no name words", top=2.0)
    box._element.find(qn("p:nvSpPr")).find(qn("p:cNvPr")).attrib.pop("name")


def damage_no_sp_pr(case, slide):
    box = add_box(slide, "no spPr words", top=2.0)
    box._element.remove(box._element.find(qn("p:spPr")))


def damage_transform_without_x(case, slide):
    box = add_box(slide, "no x words", top=2.0)
    box._element.find(qn("p:spPr")).find(qn("a:xfrm")).find(qn("a:off")).attrib.pop("x")


def damage_no_blip_fill(case, slide):
    picture = add_picture(slide, case.png("p.png", 20, 10), top=3.0, width=2.0, height=1.0, alt="A picture")
    picture._element.remove(picture._element.find(qn("p:blipFill")))


def damage_bad_crop(case, slide):
    picture = add_picture(slide, case.png("q.png", 20, 10), top=3.0, width=2.0, height=1.0, alt="A picture")
    picture._element.find(qn("p:blipFill")).append(parse_xml('<a:srcRect %s l="wide"/>' % nsdecls("a")))


def damage_non_numeric_size(case, slide):
    box = add_box(slide, "bad size words", top=2.0)
    run_props(box.text_frame.paragraphs[0].runs[0], '<a:rPr %s sz="big" b="maybe"/>' % nsdecls("a"))


def damage_run_without_text(case, slide):
    box = add_box(slide, "no text element words", top=2.0)
    run = box.text_frame.paragraphs[0].runs[0]
    run._r.remove(run._r.find(qn("a:t")))


def damage_ragged_table(case, slide):
    table = slide.shapes.add_table(2, 3, Inches(1), Inches(3), Inches(4), Inches(1))
    table.table.cell(0, 0).text = "header words"
    row = table._element.findall(".//" + qn("a:tr"))[1]
    row.remove(row.findall(qn("a:tc"))[-1])


def damage_no_graphic(case, slide):
    table = slide.shapes.add_table(1, 1, Inches(1), Inches(3), Inches(2), Inches(1))
    table._element.remove(table._element.find(qn("a:graphic")))


def damage_no_table_element(case, slide):
    table = slide.shapes.add_table(1, 1, Inches(1), Inches(3), Inches(2), Inches(1))
    graphic_data = table._element.find(qn("a:graphic")).find(qn("a:graphicData"))
    graphic_data.remove(graphic_data.find(qn("a:tbl")))


def damage_blip_without_embed(case, slide):
    picture = add_picture(slide, case.png("r.png", 20, 10), top=3.0, width=2.0, height=1.0, alt="A picture")
    picture._element.find(qn("p:blipFill")).find(qn("a:blip")).attrib.pop(qn("r:embed"))


def damage_chart_without_plot_area(case, slide):
    frame = add_chart(slide)
    chart = frame.chart._chartSpace.find(qn("c:chart"))
    chart.remove(chart.find(qn("c:plotArea")))


def damage_chart_with_bad_size(case, slide):
    frame = add_chart(slide, labels=True)
    for props in frame.chart._chartSpace.iter(qn("a:defRPr")):
        props.set("sz", "huge")


def damage_chart_with_scheme_colors(case, slide):
    frame = add_chart(slide, labels=True)
    for srgb in list(frame.chart._chartSpace.iter(qn("a:srgbClr"))):
        srgb.tag = qn("a:schemeClr")
        srgb.set("val", "accent1")


def damage_group_without_transform(case, slide):
    one = add_box(slide, "one", name="One", top=2.0)
    group = slide.shapes.add_group_shape([one])
    properties = group._element.find(qn("p:grpSpPr"))
    properties.remove(properties.find(qn("a:xfrm")))


def damage_invalid_placeholder_type(case, slide):
    slide.shapes.title._element.find(qn("p:nvSpPr")).find(qn("p:nvPr")).find(qn("p:ph")).set("type", "bogus")


DAMAGES = [
    ("a shape with no nvSpPr", damage_no_nv_sp_pr),
    ("a cNvPr with no name", damage_no_name),
    ("a shape with no spPr", damage_no_sp_pr),
    ("a transform with no x", damage_transform_without_x),
    ("a picture with no blipFill", damage_no_blip_fill),
    ("a picture with a bad crop", damage_bad_crop),
    ("a run with a non-numeric sz", damage_non_numeric_size),
    ("a run with no text element", damage_run_without_text),
    ("a table with a ragged row", damage_ragged_table),
    ("a graphic frame with no a:graphic", damage_no_graphic),
    ("a table frame with no a:tbl", damage_no_table_element),
    ("a picture whose blip has no embed", damage_blip_without_embed),
    ("a chart with no plot area", damage_chart_without_plot_area),
    ("a chart with a bad font size", damage_chart_with_bad_size),
    ("a chart with scheme colors", damage_chart_with_scheme_colors),
    ("a group with no transform", damage_group_without_transform),
    ("a placeholder with an invalid type", damage_invalid_placeholder_type),
]
