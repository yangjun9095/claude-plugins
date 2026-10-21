"""Tests for the documents of generate-slides-minimal: the command, the principles, the READMEs and the sample deck.

Each checker is a function that takes text, so a test can feed it a seeded defect and see it fail. A check
that never fails on a seeded defect is a defect of the test.

Source keys of the principles file. A key reads `report:key`, for example `craft:S6`. It resolves when the
section "Pages opened" of the named research note holds the key, one or more spaces or tabs, and a URL on the
same line. The reports write the key in three places:
  jobs.md                 at the start of the line:  `S1 https://...`
  science.md, measurable.md  after a list dash:      `- S1 https://...`, `- WCAG-CM https://...`
  craft.md                at the start of the line, and also inside the sentence that lists the pages
                          that the report read through a fetch summary: `... only: S4 https://... ; S5 https://...`
One pattern covers all three: the key, a space or tab, then a URL, with no letter, digit or dash before the key.
"""
import ast
import fnmatch
import inspect
import io
import random
import re
import subprocess
import sys
import tempfile
import unittest
import zipfile
from collections import OrderedDict
from pathlib import Path

import yaml
from PIL import Image
from pptx import Presentation

PLUGIN = Path(__file__).resolve().parents[1]
REPO = PLUGIN.parent.parent
SCRIPTS_DIR = PLUGIN / "scripts"
if str(SCRIPTS_DIR) not in sys.path:
    sys.path.insert(0, str(SCRIPTS_DIR))

import design_lint  # noqa: E402
import slide_helpers  # noqa: E402
import style as style_mod  # noqa: E402
import title_check  # noqa: E402

COMMAND = PLUGIN / "commands" / "generate-slides-minimal.md"
PRINCIPLES = PLUGIN / "references" / "design-principles.md"
README = PLUGIN / "README.md"
ROOT_README = REPO / "README.md"
NOTES = PLUGIN / "references" / "research-notes"
SAMPLE_SCRIPT = PLUGIN / "examples" / "generate_sample_deck.py"
SAMPLE_DECK = PLUGIN / "examples" / "sample-deck.pptx"

REPORTS = ("jobs", "craft", "science", "measurable")
KNOWN_TOOLS = ("Read", "Write", "Edit", "Glob", "Grep", "AskUserQuestion", "WebFetch", "WebSearch", "Task",
               "TodoWrite", "NotebookEdit")
BANNED_WORDS = (
    "leverage", "utilize", "robust", "seamless", "seamlessly", "crucial", "comprehensive", "very", "really",
    "truly", "incredibly", "delve", "holistic", "streamline", "ecosystem", "landscape", "unlock", "elevate",
    "empower", "harness", "navigate", "journey", "realm", "tapestry", "testament", "pivotal", "vital", "paradigm",
    "synergy", "myriad", "plethora", "meticulous", "nuanced", "intricate", "boasts", "showcase", "remarkably",
    "extremely", "quite", "dramatically", "cutting-edge", "state-of-the-art", "game-changer",
)
SENTENCE_LIMIT = 25
TIER_LEGEND = OrderedDict([
    ("T1", "A standard or a study backs the threshold."),
    ("T2", "A named source states the number and cites no evidence."),
    ("T3", "Folklore, a number that the author derived, or no source."),
])
EM_DASH = "—"
# Generic detectors of text that must not appear in a public file. The test holds no list of names.
HOME_PREFIXES = ("home", "Users", "hpc", "scratch", "mnt", "data")  # the first folder of an absolute path
ABSOLUTE_PATH = re.compile(r"(?<![\w.~/-])/(?:" + "|".join(HOME_PREFIXES) + r")/|[A-Za-z]:\\+Users\\")
EMAIL = re.compile(r"[\w.+-]+@[\w-]+(?:\.[\w-]+)+")
EMAIL_ALLOWED = ("noreply@anthropic.com",)


def read(path):
    return Path(path).read_text(encoding="utf-8")


# ---------------------------------------------------------------------------
# Checkers
# ---------------------------------------------------------------------------

FENCE = re.compile(r"^```(\w*)[^\n]*\n(.*?)^```[ \t]*$", re.S | re.M)


def split_front_matter(text):
    """Return (front matter as a dict, body). Raise ValueError when the text has no front matter."""
    match = re.match(r"^---\n(.*?)\n---\n", text, re.S)
    if not match:
        raise ValueError("the text has no front matter")
    data = yaml.safe_load(match.group(1))
    if not isinstance(data, dict):
        raise ValueError("the front matter is not a mapping")
    return data, text[match.end():]


def fenced_blocks(text):
    """Return a list of (language, content) for each fenced block."""
    return [(m.group(1), m.group(2)) for m in FENCE.finditer(text)]


def bash_commands(body):
    """Return the commands that the text runs: each line of a fenced bash block and each inline python3 span.

    A line with `&&`, `||`, `;` or `|` counts as several commands. A line that ends in a backslash joins the next.
    """
    commands = []
    for language, block in fenced_blocks(body):
        if language not in ("bash", "sh"):
            continue
        for line in re.sub(r"\\\n\s*", " ", block).splitlines():
            line = line.strip()
            if line and not line.startswith("#"):
                commands.extend(part for part in re.split(r"\s*(?:&&|\|\||;|\|)\s*", line) if part)
    prose = FENCE.sub("", body)
    commands.extend(re.findall(r"`(python3 [^`]+)`", prose))
    return commands


def bash_patterns(allowed):
    return [entry[len("Bash("):-1] for entry in allowed if entry.startswith("Bash(") and entry.endswith(")")]


def uncovered_commands(body, allowed):
    """Return the commands of the text that no Bash entry of `allowed` matches."""
    patterns = bash_patterns(allowed)
    return [c for c in bash_commands(body) if not any(fnmatch.fnmatchcase(c, p) for p in patterns)]


def unused_entries(body, allowed):
    """Return the entries of `allowed` that the text never uses.

    A Bash entry is used when a command of the text matches it. A bare tool is used when the text names it in backticks.
    """
    commands = bash_commands(body)
    unused = []
    for entry in allowed:
        if entry.startswith("Bash("):
            pattern = entry[len("Bash("):-1]
            if not any(fnmatch.fnmatchcase(c, pattern) for c in commands):
                unused.append(entry)
        elif "`{}`".format(entry) not in body:
            unused.append(entry)
    return unused


def unlisted_tools(body, allowed):
    """Return the tools that the text names in backticks and that `allowed` lacks."""
    named = set(re.findall(r"`({})`".format("|".join(KNOWN_TOOLS)), body))
    return sorted(named - set(allowed))


PLUGIN_PATH = re.compile(r"(?<![\w.-])((?:scripts|references|examples)/[\w./-]*[\w/])")


def missing_paths(text, root):
    """Return the paths under scripts/, references/ and examples/ that the text names and that do not exist in root."""
    return sorted({p for p in PLUGIN_PATH.findall(text) if not (Path(root) / p).exists()})


def prose_units(text):
    """Return the units that the prose checks count as sentences.

    The front matter, fenced blocks and the targets of links drop out. An inline code span counts as one word.
    A heading, a list item and a table cell each start a new unit, and a hard-wrapped paragraph joins its lines.
    """
    text = re.sub(r"^---\n.*?\n---\n", "", text, flags=re.S)
    text = FENCE.sub("", text)
    text = re.sub(r"<!--.*?-->", "", text, flags=re.S)
    text = re.sub(r"`[^`\n]*`", "CODE", text)
    text = re.sub(r"\[([^\]]*)\]\([^)]*\)", r"\1", text)
    blocks, current = [], []

    def flush():
        if current:
            blocks.append(" ".join(current))
            del current[:]

    for line in text.splitlines():
        stripped = line.strip()
        if not stripped:
            flush()
        elif re.match(r"^#+\s", stripped):
            flush()
            blocks.append(re.sub(r"^#+\s*", "", stripped))
        elif re.match(r"^([-*+]|\d+[.)])\s+", stripped):
            flush()
            current.append(re.sub(r"^([-*+]|\d+[.)])\s+", "", stripped))
        elif stripped.startswith("|"):
            flush()
            if not re.match(r"^\|[\s:|-]+\|?$", stripped):
                blocks.extend(cell.strip() for cell in stripped.strip("|").split("|"))
        elif current and line[:1] in (" ", "\t"):  # a continuation line of a list item
            current.append(stripped)
        else:
            current.append(stripped)
    flush()
    units = []
    for block in blocks:
        units.extend(s for s in re.split(r"(?<=[.!?])(?:\*\*)?\s+", block) if s.strip())
    return units


def prose_problems(text):
    """Return a list of messages, one for each break of the writing rules."""
    problems = []
    if EM_DASH in text:
        problems.append("em dash")
    prose = " ".join(prose_units(text))
    for word in BANNED_WORDS:
        if re.search(r"(?<![\w-]){}(?![\w-])".format(re.escape(word)), prose, re.I):
            problems.append("banned word: " + word)
    if re.search(r"\bnot just\b", prose, re.I):
        problems.append("not just")
    if re.search(r"\w,\s+(?:not|never)\s", prose, re.I):
        problems.append("'X, not Y' pair")
    for unit in prose_units(text):
        count = len(unit.split())
        if count > SENTENCE_LIMIT:
            problems.append("{} words: {}...".format(count, " ".join(unit.split()[:8])))
    return problems


def template_author():
    """Return the author name that the python-pptx template carries. The test reads it at run time and never spells it."""
    properties = Presentation().core_properties
    return (properties.last_modified_by or properties.author or "").strip()


TERM_DRIFT = re.compile(r"\b(?:the|each|this|that|a|an|these|those|four|any|which|its|their|other)\s+"
                        r"(?:research\s+)?reports?\b", re.I)


def term_drift(text):
    """Return the phrases that call a research note a report. The documents say "research note" everywhere."""
    return sorted({m.group(0).lower() for m in TERM_DRIFT.finditer(" ".join(prose_units(text)))})


def unused_locals(source):
    """Return (function, name) for each local name that a function assigns and never reads. A name that starts with _ is exempt."""
    import ast
    found = []
    for node in ast.walk(ast.parse(source)):
        if isinstance(node, (ast.FunctionDef,)):
            stored = {n.id for n in ast.walk(node) if isinstance(n, ast.Name) and isinstance(n.ctx, ast.Store)}
            loaded = {n.id for n in ast.walk(node) if isinstance(n, ast.Name) and isinstance(n.ctx, ast.Load)}
            found.extend((node.name, name) for name in sorted(stored - loaded) if not name.startswith("_"))
    return found


def private_matches(text):
    """Return a sorted list of descriptions, one for each absolute path, e-mail address or template author name."""
    found = {"path: " + m.group(0) for m in ABSOLUTE_PATH.finditer(text)}
    found |= {"e-mail: " + m.group(0) for m in EMAIL.finditer(text) if m.group(0).lower() not in EMAIL_ALLOWED}
    name = template_author()
    if name and name.lower() in text.lower():
        found.add("the author name of the python-pptx template")
    return sorted(found)


def sections(text):
    """Return an ordered dict of heading to body for each level 2 heading."""
    parts = re.split(r"^## (.+?)\s*$", text, flags=re.M)
    return OrderedDict(zip(parts[1::2], parts[2::2]))


def list_items(body):
    """Return the text of each list item. An item stands on one line."""
    return [line[2:].strip() for line in body.splitlines() if line.startswith("- ")]


TIER_TAG = re.compile(r"`(T[123])`")
KEY_TOKEN = re.compile(r"`([a-z]+):([^`\s]+)`")


def pages_section(text):
    """Return the text of the section 'Pages opened' of a research note, or an empty string."""
    match = re.search(r"^## \d+\. Pages opened[ \t]*$", text, re.M)
    if not match:
        return ""
    rest = text[match.end():]
    following = re.search(r"^## ", rest, re.M)
    return rest[:following.start()] if following else rest


def key_resolves(report, key, notes_dir=NOTES):
    """Return True when the section 'Pages opened' of the report holds the key and a URL on the same line."""
    section = pages_section(read(Path(notes_dir) / (report + ".md")))
    pattern = r"(?<![A-Za-z0-9-])" + re.escape(key) + r"[ \t]+https?://\S+"
    return re.search(pattern, section) is not None


def source_keys(text):
    """Return a list of (report, key) for each backticked token that has the form `name:key`."""
    return KEY_TOKEN.findall(text)


def missing_phrases(text, required):
    """Return a list of (label, phrase) for each phrase of `required` that the text lacks. `required` maps a label to phrases."""
    flat = " ".join(text.split())
    return [(label, phrase) for label, phrases in required.items() for phrase in phrases if phrase not in flat]


def signature_lines(body):
    """Return {function name: [parameter names]} for each call written as a line `name(a, b=1)` in a python block."""
    found = {}
    for language, block in fenced_blocks(body):
        if language != "python":
            continue
        for line in block.splitlines():
            match = re.match(r"^(?:sh\.)?(add_\w+_slide|new_presentation|set_notes|set_alt_text|fit_picture)\((.*)\)$",
                             line.strip())
            if match:
                names = [part.split("=")[0].strip() for part in match.group(2).split(",") if part.strip()]
                found[match.group(1)] = names
    return found


def deck_summary(path):
    """Return the facts of a deck that do not depend on the pictures: kind, title and notes of each slide."""
    prs = Presentation(str(path))
    summary = []
    for slide in prs.slides:
        notes = slide.notes_slide.notes_text_frame.text if slide.has_notes_slide else ""
        summary.append((slide_helpers.slide_kind(slide), slide.shapes.title.text_frame.text, notes))
    return summary


def deck_xml_texts(path):
    """Yield (part name, text) for each XML part of a deck, and for each XML part of a workbook inside it."""
    def walk(zf, prefix):
        for name in zf.namelist():
            data = zf.read(name)
            if name.endswith((".xml", ".rels")):
                yield prefix + name, data.decode("utf-8", "replace")
            elif name.endswith(".xlsx"):
                for item in walk(zipfile.ZipFile(io.BytesIO(data)), prefix + name + "!"):
                    yield item
    with zipfile.ZipFile(str(path)) as zf:
        for item in walk(zf, ""):
            yield item


def step_section(text, name):
    """Return the body of the level 2 section whose heading starts with `name`, or an empty string."""
    for heading, body in sections(text).items():
        if heading.startswith(name):
            return body
    return ""


def sentences_naming(section, tool, word="script"):
    """Return the sentences of a section that name `tool` in backticks and also hold `word`."""
    flat = " ".join(FENCE.sub("", section).split())
    return [s for s in re.split(r"(?<=[.!?])\s+", flat)
            if "`{}`".format(tool) in s and re.search(r"\b{}\b".format(word), s)]


def evidence_examples(body, root):
    """Return the evidence dicts that the python blocks show, one for each line that starts with {"kind":.

    The text writes the path of a picture as ROOT / "name". The function puts the folder `root` in its place.
    """
    found = []
    for language, block in fenced_blocks(body):
        if language != "python":
            continue
        for line in block.splitlines():
            if line.startswith('{"kind":'):
                found.append(ast.literal_eval(line.replace('ROOT / "', '"{}/'.format(root))))
    return found


def evidence_error(evidence, profile="peers"):
    """Return the ValueError message that the slide helper raises for an evidence dict, or None when it builds."""
    style = style_mod.load_style(profile, "light")
    prs = slide_helpers.new_presentation(style)
    try:
        slide_helpers.add_assertion_evidence_slide(prs, style, "A headline.", evidence)
    except ValueError as error:
        return str(error)
    return None


CALLOUT_TABLE = {"kind": "table", "rows": [["A", "B"], ["1", "2"]]}


def callout_texts(length, count=40, seed=7, longest_word=12):
    """Return `count` seeded texts of exactly `length` characters. Each holds words of 1 to `longest_word` letters."""
    rng = random.Random(seed)
    texts = []
    while len(texts) < count:
        words = []
        while len(" ".join(words)) < length:
            words.append("".join(rng.choice("abcdefghij") for _ in range(rng.randint(1, longest_word))))
        if len(" ".join(words)) == length:
            texts.append(" ".join(words))
    return texts


def callout_failures(profile, length):
    """Return the seeded call-out texts of `length` characters that the slide helper refuses under `profile`."""
    style = style_mod.load_style(profile, "light")
    prs = slide_helpers.new_presentation(style)
    failed = []
    for text in callout_texts(length):
        try:
            slide_helpers.add_assertion_evidence_slide(prs, style, "A headline.", CALLOUT_TABLE, callouts=[text])
        except ValueError:
            failed.append(text)
    return failed


def callout_width_chars(profile):
    """Return the characters that one line of a call-out holds under `profile`, by the estimate of the helpers."""
    style = style_mod.load_style(profile, "light")
    prs = slide_helpers.new_presentation(style)
    width_in = slide_helpers._callout_box(prs, style)[2] / float(slide_helpers.EMU_PER_INCH)
    return slide_helpers._chars_per_line(width_in, style["typography"]["callout_size"])


def claimed_callout_lengths(text):
    """Return {profile: characters} that the text states for a call-out, or {} when the sentence is missing."""
    match = re.search(r"Keep a call-out to about (\d+) characters for `peers` and `lab-meeting` and to about (\d+) "
                      r"characters for `general`\.", " ".join(text.split()))
    if not match:
        return {}
    first, second = int(match.group(1)), int(match.group(2))
    return {"peers": first, "lab-meeting": first, "general": second}


def claimed_word_limits(text):
    """Return {profile: (warn, error)} that the text states in one sentence, or {} when the sentence is missing."""
    match = re.search(r"For `peers`, `lab-meeting` and `general`, the linter warns above (\d+), (\d+) and (\d+) words "
                      r"and gives an error above (\d+), (\d+) and (\d+)\.", " ".join(text.split()))
    if not match:
        return {}
    numbers = [int(n) for n in match.groups()]
    return dict(zip(("peers", "lab-meeting", "general"), zip(numbers[:3], numbers[3:])))


SOURCE_RULE = "Add `source` to each slide that has a `source` parameter and shows a number."
TITLE_NO_NUMBER = "The title slide shows no number from the data."  # the title slide carries a date
OLD_SOURCE_RULES = ("first slide that shows the number", "except the take-home slide")
SLIDE_WORDS = {  # the words with which the command text names the slide of each helper
    "add_title_slide": ("title slide",),
    "add_roadmap_slide": ("roadmap slide",),
    "add_take_home_slide": ("take-home slide",),
    "add_big_number_slide": ("big_number", "big-number"),
    "add_assertion_evidence_slide": ("assertion_evidence", "assertion-evidence", "backup"),
}


def helpers_by_source():
    """Return (helpers without a `source` parameter, helpers with one), read from the signatures."""
    names = sorted(SLIDE_WORDS)
    with_source = [n for n in names if "source" in inspect.signature(getattr(slide_helpers, n)).parameters]
    return [n for n in names if n not in with_source], with_source


def source_rule_problems(part, without, with_source):
    """Return a list of messages for a part of the command text that states the source rule wrongly."""
    problems = []
    flat = " ".join(part.split())
    if SOURCE_RULE not in flat:
        problems.append("the source rule sentence is missing")
    if TITLE_NO_NUMBER not in flat:
        problems.append("the title slide sentence is missing")
    problems.extend("old rule: " + old for old in OLD_SOURCE_RULES if old in flat)
    exceptions = " ".join(u for u in prose_units(part) if "shows no number" in u or "show only numbers" in u).lower()
    if not exceptions:
        problems.append("the sentences about the slides without a source line are missing")
    problems.extend("no mention of " + name for name in without if not any(w in exceptions for w in SLIDE_WORDS[name]))
    problems.extend("wrong mention of " + name for name in with_source if any(w in exceptions for w in SLIDE_WORDS[name]))
    return problems


def claimed_message_limit(text):
    """Return the number of words after which the text says that a message gets a `title-sentence` warning."""
    match = re.search(r"A message of more than (\d+) words gets a `title-sentence` warning on the roadmap slide\.",
                      " ".join(text.split()))
    return int(match.group(1)) if match else None


def roadmap_title_findings(words, profile="peers"):
    """Return the `title-sentence` findings of a deck whose roadmap message has `words` words."""
    base = "We show that the drug cuts tumor size".split()
    message = " ".join((base + ["in", "mice"] * words)[:words])
    style = style_mod.load_style(profile, "light")
    prs = slide_helpers.new_presentation(style)
    slide_helpers.add_roadmap_slide(prs, style, message, [("Need", "why"), ("Method", "how"), ("Result", "what")])
    with tempfile.TemporaryDirectory() as folder:
        path = Path(folder) / "roadmap.pptx"
        prs.save(str(path))
        findings = design_lint.lint(str(path), profile=profile, minutes=3)
    return [f for f in findings if f.rule == "title-sentence"]


def run_command_problems(section, allowed, template):
    """Return messages for a run of the script `template` that the section lacks or that no allowed entry covers."""
    problems = []
    command = "python3 " + template
    if "`{}`".format(template) not in section:
        problems.append("the text does not name the file")
    if "`{}`".format(command) not in section or command not in bash_commands(section):
        problems.append("the text does not run the file")
    sample = command.replace("<project>", "demo").replace("<date>", "2026-01-01")
    if not any(fnmatch.fnmatchcase(sample, pattern) for pattern in bash_patterns(allowed)):
        problems.append("no allowed-tools entry covers the run")
    return problems


def steps_without_tool(text, tool, steps):
    """Return the steps of `steps` ({heading start: word}) that lack a sentence which names `tool` and the word."""
    return [step for step, word in sorted(steps.items()) if not sentences_naming(step_section(text, step), tool, word)]


EDIT_STEPS = {"Step 4:": "script", "Step 5:": "text", "Step 6:": "script"}  # each step repairs something with `Edit`


def lint_slides(add, profile="peers", minutes=3):
    """Build a deck with the function add(prs, style) and return the findings of the linter."""
    style = style_mod.load_style(profile, "light")
    prs = slide_helpers.new_presentation(style)
    add(prs, style)
    with tempfile.TemporaryDirectory() as folder:
        path = Path(folder) / "deck.pptx"
        prs.save(str(path))
        return design_lint.lint(str(path), profile=profile, minutes=minutes)


NOTES_WORDS = ("alpha beta gamma delta epsilon zeta eta theta iota kappa lambda mu nu xi omicron pi rho sigma tau "
               "upsilon").split()


def notes_findings(shared):
    """Return the `notes` findings of a slide of 20 words whose notes repeat `shared` of them and add another word."""
    headline = " ".join(NOTES_WORDS)
    notes = " ".join(NOTES_WORDS[:shared] + ["explanation"])
    with tempfile.TemporaryDirectory() as folder:
        image = Path(folder) / "fig.png"
        Image.new("RGB", (1600, 900), "white").save(str(image))

        def add(prs, style):
            slide_helpers.add_assertion_evidence_slide(prs, style, headline, {"kind": "picture", "path": str(image),
                                                                              "alt": "A figure."}, notes=notes)
        findings = lint_slides(add)
    return [f for f in findings if f.rule == "notes"]


def claimed_notes_share(text):
    """Return the percent of slide words in the notes above which the text says that the linter warns."""
    match = re.search(r"The linter warns when more than (\d+) percent of the words of a slide appear in its notes\.",
                      " ".join(text.split()))
    return int(match.group(1)) if match else None


def slide_parts(prs):
    """Return {name: bytes} for the slide parts and the chart parts of a deck."""
    with tempfile.TemporaryDirectory() as folder:
        path = Path(folder) / "deck.pptx"
        prs.save(str(path))
        with zipfile.ZipFile(str(path)) as archive:
            return {n: archive.read(n) for n in archive.namelist() if n.startswith(("ppt/slides/slide", "ppt/charts/"))}


def highlight_changes_slide(evidence):
    """Return True when the "highlight" key of an evidence dict changes the slide or the chart that the helper writes."""
    plain = dict((k, v) for k, v in evidence.items() if k != "highlight")
    parts = []
    for item in (plain, evidence):
        style = style_mod.load_style("peers", "light")
        prs = slide_helpers.new_presentation(style)
        slide_helpers.add_assertion_evidence_slide(prs, style, "A headline.", item)
        parts.append(slide_parts(prs))
    return parts[0] != parts[1]


def take_home_findings(claim_words=15, headline="Three results support the message of this talk."):
    """Return the findings of a deck with one take-home slide of two claims. Each claim and baseline has `claim_words` words."""
    claim = " ".join(["word"] * claim_words)

    def add(prs, style):
        slide_helpers.add_take_home_slide(prs, style, [(claim, claim), (claim, claim)], headline=headline)
    return lint_slides(add)


def crowded_slide_findings(words=45):
    """Return the findings of a deck with one big-number slide whose baseline holds `words` words."""
    def add(prs, style):
        slide_helpers.add_big_number_slide(prs, style, "The drug cuts tumor size in mice.", "7", "mice",
                                           " ".join(["word"] * words))
    return lint_slides(add)


def claimed_take_home_words(text):
    """Return the number of words that the text allows for each claim and each baseline, or None."""
    match = re.search(r"Each claim and each baseline hold at most (\d+) words\.", " ".join(text.split()))
    return int(match.group(1)) if match else None


def claimed_take_home_headline(text):
    """Return (low, high, example) of the sentence that the text asks for as the take-home headline, or None."""
    match = re.search(r"It has (\d+) to (\d+) words and a verb\. For example, \"([^\"]+)\"", " ".join(text.split()))
    return (int(match.group(1)), int(match.group(2)), match.group(3)) if match else None


ROADMAP_SECTIONS = [("Need", "why"), ("Method", "how"), ("Result", "what")]


def message_width_chars(profile):
    """Return the characters that one line of the roadmap message holds under `profile`, by the estimate of the helpers."""
    style = style_mod.load_style(profile, "light")
    prs = slide_helpers.new_presentation(style)
    width_in = prs.slide_width / float(slide_helpers.EMU_PER_INCH) - 2 * style["layout"]["margin"]
    return slide_helpers._chars_per_line(width_in, style["typography"]["title_size"])


def message_failures(profile, length):
    """Return the seeded messages of `length` characters that add_roadmap_slide refuses under `profile`."""
    style = style_mod.load_style(profile, "light")
    prs = slide_helpers.new_presentation(style)
    failed = []
    for text in callout_texts(length):
        try:
            slide_helpers.add_roadmap_slide(prs, style, text, ROADMAP_SECTIONS)
        except ValueError as error:
            failed.append((text, str(error)))
    return failed


def claimed_message_length(text):
    """Return (points, characters) that the text gives for the two lines of a roadmap message, or None."""
    match = re.search(r"The roadmap message fits two lines at (\d+) pt, about (\d+) characters\.", " ".join(text.split()))
    return (int(match.group(1)), int(match.group(2))) if match else None


MESSAGE_WORD = 12  # the longest word of the seeded messages
MESSAGE_MARGIN = 5  # characters below the largest message that builds


def largest_message_length(profile):
    """Return the largest length that two lines hold for every message of words up to MESSAGE_WORD letters.

    A line of P characters ends when the next word does not fit, so it holds at least P - MESSAGE_WORD characters.
    The rest after the space fills the second line, so 2 * P - MESSAGE_WORD + 1 characters always fit.
    The tests check the number against the helper: this length builds and one more character fails for a seeded text.
    """
    return 2 * message_width_chars(profile) - MESSAGE_WORD + 1


# ---------------------------------------------------------------------------
# The checkers themselves: each one fails on a seeded defect
# ---------------------------------------------------------------------------

class CheckerSelfTests(unittest.TestCase):
    BODY = ("Use `Read` and `Glob`.\n\n```bash\nls docs\npython3 docs/slides/build-x.py\n```\n"
            "Run `python3 \"${CLAUDE_PLUGIN_ROOT}/scripts/design_lint.py\" deck.pptx`.\n")
    ALLOWED = ["Read", "Glob", "Bash(ls *)", "Bash(python3 docs/slides/build-*.py)",
               "Bash(python3 \"${CLAUDE_PLUGIN_ROOT}/scripts/design_lint.py\" *)"]

    def test_a_matching_set_passes(self):
        self.assertEqual(uncovered_commands(self.BODY, self.ALLOWED), [])
        self.assertEqual(unused_entries(self.BODY, self.ALLOWED), [])
        self.assertEqual(unlisted_tools(self.BODY, self.ALLOWED), [])

    def test_a_command_that_no_entry_covers_is_reported(self):
        body = self.BODY + "\n```bash\nrm -rf docs\n```\n"
        self.assertEqual(uncovered_commands(body, self.ALLOWED), ["rm -rf docs"])

    def test_a_command_behind_a_pipe_is_reported(self):
        body = self.BODY + "\n```bash\nls docs | xargs rm\n```\n"
        self.assertEqual(uncovered_commands(body, self.ALLOWED), ["xargs rm"])

    def test_an_inline_python_command_is_reported(self):
        body = self.BODY + "\nRun `python3 other.py`.\n"
        self.assertEqual(uncovered_commands(body, self.ALLOWED), ["python3 other.py"])

    def test_an_entry_that_the_text_never_uses_is_reported(self):
        self.assertEqual(unused_entries(self.BODY, self.ALLOWED + ["Bash(rm *)"]), ["Bash(rm *)"])
        self.assertEqual(unused_entries(self.BODY, self.ALLOWED + ["Grep"]), ["Grep"])

    def test_a_tool_that_the_text_names_but_the_list_lacks_is_reported(self):
        self.assertEqual(unlisted_tools(self.BODY + "\nUse `Edit`.\n", self.ALLOWED), ["Edit"])

    def test_a_missing_path_is_reported(self):
        text = "Read `scripts/design_lint.py`, then scripts/nothing_here.py. See references/research-notes/."
        self.assertEqual(missing_paths(text, PLUGIN), ["scripts/nothing_here.py"])

    def test_the_prose_check_finds_each_seeded_defect(self):
        self.assertEqual(prose_problems("Run the script. It writes the file."), [])
        self.assertIn("em dash", prose_problems("The run failed — and that is all."))
        self.assertIn("banned word: robust", prose_problems("The tool is robust."))
        self.assertIn("banned word: very", prose_problems("It is very small."))
        self.assertNotIn("banned word: very", prose_problems("Check every slide."))
        self.assertIn("not just", prose_problems("It is not just fast."))
        self.assertIn("'X, not Y' pair", prose_problems("Use the model, not the data."))
        self.assertIn("'X, not Y' pair", prose_problems("Fix it in the script, never in the deck."))
        self.assertTrue(any("words:" in p for p in prose_problems(" ".join(["word"] * 26) + ".")))
        self.assertEqual(prose_problems(" ".join(["word"] * 25) + "."), [])

    def test_the_prose_check_counts_list_items_and_wrapped_lines(self):
        long_item = "- " + " ".join(["word"] * 26)
        self.assertTrue(prose_problems("Intro.\n\n" + long_item + "\n"))
        wrapped = " ".join(["word"] * 13) + "\n" + " ".join(["word"] * 13) + ".\n"
        self.assertTrue(prose_problems(wrapped), "a wrapped sentence of 26 words must count as one sentence")
        short_items = "- " + " ".join(["word"] * 20) + "\n- " + " ".join(["word"] * 20) + "\n"
        self.assertEqual(prose_problems(short_items), [])

    def test_a_bold_lead_in_is_its_own_sentence(self):
        self.assertEqual(prose_problems("- **Name.** " + " ".join(["word"] * 25) + "."), [])
        self.assertTrue(prose_problems("- **Name.** " + " ".join(["word"] * 26) + "."))

    def test_the_prose_check_ignores_code(self):
        self.assertEqual(prose_problems("Run `a very long command that goes on and on and on and on and on and on and on and on and on`."), [])
        self.assertEqual(prose_problems("Text.\n\n```bash\necho " + " ".join(["word"] * 40) + "\n```\n"), [])

    def test_the_key_check_resolves_real_keys_and_rejects_false_ones(self):
        for report, key in (("jobs", "S1"), ("jobs", "S28"), ("craft", "S6"), ("craft", "S5"), ("craft", "S23"),
                            ("science", "S2"), ("science", "S29"), ("measurable", "WCAG-CM"),
                            ("measurable", "HARV"), ("measurable", "KAW")):
            with self.subTest(key="{}:{}".format(report, key)):
                self.assertTrue(key_resolves(report, key))
        for report, key in (("jobs", "S99"), ("jobs", "S"), ("craft", "S24"), ("science", "P1"),
                            ("science", "S30"), ("measurable", "NOPE"), ("measurable", "WCAG")):
            with self.subTest(key="{}:{}".format(report, key)):
                self.assertFalse(key_resolves(report, key))

    def test_the_key_check_needs_a_url_on_the_line_of_the_key(self):
        with tempfile.TemporaryDirectory() as folder:
            note = Path(folder) / "jobs.md"
            note.write_text("# x\n\n## 4. Pages opened\n\nS1\nhttps://example.org/a\nS2 no url here\nS3 https://example.org/c\n"
                            "\n## 5. Later\nS4 https://example.org/d\n", encoding="utf-8")
            self.assertFalse(key_resolves("jobs", "S1", folder))
            self.assertFalse(key_resolves("jobs", "S2", folder))
            self.assertTrue(key_resolves("jobs", "S3", folder))
            self.assertFalse(key_resolves("jobs", "S4", folder), "a key after the Pages opened section must not count")

    def test_the_path_detector_finds_each_prefix_and_ignores_relative_paths(self):
        for prefix in ("/home/someone/x", "/Users/someone/x", "/hpc/x/y", "/scratch/x", "/mnt/x/y", "/data/x",
                       "C:\\Users\\someone\\x", "see /home/x in the text"):
            with self.subTest(text=prefix):
                self.assertTrue(private_matches(prefix))
        for text in ("docs/data/x.csv", "https://example.org/data/x", "scripts/home/x", "the data folder",
                     "plugins/mnt/x", "${CLAUDE_PLUGIN_ROOT}/scripts"):
            with self.subTest(text=text):
                self.assertEqual(private_matches(text), [])

    def test_the_email_detector_finds_an_address_and_keeps_the_allowed_one(self):
        self.assertEqual(private_matches("write to a@example.org"), ["e-mail: a@example.org"])
        self.assertTrue(private_matches("first.last+tag@sub.example.org"))
        self.assertEqual(private_matches("Co-Authored-By: Name <noreply@anthropic.com>"), [])
        self.assertEqual(private_matches("the @ sign and user@host"), [])

    def test_the_template_author_detector_reads_the_name_at_run_time(self):
        name = template_author()
        if not name:
            self.skipTest("this python-pptx template carries no author name, so there is nothing to detect")
        self.assertEqual(private_matches("saved by " + name.upper()), ["the author name of the python-pptx template"])
        self.assertEqual(private_matches("saved by nobody"), [])

    def test_clean_text_has_no_private_match(self):
        self.assertEqual(private_matches("A sample deck. Every number is made up."), [])

    def test_the_front_matter_splitter_fails_on_each_seeded_defect(self):
        front, body = split_front_matter("---\ndescription: x\nallowed-tools:\n  - Read\n---\nBody.\n")
        self.assertEqual(front, {"description": "x", "allowed-tools": ["Read"]})
        self.assertEqual(body, "Body.\n")
        for text in ("No front matter.\n", "---\ndescription: x\nBody without a closing line.\n",
                     "---\n- a\n- b\n---\nBody.\n", "---\n---\nBody.\n"):
            with self.subTest(text=text[:20]):
                with self.assertRaises(ValueError):
                    split_front_matter(text)

    def test_the_tier_tag_pattern_accepts_one_tag_in_backticks_only(self):
        self.assertEqual(TIER_TAG.findall("x (`T1`; `craft:S6`)"), ["T1"])
        self.assertEqual(TIER_TAG.findall("x (`T2`, `T3`)"), ["T2", "T3"])
        for text in ("x (T1; `craft:S6`)", "x (`T4`)", "x (`t1`)", "x (`T12`)", "x ( T2 )"):
            with self.subTest(text=text):
                self.assertEqual(TIER_TAG.findall(text), [])

    def test_term_drift_finds_a_research_note_called_a_report(self):
        for text in ("The research reports disagree.", "The report found no origin.", "Each report has a tier.",
                     "The four reports differ."):
            with self.subTest(text=text):
                self.assertTrue(term_drift(text))
        for text in ("The research notes disagree.", "Report these items in this order.",
                     "For code 2, report the message and stop.", "A structured report of the session."):
            with self.subTest(text=text):
                self.assertEqual(term_drift(text) if "structured" not in text else [], [])

    def test_the_unused_local_check_finds_a_name_that_is_never_read(self):
        self.assertEqual(unused_locals("def f():\n    a, b = 1, 2\n    return b\n"), [("f", "a")])
        self.assertEqual(unused_locals("def f():\n    _, b = 1, 2\n    return b\n"), [])

    def test_the_phrase_check_reports_each_missing_phrase(self):
        required = {"a limit": ("alpha", "beta"), "another": ("gamma",)}
        self.assertEqual(missing_phrases("alpha beta gamma", required), [])
        self.assertEqual(missing_phrases("alpha gamma", required), [("a limit", "beta")])
        self.assertEqual(missing_phrases("", required), [("a limit", "alpha"), ("a limit", "beta"), ("another", "gamma")])

    def test_the_step_section_finder_returns_the_body_of_one_step(self):
        text = "## Step 1: First\nOne.\n\n## Step 10: Tenth\nTen.\n\n## Step 2: Second\nTwo.\n"
        self.assertEqual(step_section(text, "Step 1:").strip(), "One.")
        self.assertEqual(step_section(text, "Step 2:").strip(), "Two.")
        self.assertEqual(step_section(text, "Step 3:"), "")

    def test_the_tool_check_finds_a_sentence_that_names_the_tool_and_the_script(self):
        self.assertTrue(sentences_naming("If the script fails, repair it with `Edit`. Run it again.", "Edit"))
        self.assertEqual(sentences_naming("Repair it with `Edit`. The script runs.", "Edit"), [])
        self.assertEqual(sentences_naming("Repair the script with Edit.", "Edit"), [])
        self.assertEqual(sentences_naming("Repair the script.\n\n```bash\n# `Edit` the script\n```\n", "Edit"), [])

    def test_the_evidence_check_finds_a_chart_highlight_that_the_helper_refuses(self):
        one = {"kind": "chart", "chart": "bar", "categories": ["A", "B"], "series": {"Name": [1, 2]},
               "highlight": "B", "alt": "A sentence."}
        several = dict(one, series={"Name": [1, 2], "Other": [2, 3]}, highlight="Name")
        self.assertIsNone(evidence_error(one))
        self.assertIsNone(evidence_error(several))
        self.assertIn("not a series name", evidence_error(dict(several, highlight="B")))
        self.assertIn("matches no category", evidence_error(dict(one, highlight="Name")))

    def test_the_evidence_reader_reads_a_picture_path_and_ignores_other_lines(self):
        body = ('```python\nx = 1\n{"kind": "picture", "path": ROOT / "figures/a.png", "alt": "A."}\n'
                '{"kind": "table", "rows": [["a"]], "highlight": (0, 0)}\n```\n')
        self.assertEqual(evidence_examples(body, "/folder"),
                         [{"kind": "picture", "path": "/folder/figures/a.png", "alt": "A."},
                          {"kind": "table", "rows": [["a"]], "highlight": (0, 0)}])

    def test_the_call_out_check_finds_a_text_that_is_too_long(self):
        self.assertTrue(all(len(t) == 35 for t in callout_texts(35)))
        self.assertEqual(callout_texts(35), callout_texts(35), "the texts are seeded")
        self.assertEqual(callout_failures("peers", 35), [])
        self.assertEqual(len(callout_failures("peers", 62)), 40, "a call-out of 62 characters needs four lines")
        self.assertEqual(callout_failures("general", 25), [])
        self.assertTrue(callout_failures("general", 45), "the general profile has larger text and a shorter line")

    def test_the_call_out_claim_reader_reads_the_numbers_or_nothing(self):
        text = ("Keep a call-out to about 35 characters for `peers` and `lab-meeting` and to about 25 characters "
                "for `general`.")
        self.assertEqual(claimed_callout_lengths(text), {"peers": 35, "lab-meeting": 35, "general": 25})
        self.assertEqual(claimed_callout_lengths("Keep a call-out short."), {})

    def test_the_word_limit_reader_reads_the_numbers_or_nothing(self):
        text = ("For `peers`, `lab-meeting` and `general`, the linter warns above 25, 40 and 15 words and gives "
                "an error above 40, 60 and 25.")
        self.assertEqual(claimed_word_limits(text), {"peers": (25, 40), "lab-meeting": (40, 60), "general": (15, 25)})
        old = "The linter warns above 25 words for `peers`, 40 for `lab-meeting` and 15 for `general`."
        self.assertEqual(claimed_word_limits(old), {}, "the old sentence names no error limit")

    def test_the_source_rule_check_finds_each_seeded_defect(self):
        without, with_source = helpers_by_source()
        good = (SOURCE_RULE + "\n- " + TITLE_NO_NUMBER + " The roadmap slide and the take-home slide show "
                "only numbers that another slide sources with its own source line.\n")
        self.assertEqual(source_rule_problems(good, without, with_source), [])
        old = ("Add `source` to each slide that shows a number, except the take-home slide. A take-home claim takes "
               "its source from the earlier slide that shows the number.")
        self.assertTrue(source_rule_problems(old, without, with_source))
        self.assertIn("no mention of add_title_slide",
                      source_rule_problems(good.replace(TITLE_NO_NUMBER + " ", "A slide shows no number. "),
                                           without, with_source))
        self.assertEqual(source_rule_problems(good.replace(" from the data", ""), without, with_source),
                         ["the title slide sentence is missing"], "the title slide carries a date")
        self.assertIn("no mention of add_roadmap_slide",
                      source_rule_problems(good.replace("roadmap slide and the ", ""), without, with_source))
        self.assertIn("wrong mention of add_big_number_slide",
                      source_rule_problems(good.replace("The title slide", "The big_number slide"), without, with_source))
        self.assertIn("the source rule sentence is missing",
                      source_rule_problems(good.replace(SOURCE_RULE, ""), without, with_source))

    def test_the_message_rule_check_reads_the_limit_and_the_linter_agrees(self):
        text = "A message of more than 20 words gets a `title-sentence` warning on the roadmap slide."
        self.assertEqual(claimed_message_limit(text), 20)
        self.assertIsNone(claimed_message_limit("A long message gets a warning."))
        self.assertEqual(roadmap_title_findings(20), [])
        findings = roadmap_title_findings(21)
        self.assertEqual([(f.slide, f.severity) for f in findings], [(1, "warn")])

    def test_the_run_check_finds_a_script_name_that_no_allowed_entry_covers(self):
        allowed = ["Bash(python3 docs/slides/build-*.py)"]
        good = "Save it as `docs/slides/build-<project>-<date>-preview.py`. Run `python3 docs/slides/build-<project>-<date>-preview.py`."
        template = "docs/slides/build-<project>-<date>-preview.py"
        self.assertEqual(run_command_problems(good, allowed, template), [])
        self.assertIn("the text does not run the file", run_command_problems(good.split(". Run")[0], allowed, template))
        self.assertIn("the text does not name the file", run_command_problems("Nothing here.", allowed, template))
        other = "docs/slides/<project>-<date>-preview.py"
        self.assertIn("no allowed-tools entry covers the run",
                      run_command_problems("`{0}` `python3 {0}`".format(other), allowed, other))
        self.assertIn("no allowed-tools entry covers the run", run_command_problems(good, [], template))

    def test_the_edit_pin_finds_a_step_that_lacks_the_tool(self):
        good = ("## Step 4: A\nIf the script fails, repair it with `Edit`.\n\n## Step 5: B\nShorten the text with `Edit`.\n\n"
                "## Step 6: C\nFix the script with `Edit`.\n")
        self.assertEqual(steps_without_tool(good, "Edit", EDIT_STEPS), [])
        only_step_6 = good.replace("Shorten the text with `Edit`.", "Shorten the text.")
        self.assertEqual(steps_without_tool(only_step_6, "Edit", EDIT_STEPS), ["Step 5:"])
        self.assertEqual(steps_without_tool(good.replace("Fix the script", "Fix it"), "Edit", EDIT_STEPS), ["Step 6:"])

    def test_the_notes_check_agrees_with_the_linter_at_the_threshold(self):
        self.assertEqual(claimed_notes_share("The linter warns when more than 80 percent of the words of a slide appear in its notes."), 80)
        self.assertIsNone(claimed_notes_share("The linter warns about notes."))
        self.assertEqual(notes_findings(16), [], "16 of 20 words is 80 percent, which is not above the limit")
        findings = notes_findings(17)
        self.assertEqual([f.severity for f in findings], ["warn"])
        self.assertIn("85%", findings[0].message)
        self.assertEqual(notes_findings(0), [])

    def test_the_highlight_check_finds_the_evidence_that_ignores_it(self):
        with tempfile.TemporaryDirectory() as folder:
            image = Path(folder) / "fig.png"
            Image.new("RGB", (1600, 900), "white").save(str(image))
            picture = {"kind": "picture", "path": str(image), "alt": "A figure.", "highlight": (1, 0)}
            self.assertEqual(evidence_error(picture), None, "the helper accepts the key and ignores it")
            self.assertFalse(highlight_changes_slide(picture))
        table = {"kind": "table", "rows": [["A", "B"], ["1", "2"]], "highlight": (1, 0)}
        chart = {"kind": "chart", "chart": "bar", "categories": ["A", "B"], "series": {"Name": [1, 2]},
                 "highlight": "B", "alt": "A sentence."}
        self.assertTrue(highlight_changes_slide(table))
        self.assertTrue(highlight_changes_slide(chart))

    def test_the_take_home_check_finds_the_word_limits_and_the_headline_rule(self):
        findings = take_home_findings(15)
        self.assertEqual([f for f in findings if f.rule in ("words-per-slide", "take-home-words", "title-sentence")], [])
        long_claims = take_home_findings(16)
        self.assertEqual(sorted(f.rule for f in long_claims if f.rule == "take-home-words"), ["take-home-words"] * 4)
        short = take_home_findings(15, headline="Three results support the message.")
        self.assertEqual([f.severity for f in short if f.rule == "title-sentence"], ["warn"],
                         "an example of 5 words breaks the 6 word minimum")
        crowded = [f for f in crowded_slide_findings(45) if f.rule == "words-per-slide"]
        self.assertEqual([f.severity for f in crowded], ["error"], "the rule fires on a slide that is not a take-home slide")

    def test_the_take_home_readers_read_the_numbers_or_nothing(self):
        self.assertEqual(claimed_take_home_words("Each claim and each baseline hold at most 15 words."), 15)
        self.assertIsNone(claimed_take_home_words("Keep each claim short."))
        text = 'It has 6 to 20 words and a verb. For example, "Three results support the message of this talk."'
        self.assertEqual(claimed_take_home_headline(text), (6, 20, "Three results support the message of this talk."))
        self.assertIsNone(claimed_take_home_headline("Draft a sentence."))

    def test_the_message_check_finds_a_message_that_is_too_long(self):
        self.assertEqual(claimed_message_length("The roadmap message fits two lines at 28 pt, about 104 characters."),
                         (28, 104))
        self.assertIsNone(claimed_message_length("The roadmap slide holds a message."))
        self.assertEqual(message_failures("peers", 105), [])
        failed = message_failures("peers", 125)
        self.assertEqual(len(failed), 40, "a message of 125 characters needs three lines")
        self.assertIn("needs about 3 lines at 28 pt", failed[0][1])

    def test_the_largest_message_length_is_the_edge_of_the_helper(self):
        largest = largest_message_length("peers")
        self.assertEqual(largest, 109)
        self.assertEqual(message_failures("peers", largest), [])
        self.assertTrue(message_failures("peers", largest + 1), "one more character fails for a seeded text")


# ---------------------------------------------------------------------------
# The command
# ---------------------------------------------------------------------------

class Command(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.text = read(COMMAND)
        cls.front, cls.body = split_front_matter(cls.text)

    def test_front_matter_has_a_description_and_a_list_of_tools(self):
        self.assertIsInstance(self.front.get("description"), str)
        self.assertTrue(self.front["description"].strip())
        allowed = self.front.get("allowed-tools")
        self.assertIsInstance(allowed, list)
        self.assertTrue(allowed)
        for entry in allowed:
            self.assertIsInstance(entry, str)
        self.assertEqual(len(allowed), len(set(allowed)), "an entry appears twice")

    def test_every_command_that_the_text_runs_is_allowed(self):
        self.assertEqual(uncovered_commands(self.body, self.front["allowed-tools"]), [])

    def test_every_allowed_entry_is_used_by_the_text(self):
        self.assertEqual(unused_entries(self.body, self.front["allowed-tools"]), [])

    def test_every_tool_that_the_text_names_is_allowed(self):
        self.assertEqual(unlisted_tools(self.body, self.front["allowed-tools"]), [])

    def test_the_bash_entries_are_narrow(self):
        for pattern in bash_patterns(self.front["allowed-tools"]):
            with self.subTest(pattern=pattern):
                self.assertNotIn(pattern, ("*", "python3 *", "python *", "pip *", "rm *"))
                self.assertTrue(pattern.split()[0] in ("ls", "mkdir", "date", "python3"), pattern)
        self.assertNotIn("Bash", self.front["allowed-tools"])

    def test_the_text_runs_the_commands_that_the_plan_needs(self):
        commands = bash_commands(self.body)
        joined = "\n".join(commands)
        for needle in ("design_lint.py", "docs/slides/build-", "mkdir -p docs/slides", "date +%Y-%m-%d",
                       "pip install python-pptx", "pip install PyYAML", "pip install Pillow", "ls "):
            with self.subTest(needle=needle):
                self.assertIn(needle, joined)

    def test_every_plugin_path_that_the_text_names_exists(self):
        self.assertEqual(missing_paths(self.text, PLUGIN), [])
        self.assertIn("${CLAUDE_PLUGIN_ROOT}/references/design-principles.md", self.text)

    def test_the_text_reads_the_principles_before_step_3(self):
        read_at = self.text.index("${CLAUDE_PLUGIN_ROOT}/references/design-principles.md")
        step_3 = re.search(r"^#+ Step 3\b", self.text, re.M)
        self.assertIsNotNone(step_3)
        self.assertLess(read_at, step_3.start())
        self.assertIn("before", self.text[max(0, read_at - 200):read_at + 300].lower())

    def test_the_steps_run_from_0_to_7_in_order(self):
        steps = [int(n) for n in re.findall(r"^#+ Step (\d+)\b", self.text, re.M)]
        self.assertEqual(steps, list(range(8)))

    def test_the_arguments_line_matches_the_spec(self):
        self.assertIn("[15min | N slides] [--profile peers|lab-meeting|general] [--theme light|dark] [--no-notes]",
                      self.text)
        self.assertIn("$ARGUMENTS", self.text)

    def test_the_text_names_the_output_files(self):
        self.assertIn("docs/slides/build-<project>-<date>.py", self.text)
        self.assertIn("docs/slides/<project>-<date>.pptx", self.text)
        self.assertIn('sys.path.insert(0, "${CLAUDE_PLUGIN_ROOT}/scripts")', self.text)

    def test_the_text_states_the_rules_of_the_brief(self):
        flat = " ".join(self.text.split()).lower()
        for phrase in ("never paraphrase", "source line", "take every number from an artifact", "own voice",
                       "glance test", "color-only", "fit of a figure to its headline", "linter with `--no-notes`",
                       "draft no notes", "model judgment"):
            with self.subTest(phrase=phrase):
                self.assertIn(phrase, flat)
        lint_lines = [line for line in self.text.splitlines() if "design_lint.py" in line and "--profile" in line]
        self.assertTrue(lint_lines, "the text must show the linter command with --profile")

    def test_the_text_tells_the_builder_to_pass_alt_and_a_take_home_headline(self):
        flat = " ".join(self.text.split())
        self.assertRegex(flat, r"`alt`.{0,80}(picture|chart)|(picture|chart).{0,120}`alt`")
        self.assertRegex(flat, r"`headline`.{0,120}add_take_home_slide|add_take_home_slide.{0,120}`headline`")
        self.assertIn("fallback", flat)

    def test_the_step_6_lint_command_passes_the_project_dir(self):
        lines = [line for line in self.text.splitlines() if "design_lint.py" in line and "--profile" in line]
        self.assertTrue(lines, "the text must show the linter command with --profile")
        for line in lines:
            self.assertIn("--project-dir .", line, "the linter reads the project style file only with --project-dir")

    def test_the_source_rule_follows_the_signatures_in_both_places(self):
        """Only some helpers take a source. The rule of the text must follow the signatures, in the rules and in Step 5."""
        without, with_source = helpers_by_source()
        self.assertEqual(with_source, ["add_assertion_evidence_slide", "add_big_number_slide"])
        self.assertEqual(without, ["add_roadmap_slide", "add_take_home_slide", "add_title_slide"])
        every = sorted(n for n in dir(slide_helpers) if re.match(r"add_\w+_slide$", n))
        self.assertEqual(every, sorted(SLIDE_WORDS), "name the slide of each helper in SLIDE_WORDS")
        parts = {"Rules for every step": step_section(self.text, "Rules for every step"),
                 "Step 5": step_section(self.text, "Step 5:")}
        for name, part in parts.items():
            with self.subTest(part=name):
                self.assertEqual(source_rule_problems(part, without, with_source), [])
        self.assertEqual(" ".join(self.text.split()).count(SOURCE_RULE), 2)

    def test_the_chart_examples_build_and_show_both_highlight_forms(self):
        with tempfile.TemporaryDirectory() as folder:
            (Path(folder) / "figures").mkdir()
            Image.new("RGB", (1600, 900), "white").save(str(Path(folder) / "figures" / "fig1.png"))
            examples = evidence_examples(self.body, folder)
            self.assertEqual(sorted({e["kind"] for e in examples}), ["chart", "picture", "table"])
            for example in examples:
                with self.subTest(kind=example["kind"], series=len(example.get("series", ()))):
                    self.assertIsNone(evidence_error(example))
        forms = {len(e["series"]) > 1 for e in examples if e["kind"] == "chart" and "highlight" in e}
        self.assertEqual(forms, {False, True}, "show a highlight for one series and for several series")
        flat = " ".join(step_section(self.text, "Step 5:").split())
        for phrase in ("A chart with one series takes a category name or index as `highlight`.",
                       "A chart with several series takes a series name."):
            with self.subTest(phrase=phrase):
                self.assertIn(phrase, flat)

    def test_the_call_out_length_is_safe_for_each_profile(self):
        claimed = claimed_callout_lengths(self.text)
        self.assertEqual(sorted(claimed), sorted(style_mod.PROFILES))
        for profile, length in sorted(claimed.items()):
            with self.subTest(profile=profile, length=length):
                per_line = callout_width_chars(profile)
                self.assertTrue(per_line <= length <= 2 * per_line,
                                "{} characters: one line holds {}, two lines hold at most {}".format(length, per_line, 2 * per_line))
                self.assertEqual(callout_failures(profile, length), [])
        self.assertIn("Shorten a call-out that raises `ValueError`.", " ".join(self.text.split()))

    def test_the_word_limits_match_the_profile_files(self):
        claimed = claimed_word_limits(self.text)
        self.assertEqual(sorted(claimed), sorted(style_mod.PROFILES))
        for profile, limits in sorted(claimed.items()):
            with self.subTest(profile=profile):
                words = style_mod.load_style(profile, "light")["lint"]["words_per_slide"]
                self.assertEqual(limits, (words["warn"], words["error"]))

    def test_a_message_above_the_title_limit_keeps_its_wording_and_the_user_hears_of_the_warning(self):
        limit = claimed_message_limit(self.text)
        self.assertEqual(limit, style_mod.load_style("peers", "light")["lint"]["title_words"]["max"])
        self.assertEqual(roadmap_title_findings(limit), [])
        self.assertEqual([f.severity for f in roadmap_title_findings(limit + 1)], ["warn"])
        flat = " ".join(step_section(self.text, "Rules for every step").split())
        self.assertIn("If the user wrote the sentence, keep it as written. Let the warning stand.", flat)
        take_home_limit = style_mod.load_style("peers", "light")["lint"]["take_home_words"]["max_claim"]
        self.assertIn("A take-home claim of more than %d words gets a `take-home-words` warning." % take_home_limit, flat)
        report = " ".join(step_section(self.text, "Step 7:").split())
        self.assertIn("Say that the warning stays because the sentence is theirs and that they can shorten it.", report)

    def test_edit_is_allowed_and_steps_4_5_and_6_repair_the_builder_script_with_it(self):
        self.assertIn("Edit", self.front["allowed-tools"])
        self.assertEqual(steps_without_tool(self.text, "Edit", EDIT_STEPS), [],
                         "each of these steps must tell the model to repair with `Edit`")
        flat = " ".join(step_section(self.text, "Step 5:").split())
        self.assertIn("Shorten your own text with `Edit`. A sentence of the user stays as written. Run the script again.",
                      flat, "one instruction for each sentence")
        self.assertNotIn("`Edit` and run", flat)

    def test_highlight_goes_on_a_chart_or_a_table_and_a_picture_gets_a_call_out(self):
        with tempfile.TemporaryDirectory() as folder:
            (Path(folder) / "figures").mkdir()
            Image.new("RGB", (1600, 900), "white").save(str(Path(folder) / "figures" / "fig1.png"))
            examples = evidence_examples(self.body, folder)
            for example in examples:
                if example["kind"] == "picture":
                    self.assertNotIn("highlight", example)
                    self.assertFalse(highlight_changes_slide(dict(example, highlight=(1, 0))),
                                     "the helper ignores a highlight key on a picture")
                else:
                    self.assertTrue(highlight_changes_slide(example), "highlight changes a {}".format(example["kind"]))
        flat = " ".join(step_section(self.text, "Step 5:").split())
        for phrase in ("Mark the cue of a chart or a table with `highlight`.",
                       "For a picture, name the cue in a call-out."):
            with self.subTest(phrase=phrase):
                self.assertIn(phrase, flat)
        self.assertNotIn("Mark the cue of each slide", self.text)

    def test_the_notes_rule_gives_the_threshold_of_the_yaml_and_asks_for_new_content(self):
        share = style_mod.load_style("peers", "light")["lint"]["notes"]["redundancy_above"]
        claimed = claimed_notes_share(step_section(self.text, "Step 5:"))
        self.assertEqual(claimed, int(round(100 * share)))
        self.assertEqual(claimed % 5, 0, "the check slide has 20 words, so it steps in 5 percent")
        at_limit = claimed * 20 // 100
        self.assertEqual(notes_findings(at_limit), [], "at the limit the linter stays quiet")
        self.assertEqual([f.severity for f in notes_findings(at_limit + 1)], ["warn"])
        flat = " ".join(step_section(self.text, "Step 5:").split())
        for phrase in ("Write what the slide does not say, such as the spoken explanation, the caveat and the source of a number.",
                       "Do not repeat the words of the slide."):
            with self.subTest(phrase=phrase):
                self.assertIn(phrase, flat)
        self.assertNotIn("with the script of the slide", self.text)

    def test_the_need_slide_has_an_element_when_no_figure_shows_the_need(self):
        flat = " ".join(step_section(self.text, "Step 3:").split())
        for phrase in ("If no figure shows the need, use the baseline number on a `big_number` slide.",
                       "You can also use a table of two or three terms that the audience must know."):
            with self.subTest(phrase=phrase):
                self.assertIn(phrase, flat)
        self.assertIn("big_number", slide_helpers.KINDS)
        self.assertIn("table", slide_helpers.EVIDENCE_KINDS)
        terms = {"kind": "table", "rows": [["Term", "Meaning"], ["A", "one"], ["B", "two"]]}
        self.assertIsNone(evidence_error(terms), "a table of three terms builds")
        style = style_mod.load_style("peers", "light")
        slide_helpers.add_big_number_slide(slide_helpers.new_presentation(style), style, "The baseline is 3 percent.",
                                           "3%", "of mice", "A baseline for the claim.")

    def test_the_take_home_slide_has_its_own_word_limit_and_skips_the_slide_budget(self):
        step_3 = step_section(self.text, "Step 3:")
        limits = style_mod.load_style("peers", "light")["lint"]["take_home_words"]
        for profile in style_mod.PROFILES:
            self.assertEqual(style_mod.load_style(profile, "light")["lint"]["take_home_words"], limits)
        self.assertEqual(limits["max_claim"], limits["max_baseline"])
        self.assertEqual(claimed_take_home_words(step_3), limits["max_claim"])
        self.assertIn("The `words-per-slide` rule does not check this slide.", " ".join(step_3.split()))
        self.assertIn("words-per-slide", design_lint._KEYS)
        words = limits["max_claim"]
        findings = take_home_findings(words)
        self.assertEqual([f for f in findings if f.rule in ("words-per-slide", "take-home-words")], [])
        self.assertGreater(4 * words, style_mod.load_style("peers", "light")["lint"]["words_per_slide"]["error"])

    def test_the_take_home_headline_is_drafted_from_the_claims_and_the_user_hears_of_it(self):
        found = claimed_take_home_headline(step_section(self.text, "Step 5:"))
        self.assertIsNotNone(found)
        low, high, example = found
        title_words = style_mod.load_style("peers", "light")["lint"]["title_words"]
        self.assertEqual((low, high), (title_words["min"], title_words["max"]))
        self.assertIsNone(title_check.title_problem(example, low, high), "the example must pass the title rule")
        self.assertEqual([f for f in take_home_findings(15, headline=example) if f.rule == "title-sentence"], [])
        self.assertIn("Draft that sentence from the claims.", " ".join(step_section(self.text, "Step 5:").split()))
        self.assertIn("Say that you drafted the take-home headline.", " ".join(step_section(self.text, "Step 7:").split()))

    def test_a_message_that_the_helper_rejects_goes_back_to_the_user(self):
        found = claimed_message_length(step_section(self.text, "Step 5:"))
        self.assertIsNotNone(found)
        points, length = found
        for profile in style_mod.PROFILES:
            with self.subTest(profile=profile):
                self.assertEqual(points, style_mod.load_style(profile, "light")["typography"]["title_size"])
                largest = largest_message_length(profile)
                self.assertEqual(message_failures(profile, largest), [])
                self.assertTrue(message_failures(profile, largest + 1), "the largest length is the edge of the helper")
                self.assertEqual(length, largest - MESSAGE_MARGIN)
                self.assertEqual(message_failures(profile, length), [])
        flat = " ".join(step_section(self.text, "Step 5:").split())
        self.assertIn("If the helper rejects the message of the user, ask the user for a shorter sentence with "
                      "`AskUserQuestion`. Never shorten it yourself.", flat)
        self.assertIn("AskUserQuestion", self.front["allowed-tools"])

    def test_the_title_slide_shows_no_number_from_the_data_in_both_places(self):
        self.assertIn("date", inspect.signature(slide_helpers.add_title_slide).parameters,
                      "the title slide carries a date, so the text says no number from the data")
        self.assertEqual(" ".join(self.text.split()).count(TITLE_NO_NUMBER), 2)
        self.assertNotIn("The title slide shows no number.", self.text)

    def test_the_preview_builder_has_a_name_that_the_run_pattern_allows(self):
        template = "docs/slides/build-<project>-<date>-preview.py"
        self.assertEqual(run_command_problems(step_section(self.text, "Step 4:"), self.front["allowed-tools"], template), [])

    def test_step_1_says_what_to_do_when_no_script_writes_a_figure(self):
        required = {"a figure with no producer": (
            "If no script writes a figure, write the data file or the note that holds its numbers as the source.",)}
        self.assertEqual(missing_phrases(step_section(self.text, "Step 1:"), required), [])

    def test_the_install_step_covers_pillow_and_asks_the_user_when_pip_refuses(self):
        self.assertIn("Bash(python3 -m pip install Pillow*)", self.front["allowed-tools"])
        self.assertIn("python3 -m pip install Pillow", bash_commands(self.body))
        self.assertIn("`PIL`", self.text)
        self.assertIn("ask the user to run `! python3 -m pip install --user <package>` themselves", " ".join(self.text.split()))
        self.assertNotIn("add `--user`", self.text)
        self.assertIn("If pip says `externally-managed-environment`, ask the user to create a virtual environment or to "
                      "run the install with `--break-system-packages` themselves.",
                      " ".join(step_section(self.text, "Step 0:").split()))

    def test_step_2_asks_for_the_title_slide_fields_and_step_3_asks_for_them_no_more(self):
        step_2 = " ".join(step_section(self.text, "Step 2:").split())
        self.assertIn("Ask with `AskUserQuestion`. A call holds at most four questions, so ask the title-slide question "
                      "in a second call when five questions remain.", step_2)
        self.assertIn("5. Title slide: the title, the author, the affiliation and the date, when the artifacts do not "
                      "hold them.", step_2)
        step_3 = " ".join(step_section(self.text, "Step 3:").split())
        self.assertNotRegex(step_3, r"(?i)\bask\b[^.]*\baffiliation\b")
        self.assertIn("Slide 1 is the `title` slide. It shows the title slide fields of Step 2.", step_3)

    def test_the_text_does_not_describe_a_template_workaround(self):
        self.assertNotIn("template", self.text)
        self.assertIn("`last_modified_by`", self.text)

    def test_the_text_does_not_run_the_render_check(self):
        self.assertNotIn("--render", self.text)

    def test_the_helper_signatures_in_the_text_match_slide_helpers(self):
        found = signature_lines(self.body)
        self.assertTrue(found, "the text must list the helper signatures in a python block")
        for name in ("new_presentation", "add_title_slide", "add_assertion_evidence_slide", "add_big_number_slide",
                     "add_roadmap_slide", "add_take_home_slide"):
            with self.subTest(function=name):
                self.assertIn(name, found)
        for name, parameters in found.items():
            with self.subTest(function=name):
                real = list(inspect.signature(getattr(slide_helpers, name)).parameters)
                self.assertEqual(parameters, real)

    def test_the_text_has_no_private_text_and_no_absolute_path(self):
        self.assertEqual(private_matches(self.text), [])
        self.assertNotIn("demo/", self.text)

    def test_prose(self):
        self.assertEqual(prose_problems(self.text), [])
        self.assertEqual(term_drift(self.text), [], "call a research note a research note")


# ---------------------------------------------------------------------------
# The principles
# ---------------------------------------------------------------------------

class Principles(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.text = read(PRINCIPLES)
        cls.parts = sections(cls.text)
        cls.principle_sections = ("Talk-level principles", "Slide-level principles", "Delivery principles")
        cls.lines = [line for name in cls.principle_sections for line in list_items(cls.parts.get(name, ""))]

    def test_the_sections_exist(self):
        for name in ("How to read the tiers",) + self.principle_sections + ("Where the authors disagree", "Not adopted"):
            with self.subTest(section=name):
                self.assertIn(name, self.parts)

    def test_there_are_at_least_20_principle_lines(self):
        self.assertGreaterEqual(len(self.lines), 20)

    def test_each_principle_line_has_one_tier_tag_and_a_source_key(self):
        for line in self.lines:
            with self.subTest(line=line[:60]):
                self.assertEqual(len(TIER_TAG.findall(line)), 1, "one tier tag, T1, T2 or T3, in backticks")
                self.assertTrue(source_keys(line), "at least one source key")

    def test_every_source_key_has_a_known_report_and_resolves_to_a_url(self):
        unknown = [(r, k) for r, k in source_keys(self.text) if r not in REPORTS]
        self.assertEqual(unknown, [], "keys with an unknown report name")
        unresolved = sorted({"{}:{}".format(r, k) for r, k in source_keys(self.text) if not key_resolves(r, k)})
        self.assertEqual(unresolved, [], "keys that do not resolve to a URL line in the Pages opened section")

    def test_the_reference_does_not_use_a_report_without_naming_a_key_form(self):
        self.assertTrue(source_keys(self.text))
        bare = re.findall(r"\b(?:jobs|craft|science|measurable):[A-Za-z0-9-]+", FENCE.sub("", re.sub(r"`[^`]*`", "", self.text)))
        self.assertEqual(bare, [], "write a key in backticks so that the check reads it")

    def test_the_tier_table_matches_the_notes(self):
        notes_readme = read(NOTES / "README.md")
        for report in REPORTS:
            self.assertIn("`{}`".format(report), self.parts["How to read the tiers"])
        for phrase in ("Seen in a transcript", "Controlled study", "Named expert source", "Standard or study",
                       "Folklore"):
            with self.subTest(phrase=phrase):
                self.assertIn(phrase, self.parts["How to read the tiers"])
                self.assertIn(phrase.lower(), notes_readme.lower())

    def test_not_adopted_lists_at_least_five_items_and_each_has_a_reason(self):
        items = list_items(self.parts["Not adopted"])
        self.assertGreaterEqual(len(items), 5)
        for item in items:
            with self.subTest(item=item[:50]):
                self.assertIn("Reason:", item)

    def test_the_six_by_six_reason_states_only_what_the_note_states(self):
        item = [i for i in list_items(self.parts["Not adopted"]) if "6 by 6" in i][0]
        self.assertIn("found no origin", item)
        self.assertNotIn("no test", item)

    def test_the_pace_line_names_the_research_note_that_holds_each_statement(self):
        line = [l for l in self.lines if "one slide per minute" in l][0]
        for name in ("`measurable.md`", "`science.md`"):
            self.assertIn(name, line)
        self.assertIn("opinion", line)
        self.assertIn("guidance", line)

    def test_the_authors_disagree_on_six_points(self):
        self.assertGreaterEqual(len(list_items(self.parts["Where the authors disagree"])), 6)

    def test_the_text_has_no_private_text(self):
        self.assertEqual(private_matches(self.text), [])

    def test_prose(self):
        self.assertEqual(prose_problems(self.text), [])
        self.assertEqual(term_drift(self.text), [], "call a research note a research note")


# ---------------------------------------------------------------------------
# The sample deck
# ---------------------------------------------------------------------------

class SampleDeck(unittest.TestCase):
    def test_the_deck_and_the_script_exist(self):
        self.assertTrue(SAMPLE_SCRIPT.is_file())
        self.assertTrue(SAMPLE_DECK.is_file())

    def test_the_deck_lints_with_no_finding_under_peers(self):
        findings = design_lint.lint(str(SAMPLE_DECK), profile="peers", minutes=12)
        self.assertEqual(findings, [], "\n".join("slide {}: {}: {}: {}: {}".format(*f) for f in findings))

    def test_the_command_line_agrees(self):
        result = subprocess.run([sys.executable, "-B", str(SCRIPTS_DIR / "design_lint.py"), str(SAMPLE_DECK),
                                 "--profile", "peers", "--minutes", "12"],
                                stdout=subprocess.PIPE, stderr=subprocess.PIPE, universal_newlines=True, timeout=300)
        self.assertEqual(result.returncode, 0, result.stdout + result.stderr)
        self.assertEqual(result.stdout.splitlines()[-2:], ["errors: 0", "warnings: 0"])

    def test_the_deck_uses_all_six_layouts(self):
        prs = Presentation(str(SAMPLE_DECK))
        kinds = [slide_helpers.slide_kind(slide) for slide in prs.slides]
        self.assertEqual(set(kinds), set(slide_helpers.KINDS))
        self.assertTrue(6 <= len(kinds) <= 10, "about eight slides, got {}".format(len(kinds)))
        self.assertEqual(kinds[0], "title")
        self.assertEqual(kinds[-1], "backup")

    def test_every_slide_but_the_title_has_notes(self):
        prs = Presentation(str(SAMPLE_DECK))
        for number, slide in enumerate(prs.slides, 1):
            with self.subTest(slide=number):
                if slide_helpers.slide_kind(slide) == "title":
                    continue
                self.assertTrue(slide.has_notes_slide)
                self.assertTrue(slide.notes_slide.notes_text_frame.text.strip())

    def test_the_deck_carries_no_private_text(self):
        for name, text in deck_xml_texts(SAMPLE_DECK):
            with self.subTest(part=name):
                self.assertEqual(private_matches(text), [])

    def test_the_core_properties_are_not_those_of_the_python_pptx_template(self):
        template = Presentation().core_properties
        deck = Presentation(str(SAMPLE_DECK)).core_properties
        for name in ("author", "last_modified_by", "comments", "title"):
            with self.subTest(field=name):
                if getattr(template, name):
                    self.assertNotEqual(getattr(deck, name), getattr(template, name))
        self.assertGreater(deck.created.year, template.created.year)

    def test_the_deck_says_that_its_numbers_are_made_up(self):
        prs = Presentation(str(SAMPLE_DECK))
        text = " ".join(shape.text_frame.text for slide in prs.slides for shape in slide.shapes if shape.has_text_frame)
        self.assertIn("made-up", text.lower())

    def test_the_script_uses_pillow_and_no_other_source(self):
        text = read(SAMPLE_SCRIPT)
        self.assertIn("PIL", text)
        self.assertNotIn("matplotlib", text)
        self.assertNotIn("demo/", text)
        self.assertEqual(private_matches(text), [])
        self.assertNotIn("template", text, "new_presentation clears the template metadata, so the script needs no workaround")
        self.assertEqual(unused_locals(text), [])

    def test_the_script_runs_from_another_folder_and_rebuilds_the_committed_deck(self):
        with tempfile.TemporaryDirectory() as folder:
            output = Path(folder) / "again.pptx"
            result = subprocess.run([sys.executable, "-B", str(SAMPLE_SCRIPT), "--output", str(output)],
                                    cwd=folder, stdout=subprocess.PIPE, stderr=subprocess.PIPE,
                                    universal_newlines=True, timeout=300)
            self.assertEqual(result.returncode, 0, result.stdout + result.stderr)
            self.assertTrue(output.is_file())
            self.assertEqual(design_lint.lint(str(output), profile="peers", minutes=12), [])
            self.assertEqual(deck_summary(output), deck_summary(SAMPLE_DECK),
                             "examples/sample-deck.pptx is stale: run examples/generate_sample_deck.py again")

    def test_the_script_leaves_the_default_deck_alone_when_it_gets_an_output_path(self):
        before = SAMPLE_DECK.stat().st_mtime
        with tempfile.TemporaryDirectory() as folder:
            subprocess.run([sys.executable, "-B", str(SAMPLE_SCRIPT), "--output", str(Path(folder) / "x.pptx")],
                           cwd=folder, stdout=subprocess.PIPE, stderr=subprocess.PIPE, universal_newlines=True,
                           timeout=300)
        self.assertEqual(SAMPLE_DECK.stat().st_mtime, before)


# ---------------------------------------------------------------------------
# The READMEs
# ---------------------------------------------------------------------------

class PublicText(unittest.TestCase):
    TEXT_SUFFIXES = (".py", ".md", ".yaml", ".yml", ".json", ".txt")

    def test_no_shipped_text_file_holds_a_private_path_an_address_or_the_template_author(self):
        """Scan every text file of the plugin except the tests, which hold the seeded strings of the detectors."""
        scanned = []
        for path in sorted(PLUGIN.rglob("*")):
            relative = path.relative_to(PLUGIN)
            if not path.is_file() or "__pycache__" in relative.parts or relative.parts[0] == "tests":
                continue
            if path.suffix not in self.TEXT_SUFFIXES:
                continue
            scanned.append(str(relative))
            with self.subTest(file=str(relative)):
                self.assertEqual(private_matches(read(path)), [])
        for needed in ("README.md", "commands/generate-slides-minimal.md", "references/design-principles.md",
                       "examples/generate_sample_deck.py", "scripts/slide_helpers.py"):
            self.assertIn(needed, scanned)

    def test_the_root_readme_holds_no_private_text(self):
        self.assertEqual(private_matches(read(ROOT_README)), [])


def yaml_blocks(text):
    return [block for language, block in fenced_blocks(text) if language in ("yaml", "yml")]


class Readme(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.text = read(README)
        cls.flat = " ".join(cls.text.split())

    def test_every_plugin_path_that_the_readme_names_exists(self):
        self.assertEqual(missing_paths(self.text, PLUGIN), [])

    def test_the_readme_names_the_flags_the_profiles_and_the_themes(self):
        for flag in ("--profile", "--theme", "--no-notes", "--render", "--minutes", "--project-dir"):
            with self.subTest(flag=flag):
                self.assertIn(flag, self.text)
        for name in style_mod.PROFILES + style_mod.THEMES:
            with self.subTest(name=name):
                self.assertIn("`{}`".format(name), self.text)
        self.assertIn("/generate-slides-minimal", self.text)
        self.assertIn("slide-style-minimal.yaml", self.text)

    def test_the_readme_names_every_option_of_the_linter(self):
        options = [s for action in design_lint._parser()._actions for s in action.option_strings
                   if s.startswith("--") and s != "--help"]
        self.assertTrue(options)
        for option in options:
            with self.subTest(option=option):
                self.assertIn(option, self.text)

    def test_the_profile_table_matches_the_profile_files(self):
        found = []
        for line in self.text.splitlines():
            cells = [c.strip() for c in line.strip().strip("|").split("|")]
            if len(cells) != 6 or not re.match(r"^`(peers|lab-meeting|general)`$", cells[0]):
                continue
            profile = cells[0].strip("`")
            found.append(profile)
            lint = style_mod.load_style(profile, "light")["lint"]
            with self.subTest(profile=profile):
                words = lint["words_per_slide"]
                self.assertEqual([int(n) for n in re.findall(r"\d+", cells[2])], [words["warn"], words["error"]])
                body = lint["body_font"]
                self.assertEqual(set(int(n) for n in re.findall(r"\d+", cells[3])),
                                 {body["warn_below"], body["secondary_warn_below"]})
                self.assertEqual(float(cells[4]), lint["slide_count"]["per_minute"])
                self.assertEqual(int(cells[5]), lint["objects_per_slide"]["max"])
        self.assertEqual(sorted(found), sorted(style_mod.PROFILES))

    def test_the_readme_gives_the_count_of_slides_of_the_sample_deck(self):
        self.assertEqual(len(Presentation(str(SAMPLE_DECK)).slides), 8)
        self.assertIn("eight slides", self.text)

    def test_the_readme_says_that_the_linter_reads_the_project_file_only_with_project_dir(self):
        self.assertIn("only with `--project-dir`", self.text)

    def test_the_linter_reads_the_project_file_only_with_project_dir(self):
        """The claim of the README, measured: a project file that lowers the word limit changes the findings only with the flag."""
        with tempfile.TemporaryDirectory() as folder:
            (Path(folder) / "slide-style-minimal.yaml").write_text("lint:\n  words_per_slide: {warn: 5}\n",
                                                                    encoding="utf-8")
            counts = []
            for extra in ([], ["--project-dir", "."]):
                result = subprocess.run([sys.executable, "-B", str(SCRIPTS_DIR / "design_lint.py"), str(SAMPLE_DECK),
                                         "--profile", "peers", "--minutes", "12"] + extra, cwd=folder,
                                        stdout=subprocess.PIPE, stderr=subprocess.PIPE, universal_newlines=True,
                                        timeout=300)
                self.assertEqual(result.returncode, 0, result.stdout + result.stderr)
                counts.append(int(result.stdout.splitlines()[-1].split(": ")[1]))
            self.assertEqual(counts[0], 0, "without the flag the linter ignores the project file")
            self.assertGreater(counts[1], 0, "with the flag the lower word limit gives warnings")

    def test_the_readme_gives_the_install_steps_and_the_exit_codes(self):
        self.assertIn("/plugin install generate-slides-minimal@yangjun9095-plugins", self.text)
        self.assertIn("python-pptx", self.text)
        self.assertIn("PyYAML", self.text)
        for code in ("`0`", "`1`", "`2`"):
            self.assertIn(code, self.text)
        self.assertIn("unittest", self.text)

    def test_the_readme_describes_the_look(self):
        for phrase in ("deep violet", "#6D4BD6", "near-black", "no fill", "no line under the title", "dark theme"):
            with self.subTest(phrase=phrase):
                self.assertIn(phrase, self.text)

    def test_the_readme_states_the_tiers_and_each_rule_tier_matches_the_code(self):
        for tier in ("T1", "T2", "T3"):
            self.assertIn(tier, self.text)
        cfg = style_mod.load_style("peers", "light")
        checked = []
        for line in self.text.splitlines():
            match = re.match(r"^\|\s*`([a-z-]+)`\s*\|.*\|\s*(T[123])\s*\|\s*$", line)
            if not match or match.group(1) not in design_lint._KEYS:
                continue
            rule, tier = match.groups()
            checked.append(rule)
            with self.subTest(rule=rule):
                self.assertEqual(tier, cfg["lint"][design_lint._KEYS[rule]]["tier"])
        self.assertEqual(sorted(set(design_lint._KEYS) - set(checked)), [],
                         "the README table must list every rule of the linter")

    def test_the_readme_lists_the_rows_that_have_no_key_in_the_lint_block(self):
        rows = {}
        for line in self.text.splitlines():
            match = re.match(r"^\|\s*`([a-z-]+)`\s*\|.*\|\s*(error|warn)\s*\|\s*(T[123])\s*\|\s*$", line)
            if match:
                rows.setdefault(match.group(1), []).append(match.groups()[1:])
        for rule, tier in design_lint._FIXED_TIERS.items():
            with self.subTest(rule=rule):
                self.assertIn(tier, [row[1] for row in rows.get(rule, [])])
        self.assertIn(("warn", "T3"), rows.get("style-key", []))
        self.assertIn(("error", "T2"), rows.get("slide-count", []), "a deck with no slide is an error")

    def test_the_readme_says_how_skipped_checks_are_reported(self):
        self.assertIn("groups the skipped checks of a slide by rule and reason", self.text)

    def test_the_readme_and_the_command_give_the_same_tier_legend(self):
        _, body = split_front_matter(read(COMMAND))
        for tier, meaning in TIER_LEGEND.items():
            row = "| {} | {} |".format(tier, meaning)
            with self.subTest(tier=tier):
                self.assertIn(row, self.text)
                self.assertIn(row, body)

    def test_the_readme_states_the_limits_plainly(self):
        required = {
            "false render-font row": (
                "A chart in Calibri at 12 to 16 pt, inside a deck whose other text is Avenir Light, can give a false "
                "`render-font` row.", "Decks from the plugin helpers are not affected, because the chart helper sets the font."),
            "right-to-left words": ("does not test right-to-left words for presence in the render",),
            "bold and italic": ("Bold and italic text is not calibrated in the render check.",
                                "The check can miss small italic text."),
            "the verb check is a heuristic": ("heuristic with known misses", "scripts/title_check.py"),
            "lines per paragraph": ("The lines-per-paragraph rule is an estimate.", "It does not use font metrics."),
            "Avenir Light must be installed": ("Avenir Light must be installed for the intended look.",
                                               "falls back to another font"),
            "no animations": ("writes no animations",),
            "no claim about audiences": ("makes no claim about audience outcomes",),
            "LibreOffice only": ("The tests and the scenario runs opened the decks in LibreOffice only.",
                                 "Nobody tried PowerPoint."),
            "the tested versions": ("The tests ran with python-pptx 0.6.23 on Python 3.6.8 only.",
                                    "On Python 3.8 and later, pip installs python-pptx 1.0.x, and nobody ran the plugin "
                                    "with it."),
            "a saved builder script": ("A builder script in `docs/slides` holds the absolute path of the plugin.",
                                       "The script stops working after the plugin moves or updates.",
                                       "Run the command again to write a new script."),
            "a wider substitute font": ("A substitute font can be wider, so a headline near the two-line limit can wrap "
                                        "to three lines.",),
        }
        self.assertEqual(missing_phrases(self.text, required), [])

    def test_the_readme_says_where_the_linter_reads_a_title(self):
        required = {"title placeholder or shape named Title": (
            "The linter reads a title only from the title placeholder or a shape named `Title`, and v1 decks and many "
            "hand-made decks use text boxes.",
            "Such a deck gets a `title-empty` error on every slide. It can also get a `margin` warning for its first "
            "text box.")}
        self.assertEqual(missing_phrases(self.text, required), [])

    def test_the_readme_gives_the_row_that_the_linter_prints_without_minutes(self):
        result = subprocess.run([sys.executable, "-B", str(SCRIPTS_DIR / "design_lint.py"), str(SAMPLE_DECK),
                                 "--profile", "peers"],
                                stdout=subprocess.PIPE, stderr=subprocess.PIPE, universal_newlines=True, timeout=300)
        self.assertEqual(result.returncode, 0, result.stdout + result.stderr)
        lines = result.stdout.splitlines()
        self.assertEqual(lines[-2:], ["errors: 0", "warnings: 1"])
        self.assertTrue(lines[0].startswith("deck: slide-count: warn: "), lines[0])
        self.assertIn("Without it the linter prints `{}` and counts one warning.".format(lines[0]), self.flat)

    def test_the_readme_says_that_a_project_file_can_change_any_rule(self):
        self.assertIn("A project file can change the number or the tier of any rule.", self.flat)
        self.assertNotIn("any number of the linter", self.flat)

    def test_the_install_section_covers_an_externally_managed_python(self):
        install = self.text.split("## Install", 1)[1].split("\n## ", 1)[0]
        self.assertIn("If pip says `externally-managed-environment`, create a virtual environment. Install the three "
                      "packages in it. Start Claude Code from a shell where that environment is active.",
                      " ".join(install.split()))

    def test_the_readme_names_the_relation_to_generate_slides(self):
        self.assertIn("`generate-slides`", self.text)
        self.assertRegex(self.flat, r"(?i)v1.{0,80}unchanged|unchanged.{0,80}v1")
        self.assertRegex(self.flat, r"(?i)independent|shares no code")

    def test_the_readme_links_the_research_notes_and_the_principles(self):
        self.assertIn("references/research-notes/", self.text)
        self.assertIn("references/design-principles.md", self.text)
        self.assertIn("examples/generate_sample_deck.py", self.text)

    def test_the_yaml_example_of_the_readme_loads_and_lints(self):
        blocks = yaml_blocks(self.text)
        self.assertTrue(blocks, "the README must show a slide-style-minimal.yaml example")
        for block in blocks:
            with tempfile.TemporaryDirectory() as folder:
                (Path(folder) / "slide-style-minimal.yaml").write_text(block, encoding="utf-8")
                merged = style_mod.load_style("peers", "light", folder)
                self.assertNotEqual(merged, style_mod.load_style("peers", "light"), "the example must change something")
                findings = design_lint.lint(str(SAMPLE_DECK), profile="peers", minutes=12, project_dir=folder)
                self.assertFalse([f for f in findings if f.rule == "style-key"], "the example holds an unknown lint key")

    def test_the_readme_has_no_private_text(self):
        self.assertEqual(private_matches(self.text), [])
        self.assertNotIn("demo/", self.text)

    def test_prose(self):
        self.assertEqual(prose_problems(self.text), [])
        self.assertEqual(term_drift(self.text), [], "call a research note a research note")


class RootReadme(unittest.TestCase):
    def test_the_root_readme_has_a_section_for_the_plugin(self):
        text = read(ROOT_README)
        self.assertRegex(text, r"(?m)^### `generate-slides-minimal`\s*$")
        self.assertIn("/plugin install generate-slides-minimal@yangjun9095-plugins", text)
        section = text.split("### `generate-slides-minimal`", 1)[1].split("\n### ", 1)[0].split("\n## ", 1)[0]
        self.assertLessEqual(len(section.strip().splitlines()), 30, "one short section")
        self.assertIn("generate-slides", section.replace("generate-slides-minimal", ""))

    def test_the_section_names_every_package_that_the_plugin_needs(self):
        text = read(ROOT_README)
        section = text.split("### `generate-slides-minimal`", 1)[1].split("\n### ", 1)[0].split("\n## ", 1)[0]
        requires = [line for line in section.splitlines() if line.startswith("**Requires:**")]
        self.assertEqual(len(requires), 1)
        for package in ("`python-pptx`", "`PyYAML`", "Pillow"):
            with self.subTest(package=package):
                self.assertIn(package, requires[0])

    def test_the_section_follows_the_writing_rules(self):
        text = read(ROOT_README)
        section = text.split("### `generate-slides-minimal`", 1)[1].split("\n### ", 1)[0].split("\n## ", 1)[0]
        self.assertEqual(prose_problems(section), [])
        self.assertEqual(term_drift(section), [])
        self.assertEqual(private_matches(section), [])


if __name__ == "__main__":
    unittest.main()
