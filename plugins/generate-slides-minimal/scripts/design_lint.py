"""Design linter of generate-slides-minimal.

Usage: python3 design_lint.py DECK.pptx [--profile peers|lab-meeting|general] [--minutes N]
           [--theme light|dark] [--no-notes] [--render] [--project-dir DIR]

The linter reads a finished .pptx deck and checks it against the rules of the spec. Each
finding has the evidence tier of its threshold: T1 (a standard or a study), T2 (a named
source with no evidence) or T3 (folklore, or derived by the author). The tier comes from the
`lint` block of the style, so a project file can change both a number and its tier.

The theme has no effect on the findings, because no rule reads a color from the style.
The option loads the theme file, so a broken theme file is an error.

With --render, the linter also runs the render check of render_check.py. Its rows join the findings.
The check needs soffice and pdftotext on the PATH. Without them, the output ends with one line:
render check skipped: soffice or pdftotext not found. The check uses the source_font floor of the style.

Output: one row per finding as `slide N: rule: severity: tier: message`. A deck-level row
starts with `deck:`. Two lines follow: `errors: E` and `warnings: W`. The exit code is 0 when
no finding is an error, 1 when one is, and 2 for a usage or read error.

Besides the rules of the spec, the linter writes two kinds of note. `unreadable-shape` says that a shape could not be read.
`style-key` says that a project file holds a lint key that no rule reads, which is often a typo.

The linter never guesses a value. It skips a rule for a value that it cannot resolve, or that
only a:pPr/a:defRPr sets (PowerPoint may ignore it). One warning per slide and rule says so.
The warning carries the id of the skipped rule. A slide with a shape that python-pptx cannot
read gets a warning with the rule id `unreadable-shape` and the count of the shapes.
A slide with a kind (p:cSld/@name, written by slide_helpers) gets the layout rules.
These use the shape names Callout, Subtitle, Baseline, Number, Claim and Section name.
A take-home slide with no Baseline shape holds each baseline as paragraph 1 of its Claim shape.
The rules read paragraph 0 as the claim and the paragraphs after it as the baseline.
A take-home slide with a Baseline shape reads each Claim shape as claim text only.
A slide with no kind gets the generic rules only. A kind-less slide with a centered title
placeholder (ctrTitle) counts as a title slide.

Limits:
  - Text sizes, colors and fonts come from pptx_resolve, with the limits of that module.
  - Bullets of a body placeholder count as set unless a paragraph sets a:buNone.
  - Text with no fill of its own takes the color of the nearest filled shape below it that holds
    its box: a slide shape drawn earlier, then a layout shape, then a master shape. With no such shape,
    the slide background shows. A picture, a gradient or a fill that cannot be read gives a skip.
    A chart frame or a table frame has no fill here, so text over one sees what lies below the frame.
  - A chart with no fill element counts as transparent. Chart text with no size or color of
    its own, and no chart default, is skipped.
  - The margin rule checks all four sides of a body shape. The title and the source line are checked
    on the left and right sides only, because the look puts the title 0.3 in from the top and the
    source line 0.25 in from the bottom.
  - A box ignores rotation, so inside-slide and margin skip a rotated shape and say so.
  - The verb check uses two word lists and the words around a word (see title_check.py).
    It can miss a verb, and it can pass a noun phrase that holds a word like a verb.
  - Lines per paragraph is an estimate of 0.5 em per character. It uses no font file.
"""
import argparse
import math
import re
import sys
from collections import OrderedDict, namedtuple

from pptx import Presentation
from pptx.oxml.ns import qn
from pptx.shapes.group import GroupShape

import design_chart
import design_color
import pptx_resolve as rs
import style
import title_check

Finding = namedtuple("Finding", "slide rule severity tier message")

# The kinds that slide_helpers writes to p:cSld/@name. The linter does not import the helpers.
KINDS = ("title", "assertion_evidence", "backup", "big_number", "roadmap", "take_home")

EMU = 914400.0
CHAR_EM = 0.5  # width of an average character in em, the estimate of slide_helpers
LARGE_TEXT_RATIO = 3.0  # WCAG 2.1 ratio for large text
BOLD_LARGE_PT = 14  # WCAG 2.1: bold text from this size counts as large
FULL_SLIDE_SHARE = 0.95  # a picture this large is a background picture
SOURCE_ZONE = 0.85  # a text box below this share of the slide height can be the source line
SOURCE_HEIGHT = 0.12  # and it must be no taller than this share

# The key of the lint block that holds the tier of each rule.
_KEYS = {
    "title-sentence": "title_words", "title-unique": "title_unique", "words-per-slide": "words_per_slide",
    "take-home-words": "take_home_words", "bullets": "bullets_per_list", "objects": "objects_per_slide",
    "lines-per-paragraph": "lines_per_paragraph", "body-font": "body_font", "title-font": "title_font",
    "fonts-per-slide": "fonts_per_slide", "text-contrast": "text_contrast", "mark-contrast": "mark_contrast",
    "red-green": "red_green", "inside-slide": "inside_slide", "margin": "margin", "aspect": "aspect",
    "picture-ppi": "picture_ppi", "alt-text": "alt_text", "slide-count": "slide_count",
    "take-homes": "take_homes", "notes": "notes",
}
# Two findings have no key in the lint block.
_FIXED_TIERS = {"title-empty": "T2", "unreadable-shape": "T3"}

# Shapes of slide_helpers that are parts of another object. A slide of a known kind counts them once.
_PARTS = frozenset(("Accent mark", "Callout rule", "Divider", "Section mark", "Section line", "Claim mark",
                    "Number label"))
_SECONDARY_NAMES = frozenset(("Callout", "Subtitle"))
_CHROME = frozenset(("dt", "ftr", "sldNum"))
_CAUSE = {"unresolved": "unresolved", "paragraph": "comes from paragraph defaults, which PowerPoint may ignore"}


# ---------------------------------------------------------------------------
# Reading a slide
# ---------------------------------------------------------------------------

def _clean(text, limit=60):
    """Return text on one line, cut to `limit` characters."""
    text = " ".join(str(text).split())
    return text if len(text) <= limit else text[:limit - 3] + "..."


def _name(shape):
    try:
        return _clean(shape.name or "")
    except rs.READ_ERRORS:
        return ""


def _flag(shape, attribute):
    try:
        return bool(getattr(shape, attribute, False))
    except rs.READ_ERRORS:
        return False


def _texts(shape):
    """Return the paragraph texts of a shape or a table. Charts hold their text elsewhere."""
    element = getattr(shape, "_element", None)
    return [] if element is None else [rs.paragraph_text(p) for p in element.iter(qn("a:p"))]


class _Leaf(object):
    """One shape of a slide, with the facts that the rules read. A group is expanded into its children."""

    def __init__(self, shape, group):
        element = shape._element
        self.shape, self.group, self.name = shape, group, _name(shape)
        self.box = rs.shape_box(shape)
        self.is_picture = element.tag == qn("p:pic")
        self.is_line = element.tag == qn("p:cxnSp")
        self.is_chart = _flag(shape, "has_chart")
        self.is_table = _flag(shape, "has_table")
        ph = rs.placeholder_element(shape)
        self.placeholder = None if ph is None else ph.get("type", "obj")
        cnv = rs.find_element(element, "p:nvSpPr", "p:cNvSpPr")
        self.is_autoshape = element.tag == qn("p:sp") and ph is None and (cnv is None or cnv.get("txBox") != "1")
        self.words = rs.word_count(shape)
        self.role = None  # "title", "source" or "chrome"
        self.index = 0  # the place in the drawing order of the slide


def _leaves(container):
    """Return (the leaf shapes of a slide, a layout or a master, the number of shapes that python-pptx could not read)."""
    leaves, unreadable = [], [0]

    def visit(parent, group):
        shapes, bad = rs.safe_shapes_report(parent)
        unreadable[0] += bad
        for shape in shapes:
            if isinstance(shape, GroupShape):
                visit(shape, shape._element if group is None else group)
            else:
                leaves.append(_Leaf(shape, group))

    visit(container, None)
    return leaves, unreadable[0]


_Run = namedtuple("_Run", "leaf cell paragraph run resolved cls label")


def _paragraph_index(p_el):
    """Return the place of an a:p element among the paragraphs of its text body."""
    parent = p_el.getparent()
    return 0 if parent is None else parent.findall(qn("a:p")).index(p_el)


def _is_claim_baseline(ctx, leaf, p_el):
    """Return True for a paragraph after paragraph 0 of a Claim shape on a take-home slide with no Baseline shape."""
    return ctx.claim_baselines and leaf.name == "Claim" and _paragraph_index(p_el) >= 1


def _runs(ctx):
    """Return the text runs of the slide with their resolved font values, their class and their label.

    The class is title, source, chrome, number, secondary or primary. A run with no text is left out.
    The label names the shape in a finding. The baseline of a Claim shape gets "Claim (baseline)".
    """
    runs = []
    for leaf in ctx.leaves:
        if leaf.is_chart or leaf.is_picture:
            continue
        if leaf.is_table:
            units = [(cell, paragraph, run) for cell in rs.table_cells(leaf.shape)
                     if rs.find_element(cell._tc, "a:txBody") is not None
                     for paragraph in cell.text_frame.paragraphs for run in paragraph.runs]
        else:
            units = [(None, paragraph, run) for paragraph, run in rs.iter_text_runs(leaf.shape)]
        for cell, paragraph, run in units:
            if not "".join(t.text or "" for t in run._r.findall(qn("a:t"))).strip():
                continue
            resolved = rs.resolve_run(ctx.prs, ctx.slide, leaf.shape, paragraph, run)
            baseline = cell is None and _is_claim_baseline(ctx, leaf, paragraph._p)
            runs.append(_Run(leaf, cell, paragraph, run, resolved, _text_class(ctx, leaf, cell, baseline),
                             leaf.name + " (baseline)" if baseline else leaf.name))
    return runs


def _text_class(ctx, leaf, cell, baseline):
    """Return the class of a text run. `baseline` is True for the baseline paragraph of a Claim shape."""
    if leaf.role is not None:
        return leaf.role
    if ctx.kind == "big_number" and leaf.name == "Number":
        return "number"
    if cell is not None or baseline:
        return "secondary"
    if ctx.kind is not None and (leaf.name in _SECONDARY_NAMES or (leaf.name == "Baseline" and ctx.kind == "take_home")):
        return "secondary"
    return "primary"


def _trusted(run, attribute):
    """Return (value, cause) of a resolved attribute of a run.

    The value is None when the resolver could not resolve it ("unresolved"), or when only
    a:pPr/a:defRPr sets it ("paragraph"). Otherwise the cause is None.
    """
    value = getattr(run.resolved, attribute)
    if value is None:
        return None, "unresolved"
    if run.resolved.sources.get(attribute) == rs.SOURCE_PARAGRAPH:
        return None, "paragraph"
    return value, None


def _kind(slide):
    name = rs.find_element(slide._element, "p:cSld")
    kind = None if name is None else name.get("name")
    return kind if kind in KINDS else None


class _Slide(object):
    """The facts about one slide that the rules read, and the findings so far."""

    def __init__(self, prs, slide, number, cfg, cache):
        self.prs, self.slide, self.number, self.cache = prs, slide, number, cache
        self.cfg, self.limits = cfg, cfg["lint"]
        self.kind = _kind(slide)
        self.findings = []
        self.skips = OrderedDict()
        self.leaves, self.unreadable = _leaves(slide)
        for position, leaf in enumerate(self.leaves):
            leaf.index = position
        self._base, self._behind = None, {}
        # Decks built before the Claim shape held its baseline have a Baseline shape. Their rules stay as they were.
        self.claim_baselines = self.kind == "take_home" and not any(leaf.name == "Baseline" for leaf in self.leaves)
        self.title = rs.find_title(slide)
        ph = None if self.title is None else rs.placeholder_element(self.title)
        # A slide with no kind and a centered title placeholder (ctrTitle) is a title slide.
        self.is_title = self.kind == "title" or (self.kind is None and ph is not None and ph.get("type") == "ctrTitle")
        self.is_main = not (self.is_title or self.kind == "backup")
        self.width, self.height = prs.slide_width, prs.slide_height
        self.background = rs.background_color(prs, slide)
        self._mark_roles()
        self.runs = _runs(self)
        self.charts = [design_chart.Chart(leaf.shape, leaf.name, self.background) for leaf in self.leaves if leaf.is_chart]

    def _mark_roles(self):
        named = False
        for leaf in self.leaves:
            if self.title is not None and leaf.shape._element is self.title._element:
                leaf.role = "title"
            elif leaf.name == "Source":
                leaf.role, named = "source", True
            elif leaf.placeholder in _CHROME:
                leaf.role = "chrome"
        if named or not self.height:
            return
        # With no shape named Source, the smallest text box at the bottom of the slide is the source line.
        found = [leaf for leaf in self.leaves
                 if leaf.role is None and leaf.words and leaf.box is not None and not leaf.is_table
                 and leaf.box[1] >= SOURCE_ZONE * self.height and leaf.box[3] <= SOURCE_HEIGHT * self.height]
        if found:
            min(found, key=lambda leaf: leaf.box[2] * leaf.box[3]).role = "source"

    def add(self, rule, severity, message, key=None):
        tier = _FIXED_TIERS[rule] if rule in _FIXED_TIERS else self.limits[key or _KEYS[rule]]["tier"]
        self.findings.append(Finding(self.number, rule, severity, tier, message))

    def skip(self, rule, reason, name, key=None):
        """Note that a rule skipped a value. flush_skips() writes one warning per rule and reason."""
        self.skips.setdefault((rule, reason, key), []).append(name or "unnamed shape")

    def flush_skips(self):
        for (rule, reason, key), names in self.skips.items():
            unique = list(OrderedDict.fromkeys(names))
            shown = ", ".join(unique[:3]) + (" and {} more".format(len(unique) - 3) if len(unique) > 3 else "")
            self.add(rule, "warn", "{}; rule skipped ({})".format(reason, shown), key)


def _contains(outer, inner):
    return (outer[0] <= inner[0] and outer[1] <= inner[1]
            and outer[0] + outer[2] >= inner[0] + inner[2] and outer[1] + outer[3] >= inner[1] + inner[3])


def _hides(element):
    """Return True when a slide or a layout element sets showMasterSp to false."""
    return element.get("showMasterSp") in ("0", "false")


def _container_leaves(ctx, container):
    """Return the drawn leaves of a layout or a master, nearest first. The cache of the run holds them for every slide.

    A group is expanded into its children, with the transform of the group. A placeholder is not drawn on a slide.
    """
    if container.part not in ctx.cache:
        leaves, _ = _leaves(container)
        ctx.cache[container.part] = [leaf for leaf in reversed(leaves) if rs.placeholder_element(leaf.shape) is None]
    return ctx.cache[container.part]


def _base_leaves(ctx):
    """Return the leaves of the layout and then the master that the slide draws, nearest first.

    They lie below every shape of the slide. A placeholder is not drawn on the slide, so it is left out.
    A slide that hides background graphics draws neither. A layout that hides them draws no master shape.
    """
    if ctx._base is None:
        ctx._base = []
        try:
            layout = ctx.slide.slide_layout
            levels = [layout, layout.slide_master]
        except rs.READ_ERRORS:
            levels = []
        for number, container in enumerate(levels):
            if _hides(ctx.slide._element) or (number and _hides(levels[0]._element)):
                break
            ctx._base += _container_leaves(ctx, container)
    return ctx._base


def _fill_behind(ctx, leaf):
    """Return what lies behind a leaf: ("solid", color), ("unknown", None), ("picture", None) or None.

    None means that no shape lies behind it, so the slide background shows. The nearest shape below
    the leaf whose box holds the leaf and whose fill is not "none" decides. A filled shape with no
    usable box (a group with a zero child extent, for example) might hold the leaf, so it gives "unknown".
    """
    if leaf.index not in ctx._behind:
        found = None
        if leaf.box is not None:
            for other in list(reversed(ctx.leaves[:leaf.index])) + _base_leaves(ctx):
                if other.is_line or (other.box is not None and not _contains(other.box, leaf.box)):
                    continue
                if other.is_picture:
                    found = ("unknown", None) if other.box is None else ("picture", None)
                    break
                state, color = design_chart.fill_and_color(ctx.slide, other.shape, False)
                if state != "none":
                    # A shape with a box that python-pptx cannot read may lie behind the text: the fill is unknown.
                    found = ("unknown", None) if other.box is None else (state, color)
                    break
        ctx._behind[leaf.index] = found
    return ctx._behind[leaf.index]


# ---------------------------------------------------------------------------
# Rules
# ---------------------------------------------------------------------------

def _check_unreadable(ctx):
    if ctx.unreadable:
        ctx.add("unreadable-shape", "warn",
                "{} shape(s) could not be read, so no rule checked them; open the slide and check it".format(
                    ctx.unreadable))


def _check_title(ctx, seen):
    text = " ".join(_texts(ctx.title)).strip() if ctx.title is not None else ""
    if not text:
        if ctx.title is None and ctx.unreadable:
            return  # the title can be one of the shapes that python-pptx could not read
        ctx.add("title-empty", "error",
                "no title placeholder or shape named Title holds text; add a headline that states the claim")
        return
    key = " ".join(text.lower().split())
    if key in seen:
        ctx.add("title-unique", "warn", "the title repeats the title of slide {}; give each slide its own claim"
                .format(seen[key]))
    else:
        seen[key] = ctx.number
    if ctx.is_title:
        return
    limits = ctx.limits["title_words"]
    problem = title_check.title_problem(text, limits["min"], limits["max"])
    if problem:
        ctx.add("title-sentence", "warn", problem)


def _check_words(ctx):
    if ctx.kind == "take_home":
        return
    limits = ctx.limits["words_per_slide"]
    count = sum(leaf.words for leaf in ctx.leaves if leaf.role is None)
    if count > limits["error"]:
        ctx.add("words-per-slide", "error", "{} words on the slide; the limit is {}. Cut the text or move it to "
                "the notes".format(count, limits["error"]))
    elif count > limits["warn"]:
        ctx.add("words-per-slide", "warn", "{} words on the slide; keep at most {}. Cut the text or move it to "
                "the notes".format(count, limits["warn"]))


def _paragraph_words(leaf):
    """Return the word count of each paragraph of a shape, in order. A shape with no text body gives []."""
    body = rs.find_element(leaf.shape._element, "p:txBody")
    return [] if body is None else [len(rs.paragraph_text(p).split()) for p in body.findall(qn("a:p"))]


def _take_home_parts(ctx):
    """Return [(label, lint key, words)] for the claims of a take-home slide, then for the baselines.

    With a Baseline shape on the slide, each Claim shape is a claim and each Baseline shape a baseline.
    With none, paragraph 0 of a Claim shape is the claim and the paragraphs after it are its baseline.
    """
    claims = [leaf for leaf in ctx.leaves if leaf.name == "Claim"]
    if not ctx.claim_baselines:
        baselines = [leaf for leaf in ctx.leaves if leaf.name == "Baseline"]
        return ([("claim", "max_claim", leaf.words) for leaf in claims]
                + [("baseline", "max_baseline", leaf.words) for leaf in baselines])
    counts = [_paragraph_words(leaf) or [leaf.words] for leaf in claims]
    return ([("claim", "max_claim", count[0]) for count in counts]
            + [("baseline", "max_baseline", sum(count[1:])) for count in counts if len(count) > 1])


def _check_take_home_words(ctx):
    if ctx.kind != "take_home":
        return
    limits = ctx.limits["take_home_words"]
    for label, key, count in _take_home_parts(ctx):
        if count > limits[key]:
            ctx.add("take-home-words", "warn", "a {} has {} words; keep at most {}".format(label, count, limits[key]))


def _bulleted(p_el, placeholder):
    p_pr = p_el.find(qn("a:pPr"))
    for child in [] if p_pr is None else p_pr:
        if child.tag == qn("a:buNone"):
            return False
        if child.tag in (qn("a:buChar"), qn("a:buAutoNum"), qn("a:buBlip")):
            return True
    return placeholder in ("obj", "body")


def _check_bullets(ctx):
    lists = []
    for leaf in ctx.leaves:
        if leaf.role is not None or leaf.is_table or leaf.is_chart or leaf.is_picture:
            continue
        body = rs.find_element(leaf.shape._element, "p:txBody")
        paragraphs = [] if body is None else [p for p in body.findall(qn("a:p")) if rs.paragraph_text(p).strip()]
        count = sum(1 for p in paragraphs if _bulleted(p, leaf.placeholder))
        if count >= 2:
            lists.append(count)
    longest, limit = max(lists or [0]), ctx.limits["bullets_per_list"]["max"]
    if longest > limit:
        ctx.add("bullets", "warn", "a list has {} bullets; keep at most {} per list".format(longest, limit))
    if len(lists) > ctx.limits["lists_per_slide"]["max"]:
        ctx.add("bullets", "warn", "the slide has {} bulleted lists; keep at most {}".format(
            len(lists), ctx.limits["lists_per_slide"]["max"]), key="lists_per_slide")


def _check_objects(ctx):
    units = set()
    for leaf in ctx.leaves:
        if leaf.role is not None or leaf.is_line:
            continue
        if ctx.kind is not None and (leaf.name in _PARTS or (leaf.name == "Baseline" and ctx.kind == "take_home")):
            continue
        units.add(id(leaf.group) if leaf.group is not None else id(leaf))
    limit = ctx.limits["objects_per_slide"]["max"]
    if len(units) > limit:
        ctx.add("objects", "warn", "the slide has {} objects; keep at most {}".format(len(units), limit))


def _estimated_lines(text, size_pt, width_in):
    """Estimate the lines of a paragraph at CHAR_EM per character. A line break starts a new line."""
    return sum(max(1, int(math.ceil(len(part) * CHAR_EM * size_pt / 72.0 / width_in - 1e-9)))
               for part in text.split("\n"))


def _inset(body_pr, name):
    """Return an inset of a text body in EMU. The default is 0.1 in."""
    raw = body_pr.get(name)
    return int(raw) if raw is not None and re.fullmatch(r"-?[0-9]+", raw) else 91440


def _check_lines(ctx):
    limit = ctx.limits["lines_per_paragraph"]["max"]
    paragraphs = OrderedDict()
    for run in ctx.runs:
        if run.cls not in ("source", "chrome") and not run.leaf.is_table:
            paragraphs.setdefault((run.leaf, run.paragraph._p), []).append(run)
    worst = OrderedDict()
    for (leaf, p_el), runs in paragraphs.items():
        sizes = [_trusted(run, "size_pt")[0] for run in runs]
        body_pr = rs.find_element(leaf.shape._element, "p:txBody", "a:bodyPr")
        if None in sizes or leaf.box is None or body_pr is None or body_pr.get("wrap") == "none":
            continue
        width = (leaf.box[2] - _inset(body_pr, "lIns") - _inset(body_pr, "rIns")) / EMU
        if width > 0:
            lines = _estimated_lines(rs.paragraph_text(p_el), max(sizes), width)
            key = (leaf, runs[0].label or leaf.name or "a shape")
            if lines > limit and lines > worst.get(key, 0):
                worst[key] = lines
    for (leaf, label), lines in worst.items():
        ctx.add("lines-per-paragraph", "warn", "{} has a paragraph of an estimated {} lines; keep at most {} "
                "(the estimate uses 0.5 em per character)".format(label, lines, limit))


def _lowest(found, key, size, label):
    if key not in found or size < found[key][0]:
        found[key] = (size, label)


def _check_body_font(ctx):
    limits = ctx.limits["body_font"]
    found = {}
    for run in ctx.runs:
        if run.cls in ("title", "chrome", "number"):
            continue
        size, cause = _trusted(run, "size_pt")
        if cause:
            ctx.skip("body-font", "size " + _CAUSE[cause], run.leaf.name)
        else:
            _lowest(found, run.cls, size, run.label or "a shape")
    for chart in ctx.charts:
        for size, _ in chart.texts:
            if size is None:
                ctx.skip("body-font", "size unresolved", chart.name + " (chart text)")
            else:
                _lowest(found, "secondary", size, chart.name + " (chart text)")
    source = ctx.limits["source_font"]["warn_below"]
    if "source" in found and found["source"][0] < source:
        ctx.add("body-font", "warn", "the source line is {:g} pt; use at least {:g} pt".format(
            found["source"][0], source), key="source_font")
    for cls, floor_key, noun in (("primary", "warn_below", "body text"),
                                 ("secondary", "secondary_warn_below", "secondary text such as call-outs, "
                                  "tables and chart text")):
        if cls not in found:
            continue
        size, label = found[cls]
        if size < limits["error_below"]:
            ctx.add("body-font", "error", "{} is {:g} pt; use at least {:g} pt".format(
                label, size, limits["error_below"]))
        elif size < limits[floor_key]:
            ctx.add("body-font", "warn", "{} is {:g} pt; use at least {:g} pt for {}".format(
                label, size, limits[floor_key], noun))


def _check_title_font(ctx):
    if ctx.kind == "title":
        return
    sizes = []
    for run in ctx.runs:
        if run.cls == "title":
            size, cause = _trusted(run, "size_pt")
            if cause:
                ctx.skip("title-font", "size " + _CAUSE[cause], run.leaf.name)
            else:
                sizes.append(size)
    floor = ctx.limits["title_font"]["warn_below"]
    if sizes and min(sizes) < floor:
        ctx.add("title-font", "warn", "the title is {:g} pt; use at least {:g} pt".format(min(sizes), floor))


def _check_fonts_per_slide(ctx):
    families, sizes = set(), set()
    for run in ctx.runs:
        if run.cls in ("source", "chrome"):
            continue
        family, cause = _trusted(run, "family")
        if cause:
            ctx.skip("fonts-per-slide", "font family " + _CAUSE[cause], run.leaf.name)
        else:
            families.add(family.lower())
        size, _ = _trusted(run, "size_pt")
        if size is not None:
            sizes.add(round(size, 2))
    limits, problems = ctx.limits["fonts_per_slide"], []
    if len(families) > limits["families"]:
        problems.append("{} font families; keep at most {}".format(len(families), limits["families"]))
    if len(sizes) > limits["sizes"]:
        problems.append("{} sizes; keep at most {}".format(len(sizes), limits["sizes"]))
    if problems:
        ctx.add("fonts-per-slide", "warn", "the slide uses " + " and ".join(problems))


def _contrast(ctx, found, label, color, background, size, bold):
    """Record the contrast of one text color against a background. One finding per label and color pair."""
    limits = ctx.limits["text_contrast"]
    ratio = design_color.contrast_ratio(color, background)
    large = size is not None and (size >= limits["large_text_pt"] or (size >= BOLD_LARGE_PT and bold))
    if ratio < limits["error_below"]:
        severity, floor = "error", limits["error_below"]
    elif size is not None and ratio < (LARGE_TEXT_RATIO if large else limits["warn_below"]):
        severity, floor = "warn", LARGE_TEXT_RATIO if large else limits["warn_below"]
    else:
        return
    found[(label, color, background)] = (severity, "{}: text color {} on {} has a contrast of {:.2f}:1; use at "
                                         "least {:g}:1".format(label, color, background, ratio, floor))


def _text_background(ctx, run):
    """Return the hex color behind a run, or None after a note on the slide.

    The fill of the shape or the table cell comes first. Without one, the nearest filled shape
    below the text decides, and the slide background shows when no shape lies below.
    """
    leaf = run.leaf
    state, color = design_chart.fill_and_color(ctx.slide, leaf.shape if run.cell is None else run.cell,
                                               run.cell is not None)
    if state == "solid":
        return color
    if state == "unknown":
        ctx.skip("text-contrast", "fill unknown", leaf.name)
        return None
    behind = _fill_behind(ctx, leaf)
    if behind is not None and behind[0] == "solid":
        return behind[1]
    if behind is not None:
        ctx.skip("text-contrast", "text over a picture has an unknown background" if behind[0] == "picture"
                 else "text over a filled shape with an unknown fill", leaf.name)
    elif ctx.background is None:
        ctx.skip("text-contrast", "slide background unresolved", leaf.name)
    else:
        return ctx.background
    return None


def _check_text_contrast(ctx):
    found = OrderedDict()
    for run in ctx.runs:
        color, cause = _trusted(run, "color")
        if cause:
            ctx.skip("text-contrast", "color " + _CAUSE[cause], run.leaf.name)
            continue
        background = _text_background(ctx, run)
        if background is not None:
            size, _ = _trusted(run, "size_pt")
            bold = run.resolved.bold is True and run.resolved.sources.get("bold") != rs.SOURCE_PARAGRAPH
            _contrast(ctx, found, run.label or "a shape", color, background, size, bold)
    for chart in ctx.charts:
        for size, color in chart.texts:
            if color is None:
                ctx.skip("text-contrast", "chart text color unresolved", chart.name)
            elif chart.background is None:
                ctx.skip("text-contrast", chart.why, chart.name)
            else:
                _contrast(ctx, found, chart.name + " (chart text)", color, chart.background, size, False)
    for severity, message in found.values():
        ctx.add("text-contrast", severity, message)


def _mark_candidates(ctx):
    """Yield (label, color) for each filled autoshape in the evidence area that holds no text.

    A shape whose box holds text, its own or that of another shape, is a label or a card. Its fill
    is a backdrop.
    """
    layout = ctx.cfg["layout"]
    top, bottom = layout["content_top"] * EMU, layout["content_bottom"] * EMU
    for leaf in ctx.leaves:
        if leaf.is_autoshape and leaf.box is not None:
            middle = leaf.box[1] + leaf.box[3] / 2.0
            holds_text = any(other.words and other.box is not None and _contains(leaf.box, other.box)
                             for other in ctx.leaves)
            if top <= middle <= bottom and not holds_text and rs.fill_state(leaf.shape) == "solid":
                yield leaf.name or "a shape", rs.fill_color(leaf.shape)


def _check_mark_contrast(ctx):
    floor = ctx.limits["mark_contrast"]["warn_below"]
    for chart in ctx.charts:
        for series in chart.series:
            if series.labelled:
                continue  # a mark with a data label is exempt
            if not series.colors:
                ctx.skip("mark-contrast", "series color comes from the theme", chart.name)
            elif chart.background is None:
                ctx.skip("mark-contrast", chart.why, chart.name)
            else:
                for color in series.colors:
                    ratio = design_color.contrast_ratio(color, chart.background)
                    if ratio < floor:
                        ctx.add("mark-contrast", "warn", "series {} of {} has the color {} with a contrast of "
                                "{:.2f}:1 against {}; use at least {:g}:1 or add data labels".format(
                                    series.name, chart.name, color, ratio, chart.background, floor))
    marks = list(_mark_candidates(ctx))
    if marks and ctx.background is None:
        ctx.skip("mark-contrast", "slide background unresolved", marks[0][0])
        return
    for label, color in marks:
        ratio = design_color.contrast_ratio(color, ctx.background)
        if ratio < floor:
            ctx.add("mark-contrast", "warn", "{} has the color {} with a contrast of {:.2f}:1 against the "
                    "background; use at least {:g}:1".format(label, color, ratio, floor))


def _red_green_pair(colors):
    for index, first in enumerate(colors):
        for second in colors[index + 1:]:
            if design_color.red_green_conflict(first, second):
                return first, second
    return None


def _check_red_green(ctx):
    groups = [(chart.name, [c for s in chart.series for c in s.colors]) for chart in ctx.charts]
    shapes = [rs.fill_color(leaf.shape) for leaf in ctx.leaves
              if not leaf.is_picture and not leaf.is_chart and rs.fill_state(leaf.shape) == "solid"]
    groups.append(("the filled shapes of the slide", shapes))
    for label, colors in groups:
        pair = _red_green_pair(list(OrderedDict.fromkeys(c for c in colors if c)))
        if pair:
            ratio = design_color.contrast_ratio(*pair)
            ctx.add("red-green", "warn", "{}: the colors {} and {} form a red and green pair that the hue heuristic "
                    "flags (luminance ratio {:.2f}:1); change one color or add a second cue".format(
                        label, pair[0], pair[1], ratio))


def _rotated(shape):
    """Return True when the shape or a group around it has a rotation that is not zero."""
    element = shape._element
    while element is not None and element.tag != qn("p:spTree"):
        for path in (("p:spPr", "a:xfrm"), ("p:grpSpPr", "a:xfrm"), ("p:xfrm",)):
            rotation = rs.find_element(element, *path)
            if rotation is not None and re.fullmatch(r"-?[0-9]+", rotation.get("rot", "0")) and int(rotation.get("rot", "0")):
                return True
        element = element.getparent()
    return False


def _check_geometry(ctx):
    """The inside-slide and margin rules, from the boxes of the leaf shapes.

    A box ignores rotation, so a rotated shape is skipped with one note for the slide.
    """
    if not ctx.width or not ctx.height:
        return
    margin = ctx.limits["margin"]["warn_below_in"] * EMU
    for leaf in ctx.leaves:
        if leaf.box is None:
            continue
        if _rotated(leaf.shape):
            ctx.skip("inside-slide", "inside-slide and margin: shape is rotated and its box ignores rotation", leaf.name)
            continue
        left, top, width, height = leaf.box
        over = [(left * -1, "left"), (top * -1, "top"), (left + width - ctx.width, "right"),
                (top + height - ctx.height, "bottom")]
        outside = max(over)
        if outside[0] > 0:
            ctx.add("inside-slide", "error", "{} extends {:.2f} in beyond the {} edge of the slide; move or "
                    "resize it".format(leaf.name or "a shape", outside[0] / EMU, outside[1]))
            continue
        if leaf.is_picture and width * height >= FULL_SLIDE_SHARE * ctx.width * ctx.height:
            continue  # a background picture runs to the edge
        gaps = [(left, "left"), (ctx.width - left - width, "right")]
        if leaf.role is None:  # the look puts the title near the top and the source line near the bottom
            gaps += [(top, "top"), (ctx.height - top - height, "bottom")]
        side = min(gaps)
        if side[0] < margin:
            ctx.add("margin", "warn", "{} is {:.2f} in from the {} edge; keep at least {:g} in".format(
                leaf.name or "a shape", side[0] / EMU, side[1], ctx.limits["margin"]["warn_below_in"]))


def _check_picture_ppi(ctx):
    floor = ctx.limits["picture_ppi"]["warn_below"]
    for leaf in ctx.leaves:
        if not leaf.is_picture:
            continue
        ppi = rs.effective_ppi(leaf.shape)
        if ppi is None:
            ctx.skip("picture-ppi", "picture resolution unknown", leaf.name)
        elif ppi < floor:
            ctx.add("picture-ppi", "warn", "{} shows at {:.1f} ppi; use an image with at least {:g} ppi at this "
                    "size".format(leaf.name or "a picture", ppi, floor))


_FILE_NAME = re.compile(r"\.(png|jpe?g|gif|bmp|tiff?|svg|emf|wmf|webp)$", re.IGNORECASE)
_GENERIC_ALT = re.compile(r"^(picture|image|chart|graphic|figure|photo|diagram|object|content placeholder|"
                          r"picture placeholder)( ?\d+)?$", re.IGNORECASE)


def _check_alt_text(ctx):
    for leaf in ctx.leaves:
        if not (leaf.is_picture or leaf.is_chart):
            continue
        props = (leaf.shape._element.xpath("./*[1]/p:cNvPr") or [None])[0]
        if props is not None and any(isinstance(el.tag, str) and el.tag.endswith("}decorative")
                                     and el.get("val") in ("1", "true") for el in props.iter()):
            continue  # PowerPoint: "Mark as decorative"
        text = " ".join(((props.get("descr") if props is not None else "") or "").split())
        if not text:
            problem = "has no alt text"
        elif _FILE_NAME.search(text):
            problem = "has a file name as alt text"
        elif _GENERIC_ALT.match(text) or text.lower() == leaf.name.lower():
            problem = "has a generic default as alt text"
        else:
            continue
        ctx.add("alt-text", "warn", "{} {}; describe what the {} shows".format(
            leaf.name or "a shape", problem, "chart" if leaf.is_chart else "picture"))


def _check_take_homes(ctx):
    rules = {"take_home": ("Claim", "claims", "take_homes"), "roadmap": ("Section name", "sections",
                                                                          "roadmap_sections")}
    if ctx.kind not in rules:
        return
    name, noun, key = rules[ctx.kind]
    count = sum(1 for leaf in ctx.leaves if leaf.name == name)
    if count > ctx.limits[key]["max"]:
        ctx.add("take-homes", "warn", "the slide has {} {}; keep at most {}".format(count, noun, ctx.limits[key]["max"]),
                key=key)


def _tokens(texts):
    return re.findall(r"\w+", " ".join(texts).lower())


def _notes_text(slide):
    try:
        if not slide.has_notes_slide:  # reading slide.notes_slide would create a notes slide
            return ""
        frame = slide.notes_slide.notes_text_frame
        return "" if frame is None else frame.text
    except rs.READ_ERRORS:
        return ""


def _check_notes(ctx):
    if not ctx.is_main:
        return
    text = _notes_text(ctx.slide)
    if not text.strip():
        ctx.add("notes", "warn", "the slide has no speaker notes; add a script that you speak in your own words")
        return
    words = _tokens(t for leaf in ctx.leaves if leaf.role not in ("source", "chrome") for t in _texts(leaf.shape))
    shared = set(_tokens([text]))
    if words:
        share = sum(1 for word in words if word in shared) / float(len(words))
        if share > ctx.limits["notes"]["redundancy_above"]:
            ctx.add("notes", "warn", "{:.0f}% of the words of the slide appear in the notes; make the notes add "
                    "to the slide".format(100 * share))


# ---------------------------------------------------------------------------
# The deck
# ---------------------------------------------------------------------------

def _open(path):
    try:
        return Presentation(str(path))
    except Exception as error:  # python-pptx and zipfile raise many types for a bad file
        raise ValueError("cannot read the deck {}: {}".format(path, _clean(error, 200)))


def _slides(prs):
    """Return ([(number, slide)], [numbers of the slides that python-pptx could not read])."""
    try:
        collection = prs.slides
        count = len(collection)
    except rs.READ_ERRORS:
        raise ValueError("the deck has no readable list of slides")
    slides, unreadable = [], []
    for index in range(count):
        try:
            slides.append((index + 1, collection[index]))
        except (KeyError, ValueError):
            unreadable.append(index + 1)
    return slides, unreadable


def _check_aspect(prs, cfg, findings):
    limits = cfg["lint"]["aspect"]
    width, height = prs.slide_width, prs.slide_height
    if not width or not height or width <= 0 or height <= 0:
        findings.append(Finding(0, "aspect", "error", limits["tier"], "the deck has no usable slide size"))
    elif abs(width / float(height) - limits["ratio"]) > limits["tolerance"]:
        findings.append(Finding(0, "aspect", "error", limits["tier"],
                                "the slide ratio is {:.3f}; use 16:9 ({:.3f})".format(width / float(height),
                                                                                        limits["ratio"])))


_SLIDE_CHECKS = (_check_words, _check_take_home_words, _check_bullets, _check_objects,
                 _check_lines, _check_body_font, _check_title_font, _check_fonts_per_slide, _check_text_contrast,
                 _check_mark_contrast, _check_red_green, _check_geometry, _check_picture_ppi, _check_alt_text,
                 _check_take_homes)


# The numbers that the rules read from the style. A project file can set any of them to a bad value.
_NUMBERS = (
    ("lint", "title_words", ("min", "max")), ("lint", "words_per_slide", ("warn", "error")),
    ("lint", "bullets_per_list", ("max",)), ("lint", "lists_per_slide", ("max",)),
    ("lint", "objects_per_slide", ("max",)), ("lint", "lines_per_paragraph", ("max",)),
    ("lint", "body_font", ("error_below", "warn_below", "secondary_warn_below")),
    ("lint", "take_home_words", ("max_claim", "max_baseline")), ("lint", "title_font", ("warn_below",)),
    ("lint", "source_font", ("warn_below",)), ("lint", "fonts_per_slide", ("families", "sizes")),
    ("lint", "text_contrast", ("error_below", "warn_below", "large_text_pt")),
    ("lint", "mark_contrast", ("warn_below",)), ("lint", "margin", ("warn_below_in",)),
    ("lint", "aspect", ("ratio", "tolerance")), ("lint", "picture_ppi", ("warn_below",)),
    ("lint", "slide_count", ("per_minute",)), ("lint", "take_homes", ("max",)),
    ("lint", "roadmap_sections", ("max",)), ("lint", "notes", ("redundancy_above",)),
    ("layout", None, ("content_top", "content_bottom")),
)
_TIERS = ("T1", "T2", "T3")
# The lint keys that the linter reads. Each one needs a tier.
_TIER_KEYS = frozenset(_KEYS.values()) | {"source_font", "lists_per_slide", "roadmap_sections"}


def _check_style(cfg):
    """Raise ValueError that names the key path of each number or tier that the rules cannot use."""
    for section, key, names in _NUMBERS:
        table = cfg.get(section, {}) if key is None else cfg.get(section, {}).get(key, {})
        for name in names:
            value = table.get(name) if isinstance(table, dict) else None
            if isinstance(value, bool) or not isinstance(value, (int, float)) or not math.isfinite(value):
                raise ValueError("{} must be a finite number, got {!r}".format(
                    ".".join(part for part in (section, key, name) if part), value))
    for key in sorted(_TIER_KEYS):  # a key that the linter does not read is not checked
        entry = cfg["lint"].get(key)
        if not isinstance(entry, dict) or entry.get("tier") not in _TIERS:
            raise ValueError("lint.{}.tier must be one of {}, got {!r}".format(
                key, ", ".join(_TIERS), entry.get("tier") if isinstance(entry, dict) else entry))


def lint(path, profile="peers", minutes=None, theme="light", notes=True, project_dir=None):
    """Return the findings for a deck as a list of Finding, ordered by slide (0 is the deck).

    Raise ValueError for an unknown profile or theme, a style file that cannot be read,
    a style number or tier that the rules cannot use, a minutes value that is not a finite
    number above 0, and a deck that cannot be read.
    """
    cfg = style.load_style(profile, theme, project_dir)
    _check_style(cfg)
    if minutes is not None and (isinstance(minutes, bool) or not isinstance(minutes, (int, float))
                                or not math.isfinite(minutes) or not minutes > 0):
        raise ValueError("minutes must be a finite number above 0, got {!r}".format(minutes))
    prs = _open(path)
    slides, unreadable = _slides(prs)
    findings = []
    _check_aspect(prs, cfg, findings)
    # A typo of a known key would otherwise change nothing. YAML can give keys of mixed types, so they sort as text.
    for key in sorted(set(cfg["lint"]) - _TIER_KEYS, key=str):
        findings.append(Finding(0, "style-key", "warn", "T3", "ignored unknown lint key '{}'".format(_clean(key))))
    tier = cfg["lint"]["slide_count"]["tier"]
    for number in unreadable:
        findings.append(Finding(number, "unreadable-shape", "warn", "T3", "the slide could not be read"))
    if not slides and not unreadable:
        findings.append(Finding(0, "slide-count", "error", tier, "the deck has no slides"))
        return findings
    shown = [(number, slide) for number, slide in slides if slide._element.get("show") not in ("0", "false")]
    if not shown and not unreadable:  # a hidden slide is not shown, so no rule applies to it
        findings.append(Finding(0, "slide-count", "error", tier, "the deck has no visible slides"))
        return findings
    seen, main, cache = {}, 0, {}  # the cache holds the leaves of each layout and master for this run
    for number, slide in shown:
        ctx = _Slide(prs, slide, number, cfg, cache)
        _check_unreadable(ctx)
        _check_title(ctx, seen)
        for check in _SLIDE_CHECKS:
            check(ctx)
        if notes:
            _check_notes(ctx)
        ctx.flush_skips()
        findings.extend(ctx.findings)
        if ctx.is_main:
            main += 1
    findings.extend(_check_count(cfg, main, minutes))
    findings.sort(key=lambda finding: finding.slide)
    return findings


def _check_count(cfg, main, minutes):
    limits = cfg["lint"]["slide_count"]
    if minutes is None:
        return [Finding(0, "slide-count", "warn", limits["tier"], "no --minutes given; rule skipped")]
    allowed = limits["per_minute"] * minutes
    if main > allowed:
        return [Finding(0, "slide-count", "warn", limits["tier"],
                        "{} main slides for {:g} {}; keep at most {:g}. Cut slides or move them to the "
                        "backup".format(main, minutes, "minute" if minutes == 1 else "minutes", allowed))]
    return []


# ---------------------------------------------------------------------------
# The command line
# ---------------------------------------------------------------------------

class _Usage(Exception):
    pass


class _Parser(argparse.ArgumentParser):
    def error(self, message):
        raise _Usage(message)


def _minutes(text):
    try:
        value = float(text)
    except ValueError:
        raise argparse.ArgumentTypeError("invalid number: {!r}".format(text))
    if not math.isfinite(value) or not value > 0:
        raise argparse.ArgumentTypeError("minutes must be a finite number above 0, got {!r}".format(text))
    return int(value) if value == int(value) else value


def _parser():
    parser = _Parser(prog="design_lint.py", description="Check a .pptx deck against the design rules.")
    parser.add_argument("deck", help="the .pptx file")
    parser.add_argument("--profile", default="peers", choices=style.PROFILES)
    parser.add_argument("--minutes", type=_minutes, help="the length of the talk, for the slide-count rule")
    parser.add_argument("--theme", default="light", choices=style.THEMES,
                        help="the style theme; it has no effect on the findings, because no rule reads a color from the style")
    parser.add_argument("--no-notes", action="store_true", help="switch the notes rules off")
    parser.add_argument("--render", action="store_true", help="also run the optional render check")
    parser.add_argument("--project-dir", help="the folder with slide-style-minimal.yaml")
    return parser


def _write(stream, text):
    """Write text to a stream. A character that the stream cannot encode becomes "?" and never raises."""
    buffer = getattr(stream, "buffer", None)
    if buffer is None:
        stream.write(text)
        return
    stream.flush()
    buffer.write(text.encode(getattr(stream, "encoding", None) or "utf-8", "replace"))
    buffer.flush()


def format_finding(finding):
    """Return the output row of a finding: `slide N: rule: severity: tier: message`, or `deck: ...` for slide 0.

    The message is put on one line and cut to 400 characters.
    """
    return "{}: {}: {}: {}: {}".format("slide %d" % finding.slide if finding.slide else "deck", finding.rule,
                                       finding.severity, finding.tier, _clean(finding.message, 400))


def _with_render(findings, args):
    """Return (findings, note) after the optional render check.

    The note is a line for the end of the output, or None. render_check imports Finding from this
    module, so the import waits until the option is on.
    """
    import render_check
    if not render_check.available():
        return findings, "render check skipped: soffice or pdftotext not found"
    min_pt = style.load_style(args.profile, args.theme, args.project_dir)["lint"]["source_font"]["warn_below"]
    rows = render_check.render_findings(args.deck, min_pt)
    return sorted(list(findings) + list(rows), key=lambda finding: finding.slide), None


def main(argv=None):
    """Run the linter. Return the exit code: 0 no error, 1 at least one error, 2 usage or read error."""
    try:
        args = _parser().parse_args(argv)
    except _Usage as error:
        _write(sys.stderr, "design_lint.py: error: {}\n".format(_clean(error, 300)))
        return 2
    except SystemExit as stop:  # --help
        return stop.code or 0
    try:
        findings = lint(args.deck, args.profile, args.minutes, args.theme, not args.no_notes, args.project_dir)
        note = None
        if args.render:
            findings, note = _with_render(findings, args)
    except (ValueError, OSError) as error:  # a usage or read error has a message by design, so no type name
        _write(sys.stderr, "design_lint.py: error: {}\n".format(_clean(error, 300)))
        return 2
    except Exception as error:  # a defect that the linter did not expect: report it with its type, do not hide it
        _write(sys.stderr, "design_lint.py: error: {}: {}\n".format(type(error).__name__, _clean(error, 300)))
        return 2
    errors = sum(1 for finding in findings if finding.severity == "error")
    lines = [format_finding(finding) for finding in findings]
    lines += ["errors: {}".format(errors), "warnings: {}".format(len(findings) - errors)]
    if note:
        lines.append(note)
    _write(sys.stdout, "\n".join(lines) + "\n")
    return 1 if errors else 0


if __name__ == "__main__":
    sys.exit(main())
