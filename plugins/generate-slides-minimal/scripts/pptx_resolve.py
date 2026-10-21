"""XML inheritance resolver of generate-slides-minimal.

python-pptx returns None for each size, family and color that a run inherits. This
module walks the inheritance chain in the XML and returns the value that the run has.
The design linter uses it to read decks that other tools built.

The chain of a run, from the highest priority down:
  1. the run properties, a:r/a:rPr
  2. the paragraph properties, a:pPr/a:defRPr (see the decision below)
  3. the list style of the text body of the shape, at the level of the paragraph
  4. the font reference of the shape style, p:style/a:fontRef
  5. the placeholder in the layout
  6. the placeholder in the master
  7. the title, body or other style of the master (p:txStyles)
  8. the default text style of the presentation

A normAutofit element with a fontScale attribute multiplies the size.

The resolver does not guess. A value that it cannot resolve is None, and its name is in
Resolved.unresolved. This happens in these cases:
  - no level of the chain sets the value (bold is the exception and then takes its base value, False)
  - a level sets a bad value, such as sz="big"
  - a color has a transform (lumMod, tint, alpha), or is a preset color, a gradient or a hyperlink
  - a placeholder has no layout part or no master part, so the chain has a hole
  - a table cell sets no bold or color itself, because its table style can set them
    (the resolver does not read tableStyles.xml)

No function raises on a malformed or foreign deck. A wrong kind of object is the one exception:
safe_shapes(42) raises the AttributeError of python-pptx. A test calls every public function on
decks with many kinds of defects to check this. No function changes the deck. python-pptx adds XML when
you read paragraph.font, paragraph.level or a text frame, so this module reads the XML directly.

python-pptx raises KeyError while it iterates the shapes of a slide that has a placeholder with
an invalid type. The linter should read shapes with safe_shapes_report(). It skips each shape
that python-pptx cannot read and counts it. The linter should add a deck-level or slide-level
note when the count is above 0. Then no slide with an unreadable placeholder looks like an empty slide.
fill_color() gives None for "no fill" and for "a fill that I cannot read". Use fill_state() to
tell them apart. A slide placeholder with no fill of its own can inherit one. Call
inherited_fill_state() and inherited_fill_color() to look in its layout and master placeholders.

Colors are "#RRGGBB" strings in upper case. Sizes are points. Boxes are EMU.

Limits:
  - normAutofit without fontScale is not modeled, because renderers shrink the text in their own way
  - baseline (superscript) and cap="small" are ignored
  - a:defPPr is ignored
  - rotation and flips are ignored
  - a:fld text counts in word_count, but iter_text_runs does not yield a field
  - a fill from p:style/a:fillRef is not resolved
  - fill_state() says unknown for a placeholder without an explicit fill, because the layout or master may fill it
  - a color with a transform is unresolved
  - table styles are not read
  - the counts of table cells and paragraphs can be low for a table style or a rich cell, and no function says so
"""
import copy
import itertools
import re

from lxml import etree
from pptx.opc.constants import RELATIONSHIP_TYPE as RT
from pptx.oxml.ns import qn
from pptx.shapes.group import GroupShape

try:
    from pptx.exc import InvalidXmlError
except ImportError:  # python-pptx before 0.6.22 names the module pptx.exceptions
    from pptx.exceptions import InvalidXmlError

EMU_PER_INCH = 914400.0

ATTRIBUTES = ("size_pt", "family", "color", "bold")

# The level that supplied a value. Resolved.sources holds one of these for each resolved value.
SOURCE_RUN = "run"
SOURCE_PARAGRAPH = "paragraph"
SOURCE_SHAPE = "shape"
SOURCE_SHAPE_STYLE = "shape_style"
SOURCE_LAYOUT = "layout"
SOURCE_MASTER = "master"
SOURCE_MASTER_STYLE = "master_style"
SOURCE_DEFAULT_STYLE = "default_style"
SOURCE_UNSET = "unset"  # no level sets bold, so bold takes its base value, False

# Decision on a:pPr/a:defRPr of a slide paragraph. python-pptx writes paragraph.font there.
# The schema allows the element. In the research render, LibreOffice 6.4 drew a run with no size
# at the defRPr size. That was 30 pt, where the placeholder style gives 32 pt.
# The research notes say that PowerPoint ignores the element on a slide paragraph.
# No PowerPoint was available to confirm it.
# The resolver follows the schema and the render, so defRPr sits between the run and the list
# style of the shape. It reports the source "paragraph" for each value that comes from there.
# The linter can pass paragraph_defaults=False to resolve_run to read the deck without defRPr.
# Many tool-built decks (python-pptx paragraph.font) set the size of their runs through defRPr.

_UNSET = object()  # this level does not set the value, so ask the next level
_BAD = object()  # this level sets a value that cannot be resolved, so stop
_BROKEN = object()  # the chain has a hole here, so every level below it is unknown
_TABLE_STYLE = object()  # the table style applies here, and it can set color and bold

# The six elements that can hold a fill in a:spPr or a:tcPr.
FILL_TAGS = frozenset(qn(name) for name in (
    "a:noFill", "a:solidFill", "a:gradFill", "a:blipFill", "a:pattFill", "a:grpFill"))
_THEME_COLORS = ("dk1", "lt1", "dk2", "lt2", "accent1", "accent2", "accent3", "accent4",
                 "accent5", "accent6", "hlink", "folHlink")
_MAPPED_COLORS = ("bg1", "tx1", "bg2", "tx2")
_TITLE_TYPES = ("title", "ctrTitle")
_OTHER_TYPES = ("dt", "ftr", "sldNum", "hdr")
_MASTER_STYLES = {"title": "p:titleStyle", "body": "p:bodyStyle", "other": "p:otherStyle"}
_HEX_DIGITS = re.compile(r"[0-9A-Fa-f]{6}")
_PARSER = etree.XMLParser(resolve_entities=False, no_network=True)
# The errors that python-pptx raises when it reads a defective part of a deck.
READ_ERRORS = (KeyError, ValueError, AttributeError, InvalidXmlError)


class Resolved(object):
    """The font values of one run after inheritance.

    size_pt is a float, family a typeface name, color a "#RRGGBB" string and bold a bool.
    A value that could not be resolved is None, and its name is in `unresolved`.
    `sources` maps the name of each resolved value to the level that set it:
    run, paragraph, shape, shape_style, layout, master, master_style, default_style or unset.
    """

    def __init__(self, size_pt=None, family=None, color=None, bold=None, unresolved=None, sources=None):
        self.size_pt = size_pt
        self.family = family
        self.color = color
        self.bold = bold
        self.unresolved = [] if unresolved is None else list(unresolved)
        self.sources = {} if sources is None else dict(sources)

    def __repr__(self):
        return "Resolved(size_pt={!r}, family={!r}, color={!r}, bold={!r}, unresolved={!r})".format(
            self.size_pt, self.family, self.color, self.bold, self.unresolved)


# ---------------------------------------------------------------------------
# The layout, the master, the theme and scheme colors
# ---------------------------------------------------------------------------

def find_element(element, *names):
    """Return the element at a path of prefixed names such as "p:cSld", or None."""
    for name in names:
        if element is None:
            return None
        element = element.find(qn(name))
    return element


def _layout_of(slide):
    try:
        return slide.slide_layout
    except (KeyError, ValueError, AttributeError):
        return None


def _master_of(layout):
    if layout is None:
        return None
    try:
        return layout.slide_master
    except (KeyError, ValueError, AttributeError):
        return None


def _theme_root(master):
    if master is None:
        return None
    try:
        return etree.fromstring(master.part.part_related_by(RT.THEME).blob, _PARSER)
    except (KeyError, ValueError, AttributeError, etree.XMLSyntaxError):
        return None


def _color_map(slide, layout, master):
    """Return the dict that maps bg1, tx1, bg2 and tx2 to theme colors, or None.

    The first override of the slide or the layout wins. Else the map of the master applies.
    """
    for owner in (slide, layout):
        override = find_element(getattr(owner, "_element", None), "p:clrMapOvr", "a:overrideClrMapping")
        if override is not None:
            return dict(override.attrib)
    clr_map = find_element(getattr(master, "_element", None), "p:clrMap")
    return None if clr_map is None else dict(clr_map.attrib)


def _hex(value):
    if value is None or _HEX_DIGITS.fullmatch(value) is None:
        return None
    return "#" + value.upper()


class _Scope(object):
    """The theme and the color map that the scheme colors of a slide resolve through."""

    def __init__(self, theme, color_map):
        self.theme = theme
        self.color_map = color_map

    @classmethod
    def of_slide(cls, slide):
        layout = _layout_of(slide)
        master = _master_of(layout)
        return cls(_theme_root(master), _color_map(slide, layout, master))

    @classmethod
    def of_shape(cls, shape):
        """Return the scope of a shape of a slide, a layout or a master. Raise a read error when it has none."""
        part = shape.part
        slide = getattr(part, "slide", None)
        if slide is not None:
            return cls.of_slide(slide)
        layout = getattr(part, "slide_layout", None)
        master = _master_of(layout) if layout is not None else getattr(part, "slide_master", None)
        if master is None:
            raise AttributeError("the shape has no slide, layout or master part")
        return cls(_theme_root(master), _color_map(None, layout, master))

    def scheme(self, name):
        """Return the hex color of a scheme color name, or None."""
        if name in _MAPPED_COLORS:
            name = (self.color_map or {}).get(name)
        if name not in _THEME_COLORS:
            return None
        slot = find_element(self.theme, "a:themeElements", "a:clrScheme", "a:" + name)
        if slot is None or len(slot) != 1:
            return None
        color = slot[0]
        if color.tag == qn("a:srgbClr"):
            return _hex(color.get("val"))
        if color.tag == qn("a:sysClr"):
            return _hex(color.get("lastClr"))
        return None

    def theme_font(self, which):
        """Return the Latin typeface of the "major" or "minor" theme font, or None."""
        latin = find_element(self.theme, "a:themeElements", "a:fontScheme", "a:{}Font".format(which), "a:latin")
        return None if latin is None else (latin.get("typeface") or None)

    def style_list(self, name):
        """Return a list of fill styles of the theme, such as "bgFillStyleLst", or None."""
        return find_element(self.theme, "a:themeElements", "a:fmtScheme", "a:" + name)


def _color_element(element, scope):
    """Return the hex color of an srgbClr or schemeClr element, or None.

    A color with a child transform, and any other kind of color, gives None.
    """
    if len(element):
        return None
    if element.tag == qn("a:srgbClr"):
        return _hex(element.get("val"))
    if element.tag == qn("a:schemeClr"):
        return scope.scheme(element.get("val"))
    return None


def _solid_color(solid_fill, scope):
    """Return the hex color of an a:solidFill element, or None."""
    children = [child for child in solid_fill if isinstance(child.tag, str)]
    if len(children) != 1:
        return None
    return _color_element(children[0], scope)


# ---------------------------------------------------------------------------
# The inheritance chain of a run
# ---------------------------------------------------------------------------

def _level(p_pr):
    """Return the level of a paragraph from 0 to 8, or None when it is not valid."""
    raw = None if p_pr is None else p_pr.get("lvl")
    if raw is None:
        return 0
    try:
        level = int(raw)
    except ValueError:
        return None
    return level if 0 <= level <= 8 else None


def _level_props(container, level):
    """Return the a:defRPr of a:lvl{n}pPr in a list style, or None."""
    return find_element(container, "a:lvl{}pPr".format(level + 1), "a:defRPr")


def placeholder_element(shape):
    """Return the p:ph element of a shape (a text shape, a picture or a graphic frame), or None."""
    element = getattr(shape, "_element", None)
    for holder in ("p:nvSpPr", "p:nvPicPr", "p:nvGraphicFramePr"):
        ph = find_element(element, holder, "p:nvPr", "p:ph")
        if ph is not None:
            return ph
    return None


def placeholder_class(ph_type):
    """Return "title", "body" or "other", the class of style that a placeholder type uses."""
    if ph_type in _TITLE_TYPES:
        return "title"
    if ph_type in _OTHER_TYPES:
        return "other"
    return "body"


def _placeholders(root):
    """Return the (p:sp, p:ph) pairs of the placeholder shapes under an element."""
    pairs = []
    for sp in root.iter(qn("p:sp")):
        ph = find_element(sp, "p:nvSpPr", "p:nvPr", "p:ph")
        if ph is not None:
            pairs.append((sp, ph))
    return pairs


def _layout_placeholder(layout, ph):
    """Return the p:sp of the layout that the slide placeholder `ph` inherits from, or None.

    The layout placeholder has the same index and the same class of type. A slide placeholder
    with no such partner skips the layout level, as python-pptx does.
    """
    kind = placeholder_class(ph.get("type", "obj"))
    for sp, other in _placeholders(layout._element):
        if other.get("idx", "0") == ph.get("idx", "0") and placeholder_class(other.get("type", "obj")) == kind:
            return sp
    return None


def _master_placeholder(master, ph_type):
    """Return the p:sp of the master that a placeholder of this type inherits from, or None.

    Date, footer, slide number and header placeholders match by exact type. The master has one
    of each, and each has its own size. Other placeholders match by class. A subtitle, an object
    or a picture placeholder reads the body placeholder.
    """
    kind = placeholder_class(ph_type)
    for sp, other in _placeholders(master._element):
        other_type = other.get("type", "obj")
        if kind == "other":
            matches = other_type == ph_type
        else:
            matches = placeholder_class(other_type) == kind
        if matches:
            return sp
    return None


def _font_ref_props(font_ref):
    """Return an a:defRPr that holds the typeface and the color of an a:fontRef."""
    props = etree.Element(qn("a:defRPr"))
    index = font_ref.get("idx")
    if index in ("major", "minor"):
        etree.SubElement(props, qn("a:latin")).set("typeface", "+mj-lt" if index == "major" else "+mn-lt")
    for child in font_ref:
        if child.tag in (qn("a:srgbClr"), qn("a:schemeClr")):
            etree.SubElement(props, qn("a:solidFill")).append(copy.deepcopy(child))
    return props


def _levels(prs, layout, master, shape, p_el, run, paragraph_defaults):
    """Return the (source, properties) pairs of the chain, highest priority first.

    A properties item is an a:rPr or a:defRPr element, or a marker for a hole in the chain.
    The a:pPr/a:defRPr level is in the chain only when `paragraph_defaults` is true.
    """
    levels = []
    r_pr = None if run is None else find_element(run._r, "a:rPr")
    if r_pr is not None:
        levels.append((SOURCE_RUN, r_pr))
    p_pr = find_element(p_el, "a:pPr")
    paragraph_props = find_element(p_pr, "a:defRPr")
    if paragraph_defaults and paragraph_props is not None:
        levels.append((SOURCE_PARAGRAPH, paragraph_props))
    level = _level(p_pr)
    if level is None:
        levels.append((None, _BROKEN))
        return levels
    text_body = p_el.getparent()
    own = _level_props(find_element(text_body, "a:lstStyle"), level)
    if own is not None:
        levels.append((SOURCE_SHAPE, own))
    default = _level_props(find_element(getattr(prs, "_element", None), "p:defaultTextStyle"), level)
    cell = text_body is not None and getattr(text_body.getparent(), "tag", None) == qn("a:tc")
    if cell:
        levels.append((None, _TABLE_STYLE))
    else:
        font_ref = find_element(getattr(shape, "_element", None), "p:style", "a:fontRef")
        if font_ref is not None:
            levels.append((SOURCE_SHAPE_STYLE, _font_ref_props(font_ref)))
        ph = placeholder_element(shape)
        if ph is not None and not _add_placeholder_levels(levels, layout, master, ph, level):
            return levels
    if default is not None:
        levels.append((SOURCE_DEFAULT_STYLE, default))
    return levels


def _add_placeholder_levels(levels, layout, master, ph, level):
    """Add the layout, master and master style levels of a placeholder.

    Return False when the chain has a hole. The hole is the last level in that case.
    """
    if layout is None:
        levels.append((None, _BROKEN))
        return False
    layout_sp = _layout_placeholder(layout, ph)
    props = _level_props(find_element(layout_sp, "p:txBody", "a:lstStyle"), level)
    if props is not None:
        levels.append((SOURCE_LAYOUT, props))
    if master is None:
        levels.append((None, _BROKEN))
        return False
    kind = placeholder_class(ph.get("type", "obj"))
    props = _level_props(find_element(_master_placeholder(master, ph.get("type", "obj")), "p:txBody", "a:lstStyle"), level)
    if props is not None:
        levels.append((SOURCE_MASTER, props))
    props = _level_props(find_element(master._element, "p:txStyles", _MASTER_STYLES[kind]), level)
    if props is not None:
        levels.append((SOURCE_MASTER_STYLE, props))
    return True


def _first(levels, extract, *, table_style_decides):
    """Return (value, source) from the first level whose extract() is not _UNSET.

    The value is _UNSET when no level sets it. It is _BAD when it cannot be resolved.
    """
    for source, props in levels:
        if props is _BROKEN:
            return _BAD, None
        if props is _TABLE_STYLE:
            if table_style_decides:
                return _BAD, None
            continue
        value = extract(props)
        if value is not _UNSET:
            return value, source
    return _UNSET, None


def _extract_size(props):
    """Return the size in hundredths of a point."""
    raw = props.get("sz")
    if raw is None:
        return _UNSET
    try:
        hundredths = int(raw)
    except ValueError:
        return _BAD
    return hundredths if hundredths > 0 else _BAD


def _extract_bold(props):
    raw = props.get("b")
    if raw is None:
        return _UNSET
    if raw in ("1", "true"):
        return True
    if raw in ("0", "false"):
        return False
    return _BAD


def _extract_family(props, scope):
    latin = props.find(qn("a:latin"))
    if latin is None:
        return _UNSET
    face = latin.get("typeface")
    if face in ("+mj-lt", "+mn-lt"):
        face = scope.theme_font("major" if face == "+mj-lt" else "minor")
    return face if face and not face.startswith("+") else _BAD


def _extract_color(props, scope):
    if props.find(qn("a:hlinkClick")) is not None:
        return _BAD
    for child in props:
        if child.tag in FILL_TAGS:
            color = _solid_color(child, scope) if child.tag == qn("a:solidFill") else None
            return _BAD if color is None else color
    return _UNSET


def _font_scale(text_body):
    """Return the font scale of a text body in 1/1000 of a percent, or None when it is malformed."""
    norm = find_element(text_body, "a:bodyPr", "a:normAutofit")
    raw = None if norm is None else norm.get("fontScale")
    if raw is None:
        return 100000
    try:
        scale = int(raw)
    except ValueError:
        return None
    return scale if scale > 0 else None


def _found(result):
    """Return (value, source) for a resolved value, else None."""
    value, source = result
    return None if value is _UNSET or value is _BAD else (value, source)


def resolve_run(prs, slide, shape, paragraph, run, paragraph_defaults=True):
    """Return the Resolved font values of a run, with all inheritance applied.

    `shape` is the shape that holds the run: a text shape, a group child or a table frame.
    Pass paragraph_defaults=False to leave out a:pPr/a:defRPr, which is the way that
    PowerPoint is said to read a deck. Compare both results to see what depends on it.
    """
    layout = _layout_of(slide)
    master = _master_of(layout)
    scope = _Scope(_theme_root(master), _color_map(slide, layout, master))
    p_el = paragraph._p
    levels = _levels(prs, layout, master, shape, p_el, run, paragraph_defaults)
    found = {}

    size, source = _first(levels, _extract_size, table_style_decides=False)
    scale = _font_scale(p_el.getparent())
    if scale is not None and _found((size, source)) is not None:
        found["size_pt"] = (size * scale / 10000000.0, source)
    found["family"] = _found(_first(levels, lambda props: _extract_family(props, scope),
                                    table_style_decides=False))
    found["color"] = _found(_first(levels, lambda props: _extract_color(props, scope),
                                   table_style_decides=True))
    bold, source = _first(levels, _extract_bold, table_style_decides=True)
    found["bold"] = (False, SOURCE_UNSET) if bold is _UNSET else _found((bold, source))

    result = Resolved()
    for name in ATTRIBUTES:
        if found.get(name) is None:
            result.unresolved.append(name)
        else:
            setattr(result, name, found[name][0])
            result.sources[name] = found[name][1]
    return result


# ---------------------------------------------------------------------------
# Text runs and words
# ---------------------------------------------------------------------------

def safe_shapes_report(container):
    """Return (shapes, unreadable_count) for a slide, a layout, a master or a group.

    python-pptx raises KeyError for a placeholder with an invalid type, for example. This function
    first reads all shapes in one pass, which costs one step per shape. When one pass raises, it reads
    the shapes one at a time with len() and the index. It skips each shape that python-pptx cannot
    read and counts it. A read by index costs a pass over all shapes in python-pptx, so only a
    container with an unreadable shape pays that price. A container whose shapes cannot be counted
    gives ([], 1). The linter should add a deck-level or slide-level note when the count is above 0.
    Then no slide with an unreadable placeholder looks like an empty slide.
    A wrong kind of object, such as a number, raises the AttributeError of python-pptx.
    """
    try:
        shapes = container.shapes
        count = len(shapes)
    except (KeyError, ValueError, InvalidXmlError):
        return [], 1
    try:
        return list(itertools.islice(shapes, count)), 0
    except (KeyError, ValueError, InvalidXmlError):
        pass  # one shape is unreadable: read them one by one and count the bad ones
    readable = []
    unreadable = 0
    for index in range(count):
        try:
            readable.append(shapes[index])
        except (KeyError, ValueError, InvalidXmlError):
            unreadable += 1
    return readable, unreadable


def safe_shapes(container):
    """Return an iterator over the shapes of a container that python-pptx can read.

    It skips the others and does not say how many. Use safe_shapes_report() when the count matters.
    """
    return iter(safe_shapes_report(container)[0])


def _shape_name(shape):
    """Return the name of a shape, or "" when the shape has no readable name."""
    try:
        return shape.name or ""
    except READ_ERRORS:
        return ""


def _has_table(shape):
    try:
        return bool(shape.has_table)
    except READ_ERRORS:
        return False


def table_cells(shape):
    """Return the cells of a table in row order, or an empty list when the shape is not a readable table."""
    try:
        return [cell for row in shape.table.rows for cell in row.cells]
    except READ_ERRORS:
        return []


def _paragraphs(shape):
    """Yield the python-pptx paragraphs of a shape, of the children of a group and of table cells.

    A shape with no text body yields nothing. python-pptx would add a text body on access.
    """
    if isinstance(shape, GroupShape):
        for child in safe_shapes(shape):
            for paragraph in _paragraphs(child):
                yield paragraph
    elif _has_table(shape):
        for cell in table_cells(shape):
            if find_element(cell._tc, "a:txBody") is not None:
                for paragraph in cell.text_frame.paragraphs:
                    yield paragraph
    elif find_element(getattr(shape, "_element", None), "p:txBody") is not None:
        for paragraph in shape.text_frame.paragraphs:
            yield paragraph


def iter_text_runs(shape):
    """Yield (paragraph, run) for each run of a shape, of a group child or of a table cell."""
    for paragraph in _paragraphs(shape):
        for run in paragraph.runs:
            yield paragraph, run


def paragraph_text(p_el):
    """Return the text of an a:p element: runs and fields joined, a line break as a newline."""
    parts = []
    for child in p_el:
        if child.tag in (qn("a:r"), qn("a:fld")):
            parts.append("".join(t.text or "" for t in child.findall(qn("a:t"))))
        elif child.tag == qn("a:br"):
            parts.append("\n")
    return "".join(parts)


def word_count(shape):
    """Return the number of words in a shape, a group or a table.

    A word is a run of characters between spaces. A line break also ends a word.
    """
    return sum(len(paragraph_text(paragraph._p).split()) for paragraph in _paragraphs(shape))


def find_title(slide):
    """Return the title placeholder of a slide, else the shape named "Title", else None.

    A slide whose shapes python-pptx cannot read gives None.
    """
    shapes = list(safe_shapes(slide))
    for shape in shapes:
        ph = placeholder_element(shape)
        if ph is not None and ph.get("type") in _TITLE_TYPES:
            return shape
    for shape in shapes:
        if _shape_name(shape) == "Title":
            return shape
    return None


# ---------------------------------------------------------------------------
# Backgrounds and fills
# ---------------------------------------------------------------------------

def _background_reference(reference, scope):
    """Return the hex color of a p:bgRef, which points into the fill styles of the theme."""
    try:
        index = int(reference.get("idx"))
    except (TypeError, ValueError):
        return None
    if index >= 1001:
        styles, number = scope.style_list("bgFillStyleLst"), index - 1000
    else:
        styles, number = scope.style_list("fillStyleLst"), index
    if styles is None or not 1 <= number <= len(styles):
        return None
    style = styles[number - 1]
    if style.tag != qn("a:solidFill"):
        return None
    colors = [child for child in style if isinstance(child.tag, str)]
    if len(colors) != 1:
        return None
    if colors[0].tag == qn("a:schemeClr") and colors[0].get("val") == "phClr" and not len(colors[0]):
        given = [child for child in reference if isinstance(child.tag, str)]
        return _color_element(given[0], scope) if len(given) == 1 else None
    return _color_element(colors[0], scope)


def _own_background(owner, scope):
    """Return the hex color or None, or _UNSET when the slide, layout or master has no p:bg."""
    background = find_element(getattr(owner, "_element", None), "p:cSld", "p:bg")
    if background is None:
        return _UNSET
    properties = background.find(qn("p:bgPr"))
    if properties is not None:
        for child in properties:
            if child.tag in FILL_TAGS:
                return _solid_color(child, scope) if child.tag == qn("a:solidFill") else None
        return None
    reference = background.find(qn("p:bgRef"))
    return None if reference is None else _background_reference(reference, scope)


def background_color(prs, slide):
    """Return the solid background color of a slide as hex, else None.

    The slide wins, then its layout, then the master. The first one with a background decides.
    A gradient, a picture or a pattern gives None, and the walk stops there.
    """
    layout = _layout_of(slide)
    scope = _Scope.of_slide(slide)
    for owner in (slide, layout, _master_of(layout)):
        if owner is not None:
            value = _own_background(owner, scope)
            if value is not _UNSET:
                return value
    return None


def _fill_properties(shape):
    """Return (the element that holds the fill, whether the object is a table cell).

    A shape holds its fill in p:spPr. A table cell holds it in a:tcPr. The element can be None.
    """
    cell = getattr(shape, "_tc", None)
    if cell is not None:
        return find_element(cell, "a:tcPr"), True
    return find_element(getattr(shape, "_element", None), "p:spPr"), False


def fill_color(shape):
    """Return the explicit solid fill of a shape or a table cell as hex, else None.

    None is ambiguous. It means "no fill". It also means "a fill that this module cannot read".
    Examples: a gradient, a pattern, a picture, a fill from p:style/a:fillRef, a color with a
    transform and a missing theme. Call fill_state() to tell these cases apart.
    """
    solid = find_element(_fill_properties(shape)[0], "a:solidFill")
    if solid is None:
        return None
    try:
        scope = _Scope.of_shape(shape)
    except (AttributeError, KeyError, ValueError):
        scope = _Scope(None, None)
    return _solid_color(solid, scope)


def _own_fill(shape):
    """Return the fill that a shape or a table cell sets itself, as (state, color), else None.

    The state is "solid", "none" or "unknown". None means that the object sets no fill.
    A p:style fill reference with an index above 0 counts as an unknown fill.
    """
    props, is_cell = _fill_properties(shape)
    for child in [] if props is None else list(props):
        if child.tag == qn("a:noFill"):
            return "none", None
        if child.tag == qn("a:solidFill"):
            color = fill_color(shape)
            return ("solid", color) if color is not None else ("unknown", None)
        if child.tag in FILL_TAGS:
            return "unknown", None
    reference = None if is_cell else find_element(shape._element, "p:style", "a:fillRef")
    if reference is None:
        return None
    try:
        index = int(reference.get("idx"))
    except (TypeError, ValueError):
        return "unknown", None
    return ("unknown", None) if index > 0 else None


def fill_state(shape):
    """Return "solid", "none" or "unknown" for the fill of a shape or a table cell.

    "solid": an explicit solid fill that fill_color() reads.
    "none": an explicit a:noFill. Also a shape with no fill element and no style fill, such as a text box.
    "unknown" in these cases:
      - a gradient, pattern or picture fill
      - a solid fill that fill_color() cannot read, for example a color with a transform
      - no explicit fill, and a p:style with an a:fillRef index above 0 (python-pptx autoshapes)
      - a table cell with no fill element, because its table style may fill it
      - a placeholder (text, picture or graphic frame) with no fill element, because its layout or master
        may fill it. inherited_fill_state() looks there.
      - an object that is not a shape
    """
    if getattr(shape, "_element", None) is None and getattr(shape, "_tc", None) is None:
        return "unknown"
    own = _own_fill(shape)
    if own is not None:
        return own[0]
    is_cell = getattr(shape, "_tc", None) is not None
    return "unknown" if is_cell or placeholder_element(shape) is not None else "none"


def _inherited_fill(slide, shape):
    """Return (state, color) of a slide shape with inheritance. See inherited_fill_state()."""
    own = _own_fill(shape)
    if own is not None:
        return own
    ph = placeholder_element(shape)
    if ph is None:
        return "none", None
    layout = _layout_of(slide)
    master = _master_of(layout)
    if layout is None or master is None:
        return "unknown", None
    levels = ((layout, _layout_placeholder(layout, ph)), (master, _master_placeholder(master, ph.get("type", "obj"))))
    for container, element in levels:
        owner = next((item for item in safe_shapes(container) if item._element is element), None)
        found = None if owner is None else _own_fill(owner)
        if found is not None:
            return found
    return "none", None


def inherited_fill_state(slide, shape):
    """Return "solid", "none" or "unknown" for the fill of a slide shape, looking through inheritance.

    A placeholder is read in this order: the shape, the layout placeholder that it inherits from and
    the master placeholder. The first level that sets a fill decides. "none" means that no level sets one.
    "unknown" means that a level sets a fill that this module cannot read, or that the slide has no
    layout or master. A shape that is not a placeholder gives the state of fill_state().
    """
    if placeholder_element(shape) is None:
        return fill_state(shape)
    return _inherited_fill(slide, shape)[0]


def inherited_fill_color(slide, shape):
    """Return the "#RRGGBB" color that inherited_fill_state() found, or None when the state is not "solid"."""
    return _inherited_fill(slide, shape)[1]


# ---------------------------------------------------------------------------
# Boxes and pictures
# ---------------------------------------------------------------------------

def _int_attributes(element, *names):
    """Return the integer attributes of an element as a tuple, or None."""
    try:
        return tuple(int(element.get(name)) for name in names)
    except (TypeError, ValueError):
        return None


def _group_scale(group):
    """Return ((off x, off y), (chOff x, chOff y), (scale x, scale y)) of a group, or None."""
    xfrm = find_element(group, "p:grpSpPr", "a:xfrm")
    parts = [find_element(xfrm, "a:" + name) for name in ("off", "ext", "chOff", "chExt")]
    if None in parts:
        return None
    off = _int_attributes(parts[0], "x", "y")
    ext = _int_attributes(parts[1], "cx", "cy")
    ch_off = _int_attributes(parts[2], "x", "y")
    ch_ext = _int_attributes(parts[3], "cx", "cy")
    if None in (off, ext, ch_off, ch_ext) or 0 in ch_ext:
        return None
    return off, ch_off, (ext[0] / float(ch_ext[0]), ext[1] / float(ch_ext[1]))


def shape_box(shape):
    """Return (left, top, width, height) of a shape in EMU, in slide coordinates.

    The transform of each parent group applies. Rotation and flips are ignored.
    The result is None when the shape or a parent group has no usable transform.
    """
    try:
        values = (shape.left, shape.top, shape.width, shape.height)
    except READ_ERRORS:
        return None
    if None in values:
        return None
    left, top, width, height = (float(value) for value in values)
    parent = shape._element.getparent()
    while parent is not None and parent.tag == qn("p:grpSp"):
        scale = _group_scale(parent)
        if scale is None:
            return None
        (off_x, off_y), (child_x, child_y), (scale_x, scale_y) = scale
        left = off_x + (left - child_x) * scale_x
        top = off_y + (top - child_y) * scale_y
        width, height = width * scale_x, height * scale_y
        parent = parent.getparent()
    return tuple(int(round(value)) for value in (left, top, width, height))


def effective_ppi(picture):
    """Return the pixels per inch of a picture on the slide, or None.

    The count uses the visible part of the image after cropping and the box on the slide.
    A picture that is stretched on one axis gives the lower of the two values.
    The result is None when the box is empty, the crop leaves nothing or the image is missing.
    """
    box = shape_box(picture)
    if box is None:
        return None
    try:
        pixels_w, pixels_h = picture.image.size
        left, right, top, bottom = (picture.crop_left, picture.crop_right,
                                    picture.crop_top, picture.crop_bottom)
    except READ_ERRORS + (OSError,):
        return None
    visible_w = pixels_w * (1 - left - right)
    visible_h = pixels_h * (1 - top - bottom)
    if box[2] <= 0 or box[3] <= 0 or visible_w <= 0 or visible_h <= 0:
        return None
    return min(visible_w / (box[2] / EMU_PER_INCH), visible_h / (box[3] / EMU_PER_INCH))
