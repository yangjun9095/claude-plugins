"""Optional render check of generate-slides-minimal.

`design_lint.py DECK --render` calls render_findings(). The check needs soffice (LibreOffice) and
pdftotext on the PATH. available() says whether both exist.

The check converts a copy of the deck to a PDF and reads the words of each page with
`pdftotext -bbox`. A hidden slide has no page, so the pages map to the visible slides in order.
The check stops with one deck-level row when the page count differs from the visible slide count.

Rules (all warn):
  render-font          T2  a word renders below the size min_pt. One row per slide.
  render-missing-text  T3  a word of the deck is not in the text of its page. One row per slide.
  render-check         T3  the check failed. One deck-level row. The check never raises for this.

Limits:
  - The height of a word is the font size times a factor that depends on the font. Avenir Light
    gives 1.37, Carlito gives 1.00. The check renders one word per font of the deck at 40 pt and
    divides each height by its factor. A word takes the smallest factor of the fonts that hold
    its text on the slide. A word with no match in the text frames, such as chart text, takes the
    smallest factor of the fonts that the deck uses, so the check can miss small chart text in
    a deck with two fonts.
  - The check measures with the fonts of this computer. A font that is not installed gives the
    size of the substitute font that LibreOffice picks.
  - A word below the floor by less than 1 percent passes, because the renderer rounds its heights.
  - The size check skips a word with no letter or digit, a word of a raised or lowered run (a
    subscript renders smaller by design), and a word with a character above U+00FF. LibreOffice
    draws a Greek letter or a rare symbol in another font, and that font has another factor.
  - The word compare skips a word with a right-to-left letter (Hebrew, Arabic), because pdftotext
    gives its letters in reverse order.
  - The word compare is a set compare per page. A word that appears twice in the deck and once
    in the PDF passes. The compare was tested on Latin text with Greek letters, and on one line
    each of Chinese and Cyrillic text. Other scripts are untested.
  - Bold and italic text is not calibrated: the calibration renders regular text. On this
    computer bold Avenir Light has the factor of regular text, and italic Avenir Light is 3.5
    percent taller, so the check errs on the quiet side. Other fonts were not measured.
  - pdftotext drops text that lies outside the page. That is how the check finds such text.
  - LibreOffice shrinks text with autofit in its own way, so PowerPoint may differ.
"""
import html
import itertools
import math
import os
import re
import shutil
import signal
import subprocess
import tempfile
import unicodedata
from collections import namedtuple
from pathlib import Path

from pptx import Presentation
from pptx.oxml.ns import qn
from pptx.shapes.group import GroupShape
from pptx.util import Inches, Pt

import pptx_resolve as rs
from design_lint import Finding

TIMEOUT_SECONDS = 120  # for each conversion and each pdftotext call
CALIBRATION_WORD = "Sample"
CALIBRATION_PT = 40
SIZE_MARGIN = 0.99  # a word counts as small below this share of the floor
JOIN_LIMIT = 4  # _page_tokens joins at most this many neighbor tokens. _is_joined checks a longer run.

_TIERS = {"render-font": "T2", "render-missing-text": "T3", "render-check": "T3"}
_PAGE = re.compile(r"<page\b[^>]*>(.*?)</page>", re.S)
_WORD = re.compile(r'<word xMin="-?[\d.]+" yMin="(-?[\d.]+)" xMax="-?[\d.]+" yMax="(-?[\d.]+)">([^<]*)</word>')
_TOKEN = re.compile(r"[^\W_]+")
_EDGE = re.compile(r"^[\W_]+|[\W_]+$")

# A slide of the deck: its number, its words as (shown, tokens), a dict of token to the fonts of
# the paragraphs that hold it (None is the theme font), and the tokens of raised or lowered runs.
_Slide = namedtuple("_Slide", "number words families shifted")


class _Failure(Exception):
    """The check could not run. The message is the reason."""


def available():
    """Return True when both soffice and pdftotext are on the PATH."""
    return shutil.which("soffice") is not None and shutil.which("pdftotext") is not None


def _row(slide, rule, message):
    return Finding(slide, rule, "warn", _TIERS[rule], message)


def _count(number, noun):
    return "{} {}{}".format(number, noun, "" if number == 1 else "s")


# ---------------------------------------------------------------------------
# Words
# ---------------------------------------------------------------------------

def _split(text):
    """Return the tokens of a text: NFKC, case folded, cut at every character that is no letter or digit.

    A soft hyphen is dropped first. A hyphen, a dash, a slash and a space all end a token.
    """
    text = unicodedata.normalize("NFKC", text).replace("\u00ad", "").casefold()
    return _TOKEN.findall(text)


def _right_to_left(word):
    """Return True when a word holds a letter that is written from right to left (Hebrew, Arabic)."""
    return any(unicodedata.bidirectional(char) in ("R", "AL") for char in word)


def _deck_tokens(word):
    """Return the tokens of a word of the deck that hold a letter. A token of digits only is left out."""
    return [token for token in _split(word) if any(char.isalpha() for char in token)]


def _page_tokens(page):
    """Return the set of tokens of a page, with the joins of up to JOIN_LIMIT neighbor tokens.

    A renderer can wrap a word at a hyphen, cut it at a baseline shift (H2O) or draw one glyph in
    another font. pdftotext then gives the word in parts. The join makes such a word match whole.
    """
    tokens = [token for _height, text in page for token in _split(text)]
    found = set(tokens)
    for start in range(len(tokens)):
        joined = tokens[start]
        for token in tokens[start + 1:start + JOIN_LIMIT]:
            joined += token
            found.add(joined)
    return found


def _page_list(page):
    """Return the tokens of a page in reading order."""
    return [token for _height, text in page for token in _split(text)]


def _is_joined(word, tokens):
    """Return True when a word equals the join of a run of neighbor tokens, of any length.

    A word of widely spaced letters comes from pdftotext as more than JOIN_LIMIT parts. The run
    starts and ends at a token edge, so the test finds no word inside a longer token or across two.
    """
    for start in range(len(tokens)):
        joined = ""
        for token in itertools.islice(tokens, start, None):
            joined += token
            if not word.startswith(joined):
                break
            if joined == word:
                return True
    return False


def _pages(text):
    """Return the pages of a `pdftotext -bbox` output as lists of (height, text) pairs of words."""
    pages = [[(float(match.group(2)) - float(match.group(1)), html.unescape(match.group(3)))
              for match in _WORD.finditer(body)] for body in _PAGE.findall(text)]
    if not pages:
        raise _Failure("the output of pdftotext has no page")
    return pages


# ---------------------------------------------------------------------------
# Reading the deck
# ---------------------------------------------------------------------------

def _leaves(container):
    """Yield the shapes of a slide or a group. The children of a group take the place of the group."""
    for shape in rs.safe_shapes(container):
        if isinstance(shape, GroupShape):
            yield from _leaves(shape)
        else:
            yield shape


def _family(prs, slide, shape, paragraph, run):
    try:
        return rs.resolve_run(prs, slide, shape, paragraph, run).family or None
    except rs.READ_ERRORS:
        return None


def _shifted(r_el):
    """Return True for a run that sits above or below the baseline."""
    properties = r_el.find(qn("a:rPr"))
    return properties is not None and properties.get("baseline") not in (None, "0")


def _paragraph_text(p_el):
    """Return the text of a paragraph from its runs, a line break as a space.

    A field is left out, because the renderer replaces its text (a date). A raised or lowered run
    is also left out, as a space. The renderer draws it small, and pdftotext puts it anywhere in the
    reading order, so it cannot take part in the compare.
    """
    parts = []
    for child in p_el:
        if child.tag == qn("a:r") and not _shifted(child):
            parts.append("".join(t.text or "" for t in child.findall(qn("a:t"))))
        elif child.tag in (qn("a:r"), qn("a:br")):
            parts.append(" ")
    return "".join(parts)


def _read_slide(prs, slide, number):
    paragraphs, shifted = [], set()  # a paragraph is (its element, the set of the fonts of its runs)
    for shape in _leaves(slide):
        for paragraph, run in rs.iter_text_runs(shape):
            if not paragraphs or paragraphs[-1][0] is not paragraph._p:
                paragraphs.append((paragraph._p, set()))
            paragraphs[-1][1].add(_family(prs, slide, shape, paragraph, run))
            if _shifted(run._r):
                shifted.update(_split(run.text))
    words, families = [], {}
    for element, fonts in paragraphs:
        for word in _paragraph_text(element).split():
            for token in _split(word):
                families.setdefault(token, set()).update(fonts)
            tokens = _deck_tokens(word)
            if tokens and not _right_to_left(word):
                words.append((_EDGE.sub("", word), tokens))
    return _Slide(number, words, families, shifted)


def _read_deck(path):
    """Return the _Slide of each visible slide, in order. A slide that python-pptx cannot read has no words."""
    try:
        prs = Presentation(str(path))
        slides = prs.slides
        count = len(slides)
    except Exception as error:  # python-pptx and zipfile raise many types for a bad file
        raise _Failure("cannot read the deck: {}".format(" ".join(str(error).split())[:120]))
    result = []
    for number in range(1, count + 1):
        try:
            slide = slides[number - 1]
        except (KeyError, ValueError):
            result.append(_Slide(number, [], {}, set()))
            continue
        if slide._element.get("show") not in ("0", "false"):
            result.append(_read_slide(prs, slide, number))
    return result


# ---------------------------------------------------------------------------
# Running the tools
# ---------------------------------------------------------------------------

def _kill(process):
    try:
        os.killpg(process.pid, signal.SIGKILL)
    except (ProcessLookupError, PermissionError):
        pass
    process.wait()


def _tail(log):
    """Return ": " and the last 200 bytes of a log file on one line, or "" when the log is empty."""
    try:
        text = " ".join(Path(str(log)).read_bytes()[-200:].decode("utf-8", "replace").split())
    except OSError:
        return ""
    return ": " + text if text else ""


def _run(command, seconds, log):
    """Run a command in its own process group. Raise _Failure for a start error, a timeout or a nonzero exit.

    The output goes to the file `log`, so a child that outlives the command cannot block this
    function. A timeout kills the group.
    """
    name = os.path.basename(command[0])
    with open(str(log), "wb") as handle:
        try:
            process = subprocess.Popen(command, stdin=subprocess.DEVNULL, stdout=handle, stderr=handle,
                                       start_new_session=True)
        except OSError as error:
            raise _Failure("cannot start {}: {}".format(name, error.strerror or error))
        try:
            code = process.wait(timeout=seconds)
        except subprocess.TimeoutExpired:
            _kill(process)
            raise _Failure("{} did not finish in {:g} s".format(name, seconds))
        except BaseException:
            _kill(process)
            raise
    if code != 0:
        raise _Failure("{} exited with code {}{}".format(name, code, _tail(log)))


def _render(deck, work, tools, seconds):
    """Convert a deck in the folder `work` to a PDF. Return its pages as `pdftotext -bbox` reads them.

    The profile of LibreOffice lies in `work`, so the profile of the user stays as it is. The child
    keeps the environment of the caller, so LibreOffice finds the fonts that the user installed.
    """
    soffice, pdftotext = tools
    pdf = work / (deck.stem + ".pdf")
    _run([soffice, "-env:UserInstallation=" + (work / "profile").as_uri(), "--headless", "--norestore",
          "--convert-to", "pdf", "--outdir", str(work), str(deck)], seconds, work / "soffice.log")
    if not pdf.is_file():
        raise _Failure("soffice wrote no PDF")
    out = work / (deck.stem + ".html")
    _run([pdftotext, "-bbox", str(pdf), str(out)], seconds, work / "pdftotext.log")
    if not out.is_file():
        raise _Failure("pdftotext wrote no text file")
    return _pages(out.read_text(encoding="utf-8", errors="replace"))


def _factors(families, work, tools, seconds):
    """Return a dict of font to (the height of a word) / (its font size), measured by one render.

    The render holds one slide per font, with a word of 40 pt. The key None is the font of the theme.
    The theme font is measured when a run resolves to no font and when the deck has no run (a chart only).
    """
    names = sorted(name for name in families if name)
    if None in families or not families:
        names.insert(0, None)
    prs = Presentation()
    for name in names:
        slide = prs.slides.add_slide(prs.slide_layouts[6])
        run = slide.shapes.add_textbox(Inches(1), Inches(1), Inches(6), Inches(2)).text_frame.paragraphs[0].add_run()
        run.text = CALIBRATION_WORD
        run.font.size = Pt(CALIBRATION_PT)
        if name:
            run.font.name = name
    deck = work / "calibration.pptx"
    prs.save(str(deck))
    pages = _render(deck, work, tools, seconds)
    if len(pages) != len(names):
        raise _Failure("cannot measure the font size: the calibration PDF has {} for {}".format(
            _count(len(pages), "page"), _count(len(names), "font")))
    factors = {}
    for name, page in zip(names, pages):
        heights = [height for height, text in page if text == CALIBRATION_WORD]
        if not heights or not heights[0] > 0:
            raise _Failure("cannot measure the font size of {}".format(repr(name) if name else "the default font"))
        factors[name] = heights[0] / CALIBRATION_PT
    return factors


# ---------------------------------------------------------------------------
# The rules
# ---------------------------------------------------------------------------

def _small_words(page, slide, factors, min_pt):
    """Return (size in pt, text) for each word of a page that renders below the floor."""
    small = []
    for height, text in page:
        tokens = _split(text)
        if not tokens or set(tokens) <= slide.shifted or any(ord(char) > 0xFF for char in text):
            continue  # a bullet, a raised or lowered run, or a glyph that the renderer may draw in another font
        fonts = set().union(*(slide.families.get(token, ()) for token in tokens)) or set(factors)
        size = height / min(factors[font] for font in fonts)
        if size < min_pt * SIZE_MARGIN:
            small.append((size, text))
    return small


def _missing_words(slide, page_tokens, tokens):
    """Return the words of the slide that the page lacks, each once, in deck order.

    A token is on the page when it is in `page_tokens` or the join of a longer run of `tokens`.
    """
    missing, seen = [], set()
    for shown, word_tokens in slide.words:
        key = tuple(word_tokens)
        if key in seen:
            continue
        if not all(token in page_tokens or _is_joined(token, tokens) for token in word_tokens):
            seen.add(key)
            missing.append(shown)
    return missing


def _failure_row(error):
    """Return the deck-level row for a _Failure, or for an OSError of a file (a full disk, no right to read)."""
    reason = "file error: {}".format(error) if isinstance(error, OSError) else error
    return _row(0, "render-check", "render check failed: {}".format(reason))


def _font_rows(slides, pages, factors, min_pt):
    rows = []
    for slide, page in zip(slides, pages):
        small = _small_words(page, slide, factors, min_pt)
        if small:
            size, word = min(small)
            rows.append(_row(slide.number, "render-font",
                             "{} rendered {} below {:g} pt. The smallest is {:.1f} pt (the word '{}'). "
                             "Check the font size and the autofit setting.".format(
                                 len(small), "word is" if len(small) == 1 else "words are", min_pt, size, word[:30])))
    return rows


def _missing_rows(slides, pages):
    rows = []
    for slide, page in zip(slides, pages):
        missing = _missing_words(slide, _page_tokens(page), _page_list(page))
        if missing:
            rows.append(_row(slide.number, "render-missing-text",
                             "{} of the deck {} not in the rendered page (for example: {}). "
                             "The text may lie outside the page or be cut off.".format(
                                 _count(len(missing), "word"), "is" if len(missing) == 1 else "are",
                                 ", ".join(word[:30] for word in missing[:3]))))
    return rows


def _check(path, min_pt, seconds):
    slides = _read_deck(path)
    if not slides:
        return []  # no visible slide, so no page
    tools = (shutil.which("soffice"), shutil.which("pdftotext"))
    if None in tools:
        raise _Failure("soffice or pdftotext not found")
    work = Path(tempfile.mkdtemp(prefix="render-check-")).resolve()
    try:
        shutil.copyfile(str(path), str(work / "deck.pptx"))
        pages = _render(work / "deck.pptx", work, tools, seconds)
        if len(pages) != len(slides):
            raise _Failure("the PDF has {} for {}".format(
                _count(len(pages), "page"), _count(len(slides), "visible slide")))
        missing = _missing_rows(slides, pages)  # this rule needs no calibration
        fonts = set().union(*(used for slide in slides for used in slide.families.values()))
        try:
            factors = _factors(fonts, work, tools, seconds)
        except (_Failure, OSError) as failure:
            return [_failure_row(failure)] + missing
    finally:
        shutil.rmtree(str(work), ignore_errors=True)
    rows = _font_rows(slides, pages, factors, min_pt) + missing
    return sorted(rows, key=lambda row: (row.slide, row.rule != "render-font"))


def render_findings(path, min_pt, timeout=TIMEOUT_SECONDS):
    """Return the Finding rows of the render check for a deck.

    `min_pt` is the smallest rendered font size in points that passes. The CLI of design_lint.py
    passes the source_font floor of the style. `timeout` is the limit in seconds for each conversion
    and each pdftotext call. A failure of the check gives one deck-level row (slide 0, render-check)
    and no exception. When only the font calibration fails, the render-missing-text rows stay.
    A min_pt that is no finite number above 0 raises ValueError.
    """
    if isinstance(min_pt, bool) or not isinstance(min_pt, (int, float)) or not math.isfinite(min_pt) \
            or not min_pt > 0:
        raise ValueError("min_pt must be a finite number above 0, got {!r}".format(min_pt))
    try:
        return _check(Path(path), min_pt, timeout)
    except (_Failure, OSError) as failure:
        return [_failure_row(failure)]
