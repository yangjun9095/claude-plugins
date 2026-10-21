"""Slide helpers of generate-slides-minimal: six layouts, speaker notes and alt text.

Derived from generate-slides v1, same author, MIT.

Each add_*_slide function takes the Presentation, the style dict from
style.load_style(), and the content of the slide. It returns the new slide.
Every color, font and size comes from the style dict. A function validates its
input before it adds a slide, so a failed call leaves the Presentation unchanged.

Text that needs more lines than its box holds raises ValueError. A headline takes
at most two lines, a callout two lines and a section name one line. The line
count is an estimate that uses no font file and errs on the wide side.

The linter reads shape names. This module sets these names: Title, Source,
Subtitle, Accent mark, Callout, Callout rule, Evidence chart, Evidence picture,
Evidence table, Number, Number label, Baseline, Divider, Section name,
Section line, Section mark, Claim and Claim mark. The Baseline shape is on the
big-number slide. On the take-home slide each baseline is paragraph 1 of its Claim shape.

Every length argument of fit_picture is a number in inches or a pptx.util length.
"""
import datetime
import math
import re
from collections import namedtuple
from pathlib import Path

from lxml import etree
from PIL import Image
from pptx import Presentation
from pptx.chart.data import CategoryChartData
from pptx.dml.color import RGBColor
from pptx.enum.chart import XL_CHART_TYPE, XL_LABEL_POSITION, XL_MARKER_STYLE
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, MSO_AUTO_SIZE, PP_ALIGN
from pptx.oxml.ns import qn
from pptx.util import Emu, Inches, Length, Pt

import design_color

EMU_PER_INCH = 914400

KINDS = ("title", "assertion_evidence", "backup", "big_number", "roadmap", "take_home")
PICTURE_FORMATS = ("BMP", "GIF", "JPEG", "PNG", "TIFF", "WMF")  # the formats python-pptx embeds
EVIDENCE_KINDS = ("picture", "table", "chart")
CHART_TYPES = ("bar", "line")

# Lengths are in inches, sizes in points. The style dict has no key for these values.
# A free-form layout places its shapes at fixed offsets from a reference content top of 2.1 inches.
# The offsets move with layout.content_top.
FREE_FORM_REFERENCE_TOP = 2.1
TITLE_SLIDE_PT = 44
TITLE_SLIDE_TOP, TITLE_SLIDE_HEIGHT = 2.1, 1.9
TITLE_SLIDE_MARK_TOP, TITLE_SLIDE_MARK_WIDTH = 4.2, 1.4
TITLE_SLIDE_SUBTITLE_TOP, TITLE_SLIDE_SUBTITLE_HEIGHT = 4.4, 1.2
NUMBER_TOP, NUMBER_HEIGHT, NUMBER_BOX_WIDTH = 2.6, 1.7, 5.4
NUMBER_LABEL_TOP, NUMBER_LABEL_HEIGHT = 4.3, 0.7
DIVIDER_LEFT, DIVIDER_TOP, DIVIDER_HEIGHT = 6.5, 2.7, 2.3
BASELINE_LEFT, BASELINE_TOP, BASELINE_HEIGHT = 6.85, 2.7, 2.3
ROADMAP_GAP = 0.3
SECTION_MARK_TOP = 2.9
SECTION_NAME_TOP, SECTION_NAME_HEIGHT = 3.15, 0.8
SECTION_LINE_TOP, SECTION_LINE_HEIGHT = 4.1, 1.8
SECTION_NAME_EXTRA_PT = 8  # a section name is this much larger than body_size
CLAIM_TOP, CLAIM_INDENT = 2.4, 0.45
CLAIM_MARK_OFFSET, CLAIM_MARK_SIZE = 0.12, 0.18
CLAIM_PAD, CLAIM_BLOCK_GAP = 0.39, 0.33  # inches below the text lines of a claim box, and between two boxes
TAKE_HOME_COUNTS = {1: "One thing", 2: "Two things", 3: "Three things"}
CALLOUT_TEXT_OFFSET = 0.25  # from the callout rule to the callout text
CALLOUT_TOP_OFFSET, CALLOUT_HEIGHT, CALLOUT_STEP = 0.1, 1.2, 1.55
SOURCE_GAP, SOURCE_HEIGHT = 0.25, 0.3  # the gap is below layout.content_bottom
TABLE_STYLE_ID = "{2D5ABB26-0587-4C30-8999-92F81FD0307C}"  # "No Style, No Grid"
TABLE_ROW_INCHES, TABLE_ROW_PAD, TABLE_CELL_MARGIN, TABLE_RULE_PT = 0.75, 0.2, 0.12, 1.5
TABLE_WEIGHT_MIN, TABLE_WEIGHT_MAX = 6, 40  # bounds of the text length that sizes a column
BAR_GAP_WIDTH = 70
CONTEXT_LINE_PT, ACCENT_LINE_PT = 2.5, 5
MARKER_PT, KEY_MARKER_PT = 9, 14
PLOT_LEFT, PLOT_TOP, PLOT_HEIGHT = 0.02, 0.05, 0.82  # fractions of the chart, for several line series
LABEL_PAD, MAX_LABEL_ROOM = 0.15, 0.4  # inches, and a fraction of the chart width
HEADLINE_LINES, CALLOUT_LINES = 2, 2  # two lines is the maximum of the lines-per-paragraph rule

# Estimates that use no font file. They err on the wide side so that a fallback font does not overflow.
CHAR_EM = 0.5  # width of an average character, in em
DIGIT_EM = 0.55
LINE_HEIGHT_EM = 1.2

# The characters below U+0020 that XML 1.0 forbids (all but tab, line feed and carriage return), and U+FFFE, U+FFFF.
_NOT_XML = re.compile("[\x00-\x08\x0b\x0c\x0e-\x1f\ufffe\uffff]")

_REQUIRED = {
    "typography": ("body_font", "title_size", "body_size", "callout_size", "source_size",
                   "big_number_size", "label_size"),
    "layout": ("slide_width", "slide_height", "margin", "title_top", "title_height", "content_top",
               "content_bottom", "evidence_width", "evidence_height", "callout_left", "rule_width"),
    "colors": ("background", "text", "muted", "accent", "accent_ink", "context"),
}


# ---------------------------------------------------------------------------
# Style access and units
# ---------------------------------------------------------------------------

def _value(style, section, key):
    try:
        return style[section][key]
    except (KeyError, TypeError):
        raise ValueError("the style has no value for {}.{}".format(section, key))


def _rgb(style, key):
    """Return the RGBColor of the style color `key`."""
    return RGBColor(*design_color.parse_hex(_value(style, "colors", key)))


def _check_style(style):
    """Raise ValueError when the style lacks a value that the layouts read."""
    for section, keys in _REQUIRED.items():
        for key in keys:
            value = _value(style, section, key)
            if section == "colors":
                design_color.parse_hex(value)


def _emu(value):
    """Return EMU for a pptx length or for a number of inches."""
    if isinstance(value, Length):
        return int(value)
    if isinstance(value, bool) or not isinstance(value, (int, float)):
        raise ValueError("a length must be a number of inches or a pptx length, got {!r}".format(value))
    return int(round(value * EMU_PER_INCH))


def _encodable(value, what):
    """Return `value` when a file can hold it. Raise ValueError otherwise.

    A lone surrogate cannot be saved. XML 1.0 forbids most control characters, such as NUL and U+000B.
    """
    try:
        value.encode("utf-8")
    except UnicodeEncodeError:
        raise ValueError("{} holds a character that cannot be saved, such as a lone surrogate".format(what))
    bad = _NOT_XML.search(value)
    if bad:
        raise ValueError("{} holds the control character U+{:04X}, which a .pptx file cannot hold. Remove it."
                         .format(what, ord(bad.group())))
    return value


def _text(value, what):
    """Return `value` when it is a non-blank string. Raise ValueError otherwise."""
    if not isinstance(value, str) or not value.strip():
        raise ValueError("{} must be a non-empty string, got {!r}".format(what, value))
    return _encodable(value, what)


def _optional_text(value, what):
    if value is None:
        return ""
    if not isinstance(value, str):
        raise ValueError("{} must be a string, got {!r}".format(what, value))
    return _encodable(value, what)


def _chars_per_line(width_in, size_pt):
    return max(1, int(width_in * 72 / (CHAR_EM * size_pt)))


def _estimate_lines(text, width_in, size_pt):
    """Estimate how many lines `text` takes in a box `width_in` inches wide. A long word breaks over lines."""
    per_line = _chars_per_line(width_in, size_pt)
    lines, used = 1, 0
    for word in text.split():
        if used and used + 1 + len(word) <= per_line:
            used += 1 + len(word)
            continue
        if used:
            lines += 1
        extra = (len(word) - 1) // per_line
        lines, used = lines + extra, len(word) - extra * per_line
    return lines


def _check_lines(text, width_in, size_pt, limit, what):
    """Raise ValueError when `text` needs more than `limit` lines."""
    lines = _estimate_lines(text, width_in, size_pt)
    if lines > limit:
        raise ValueError("{} needs about {} lines at {} pt but its box holds {}. Shorten it."
                         .format(what, lines, size_pt, limit))


def _line_height(size_pt):
    """Return the height in inches of one text line."""
    return size_pt * LINE_HEIGHT_EM / 72


def _check_headline(prs, style, text, what="the headline", size=None):
    """Raise ValueError when a headline needs more than two lines in the title box."""
    width_in = prs.slide_width / float(EMU_PER_INCH) - 2 * _value(style, "layout", "margin")
    size = _value(style, "typography", "title_size") if size is None else size
    _check_lines(text, width_in, size, HEADLINE_LINES, what)


# ---------------------------------------------------------------------------
# Presentation, kind, notes and alt text
# ---------------------------------------------------------------------------

def new_presentation(style):
    """Return an empty Presentation with the slide size of the style.

    The python-pptx template carries the name of its author, dates from 2013 and a comment. This
    function replaces them with blank fields and the current UTC time. A caller can set any field after.
    """
    prs = Presentation()
    prs.slide_width = Inches(_value(style, "layout", "slide_width"))
    prs.slide_height = Inches(_value(style, "layout", "slide_height"))
    properties = prs.core_properties
    now = datetime.datetime.utcnow()
    properties.author = ""
    properties.last_modified_by = ""
    properties.created = now
    properties.modified = now
    for name in ("title", "subject", "keywords", "comments", "category"):
        setattr(properties, name, "")
    properties.revision = 1
    return prs


def set_slide_kind(slide, kind):
    """Store the kind of a slide as the name attribute of p:cSld."""
    if kind not in KINDS:
        raise ValueError("unknown slide kind {!r}. Valid kinds: {}".format(kind, ", ".join(KINDS)))
    slide._element.cSld.set("name", kind)


def slide_kind(slide):
    """Return the kind of a slide, or None when the slide has none."""
    return slide._element.cSld.get("name")


def set_notes(slide, text):
    """Write the speaker notes. Empty or blank text writes nothing and adds no notes slide."""
    text = _optional_text(text, "the notes")
    if text.strip():
        slide.notes_slide.notes_text_frame.text = text


def set_alt_text(shape, text):
    """Write the descr attribute of the p:cNvPr of a shape."""
    if not isinstance(text, str):
        raise ValueError("alt text must be a string, got {!r}".format(text))
    _encodable(text, "the alt text")
    nodes = shape._element.xpath("./*[1]/p:cNvPr")
    if not nodes:
        raise ValueError("the shape has no p:cNvPr element to hold alt text")
    nodes[0].set("descr", text)


# ---------------------------------------------------------------------------
# Pictures
# ---------------------------------------------------------------------------

def _check_image(path):
    """Return (Path, (pixel width, pixel height)) of a readable image file."""
    if path is None:
        raise ValueError("the picture has no path")
    if str(path).strip() == "":
        raise ValueError("the picture path is empty")
    path = Path(str(path))
    if path.is_dir():
        raise ValueError("the picture path {} is a folder. Give the path of an image file.".format(path))
    if not path.exists():
        raise ValueError("the picture file {} does not exist".format(path))
    try:
        with Image.open(str(path)) as image:
            size, image_format = image.size, image.format
            image.verify()
    except Exception as error:  # Pillow raises several error types for a broken file.
        raise ValueError("cannot read the image {}: {}".format(path, error))
    if image_format not in PICTURE_FORMATS:
        raise ValueError("the image {} has the format {}. PowerPoint files take: {}."
                         .format(path, image_format, ", ".join(PICTURE_FORMATS)))
    return path, size


def fit_picture(slide, path, left, top, width, height):
    """Add a picture that keeps its aspect ratio and sits centered in the box.

    The box is (left, top, width, height) in inches or pptx lengths.
    Return (shape, effective_ppi). The ppi is the pixel width over the inches of the picture.
    """
    path, (pixels_w, pixels_h) = _check_image(path)
    box_left, box_top, box_w, box_h = (_emu(v) for v in (left, top, width, height))
    if box_w <= 0 or box_h <= 0:
        raise ValueError("the picture box needs a positive width and height")
    scale = min(box_w / pixels_w, box_h / pixels_h)
    pic_w = min(box_w, int(round(pixels_w * scale)))
    pic_h = min(box_h, int(round(pixels_h * scale)))
    x = box_left + (box_w - pic_w) // 2
    y = box_top + (box_h - pic_h) // 2
    shape = slide.shapes.add_picture(str(path), x, y, pic_w, pic_h)
    set_alt_text(shape, "")  # python-pptx writes the file name. Replace it.
    return shape, pixels_w / (pic_w / float(EMU_PER_INCH))


# ---------------------------------------------------------------------------
# Shape primitives
# ---------------------------------------------------------------------------

def _set_font(font, style, size, color_key):
    font.name = _value(style, "typography", "body_font")
    font.size = Pt(size)
    font.bold = False
    font.color.rgb = _rgb(style, color_key)


def _add_run(paragraph, style, text, size, color_key):
    run = paragraph.add_run()
    run.text = text
    _set_font(run.font, style, size, color_key)
    return run


def _fill_frame(frame, style, anchor, paragraphs):
    """Set up a text frame with no margins and no autofit, then add the paragraphs.

    `paragraphs` is a list of (text, size in pt, color key of the style).
    """
    frame.word_wrap = True
    frame.auto_size = MSO_AUTO_SIZE.NONE
    frame.vertical_anchor = anchor
    frame.margin_left = frame.margin_right = frame.margin_top = frame.margin_bottom = 0
    for index, (text, size, color_key) in enumerate(paragraphs):
        paragraph = frame.paragraphs[0] if index == 0 else frame.add_paragraph()
        paragraph.alignment = PP_ALIGN.LEFT
        _add_run(paragraph, style, text, size, color_key)


def _add_textbox(slide, style, name, box, paragraphs, anchor=MSO_ANCHOR.TOP):
    """Add a text box with no fill. `box` is (left, top, width, height) in EMU."""
    shape = slide.shapes.add_textbox(*box)
    shape.name = name
    _fill_frame(shape.text_frame, style, anchor, paragraphs)
    return shape


def _add_rect(slide, style, name, box, color_key):
    """Add a filled rectangle with no outline and no shadow."""
    shape = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, *box)
    shape.name = name
    shape.fill.solid()
    shape.fill.fore_color.rgb = _rgb(style, color_key)
    shape.line.fill.background()
    shape.shadow.inherit = False
    return shape


def _new_slide(prs, style, kind, headline, size=None, top=None, height=None):
    """Add a slide on the Title Only layout and fill in its title placeholder.

    The headline takes `size` points at `top` EMU with `height` EMU. The defaults come from the style.
    """
    slide = prs.slides.add_slide(prs.slide_layouts.get_by_name("Title Only"))
    slide.background.fill.solid()
    slide.background.fill.fore_color.rgb = _rgb(style, "background")
    set_slide_kind(slide, kind)
    margin = Inches(_value(style, "layout", "margin"))
    title = slide.shapes.title
    title.name = "Title"
    title.left, title.width = margin, prs.slide_width - 2 * margin
    title.top = Inches(_value(style, "layout", "title_top")) if top is None else top
    title.height = Inches(_value(style, "layout", "title_height")) if height is None else height
    size = _value(style, "typography", "title_size") if size is None else size
    _fill_frame(title.text_frame, style, MSO_ANCHOR.TOP, [(headline, size, "text")])
    return slide


def _add_source(prs, slide, style, text):
    """Add the source line below the content area. An empty or blank text adds nothing."""
    if not text.strip():
        return
    margin = Inches(_value(style, "layout", "margin"))
    box = (margin, Inches(_value(style, "layout", "content_bottom") + SOURCE_GAP),
           prs.slide_width - 2 * margin, Inches(SOURCE_HEIGHT))
    _add_textbox(slide, style, "Source", box, [(text, _value(style, "typography", "source_size"), "muted")])


def _pairs(items, noun, low, high, labels):
    """Return a list of (text, text) pairs. Raise ValueError for a wrong count or a bad pair.

    `labels` names the two parts of a pair in the error messages.
    """
    try:
        items = list(items)
    except TypeError:
        raise ValueError("the {} must be a list of pairs, got {!r}".format(noun, items))
    if not low <= len(items) <= high:
        raise ValueError("a slide needs {} to {} {}, got {}".format(low, high, noun, len(items)))
    pairs = []
    for item in items:
        if isinstance(item, str) or not hasattr(item, "__len__") or len(item) != 2:
            raise ValueError("each of the {} must be a ({}) pair, got {!r}".format(noun, ", ".join(labels), item))
        pairs.append((_text(item[0], "a " + labels[0]), _text(item[1], "a " + labels[1])))
    return pairs


# ---------------------------------------------------------------------------
# Layouts
# ---------------------------------------------------------------------------

def add_title_slide(prs, style, title, author="", affiliation="", date="", notes=None):
    """Add the title slide: the title, a short accent mark and up to three muted lines."""
    _check_style(style)
    title = _text(title, "the title")
    lines = [_optional_text(v, what) for v, what in ((author, "the author"), (affiliation, "the affiliation"),
                                                      (date, "the date"))]
    lines = [line for line in lines if line.strip()]
    notes = _optional_text(notes, "the notes")
    _check_headline(prs, style, title, "the title", TITLE_SLIDE_PT)
    slide = _new_slide(prs, style, "title", title, TITLE_SLIDE_PT, Inches(TITLE_SLIDE_TOP),
                       Inches(TITLE_SLIDE_HEIGHT))
    margin = Inches(_value(style, "layout", "margin"))
    _add_rect(slide, style, "Accent mark",
              (margin, Inches(TITLE_SLIDE_MARK_TOP), Inches(TITLE_SLIDE_MARK_WIDTH),
               Inches(_value(style, "layout", "rule_width"))), "accent")
    if lines:
        size = _value(style, "typography", "callout_size")
        _add_textbox(slide, style, "Subtitle",
                     (margin, Inches(TITLE_SLIDE_SUBTITLE_TOP), prs.slide_width - 2 * margin,
                      Inches(TITLE_SLIDE_SUBTITLE_HEIGHT)),
                     [(line, size, "muted") for line in lines])
    set_notes(slide, notes)
    return slide


def add_big_number_slide(prs, style, headline, number, label, baseline, source="", notes=None):
    """Add a slide with one large number, its label and a baseline that gives it meaning."""
    _check_style(style)
    headline = _text(headline, "the headline")
    number = _text(number, "the number")
    label = _text(label, "the label")
    baseline = _text(baseline, "the baseline")
    source = _optional_text(source, "the source")
    notes = _optional_text(notes, "the notes")
    _check_headline(prs, style, headline)
    big = _value(style, "typography", "big_number_size")
    limit = int(NUMBER_BOX_WIDTH * 72 / (DIGIT_EM * big))
    if len(number) > limit:
        raise ValueError("the number {!r} has {} characters but the number box holds {} at {} pt. "
                         "Shorten it, for example 1.2M.".format(number, len(number), limit, big))
    slide = _new_slide(prs, style, "big_number", headline)
    margin = Inches(_value(style, "layout", "margin"))
    dy = _value(style, "layout", "content_top") - FREE_FORM_REFERENCE_TOP
    box_w = Inches(NUMBER_BOX_WIDTH)
    _add_textbox(slide, style, "Number", (margin, Inches(NUMBER_TOP + dy), box_w, Inches(NUMBER_HEIGHT)),
                 [(number, big, "accent_ink")])
    _add_textbox(slide, style, "Number label",
                 (margin, Inches(NUMBER_LABEL_TOP + dy), box_w, Inches(NUMBER_LABEL_HEIGHT)),
                 [(label, _value(style, "typography", "label_size"), "text")])
    _add_rect(slide, style, "Divider",
              (Inches(DIVIDER_LEFT), Inches(DIVIDER_TOP + dy), Inches(_value(style, "layout", "rule_width")),
               Inches(DIVIDER_HEIGHT)), "accent")
    _add_textbox(slide, style, "Baseline",
                 (Inches(BASELINE_LEFT), Inches(BASELINE_TOP + dy), prs.slide_width - margin - Inches(BASELINE_LEFT),
                  Inches(BASELINE_HEIGHT)),
                 [(baseline, _value(style, "typography", "body_size"), "text")], anchor=MSO_ANCHOR.MIDDLE)
    _add_source(prs, slide, style, source)
    set_notes(slide, notes)
    return slide


def _check_section_names(sections, column_in, size):
    """Raise ValueError for a section name that needs more than one line of its column."""
    for name, _ in sections:
        if _estimate_lines(name, column_in, size) > 1:
            raise ValueError("the section name {!r} needs more than one line in a column of about {} characters "
                             "at {} pt. A section name takes one line. Shorten it."
                             .format(name, _chars_per_line(column_in, size), size))


def add_roadmap_slide(prs, style, message, sections, notes=None):
    """Add a roadmap slide. `sections` is a list of 3 to 5 (name, one_line) pairs.

    A section name takes one line of its column: about 9 characters for 5 sections, 12 for 4, 16 for 3.
    """
    _check_style(style)
    message = _text(message, "the message")
    sections = _pairs(sections, "sections", 3, 5, ("section name", "section line"))
    notes = _optional_text(notes, "the notes")
    _check_headline(prs, style, message)
    margin = Inches(_value(style, "layout", "margin"))
    gap = Inches(ROADMAP_GAP)
    width = (prs.slide_width - 2 * margin - gap * (len(sections) - 1)) // len(sections)
    body = _value(style, "typography", "body_size")
    _check_section_names(sections, width / float(EMU_PER_INCH), body + SECTION_NAME_EXTRA_PT)
    slide = _new_slide(prs, style, "roadmap", message)
    dy = _value(style, "layout", "content_top") - FREE_FORM_REFERENCE_TOP
    for index, (name, line) in enumerate(sections):
        left = margin + index * (width + gap)
        _add_rect(slide, style, "Section mark",
                  (left, Inches(SECTION_MARK_TOP + dy), width, Inches(_value(style, "layout", "rule_width"))), "accent")
        _add_textbox(slide, style, "Section name",
                     (left, Inches(SECTION_NAME_TOP + dy), width, Inches(SECTION_NAME_HEIGHT)),
                     [(name, body + SECTION_NAME_EXTRA_PT, "accent_ink")])
        _add_textbox(slide, style, "Section line",
                     (left, Inches(SECTION_LINE_TOP + dy), width, Inches(SECTION_LINE_HEIGHT)), [(line, body, "text")])
    set_notes(slide, notes)
    return slide


def add_take_home_slide(prs, style, claims, notes=None, headline=None):
    """Add a take-home slide. `claims` is a list of 1 to 3 (claim, baseline) pairs.

    The headline defaults to "Three things to remember from this talk", with the count of the claims.
    Each pair takes one text box named Claim. Paragraph 0 holds the claim and paragraph 1 its baseline,
    so the baseline sits directly under the claim however many lines the claim takes.
    A claim or baseline that needs a second line gets a taller box, and the next claim moves down.
    """
    _check_style(style)
    claims = _pairs(claims, "claims", 1, 3, ("claim", "baseline"))
    notes = _optional_text(notes, "the notes")
    if headline is None:
        headline = "{} to remember from this talk".format(TAKE_HOME_COUNTS[len(claims)])
    headline = _text(headline, "the headline")
    _check_headline(prs, style, headline)
    margin_in = _value(style, "layout", "margin")
    margin = Inches(margin_in)
    text_in = _value(style, "layout", "slide_width") - 2 * margin_in - CLAIM_INDENT
    body = _value(style, "typography", "body_size")
    small = _value(style, "typography", "callout_size")
    top = _value(style, "layout", "content_top") - FREE_FORM_REFERENCE_TOP + CLAIM_TOP
    plan = []  # (top, height) of each claim box in inches
    for claim, baseline in claims:
        height = (_estimate_lines(claim, text_in, body) * _line_height(body)
                  + _estimate_lines(baseline, text_in, small) * _line_height(small) + CLAIM_PAD)
        plan.append((top, height))
        top += height + CLAIM_BLOCK_GAP
    if plan[-1][0] + plan[-1][1] > _value(style, "layout", "content_bottom"):
        raise ValueError("the claims need more room than the slide has. Shorten them or use fewer claims.")
    slide = _new_slide(prs, style, "take_home", headline)
    text_left = margin + Inches(CLAIM_INDENT)
    text_width = prs.slide_width - margin - text_left
    mark = Inches(CLAIM_MARK_SIZE)
    for (claim, baseline), (top, height) in zip(claims, plan):
        _add_rect(slide, style, "Claim mark", (margin, Inches(top + CLAIM_MARK_OFFSET), mark, mark), "accent")
        _add_textbox(slide, style, "Claim", (text_left, Inches(top), text_width, Inches(height)),
                     [(claim, body, "text"), (baseline, small, "muted")])
    set_notes(slide, notes)
    return slide


# ---------------------------------------------------------------------------
# Evidence: table
# ---------------------------------------------------------------------------

def _evidence_box(style):
    """Return (left, top, width, height) of the evidence area in EMU."""
    return (Inches(_value(style, "layout", "margin")), Inches(_value(style, "layout", "content_top")),
            Inches(_value(style, "layout", "evidence_width")), Inches(_value(style, "layout", "evidence_height")))


def _is_int(value):
    return isinstance(value, int) and not isinstance(value, bool)


def _prepare_table(style, evidence):
    """Check a table evidence dict. Return a function that draws the table on a slide."""
    rows = evidence.get("rows")
    if not isinstance(rows, (list, tuple)) or not rows:
        raise ValueError("the table needs a non-empty list of rows")
    cells = []
    for i, row in enumerate(rows):
        if not isinstance(row, (list, tuple)):
            raise ValueError("each table row must be a list of cells, got {!r}".format(row))
        cells.append([_encodable("" if cell is None else str(cell), "the table cell (row {}, column {})".format(i, j))
                      for j, cell in enumerate(row)])
    n_rows, n_cols = len(cells), len(cells[0])
    if n_cols == 0:
        raise ValueError("the table needs at least one column")
    if any(len(row) != n_cols for row in cells):
        raise ValueError("every table row needs the same number of cells")
    size = _value(style, "typography", "callout_size")
    height_in = _value(style, "layout", "evidence_height")
    max_rows = int(height_in / (_line_height(size) + TABLE_ROW_PAD))
    if n_rows > max_rows:
        raise ValueError("a table of {} rows does not fit. At most {} rows fit the evidence area at {} pt."
                         .format(n_rows, max_rows, size))
    highlight = evidence.get("highlight")
    if highlight is not None:
        if not (isinstance(highlight, (list, tuple)) and len(highlight) == 2 and all(_is_int(i) for i in highlight)):
            raise ValueError("the table highlight must be a (row, column) pair or None, got {!r}".format(highlight))
        if not (0 <= highlight[0] < n_rows and 0 <= highlight[1] < n_cols):
            raise ValueError("the table highlight {!r} lies outside the table of {} rows and {} columns"
                             .format(tuple(highlight), n_rows, n_cols))

    def draw(slide):
        left, top, width, _ = _evidence_box(style)
        row_h = Inches(min(TABLE_ROW_INCHES, height_in / n_rows))
        shape = slide.shapes.add_table(n_rows, n_cols, left, top, width, row_h * n_rows)
        shape.name = "Evidence table"
        table = shape.table
        properties = table._tbl.tblPr
        properties.find(qn("a:tableStyleId")).text = TABLE_STYLE_ID
        properties.set("firstRow", "0")
        properties.set("bandRow", "0")
        weights = [max(TABLE_WEIGHT_MIN, min(TABLE_WEIGHT_MAX, max(len(row[j]) for row in cells)))
                   for j in range(n_cols)]
        widths = [width * weight // sum(weights) for weight in weights]
        widths[-1] += width - sum(widths)
        for index, column_width in enumerate(widths):
            table.columns[index].width = Emu(column_width)
        for i, row in enumerate(cells):
            table.rows[i].height = row_h
            accent_text = i == 0 or (highlight is not None and i == highlight[0])
            for j, text in enumerate(row):
                cell = table.cell(i, j)
                cell.vertical_anchor = MSO_ANCHOR.MIDDLE
                cell.margin_left = cell.margin_right = Inches(TABLE_CELL_MARGIN)
                cell.fill.background()
                paragraph = cell.text_frame.paragraphs[0]
                paragraph.alignment = PP_ALIGN.LEFT
                _add_run(paragraph, style, text, size, "accent_ink" if accent_text else "text")
                if i == 0:
                    cell._tc.get_or_add_tcPr().insert(0, _bottom_rule(style))
    return draw


def _bottom_rule(style):
    """Return an a:lnB element: the accent rule under a header cell."""
    line = etree.Element(qn("a:lnB"), w=str(int(Pt(TABLE_RULE_PT))))
    fill = etree.SubElement(line, qn("a:solidFill"))
    etree.SubElement(fill, qn("a:srgbClr"), val=str(_rgb(style, "accent")))
    return line


# ---------------------------------------------------------------------------
# Evidence: chart
# ---------------------------------------------------------------------------

def _check_highlight(highlight, categories, names):
    """Return the index of the highlighted category or series, or None.

    One series highlights a category by name or index. Several series highlight a series by name.
    """
    if highlight is None:
        return None
    if len(names) > 1:
        if isinstance(highlight, str) and highlight in names:
            return names.index(highlight)
        if isinstance(highlight, str):
            raise ValueError("the highlight {!r} is not a series name. Series: {}".format(highlight, ", ".join(names)))
        raise ValueError("with several series the highlight must be a series name, got {!r}".format(highlight))
    if isinstance(highlight, str) and highlight in categories:
        return categories.index(highlight)
    if _is_int(highlight) and 0 <= highlight < len(categories):
        return highlight
    raise ValueError("the highlight {!r} matches no category. Categories: {}".format(highlight, ", ".join(categories)))


_ChartSpec = namedtuple("_ChartSpec", "kind categories series key colors fmt vmax alt")


def _prepare_chart(style, evidence):
    """Check a chart evidence dict. Return a function that draws the chart on a slide."""
    chart_type = evidence.get("chart")
    if chart_type not in CHART_TYPES:
        raise ValueError("unknown chart type {!r}. Valid types: {}".format(chart_type, ", ".join(CHART_TYPES)))
    categories = evidence.get("categories")
    if not isinstance(categories, (list, tuple)) or not categories:
        raise ValueError("the chart needs a non-empty list of categories")
    categories = [_encodable(str(category), "the category at position {}".format(i))
                  for i, category in enumerate(categories)]
    raw = evidence.get("series")
    if not isinstance(raw, dict) or not raw:
        raise ValueError("the chart needs a non-empty dict of series, as {name: [numbers]}")
    series = []
    for name, values in raw.items():
        name = _text(name, "a series name")
        if not isinstance(values, (list, tuple)):
            raise ValueError("series {!r} must be a list of numbers, got {!r}".format(name, values))
        if len(values) != len(categories):
            raise ValueError("series {!r} has {} values but the chart has {} categories"
                             .format(name, len(values), len(categories)))
        for value in values:
            if isinstance(value, bool) or not isinstance(value, (int, float)) or not math.isfinite(value):
                raise ValueError("series {!r} holds {!r}, which is not a number".format(name, value))
        series.append((name, list(values)))
    if len(series) > 1 and "fmt" in evidence:
        raise ValueError("fmt has no effect on a chart with several series. It shows series names and no values. "
                         "Remove it.")
    key = _check_highlight(evidence.get("highlight"), categories, [name for name, _ in series])
    colors = None
    if len(series) > 1:
        colors = _palette(style, len(series)) if key is None else [
            _rgb(style, "accent") if i == key else _rgb(style, "context") for i in range(len(series))]
    fmt = evidence.get("fmt", "General")
    if not isinstance(fmt, str):
        raise ValueError("the chart number format must be a string, got {!r}".format(fmt))
    _encodable(fmt, "the chart number format")
    vmax = evidence.get("vmax")
    if vmax is not None:
        largest = max(max(values) for _, values in series)
        if isinstance(vmax, bool) or not isinstance(vmax, (int, float)) or vmax < largest:
            raise ValueError("vmax {!r} must be a number that is not below the largest value {}".format(vmax, largest))
    alt = _text(evidence.get("alt"), "the chart alt text")
    spec = _ChartSpec(chart_type, categories, series, key, colors, fmt, vmax, alt)
    return lambda slide: _draw_chart(slide, style, spec)


def _no_fill_properties():
    """Return a c:spPr element with no fill and no line."""
    properties = etree.Element(qn("c:spPr"))
    etree.SubElement(properties, qn("a:noFill"))
    etree.SubElement(etree.SubElement(properties, qn("a:ln")), qn("a:noFill"))
    return properties


def _palette(style, count):
    """Return `count` colors of colors.data_palette that show on the background."""
    background = _value(style, "colors", "background")
    floor = style.get("lint", {}).get("mark_contrast", {}).get("warn_below", 3.0)
    usable = [c for c in _value(style, "colors", "data_palette") if design_color.contrast_ratio(c, background) >= floor]
    if not usable:
        raise ValueError("no color of colors.data_palette has a contrast of {} with the background".format(floor))
    return [RGBColor(*design_color.parse_hex(usable[i % len(usable)])) for i in range(count)]


def _style_one_series(plot, style, bar, key, fmt):
    """One series: the key point is accent and the rest are context. Label every point."""
    series = plot.series[0]
    accent, context = _rgb(style, "accent"), _rgb(style, "context")
    count = len(list(plot.categories))
    colors = [accent if key is None or i == key else context for i in range(count)]
    if bar:
        series.format.line.fill.background()
        for index, color in enumerate(colors):
            fill = series.points[index].format.fill
            fill.solid()
            fill.fore_color.rgb = color
    else:
        series.format.line.color.rgb = accent if key is None else context
        series.format.line.width = Pt(CONTEXT_LINE_PT)
        for index, color in enumerate(colors):
            marker = series.points[index].marker
            marker.style = XL_MARKER_STYLE.CIRCLE
            marker.size = KEY_MARKER_PT if index == key else MARKER_PT
            marker.format.fill.solid()
            marker.format.fill.fore_color.rgb = color
            marker.format.line.color.rgb = color
    plot.has_data_labels = True
    labels = plot.data_labels
    labels.number_format = fmt
    labels.position = XL_LABEL_POSITION.OUTSIDE_END if bar else XL_LABEL_POSITION.ABOVE
    _set_font(labels.font, style, _value(style, "typography", "body_size"), "text")


def _style_many_series(plot, style, bar, key, colors):
    """Several series: no legend. Each series gets its name as a label at its last point."""
    count = len(list(plot.categories))
    size = _value(style, "typography", "callout_size")
    plot.has_data_labels = True
    plot.data_labels.show_value = False
    for index, series in enumerate(plot.series):
        if bar:
            series.format.fill.solid()
            series.format.fill.fore_color.rgb = colors[index]
            series.format.line.fill.background()
        else:
            series.format.line.color.rgb = colors[index]
            series.format.line.width = Pt(ACCENT_LINE_PT if index == key else CONTEXT_LINE_PT)
            series.marker.style = XL_MARKER_STYLE.NONE
        label = series.points[count - 1].data_label
        label.position = XL_LABEL_POSITION.OUTSIDE_END if bar else XL_LABEL_POSITION.RIGHT
        color_key = "accent_ink" if index == key else "text"
        paragraph = label.text_frame.paragraphs[0]
        _set_font(paragraph.font, style, size, color_key)  # python-pptx adds an empty a:defRPr here
        _add_run(paragraph, style, series.name, size, color_key)


def _reserve_label_room(chart, style, names, width_in):
    """Set a manual plot layout that leaves room at the right for the series names."""
    size = _value(style, "typography", "callout_size")
    room = min(MAX_LABEL_ROOM, (max(len(name) for name in names) * CHAR_EM * size / 72 + LABEL_PAD) / width_in)
    layout = etree.Element(qn("c:layout"))
    manual = etree.SubElement(layout, qn("c:manualLayout"))
    for tag, value in (("layoutTarget", "inner"), ("xMode", "edge"), ("yMode", "edge"), ("x", str(PLOT_LEFT)),
                       ("y", str(PLOT_TOP)), ("w", "{:.3f}".format(1 - 2 * PLOT_LEFT - room)),
                       ("h", str(PLOT_HEIGHT))):
        etree.SubElement(manual, qn("c:" + tag), val=value)
    chart._chartSpace.find(qn("c:chart")).find(qn("c:plotArea")).insert(0, layout)


def _draw_chart(slide, style, spec):
    left, top, width, height = _evidence_box(style)
    data = CategoryChartData()
    data.categories = spec.categories
    for name, values in spec.series:
        data.add_series(name, values)
    bar = spec.kind == "bar"
    frame = slide.shapes.add_chart(XL_CHART_TYPE.COLUMN_CLUSTERED if bar else XL_CHART_TYPE.LINE_MARKERS,
                                   left, top, width, height, data)
    frame.name = "Evidence chart"
    set_alt_text(frame, spec.alt)
    chart = frame.chart
    chart.has_legend = False
    chart.has_title = False
    chart._chartSpace.find(qn("c:chart")).addnext(_no_fill_properties())
    _set_font(chart.font, style, _value(style, "typography", "callout_size"), "text")
    plot = chart.plots[0]
    plot.vary_by_categories = False
    if bar:
        plot.gap_width = BAR_GAP_WIDTH
    axis = chart.value_axis
    axis.visible = False
    axis.has_major_gridlines = False
    axis.has_minor_gridlines = False
    if bar and all(value >= 0 for _, values in spec.series for value in values):
        axis.minimum_scale = 0
    if spec.vmax is not None:
        axis.maximum_scale = spec.vmax
    category_axis = chart.category_axis
    category_axis.has_major_gridlines = False
    category_axis.format.line.color.rgb = _rgb(style, "muted")
    _set_font(category_axis.tick_labels.font, style, _value(style, "typography", "callout_size"), "text")
    if len(spec.series) == 1:
        _style_one_series(plot, style, bar, spec.key, spec.fmt)
    else:
        _style_many_series(plot, style, bar, spec.key, spec.colors)
        if not bar:
            _reserve_label_room(chart, style, [name for name, _ in spec.series], width / float(EMU_PER_INCH))
    chart._chartSpace.find(qn("c:chart")).find(qn("c:plotArea")).append(_no_fill_properties())


# ---------------------------------------------------------------------------
# Assertion-evidence layout
# ---------------------------------------------------------------------------

def _prepare_evidence(style, evidence):
    """Check the evidence dict. Return a function that draws it on a slide."""
    kinds = ", ".join(EVIDENCE_KINDS)
    if not isinstance(evidence, dict) or "kind" not in evidence:
        raise ValueError("the evidence must be a dict with a 'kind' key. Valid kinds: {}".format(kinds))
    kind = evidence["kind"]
    if kind == "picture":
        path, _ = _check_image(evidence.get("path"))
        alt = _text(evidence.get("alt"), "the picture alt text")

        def draw(slide):
            shape, _ = fit_picture(slide, path, *_evidence_box(style))
            shape.name = "Evidence picture"
            set_alt_text(shape, alt)
        return draw
    if kind == "table":
        return _prepare_table(style, evidence)
    if kind == "chart":
        return _prepare_chart(style, evidence)
    raise ValueError("unknown evidence kind {!r}. Valid kinds: {}".format(kind, kinds))


def _callout_box(prs, style):
    """Return (rule left, text left, text width) of a callout in EMU."""
    left = Inches(_value(style, "layout", "callout_left"))
    text_left = left + Inches(CALLOUT_TEXT_OFFSET)
    return left, text_left, prs.slide_width - Inches(_value(style, "layout", "margin")) - text_left


def _callout_list(prs, style, callouts):
    """Return the callouts as a list of strings. Raise ValueError for more than two or a text that is too long."""
    if callouts is None:
        return []
    if isinstance(callouts, str):
        callouts = [callouts]
    try:
        items = list(callouts)
    except TypeError:
        raise ValueError("the callouts must be a list of strings, got {!r}".format(callouts))
    if len(items) > 2:
        raise ValueError("a slide takes at most 2 callouts, got {}".format(len(items)))
    width_in = _callout_box(prs, style)[2] / float(EMU_PER_INCH)
    for item in items:
        _check_lines(_text(item, "a callout"), width_in, _value(style, "typography", "callout_size"),
                     CALLOUT_LINES, "a callout")
    return items


def _add_callouts(prs, slide, style, callouts):
    """Add each callout as a text box with a thin accent rule at its left. Neither has a fill."""
    left, text_left, text_width = _callout_box(prs, style)
    rule_width = Inches(_value(style, "layout", "rule_width"))
    size = _value(style, "typography", "callout_size")
    top = _value(style, "layout", "content_top") + CALLOUT_TOP_OFFSET
    for text in callouts:
        _add_rect(slide, style, "Callout rule", (left, Inches(top), rule_width, Inches(CALLOUT_HEIGHT)), "accent")
        _add_textbox(slide, style, "Callout", (text_left, Inches(top), text_width, Inches(CALLOUT_HEIGHT)),
                     [(text, size, "text")])
        top += CALLOUT_STEP


def add_assertion_evidence_slide(prs, style, headline, evidence, callouts=(), source="", notes=None,
                                 backup=False):
    """Add a slide with a sentence headline over one piece of evidence.

    `evidence` is one of these dicts:
      {"kind": "picture", "path": str, "alt": str}
      {"kind": "table", "rows": [[str, ...], ...], "highlight": (row, col) or None}
      {"kind": "chart", "chart": "bar" or "line", "categories": [...], "series": {name: [numbers]},
       "highlight": a category name or index for one series, a series name for several, or None,
       "alt": str}
    A picture and a chart need a non-empty "alt". A table needs none and gets no alt text.
    The whole row of a table highlight shows accent text.
    A chart dict can add "vmax", the top of the value axis. A chart with one series can add "fmt",
    the number format of its data labels. A chart with several series shows each series name at its
    last point and no values, so "fmt" raises ValueError there.
    A slide takes at most two callouts of two lines each. backup=True gives the slide the kind "backup".
    """
    _check_style(style)
    headline = _text(headline, "the headline")
    draw = _prepare_evidence(style, evidence)
    callouts = _callout_list(prs, style, callouts)
    source = _optional_text(source, "the source")
    notes = _optional_text(notes, "the notes")
    _check_headline(prs, style, headline)
    slide = _new_slide(prs, style, "backup" if backup else "assertion_evidence", headline)
    draw(slide)
    _add_callouts(prs, slide, style, callouts)
    _add_source(prs, slide, style, source)
    set_notes(slide, notes)
    return slide
